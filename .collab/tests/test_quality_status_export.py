# -*- coding: utf-8 -*-
import io
import os
import tempfile
import unittest

import openpyxl

import database as db
import report_builder


class QualityStatusExportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        os.environ.setdefault("SECRET_KEY", "task-97-quality-status-test-only")

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.old_db_path = db.DB_PATH
        db.DB_PATH = os.path.join(self.tmp.name, "test.db")
        db.init_db()
        import app as qms_app
        self.qms = qms_app
        self.qms.app.config.update(TESTING=True, WTF_CSRF_ENABLED=False,
                                   SECRET_KEY="task-97-quality-status-test-only")

    def tearDown(self):
        db.DB_PATH = self.old_db_path
        self.tmp.cleanup()

    def _inspection(self, material_no, major="표준부품", middle="기타", *,
                    status="approved", approval_type="normal", date="2026-09-10",
                    qty=100, aql="1.5", sample=20, snapshot_set=1,
                    measured="1,2", supplier="업체", po="PO-1", remarks=(None, None, None)):
        db.upsert_material(material_no, material_no)
        db.update_material_classification(material_no, major, middle)
        conn = db.get_conn()
        conn.execute("""INSERT OR REPLACE INTO specs
            (material_no, material_name, item_name, spec_display, judge_type,
             lower_limit, upper_limit, inspect_method, aql, item_order)
            VALUES (?, ?, 'B', '외관 검사', 'ok_ng', NULL, NULL, '육안', ?, 1)""",
                     (material_no, material_no, aql))
        cur = conn.execute("""INSERT INTO inspections
            (material_no, material_name, supplier, po_number, receive_date, inspect_date,
             inspector, quantity, overall_result, status, approval_type,
             remark_inspector, remark_manager, remark_approver)
            VALUES (?, ?, ?, ?, ?, ?, '검사자', ?, ?, ?, ?, ?, ?, ?)""",
            (material_no, material_no, supplier, po, date, date, qty,
             "합격" if approval_type == "normal" else "불합격", status, approval_type,
             remarks[0], remarks[1], remarks[2]))
        inspection_id = cur.lastrowid
        conn.execute("""INSERT INTO inspection_items
            (inspection_id, item_name, measured_value, result, part_material_no,
             aql_snapshot, sample_qty_snapshot, ac_allowance_snapshot, aql_snapshot_set)
            VALUES (?, 'B', ?, '합격', ?, ?, ?, 0, ?)""",
            (inspection_id, measured, material_no, aql, sample, snapshot_set))
        conn.commit()
        conn.close()
        return inspection_id

    def _ncr(self, inspection_id, status, sample, defect, special_note=None,
             defect_description="NCR 불량 원문"):
        conn = db.get_conn()
        cur = conn.execute("""INSERT INTO ncr
            (ncr_no, inspection_id, status, sample_qty, defect_qty,
             defect_description, special_note, issued_date)
            VALUES (?, ?, ?, ?, ?, ?, ?, '2026-09-11')""",
            (f"NCR-{inspection_id}-{status}-{sample}", inspection_id, status,
             sample, defect, defect_description, special_note))
        conn.commit()
        result = cur.lastrowid
        conn.close()
        return result

    def _improvement(self, inspection_id, status, sample, defect, detail="개선요청 원문"):
        conn = db.get_conn()
        cur = conn.execute("""INSERT INTO improvement_requests
            (request_no, inspection_id, status, sample_qty, defect_qty,
             request_detail, issued_date)
            VALUES (?, ?, ?, ?, ?, ?, '2026-09-11')""",
            (f"IMP-{inspection_id}-{status}-{sample}", inspection_id, status,
             sample, defect, detail))
        conn.commit()
        result = cur.lastrowid
        conn.close()
        return result

    def test_improvement_quantity_columns_create_and_update(self):
        req_id, _ = db.create_improvement_request(
            None, "M", "자재", "업체", None, None, "10", "외관", None,
            "요청 원문", "담당자", "작성자", "2026-09-10", sample_qty=5, defect_qty=1)
        row = db.get_improvement_request(req_id)
        self.assertEqual((row["sample_qty"], row["defect_qty"]), (5, 1))
        db.update_improvement_request(
            req_id, "M", "자재", "업체", None, None, "10", "외관", None,
            "수정 원문", "담당자", "2026-09-11", sample_qty=8, defect_qty=2)
        row = db.get_improvement_request(req_id)
        self.assertEqual((row["sample_qty"], row["defect_qty"], row["status"]), (8, 2, "draft"))

    def test_improvement_quantity_parser_rejects_unparseable_digit_character(self):
        with self.qms.app.test_request_context(
                "/improvement/new", method="POST",
                data={"sample_qty": "²", "defect_qty": "0"}):
            sample_qty, defect_qty, error = self.qms._improvement_quantities_from_form()
        self.assertIsNone(sample_qty)
        self.assertIsNone(defect_qty)
        self.assertEqual(error, "표본수량은 0 이상의 정수로 입력해줘.")

    def test_improvement_quantity_parser_rejects_sqlite_integer_overflow(self):
        with self.qms.app.test_request_context(
                "/improvement/new", method="POST",
                data={"sample_qty": "999999999999999999999999999999", "defect_qty": "0"}):
            sample_qty, defect_qty, error = self.qms._improvement_quantities_from_form()
        self.assertIsNone(sample_qty)
        self.assertIsNone(defect_qty)
        self.assertEqual(error, "표본수량이 입력 가능한 범위를 초과했어.")

    def test_existing_improvement_table_receives_quantity_columns(self):
        conn = db.get_conn()
        conn.execute("DROP TABLE improvement_requests")
        conn.execute("""CREATE TABLE improvement_requests (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            request_no TEXT UNIQUE,
            status TEXT DEFAULT 'draft'
        )""")
        conn.commit()
        conn.close()

        db.init_db()
        conn = db.get_conn()
        columns = {row[1] for row in conn.execute(
            "PRAGMA table_info(improvement_requests)").fetchall()}
        conn.close()
        self.assertTrue({"sample_qty", "defect_qty"}.issubset(columns))

    def test_snapshot_is_immutable_and_old_rows_use_only_recorded_exterior_values(self):
        self._inspection("SNAP", sample=20, measured="1")
        self._inspection("OLD", sample=None, snapshot_set=0, measured="A,B,C")
        self._inspection("OUT", date="2026-10-01", sample=99)
        self._inspection("SUPER", status="superseded", sample=99)
        # 현재 규격을 체크검사로 바꿔도 과거 SNAP/OLD 표본수와 검사방법은 달라지면 안 된다.
        conn = db.get_conn()
        conn.execute("UPDATE specs SET aql='체크검사' WHERE material_no IN ('SNAP','OLD')")
        conn.commit(); conn.close()

        rows = db.quality_status_export_rows("2026-09-01", "2026-09-30")
        by_no = {r["material_no"]: r for r in rows}
        self.assertEqual(set(by_no), {"SNAP", "OLD"})
        self.assertEqual((by_no["SNAP"]["sample_qty"], by_no["SNAP"]["inspection_method"]),
                         (20, "샘플링 검사"))
        self.assertEqual((by_no["OLD"]["sample_qty"], by_no["OLD"]["sample_source"]),
                         (3, "외관 실제 기록"))

    def test_ncr_then_improvement_then_exterior_priority_and_null_fallback(self):
        special = self._inspection("SPECIAL", approval_type="special", sample=30,
                                   remarks=("검사자 원문", "검토자 원문", "승인자 원문"))
        self._ncr(special, "draft", 99, 9)
        self._ncr(special, "confirmed", 12, 2, special_note="특채사유 원문")
        self._ncr(special, "draft", 100, 10)
        self._improvement(special, "sent", 20, 3)

        fallback = self._inspection("FALLBACK", approval_type="special", sample=40)
        self._ncr(fallback, "confirmed", None, None)
        self._improvement(fallback, "sent", 7, 1)

        rows = {r["material_no"]: r for r in
                db.quality_status_export_rows("2026-09-01", "2026-09-30")}
        self.assertEqual((rows["SPECIAL"]["sample_qty"], rows["SPECIAL"]["defect_qty"]), (12, 2))
        self.assertTrue(rows["SPECIAL"]["sample_source"].startswith("NCR "))
        self.assertEqual(rows["SPECIAL"]["special_content"], "특채사유 원문\nNCR 불량 원문\n개선요청 원문")
        self.assertEqual(rows["SPECIAL"]["notes"],
                         "검사자: 검사자 원문\n검토자: 검토자 원문\n최종승인자: 승인자 원문")
        self.assertEqual((rows["FALLBACK"]["sample_qty"], rows["FALLBACK"]["defect_qty"]), (7, 1))
        self.assertTrue(rows["FALLBACK"]["sample_source"].startswith("개선요청서 "))

    def test_all_inspection_classifications_and_historic_check_snapshot(self):
        self._inspection("HOUSING", major="하우징", middle="기타", aql="체크검사", sample=5)
        self._inspection("MECH", major="기구물", middle="기타", aql="체크검사", sample=5)
        self._inspection("RES", major="전기", middle="고저항", aql="체크검사", sample=5)
        self._inspection("VI", major="전기", middle="VI", aql="체크검사", sample=5)
        self._inspection("CHECK", major="기타 소자류", middle="IC", aql="체크검사", sample=6)
        rows = {r["material_no"]: r for r in
                db.quality_status_export_rows("2026-09-01", "2026-09-30")}
        for key in ("HOUSING", "MECH", "RES", "VI"):
            self.assertEqual(rows[key]["inspection_method"], "전수 검사")
        self.assertEqual((rows["CHECK"]["inspection_method"], rows["CHECK"]["sample_qty"]),
                         ("체크 검사", 6))

    def test_unassigned_classification_uses_user_facing_label(self):
        self._inspection("UNASSIGNED", supplier="미지정업체", sample=5)
        conn = db.get_conn()
        conn.execute("DELETE FROM material_classification_map WHERE material_no = 'UNASSIGNED'")
        conn.commit()
        conn.close()

        rows = db.quality_status_export_rows("2026-09-01", "2026-09-30")
        row = next(item for item in rows if item["material_no"] == "UNASSIGNED")
        self.assertEqual((row["major_category"], row["middle_category"]), ("미지정", "미지정"))
        filtered = db.quality_status_export_rows(
            "2026-09-01", "2026-09-30",
            major_category=db.UNASSIGNED_CLASSIFICATION)
        self.assertEqual([item["material_no"] for item in filtered], ["UNASSIGNED"])

        buf = report_builder.build_quality_status_excel(
            rows, "2026-09-01", "2026-09-30", "전체 조건")
        wb = openpyxl.load_workbook(buf, data_only=True)
        detail = wb["품목상세"]
        exported = [detail.cell(r, c).value
                    for r in range(5, detail.max_row + 1)
                    for c in (3, 4)]
        self.assertIn("미지정", exported)
        self.assertNotIn(db.UNASSIGNED_CLASSIFICATION, exported)

    def test_filter_export_workbook_totals_and_five_sheet_contract(self):
        self._inspection("A", supplier="선택업체", qty=10, sample=5)
        self._inspection("A2", supplier="선택업체", middle="볼트", qty=15, sample=7)
        self._inspection("B", supplier="제외업체", qty=20, sample=8)
        rows = db.quality_status_export_rows(
            "2026-09-01", "2026-09-30", supplier=["선택업체"], states=["합격"])
        self.assertEqual([r["material_no"] for r in rows], ["A", "A2"])
        buf = report_builder.build_quality_status_excel(rows, "2026-09-01", "2026-09-30",
                                                        "협력사: 선택업체 / 판정: 합격")
        wb = openpyxl.load_workbook(buf, data_only=False)
        self.assertEqual(wb.sheetnames, ["수입검사 LIST", "협력사별 현황", "대분류별 요약", "품목상세", "분류기준"])
        detail = wb["품목상세"]
        self.assertEqual(detail["I5"].value, 10)
        self.assertEqual(detail["K5"].value, 5)
        self.assertEqual(detail["G5"].value, "A")
        self.assertEqual(wb["수입검사 LIST"]["A1"].value,
                         "수입검사 현황 LIST (2026-09-01 ~ 2026-09-30)")
        expected_panes = {
            "수입검사 LIST": "E5", "협력사별 현황": "C5", "대분류별 요약": "C5",
            "품목상세": "H5", "분류기준": None,
        }
        expected_filters = {
            "수입검사 LIST": "A4:Q6", "협력사별 현황": None,
            "대분류별 요약": None, "품목상세": "A4:V6", "분류기준": None,
        }
        for sheet_name in wb.sheetnames:
            self.assertEqual(wb[sheet_name].freeze_panes, expected_panes[sheet_name])
            self.assertEqual(wb[sheet_name].auto_filter.ref, expected_filters[sheet_name])

        category = wb["대분류별 요약"]
        self.assertEqual(
            [(category.cell(row, 1).value, category.cell(row, 2).value) for row in range(5, 9)],
            [("표준부품", "기타"), ("표준부품", "볼트"), ("표준부품 소계", None), ("총계", None)],
        )
        self.assertEqual((category["D7"].value, category["F7"].value, category["G7"].value,
                          category["I7"].value), (2, 2, 25, 12))
        self.assertEqual((category["D8"].value, category["F8"].value, category["G8"].value,
                          category["I8"].value), (2, 2, 25, 12))

    def test_login_filter_download_e2e(self):
        self._inspection("E2E-A", supplier="선택업체", qty=11, sample=5)
        self._inspection("E2E-B", supplier="제외업체", qty=22, sample=8)
        db.create_user("exporter", "pw", "출력담당", "inspect_history")
        user = db.get_user_by_username("exporter")
        client = self.qms.app.test_client()
        with client.session_transaction() as session:
            session["user_id"] = user["id"]
            session["last_seen"] = 9999999999
        response = client.get(
            "/dashboard/export.xlsx?period=custom&start=2026-09-01&end=2026-09-30&supplier=선택업체")
        self.assertEqual(response.status_code, 200)
        disposition = response.headers.get("Content-Disposition", "")
        self.assertIn("filename*=UTF-8''", disposition)
        self.assertIn("LIST_2026-09-01_2026-09-30.xlsx", disposition)
        wb = openpyxl.load_workbook(io.BytesIO(response.data), data_only=True)
        self.assertEqual(wb.sheetnames, ["수입검사 LIST", "협력사별 현황", "대분류별 요약", "품목상세", "분류기준"])
        self.assertEqual(wb["품목상세"]["G5"].value, "E2E-A")
        self.assertNotIn("E2E-B", [cell.value for cell in wb["품목상세"]["G"]])


if __name__ == "__main__":
    unittest.main()
