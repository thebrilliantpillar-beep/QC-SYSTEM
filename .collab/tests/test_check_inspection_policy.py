# -*- coding: utf-8 -*-
import os
import tempfile
import unittest

import database as db


class CheckInspectionPolicyTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        os.environ.setdefault("SECRET_KEY", "check-inspection-policy-test")

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.old_db_path = db.DB_PATH
        db.DB_PATH = os.path.join(self.tmp.name, "test.db")
        db.init_db()

        # app은 database 모듈 객체를 공유하므로 import 전/후 모두 위 임시 DB_PATH를 쓴다.
        import app as qms_app
        self.qms = qms_app
        self.qms.app.config.update(TESTING=True, WTF_CSRF_ENABLED=False)

    def tearDown(self):
        db.DB_PATH = self.old_db_path
        self.tmp.cleanup()

    def _material(self, no, major, middle, aql="4", spec_count=1):
        db.upsert_material(no, no)
        db.update_material_classification(no, major, middle)
        for idx in range(spec_count):
            db.add_spec(no, no, chr(65 + idx), "10±1", "numeric", 9, 11,
                        "버니어", aql, idx + 1)

    def _intake(self, material_no, quantity):
        conn = db.get_conn()
        cur = conn.execute("""
            INSERT INTO intake_list
                (material_no, quantity, supplier, receive_date, po_number, status)
            VALUES (?, ?, '업체', '2026-10-09', ?, '대기')
        """, (material_no, quantity, f"PO-{material_no}"))
        intake_id = cur.lastrowid
        conn.commit()
        conn.close()
        return intake_id

    def _login_session(self, client):
        db.create_user("tester", "pw", "검사자", "inspect_input,users,approve")
        user = db.get_user_by_username("tester")
        with client.session_transaction() as sess:
            sess["user_id"] = user["id"]
            sess["last_seen"] = 9999999999
        return user

    def _snapshot_inspection(self, material_no, quantity, aql, sample_qty, ac,
                             result="합격", actual_time_sec=10, created_by_user_id=None):
        spec = db.get_specs_by_material(material_no)[0]
        return db.create_inspection(
            {"material_no": material_no, "material_name": material_no, "supplier": "업체",
             "po_number": f"PO-{material_no}", "receive_date": "2026-10-09",
             "inspect_date": "2026-10-09", "inspector": "검사자", "quantity": quantity},
            [{"item_name": spec["item_name"], "measured_value": "12" if result == "불합격" else "10",
              "result": result, "part_material_no": material_no,
              "aql_snapshot": aql, "sample_qty_snapshot": sample_qty,
              "ac_allowance_snapshot": ac}],
            "합격" if result == "합격" else "검토필요",
            actual_time_sec=actual_time_sec, created_by_user_id=created_by_user_id)

    def test_exact_classification_matching_and_every_spec_gets_fixed_five(self):
        cases = (
            ("A", "기타 소자류", "아무 중분류", True),
            ("B", "조립부품류 (소모품)", "아무 중분류", True),
            ("C", "상품", "엘보/부싱 액세서리", True),
            ("D", "상품", "엘보/부싱", False),
            ("E", "조립부품류(소모품)", "아무 중분류", False),
        )
        for no, major, middle, _matched in cases:
            self._material(no, major, middle, spec_count=3)

        for no, _major, _middle, matched in cases:
            built = self.qms.build_specs_with_sample(db.get_specs_by_material(no), 1000)
            self.assertEqual(len(built), 3)
            if matched:
                self.assertTrue(all(s["aql"] == "체크검사" for s in built))
                self.assertTrue(all(s["sample_qty"] == 5 for s in built))
                self.assertTrue(all(s["ac_allowance"] == 0 for s in built))
                self.assertTrue(all(s["fixed_aql"] for s in built))
            else:
                self.assertTrue(all(s["aql"] == "4" for s in built))

    def test_check_sample_quantity_is_min_of_intake_and_five(self):
        self._material("A", "기타 소자류", "중", spec_count=1)
        specs = db.get_specs_by_material("A")
        for quantity, expected in ((1, 1), (4, 4), (5, 5), (6, 5), (99999, 5)):
            built = self.qms.build_specs_with_sample(specs, quantity)
            self.assertEqual(built[0]["sample_qty"], expected)
            self.assertEqual(built[0]["ac_allowance"], 0)

    def test_historical_snapshot_survives_category_and_spec_changes(self):
        self._material("A", "기타 소자류", "중", spec_count=1)
        spec = self.qms.build_specs_with_sample(db.get_specs_by_material("A"), 100)[0]
        inspection_id = db.create_inspection(
            {"material_no": "A", "material_name": "A", "supplier": "업체",
             "po_number": "PO", "receive_date": "2026-10-09", "inspect_date": "2026-10-09",
             "inspector": "검사자", "quantity": 100},
            [{"item_name": spec["item_name"], "measured_value": "10", "result": "합격",
              "part_material_no": "A", **self.qms._snapshot_fields(spec)}],
            "합격")

        db.update_material_classification("A", "상품", "다른 중분류")
        db.update_spec_item(spec["id"], "A", "10±1", "numeric", 9, 11, "버니어", "0.65", 1)
        _header, items = db.get_inspection(inspection_id)
        self.assertEqual((items[0]["aql"], items[0]["sample_qty"], items[0]["ac_allowance"]),
                         ("체크검사", 5, 0))
        rebuilt = self.qms.build_specs_with_sample(
            db.get_specs_by_material("A"), 100, snapshot_items=items)
        self.assertEqual((rebuilt[0]["aql"], rebuilt[0]["sample_qty"]), ("체크검사", 5))

    def test_legacy_null_snapshot_is_frozen_and_new_reinspection_uses_current_policy(self):
        self._material("A", "상품", "다른 중분류", aql=None, spec_count=1)
        spec = db.get_specs_by_material("A")[0]
        inspection_id = db.create_inspection(
            {"material_no": "A", "material_name": "A", "supplier": "업체",
             "po_number": "PO", "receive_date": "2026-10-09", "inspect_date": "2026-10-09",
             "inspector": "검사자", "quantity": 8},
            [{"item_name": "A", "measured_value": "10", "result": "합격",
              "part_material_no": "A"}], "합격")
        self.assertEqual(self.qms.ensure_inspection_aql_snapshot_migration(), 1)

        db.update_material_classification("A", "상품", "엘보/부싱 액세서리")
        db.update_spec_item(spec["id"], "A", "10±1", "numeric", 9, 11, "버니어", "4", 1)
        _header, old_items = db.get_inspection(inspection_id)
        self.assertEqual(old_items[0]["aql_snapshot_set"], 1)
        self.assertIsNone(old_items[0]["aql"])
        self.assertIsNone(old_items[0]["sample_qty"])

        new_specs = self.qms.build_specs_with_sample(db.get_specs_by_material("A"), 8)
        self.assertEqual((new_specs[0]["aql"], new_specs[0]["sample_qty"],
                          new_specs[0]["ac_allowance"]), ("체크검사", 5, 0))

    def test_normal_form_and_auto_batch_share_policy_and_render_fixed_label(self):
        self._material("NORMAL", "기타 소자류", "중", spec_count=1)
        self._material("AUTO", "조립부품류 (소모품)", "중", spec_count=1)
        normal_id = self._intake("NORMAL", 6)
        auto_id = self._intake("AUTO", 4)
        client = self.qms.app.test_client()
        self._login_session(client)

        page = client.get(f"/inspect/start/{normal_id}")
        self.assertEqual(page.status_code, 200)
        html = page.get_data(as_text=True)
        self.assertIn("체크검사", html)
        self.assertIn('class="aql-fixed"', html)
        self.assertNotIn('id="aql_', html)
        self.assertIn(f'id="sqty_{db.get_specs_by_material("NORMAL")[0]["id"]}"><b>5</b>', html)

        spec = db.get_specs_by_material("NORMAL")[0]
        form = {"inspect_date": "2026-10-09", "actual_time_sec": "10"}
        for idx in range(1, 6):
            form[f"item_{spec['id']}_{idx}"] = "10"
        response = client.post(f"/inspect/start/{normal_id}", data=form)
        self.assertEqual(response.status_code, 302)

        response = client.post("/inspect/auto-batch", data={
            "intake_ids": str(auto_id), "inspect_date": "2026-10-09",
            "time_min_m": "0", "time_min_s": "1", "time_max_m": "0", "time_max_s": "1",
        })
        self.assertEqual(response.status_code, 302)

        conn = db.get_conn()
        rows = conn.execute("SELECT id, material_no FROM inspections ORDER BY id").fetchall()
        conn.close()
        self.assertEqual({r["material_no"] for r in rows}, {"NORMAL", "AUTO"})
        snapshots = {}
        for row in rows:
            _header, items = db.get_inspection(row["id"])
            snapshots[row["material_no"]] = (items[0]["aql"], items[0]["sample_qty"],
                                               items[0]["ac_allowance"])
        self.assertEqual(snapshots["NORMAL"], ("체크검사", 5, 0))
        self.assertEqual(snapshots["AUTO"], ("체크검사", 4, 0))

        normal_inspection_id = next(r["id"] for r in rows if r["material_no"] == "NORMAL")
        detail = client.get(f"/inspection/{normal_inspection_id}")
        approval = client.get(f"/approve/{normal_inspection_id}")
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(approval.status_code, 200)
        for rendered in (detail.get_data(as_text=True), approval.get_data(as_text=True)):
            self.assertIn("체크검사", rendered)
            self.assertIn("Ac0/Re1", rendered)

    def test_report_sheet_shows_check_label_five_and_zero_defect(self):
        from openpyxl import Workbook
        import report_builder

        ws = Workbook().active
        result = {
            "item_name": "A", "aql": "체크검사", "sample_qty": 5,
            "ac_allowance": 0, "values": [10] * 5, "max": 10, "min": 10,
            "verdict": "합격", "gauge_expiry": None, "spec_display": "10±1",
            "inspect_method": "버니어", "lower_limit": 9, "upper_limit": 11,
        }
        report_builder._fill_sheet(ws, "A", "자재", {"qty": 100}, [result], "합격")
        self.assertEqual(ws["C9"].value, "[체크검사]\n무결점")
        self.assertEqual(ws["D9"].value, 5)

    def test_edit_total_time_uses_stored_snapshot_not_current_classification(self):
        self._material("EDIT", "상품", "다른 중분류", aql="4", spec_count=1)
        client = self.qms.app.test_client()
        user = self._login_session(client)
        inspection_id = self._snapshot_inspection(
            "EDIT", 100, "4", 8, 1, created_by_user_id=user["id"])
        db.update_material_classification("EDIT", "기타 소자류", "중")
        spec = db.get_specs_by_material("EDIT")[0]

        response = client.post(f"/inspection/{inspection_id}/edit", data={
            "inspect_date": "2026-10-09", "actual_time_sec": "10",
            f"item_{spec['id']}_1": "10",
        })
        self.assertEqual(response.status_code, 302)
        header, items = db.get_inspection(inspection_id)
        self.assertEqual(header["total_time_sec"], 80)  # 저장 당시 8개 × 10초
        self.assertEqual((items[0]["aql"], items[0]["sample_qty"]), ("4", 8))

    def test_approved_history_backfill_uses_snapshot_for_time_and_pass_quantity(self):
        self._material("HISTORY", "상품", "다른 중분류", aql="4", spec_count=1)
        inspection_id = self._snapshot_inspection("HISTORY", 100, "4", 8, 1)
        conn = db.get_conn()
        conn.execute("""UPDATE inspections
                           SET status='approved', approval_type='normal', total_time_sec=NULL,
                               remark_inspector='검사 수량 8개 중 1개 불량'
                         WHERE id=?""", (inspection_id,))
        conn.commit()
        conn.close()
        db.update_material_classification("HISTORY", "기타 소자류", "중")

        with self.qms.app.test_request_context("/history/approved"):
            rows, _filters = self.qms._collect_approval_history()
        row = next(r for r in rows if r["id"] == inspection_id)
        self.assertEqual(row["max_sample"], 8)
        self.assertEqual(row["pass_count"], 7)
        self.assertEqual(row["total_time_sec"], 80)

    def test_ncr_auto_sample_prefers_snapshot_and_only_legacy_falls_back(self):
        self._material("NCR", "기타 소자류", "중", aql="4", spec_count=1)
        client = self.qms.app.test_client()
        self._login_session(client)
        inspection_id = self._snapshot_inspection("NCR", 100, "4", 3, 0, result="불합격")

        response = client.get(f"/ncr/new/{inspection_id}")
        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True)
        self.assertIn('name="sample_qty" min="0" value="3"', html)

        # 스냅샷 도입 전 구형 행만 AQL 표에 따른 8개로 계산한다.
        conn = db.get_conn()
        conn.execute("""UPDATE inspection_items
                           SET aql_snapshot_set=0, aql_snapshot=NULL,
                               sample_qty_snapshot=NULL, ac_allowance_snapshot=NULL
                         WHERE inspection_id=?""", (inspection_id,))
        conn.commit()
        conn.close()
        response = client.get(f"/ncr/new/{inspection_id}")
        self.assertEqual(response.status_code, 200)
        self.assertIn('name="sample_qty" min="0" value="8"', response.get_data(as_text=True))


if __name__ == "__main__":
    unittest.main()
