# -*- coding: utf-8 -*-
# Copyright (c) 2026 윤주호. All rights reserved.
# 무단 복제·배포·수정을 금합니다.
"""DB 초기화 및 데이터 접근 — SQLite (파일 하나 = DB 전체)"""
import functools, hashlib, json, re, sqlite3, os, time, uuid
from datetime import datetime as _dt, timezone as _tz, timedelta as _td
from decimal import Decimal as _Dec, InvalidOperation as _DecInvalid

# 검사자 비고의 표준 불량 문구 파싱 — app.py의 parse_defect_counts와 동일 패턴
_DEFECT_RE = re.compile(r"검사\s*수량\s*(\d+)\s*개\s*중\s*(\d+)\s*개\s*불량")

DATA_DIR = os.environ.get("DATA_DIR", os.path.dirname(__file__))
os.makedirs(DATA_DIR, exist_ok=True)
DB_PATH = os.path.join(DATA_DIR, "iqc.db")

# ---------- 2026-09-20 확장: 출고 항목 품질확인 9종 + 전체판정 ----------
# 사용자가 준 실제 회사 "출고 내역서" 서식 헤더 순서 그대로. 배치(차수) 단위가 아니라
# 항목(S/N, outbound_items 한 행) 단위다 — 실제 첨부 엑셀을 openpyxl로 직접 읽어서 확인함.
# 2026-09-20: 5종 → 9종으로 확장 (스티커/도장/케이블타이/인디케이터 추가, 항목명 일부 변경)
OUTBOUND_CHECK_FIELDS = [
    "check_tie", "check_qr", "check_wrap", "check_rst",
    "check_sticker", "check_paint", "check_access",
    "check_cable", "check_indicator",
]
# 2026-10-07 사용자 요청: 화면(새 항목 입력 카드·저장 항목 세부·출고 승인 상세)에서는 케이블 타이
# 식별(155V)을 맨 위에 보여준다(검사자가 155V 여부부터 고르도록). OUTBOUND_CHECK_FIELDS 자체의
# 순서는 회사 "출고 내역서" 서식의 열 순서라서 바꾸지 않는다 — 엑셀(report_builder)은 그 목록을
# 그대로 쓰고, 화면만 이 표시 순서를 쓴다.
OUTBOUND_CHECK_DISPLAY_ORDER = ["check_cable"] + [
    f for f in OUTBOUND_CHECK_FIELDS if f != "check_cable"
]
OUTBOUND_CHECK_LABELS = {
    "check_tie": "볼트/너트 체결 상태 (가대/탱크 다리)",
    "check_qr": "QR 번호 부착 리셉터클 캡 결합",
    "check_wrap": "하우징 커버 포장 상태",
    "check_rst": "RST단자 나무판 결착",
    "check_sticker": "스티커 부착 상태",
    "check_paint": "도장 상태",
    "check_access": "부속품 유무 확인",
    "check_cable": "케이블 타이 식별 (155V)",
    "check_indicator": "INDICATOR 상태 확인",
}
# 2026-09-20: check_cable은 "해당없음" 포함 4종, 나머지는 3종.
# 2026-09-23: check_cable 제외 8종도 "해당없음"까지 4종(PASS/FAIL/SPECIAL/해당없음)으로 확장.
# "해당없음"은 FAIL/SPECIAL이 아니어서 자동판정에서 PASS와 동일하게 취급됨.
OUTBOUND_RESULT_VALUES = ("PASS", "FAIL", "SPECIAL", "해당없음")
# 필드별로 실제 허용되는 값(위 OUTBOUND_RESULT_VALUES는 "존재하는 값 전체"일 뿐,
# 필드마다 어떤 값이 유효한지는 이걸로 걸러야 한다 — check_cable만 3종(PASS/FAIL/해당없음,
# SPECIAL 불가), 나머지 8종은 4종(PASS/FAIL/SPECIAL/해당없음 전부 가능). 2026-09-20
# quality-watcher가 이 검증 누락을 HIGH로 지적: 이게 없으면 API를 직접 호출해
# check_tie="해당없음" 같은 값을 넣을 수 있었다.
OUTBOUND_FIELD_ALLOWED_VALUES = {
    f: ("PASS", "FAIL", "해당없음") if f == "check_cable" else ("PASS", "FAIL", "SPECIAL", "해당없음")
    for f in OUTBOUND_CHECK_FIELDS
}
# 2026-09-16: "새 항목 스캔·입력" 카드 안 좁은 폭에 넣을 축약 라벨. 정식 명칭은
# OUTBOUND_CHECK_LABELS(표 헤더·엑셀 출력용)를 계속 쓰고, 이건 카드 UI 전용 —
# 화면에선 이 짧은 텍스트를 쓰고 title 속성에 정식 명칭을 붙여 보완한다.
OUTBOUND_CHECK_SHORT_LABELS = {
    "check_tie": "체결상태",
    "check_qr": "QR 부착",
    "check_wrap": "커버포장",
    "check_rst": "RST단자",
    "check_sticker": "스티커",
    "check_paint": "도장상태",
    "check_access": "부속품",
    "check_cable": "케이블타이",
    "check_indicator": "인디케이터",
}
# 2026-09-20: 사용자가 준 실제 "출고 패킹리스트" 참고파일의 6행("검사 기준") 문구
# 그대로 — 출고 스캔 화면(새 항목 입력 카드)과 report_builder.build_outbound_excel()의
# 6행이 둘 다 이 딕셔너리 하나를 쓴다(8-1절 공용헬퍼 원칙, 문구를 두 곳에 복붙하지 않음).
OUTBOUND_CHECK_CRITERIA = {
    "check_tie": "각각 해당 위치에 체결상태 확인",
    "check_qr": "탱크에 S/N 및 QR 번호 부착 및 리셉터클 캡 장착 여부",
    "check_wrap": "커버에 하우징이 노출이 되는지 확인",
    "check_rst": "볼트가 단자에 잘 고정 되었는지",
    "check_sticker": "정격표시/ 배큠표시/오픈락/골든이글",
    "check_paint": "도장 벗겨짐/파임/오염",
    "check_access": "부속품 유무",
    "check_cable": "LIFT RING 에 노란색 케이블 타이",
    "check_indicator": "투입상태(빨강)확인",
}


def compute_outbound_item_auto_result(item):
    """item: dict(아홉 개 check_* 키 포함 — sqlite3.Row면 호출 전에 dict()로 바꿀 것).
    9개 전부 값이 있어야 계산하고, 하나라도 비어있으면(None, 아직 안 눌러봄) None을
    돌려준다("미검사").

    "해당없음"은 OUTBOUND_RESULT_VALUES에 포함돼 있어 미검사로 안 잡히고,
    FAIL/SPECIAL이 아니므로 자동판정에서 PASS와 동일하게 취급된다.

    우선순위: FAIL이 하나라도 있으면 무조건 FAIL(SPECIAL이 섞여 있어도 FAIL이 이긴다) →
    그 다음 SPECIAL이 하나라도 있으면 SPECIAL → 전부 PASS/해당없음이면 PASS.
    (2026-09-16, 사용자가 명시적으로 확정: "Special이 있어도 Fail이 들어가면 무조건
    최종 판결은 Fail" — 이 순서를 절대 바꾸지 말 것.)"""
    vals = [item.get(f) for f in OUTBOUND_CHECK_FIELDS]
    if any(v not in OUTBOUND_RESULT_VALUES for v in vals):
        return None
    if "FAIL" in vals:
        return "FAIL"
    if "SPECIAL" in vals:
        return "SPECIAL"
    return "PASS"


def outbound_item_effective_result(item):
    """수동 오버라이드가 있으면 그 값, 없으면 자동판정 그대로.
    inspection_items.override_result / app.py의 effective_result() 패턴과 동일 설계
    (CLAUDE.md 관례 — 원본 자동판정은 안 지우고 오버라이드만 별도 저장)."""
    return item.get("result_override") or compute_outbound_item_auto_result(item)
# ---------- 확장 끝 ----------

def get_ma_by_component(component_no):
    """파츠 자재번호 -> 그 파츠가 속한 MA와 그 MA의 파츠 전체를 반환.

    입고 화면에서 파츠 하나를 입력하면 같은 MA의 나머지 파츠까지 한 번에 펼치기 위한 함수.
    반환: {"ma_master": MA명, "components": [자재번호, ...]} 또는 None(= 일반 자재)

    (자재번호 하나가 여러 MA에 걸쳐 있으면 어느 MA로 펼칠지 알 수 없으므로 None을 돌려
     일반 자재로 처리한다 — 잘못된 MA로 8줄이 튀어나오는 것보다 안전하다.)
    """
    if not component_no:
        return None

    conn = get_conn()
    try:
        masters = conn.execute(
            "SELECT DISTINCT assembly_id FROM assembly_components WHERE component_no=?",
            (component_no,),
        ).fetchall()
        if len(masters) != 1:
            return None

        assembly_id = masters[0]["assembly_id"]
        master = conn.execute(
            "SELECT assembly_no FROM assembly_masters WHERE id=?", (assembly_id,)
        ).fetchone()
        if not master:
            return None

        components = [
            r["component_no"]
            for r in conn.execute(
                "SELECT component_no FROM assembly_components WHERE assembly_id=? ORDER BY component_order",
                (assembly_id,),
            ).fetchall()
        ]
        if not components:
            return None
        return {"ma_master": master["assembly_no"], "components": components}
    finally:
        conn.close()


def get_conn():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


class DbBusyError(ValueError):
    """DB 잠금이 풀리지 않아 재시도를 모두 소진했을 때."""


APPEND_ONLY_TABLES = ("lot_inspection_results", "lot_inspection_unit_events", "product_report_revisions",
                      "product_report_events", "ng_events", "ng_event_log", "unit_disposal_events")


def _install_append_only_triggers(cur, tables):
    """증거성 테이블마다 UPDATE/DELETE를 거부하는 트리거 2개(trg_<table>_no_update / _no_delete)를 만든다.
    CREATE TRIGGER IF NOT EXISTS라 멱등이고 gunicorn 워커가 동시에 init_db를 돌려도 안전하다(ALTER 아님, 기존 행 불변).
    정당한 수정 경로는 없다 — 정정은 새 행(rev+1)·새 이벤트로만 한다."""
    for tb in tables:
        for op in ("UPDATE", "DELETE"):
            cur.execute(f"CREATE TRIGGER IF NOT EXISTS trg_{tb}_no_{op.lower()} BEFORE {op} ON {tb} "
                        f"BEGIN SELECT RAISE(ABORT, '{tb}은(는) append-only라 수정할 수 없어요'); END")


def _write_retry(fn):
    """쓰기 함수 전체를 최대 3회 시도 — 'locked'/'busy' OperationalError일 때만 0.4·0.8초 쉬고 재시도.
    각 함수가 실패 시 rollback·close 하는 구조라 통째로 다시 불러도 안전하다."""
    @functools.wraps(fn)
    def wrapper(*args, **kwargs):
        for attempt in range(3):
            try:
                return fn(*args, **kwargs)
            except sqlite3.OperationalError as e:
                msg = str(e).lower()
                if "locked" not in msg and "busy" not in msg:
                    raise
                if attempt < 2:
                    time.sleep(0.4 * (attempt + 1))
        raise DbBusyError("다른 작업이 처리 중이야. 잠시 후 다시 시도해줘.")
    return wrapper


def _safe_add_column(cur, alter_sql):
    """ALTER TABLE ADD COLUMN 실행 — "컬럼 없으면 추가" 가드(existing_cols 체크)를
    통과했더라도, gunicorn이 워커 여러 개를 동시에 띄우면 각 워커가 독립 프로세스로
    똑같이 init_db()를 실행해서 두 워커가 거의 동시에 같은 컬럼을 추가하려는 경합이
    생길 수 있다(2026-09-26 실제 프로덕션 장애 — 한 워커는 성공, 다른 워커는
    "duplicate column name"으로 실패해서 부팅 자체가 안 됨). SQLite는 ADD COLUMN에
    IF NOT EXISTS 문법이 없어서, existing_cols 체크만으론 이 레이스를 못 막는다 —
    실제 실행 시점에 "이미 있다"는 에러만 무시하면 두 워커 다 안전하게 통과한다."""
    try:
        cur.execute(alter_sql)
    except sqlite3.OperationalError as e:
        if "duplicate column name" not in str(e):
            raise


def init_db():
    """테이블이 없으면 생성. 있으면 그대로 둠(재실행 안전)."""
    conn = get_conn()
    cur = conn.cursor()

    # 0. 자재 마스터 — 항목(규격)이 아직 하나도 없어도 자재 자체는 존재할 수 있게 별도 테이블로 관리
    #    (규격 개별 등록에서 자재번호+제품명만 먼저 등록하고 항목은 나중에 추가하는 경우 대응)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS materials (
            material_no TEXT PRIMARY KEY,
            material_name TEXT,
            created_at TEXT DEFAULT (datetime('now', 'localtime'))
        )
    """)
    existing_material_cols = [row[1] for row in cur.execute("PRAGMA table_info(materials)").fetchall()]
    # 기준서(SAM 양식) 정보 — 도면번호는 자재번호에서 자동 계산되므로 저장 안 함(A+P→- 규칙)
    if "drawing_version" not in existing_material_cols:
        cur.execute("ALTER TABLE materials ADD COLUMN drawing_version TEXT DEFAULT '1'")
    if "revision_date" not in existing_material_cols:
        cur.execute("ALTER TABLE materials ADD COLUMN revision_date TEXT")
    if "edition" not in existing_material_cols:
        cur.execute("ALTER TABLE materials ADD COLUMN edition INTEGER DEFAULT 1")
    if "unit" not in existing_material_cols:
        cur.execute("ALTER TABLE materials ADD COLUMN unit TEXT DEFAULT 'mm'")
    if "drawing_file" not in existing_material_cols:
        cur.execute("ALTER TABLE materials ADD COLUMN drawing_file TEXT")
    # 폐지된 커스텀 성적서 양식의 레거시 지정값 보존용. 앱은 이 값을 출력에 사용하지 않는다.
    if "custom_template_id" not in existing_material_cols:
        cur.execute("ALTER TABLE materials ADD COLUMN custom_template_id INTEGER")
    # 전수검사 — NULL이면 전수검사 없음. 열 정의는 JSON {"note":"...", "columns":[...]}
    if "full_inspect_config" not in existing_material_cols:
        cur.execute("ALTER TABLE materials ADD COLUMN full_inspect_config TEXT DEFAULT NULL")
    # 자재 분류(어셈블리/하우징/PCB 등) — FK 없이 문자열 그대로 저장(intake_list.supplier와 동일 관례)
    if "category" not in existing_material_cols:
        cur.execute("ALTER TABLE materials ADD COLUMN category TEXT")

    # 0-0-1. 자재 분류 마스터 — suppliers/gauges와 같은 독립 마스터 테이블
    cur.execute("""
        CREATE TABLE IF NOT EXISTS material_categories (
            name TEXT PRIMARY KEY,
            created_at TEXT DEFAULT (datetime('now', 'localtime'))
        )
    """)

    # 0-0-1a. 자재 대분류/중분류 — 기존 materials.category는 과거 보고서 호환용으로
    # 그대로 보존한다. materials FK를 두지 않아 미등록 자재번호의 확정 매핑도 보관한다.
    cur.execute("""
        CREATE TABLE IF NOT EXISTS material_classification_pairs (
            major_category TEXT NOT NULL,
            middle_category TEXT NOT NULL,
            source TEXT NOT NULL DEFAULT 'manual',
            created_at TEXT DEFAULT (datetime('now', 'localtime')),
            PRIMARY KEY (major_category, middle_category)
        )
    """)
    existing_class_pair_cols = [row[1] for row in cur.execute(
        "PRAGMA table_info(material_classification_pairs)").fetchall()]
    if "source" not in existing_class_pair_cols:
        _safe_add_column(cur, "ALTER TABLE material_classification_pairs "
                         "ADD COLUMN source TEXT NOT NULL DEFAULT 'manual'")
    cur.execute("""
        CREATE TABLE IF NOT EXISTS material_classification_map (
            material_no TEXT PRIMARY KEY,
            major_category TEXT NOT NULL,
            middle_category TEXT NOT NULL,
            source TEXT NOT NULL DEFAULT 'manual',
            updated_at TEXT DEFAULT (datetime('now', 'localtime')),
            FOREIGN KEY (major_category, middle_category)
                REFERENCES material_classification_pairs (major_category, middle_category)
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_material_classification_major "
                "ON material_classification_map(major_category)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_material_classification_middle "
                "ON material_classification_map(middle_category)")

    # 0-0-2. 불량 유형 마스터 — material_categories와 동일한 관례(name PK, FK 없이
    # ncr.defect_type에 문자열로만 저장). 2026-09-08 신설.
    cur.execute("""
        CREATE TABLE IF NOT EXISTS defect_types (
            name TEXT PRIMARY KEY,
            created_at TEXT DEFAULT (datetime('now', 'localtime'))
        )
    """)

    # 0-1. 폐지된 커스텀 성적서 양식의 레거시 데이터 보존용 테이블.
    #      기존 레코드 삭제·변환 없이 유지하며, 앱 화면·라우트·출력에서는 사용하지 않는다.
    cur.execute("""
        CREATE TABLE IF NOT EXISTS custom_templates (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            page_size TEXT DEFAULT 'A4',
            orientation TEXT DEFAULT 'portrait',   -- portrait / landscape
            canvas_w INTEGER DEFAULT 495,
            canvas_h INTEGER DEFAULT 700,
            layout_json TEXT DEFAULT '[]',
            created_by TEXT,
            created_at TEXT DEFAULT (datetime('now', 'localtime')),
            updated_at TEXT DEFAULT (datetime('now', 'localtime'))
        )
    """)

    # 1. 규격표 — 자재별 검사항목 하한/상한 (또는 육안판정)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS specs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            material_no TEXT NOT NULL,
            material_name TEXT,
            item_name TEXT NOT NULL,
            spec_display TEXT,          -- 화면 표기용 (예: "108.5 ± 0.8")
            judge_type TEXT NOT NULL DEFAULT 'numeric',  -- 'numeric' or 'visual'
            lower_limit REAL,
            upper_limit REAL,
            inspect_method TEXT,
            aql TEXT,
            item_order INTEGER DEFAULT 0
        )
    """)

    existing_specs_cols = [row[1] for row in cur.execute("PRAGMA table_info(specs)").fetchall()]
    if "stage_group" not in existing_specs_cols:
        # 하우징 성적서 프로젝트 2단계(2026-09-26): 1차/2차/3차 분류. NULL=하우징 무관
        # (일반 자재는 전부 NULL로 남는다), 1/2/3=하우징 항목이 속한 차수.
        _safe_add_column(cur, "ALTER TABLE specs ADD COLUMN stage_group INTEGER DEFAULT NULL")

    # 기존 specs 테이블에만 있던 자재들을 materials 테이블로 1회 백필 (재실행 안전 — INSERT OR IGNORE)
    cur.execute("""
        INSERT OR IGNORE INTO materials (material_no, material_name)
        SELECT DISTINCT material_no, material_name FROM specs
    """)

    # 1-1. 입고 리스트 — 엑셀에서 붙여넣은 입고 건 (검사 전 대기 상태)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS intake_list (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            material_no TEXT NOT NULL,
            quantity INTEGER,
            supplier TEXT,
            receive_date TEXT,
            po_number TEXT,
            status TEXT NOT NULL DEFAULT '대기',   -- 대기 / 검사완료
            created_at TEXT DEFAULT (datetime('now', 'localtime'))
        )
    """)
    existing_intake_cols = [row[1] for row in cur.execute("PRAGMA table_info(intake_list)").fetchall()]
    if "product_name" not in existing_intake_cols:
        # 제품명 — 입고리스트 표시 전용 필드. 규격표에 등록된 자재명과는 별개(자동 연동 안 함)
        cur.execute("ALTER TABLE intake_list ADD COLUMN product_name TEXT")
    if "assembly_no" not in existing_intake_cols:
        # 이 행이 조립품(MA) 파츠 자동전개로 생성됐으면 그 MA번호. 일반 입고는 NULL.
        # (예전엔 product_name에 " - "가 들어있는지로 MA 파츠 여부를 추측했는데, 우연히 제품명에
        #  " - "가 들어간 일반 자재까지 "MA 파츠"로 잘못 표시되는 오판정이 있었다 — 이제 이 컬럼으로 실제 출처를 기록한다.)
        cur.execute("ALTER TABLE intake_list ADD COLUMN assembly_no TEXT")
    if "is_priority" not in existing_intake_cols:
        # 우선검사 플래그 — 입고 등록 화면에서 체크하면 검사 대기 목록 상단에 강조 표시된다.
        # 검사완료 후엔 의미 없음(대기 상태에서만 씀), 다른 화면엔 노출 안 함(스펙 확정 사항).
        cur.execute("ALTER TABLE intake_list ADD COLUMN is_priority INTEGER NOT NULL DEFAULT 0")

    # 1-2. 과거 입고 이력 — 일일보고 엑셀 등에서 옮겨온 "이미 끝난" 입고 기록 조회 전용.
    # intake_list와 절대 혼동하지 말 것: 여기 등록해도 검사 대기 큐(상태='대기')에 안 뜨고
    # 품질현황(quality_report)에도 안 잡힌다 — 검사/승인 워크플로우와 완전히 무관한 순수 로그다.
    # FK를 일부러 안 건다(materials/intake_list/inspections 어디와도 조인 강제 없음).
    cur.execute("""
        CREATE TABLE IF NOT EXISTS intake_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            material_no TEXT NOT NULL,
            product_name TEXT,
            supplier TEXT,
            receive_date TEXT,      -- 가능하면 'YYYY-MM-DD'로 정규화해서 저장(정렬/기간검색 위해)
            po_number TEXT,
            quantity INTEGER,
            source TEXT,            -- 이 행이 어느 등록(파일+시각)에서 왔는지 — 배치 단위 되돌리기 키로도 씀
            imported_by TEXT,       -- 등록한 계정의 username
            created_at TEXT DEFAULT (datetime('now', 'localtime'))
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_intake_history_material ON intake_history(material_no)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_intake_history_supplier ON intake_history(supplier)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_intake_history_date ON intake_history(receive_date)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_intake_history_source ON intake_history(source)")

    # 2. 검사(성적서) 헤더 — 자재 입고 1건 = 성적서 1건
    cur.execute("""
        CREATE TABLE IF NOT EXISTS inspections (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            intake_id INTEGER,
            material_no TEXT NOT NULL,
            material_name TEXT,
            supplier TEXT,
            po_number TEXT,
            receive_date TEXT,
            inspect_date TEXT,
            inspector TEXT,
            quantity INTEGER,
            overall_result TEXT,        -- '합격' / '불합격'
            status TEXT NOT NULL DEFAULT 'pending',  -- pending / approved / rejected
            approver TEXT,
            approved_at TEXT,
            signature_path TEXT,
            pdf_path TEXT,
            reject_reason TEXT,
            est_time_label TEXT,        -- 실제 측정 시간 ("N시간 N분 N초") — 태블릿 스톱워치로 측정
            actual_time_sec INTEGER,    -- 실제 측정 시간(초) — 원본값. est_time_label은 이 값을 표시용으로 포맷한 것
            created_at TEXT DEFAULT (datetime('now', 'localtime'))
        )
    """)

    # 3. 검사 결과 — 성적서 항목별 측정값·판정
    cur.execute("""
        CREATE TABLE IF NOT EXISTS inspection_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            inspection_id INTEGER NOT NULL,
            item_name TEXT NOT NULL,
            measured_value TEXT,        -- 수치 또는 O/X, 콤마로 여러 샘플 가능
            max_value REAL,
            min_value REAL,
            result TEXT,                -- '합격' / '불합격'
            gauge_expiry TEXT,          -- 계측기 유효기간 (항목별, YYYY-MM-DD)
            FOREIGN KEY (inspection_id) REFERENCES inspections(id)
        )
    """)

    # 기존 DB에 gauge_expiry 컬럼이 없으면 추가 (재실행 안전 마이그레이션)
    existing_cols = [row[1] for row in cur.execute("PRAGMA table_info(inspection_items)").fetchall()]
    if "gauge_expiry" not in existing_cols:
        cur.execute("ALTER TABLE inspection_items ADD COLUMN gauge_expiry TEXT")
    if "gauge_name" not in existing_cols:
        cur.execute("ALTER TABLE inspection_items ADD COLUMN gauge_name TEXT")
    if "part_material_no" not in existing_cols:
        # 조립품 그룹 검사일 때, 이 측정값이 실제로 어느 부품(자재)에 속하는지 기록.
        # 일반(단일 자재) 검사는 항상 헤더의 material_no와 동일하게 채움.
        cur.execute("ALTER TABLE inspection_items ADD COLUMN part_material_no TEXT")
    if "aql_snapshot" not in existing_cols:
        # 검사 당시의 AQL 정책을 항목별로 고정한다. specs.aql 또는 자재 분류가 나중에
        # 바뀌어도 승인 이력·재출력 결과가 달라지면 안 된다. NULL AQL도 유효한 스냅샷이라
        # aql_snapshot_set 플래그로 "아직 이관 안 됨"과 구분한다.
        cur.execute("ALTER TABLE inspection_items ADD COLUMN aql_snapshot TEXT")
    if "sample_qty_snapshot" not in existing_cols:
        cur.execute("ALTER TABLE inspection_items ADD COLUMN sample_qty_snapshot INTEGER")
    if "ac_allowance_snapshot" not in existing_cols:
        cur.execute("ALTER TABLE inspection_items ADD COLUMN ac_allowance_snapshot INTEGER")
    if "aql_snapshot_set" not in existing_cols:
        cur.execute("ALTER TABLE inspection_items ADD COLUMN aql_snapshot_set INTEGER NOT NULL DEFAULT 0")

    existing_insp_cols = [row[1] for row in cur.execute("PRAGMA table_info(inspections)").fetchall()]
    if "est_time_label" not in existing_insp_cols:
        cur.execute("ALTER TABLE inspections ADD COLUMN est_time_label TEXT")
    if "actual_time_sec" not in existing_insp_cols:
        cur.execute("ALTER TABLE inspections ADD COLUMN actual_time_sec INTEGER")
    if "created_by_user_id" not in existing_insp_cols:
        cur.execute("ALTER TABLE inspections ADD COLUMN created_by_user_id INTEGER")
    # 비고란 — 검사자/중간관리자/최종결정권자 3명이 각자 따로 작성 (역할별 색상 구분해서 성적서에도 반영)
    if "remark_inspector" not in existing_insp_cols:
        cur.execute("ALTER TABLE inspections ADD COLUMN remark_inspector TEXT")
    if "remark_manager" not in existing_insp_cols:
        cur.execute("ALTER TABLE inspections ADD COLUMN remark_manager TEXT")
    if "remark_approver" not in existing_insp_cols:
        cur.execute("ALTER TABLE inspections ADD COLUMN remark_approver TEXT")
    # 특채 승인 여부 — 'normal'(일반 합/불 판정 그대로) / 'special'(규격 벗어나도 특채로 승인)
    if "approval_type" not in existing_insp_cols:
        cur.execute("ALTER TABLE inspections ADD COLUMN approval_type TEXT DEFAULT 'normal'")
    if "total_time_sec" not in existing_insp_cols:
        cur.execute("ALTER TABLE inspections ADD COLUMN total_time_sec INTEGER")
    # 위변조 검증 — 승인 시점의 성적서 내용을 해시로 굳혀둔다.
    #   content_hash: DB에 저장된 판정 내용(헤더+항목)의 해시. 나중에 값이 바뀌면 불일치로 잡힌다.
    #   pdf_hash    : 발행된 PDF 파일 자체의 해시. 파일이 덮어써지면 불일치로 잡힌다.
    #   ※ 해시를 넣은 시점 이후 승인분만 보호된다(소급 불가).
    if "content_hash" not in existing_insp_cols:
        cur.execute("ALTER TABLE inspections ADD COLUMN content_hash TEXT")
    if "pdf_hash" not in existing_insp_cols:
        cur.execute("ALTER TABLE inspections ADD COLUMN pdf_hash TEXT")
    if "ncr_waived" not in existing_insp_cols:
        cur.execute("ALTER TABLE inspections ADD COLUMN ncr_waived INTEGER DEFAULT 0")
    if "ncr_waived_reason" not in existing_insp_cols:
        cur.execute("ALTER TABLE inspections ADD COLUMN ncr_waived_reason TEXT")

    # 4. 사용자 계정 — 고정 역할 대신 개별 권한(콤마구분 텍스트)을 admin이 하나하나 부여/회수
    # 권한 종류: intake(입고리스트) / spec(규격관리) / inspect(검사입력,본인것만수정) /
    #           inspect_all(타인 성적서도 수정) / approve(승인·반려·특채) / output(출력) /
    #           users(계정관리, 10분 자동로그아웃 대상) / logs(활동로그 열람)
    # 비밀번호는 내부 시스템 특성상 평문 저장(관리자가 계정 발급·확인 용도)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL UNIQUE,
            password TEXT NOT NULL,
            display_name TEXT,
            permissions TEXT NOT NULL DEFAULT '',
            created_at TEXT DEFAULT (datetime('now', 'localtime'))
        )
    """)
    existing_user_cols = [row[1] for row in cur.execute("PRAGMA table_info(users)").fetchall()]
    if "permissions" not in existing_user_cols:
        cur.execute("ALTER TABLE users ADD COLUMN permissions TEXT NOT NULL DEFAULT ''")
    if "is_final_approver" not in existing_user_cols:
        # 최종결정권자 — 승인/특채/불합격 확정에 서명할 수 있는 사람. 최대 2명(MAX_FINAL_APPROVERS)
        cur.execute("ALTER TABLE users ADD COLUMN is_final_approver INTEGER NOT NULL DEFAULT 0")
    if "role" in existing_user_cols:
        # 예전 role 컬럼이 남아있으면 권한 세트로 1회 변환(이미 permissions가 채워진 계정은 건드리지 않음)
        role_to_perms = {
            "admin": "intake,spec,inspect,inspect_all,approve,output,users,logs",
            "approver": "intake,spec,inspect,inspect_all,approve,output",
            "manager": "intake,spec,inspect,inspect_all,output",
            "inspector": "inspect",
        }
        for r in cur.execute("SELECT id, role, permissions FROM users").fetchall():
            uid, old_role, perms = r
            if not perms and old_role in role_to_perms:
                cur.execute("UPDATE users SET permissions = ? WHERE id = ?",
                           (role_to_perms[old_role], uid))

    # 4-1. 조립 제품(MA) 마스터 및 파츠 분해
    cur.execute("""
        CREATE TABLE IF NOT EXISTS assembly_masters (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            assembly_no TEXT NOT NULL UNIQUE,
            assembly_name TEXT,
            component_count INTEGER DEFAULT 8,
            created_at TEXT DEFAULT (datetime('now', 'localtime'))
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS assembly_components (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            assembly_id INTEGER NOT NULL,
            component_order INTEGER NOT NULL,
            component_no TEXT NOT NULL,
            component_name TEXT,
            FOREIGN KEY(assembly_id) REFERENCES assembly_masters(id),
            UNIQUE(assembly_id, component_order)
        )
    """)

    # 4-1-1. 통합BOM 계층 정보 (2026-09-09) — assembly_masters/components와는 완전히 별개.
    #        조회/필터링 전용("자재 찾기" 화면). 재임포트 시 전량 삭제 후 재삽입한다.
    cur.execute("""
        CREATE TABLE IF NOT EXISTS material_bom_links (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            material_no        TEXT NOT NULL,
            parent_material_no TEXT,
            model_name          TEXT NOT NULL,
            level               INTEGER NOT NULL,
            kind                TEXT,
            qty_per_parent      REAL,
            qty_per_model       REAL,
            unit                TEXT,
            source_row_no       INTEGER,
            imported_at         TEXT DEFAULT (datetime('now', 'localtime'))
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_bom_material_no ON material_bom_links(material_no)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_bom_model_name ON material_bom_links(model_name)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_bom_parent ON material_bom_links(parent_material_no)")
    existing_bom_cols = [row[1] for row in cur.execute("PRAGMA table_info(material_bom_links)").fetchall()]
    if "bom_name" not in existing_bom_cols:
        # BOM 원본의 "규격사양(SPEC)" 열 — materials.material_name은 절대 안 건드리고,
        # "자재 찾기"에서 미등록 자재를 등록할 때 이름을 자동으로 채워주는 참고용 값.
        cur.execute("ALTER TABLE material_bom_links ADD COLUMN bom_name TEXT")
    if "manually_edited" not in existing_bom_cols:
        # 사용자가 "자재 찾기"에서 model_name/parent_material_no/level을 직접 수정한
        # 행 표시(2026-09-22). 이 플래그가 1인 행은 import_bom_from_excel() 재임포트
        # 시 덮어쓰기 대상에서 제외된다 — CLAUDE.md 18절 참고.
        cur.execute("ALTER TABLE material_bom_links ADD COLUMN manually_edited INTEGER NOT NULL DEFAULT 0")

    # 4-1. 검사 입력 임시저장 — 검사자가 입력하는 즉시 서버에 저장된다.
    #      예전엔 브라우저 localStorage에만 있어서 태블릿이 꺼지거나 기기를 바꾸면 날아갔다.
    cur.execute("""
        CREATE TABLE IF NOT EXISTS inspection_drafts (
            intake_id  INTEGER PRIMARY KEY,
            user_id    INTEGER,
            username   TEXT,
            payload    TEXT NOT NULL,          -- 화면 입력값 전체를 JSON으로
            updated_at TEXT DEFAULT (datetime('now', 'localtime'))
        )
    """)

    # 4-1-1. 항목 판정 수동 오버라이드 — 자동판정(AQL 허용범위 등)과 무관하게
    #        사후에 사람이 개별 항목 합격/불합격을 뒤집을 때 쓴다.
    #        inspection_items.result 원본은 절대 안 건드리고 별도 주석 레이어로만 존재한다.
    cur.execute("""
        CREATE TABLE IF NOT EXISTS inspection_item_overrides (
            id                 INTEGER PRIMARY KEY AUTOINCREMENT,
            inspection_id      INTEGER NOT NULL,
            part_material_no   TEXT NOT NULL,
            item_name          TEXT NOT NULL,
            override_result    TEXT NOT NULL,
            reason             TEXT NOT NULL,
            created_by_user_id INTEGER,
            created_by_name    TEXT,
            created_at         TEXT DEFAULT (datetime('now', 'localtime')),
            FOREIGN KEY (inspection_id) REFERENCES inspections(id),
            UNIQUE(inspection_id, part_material_no, item_name)
        )
    """)

    # 4-2. 4M 변경점 — 협력사가 사람/설비/자재/방법을 바꾼 시점 기록.
    #      변경 전후 불량률을 비교하려면 "언제 바뀌었는지"가 남아 있어야 한다(IATF 변경점 관리).
    cur.execute("""
        CREATE TABLE IF NOT EXISTS change_points (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            supplier    TEXT NOT NULL,
            material_no TEXT,                  -- 비어 있으면 그 업체 전체에 해당
            change_type TEXT NOT NULL,         -- Man / Machine / Material / Method
            change_date TEXT NOT NULL,
            description TEXT,
            reported_by TEXT,
            created_at  TEXT DEFAULT (datetime('now', 'localtime'))
        )
    """)

    # 4-3. 업체 월간 품질 성적표 — 자동 생성하되 최종결정권자 승인 전에는 발송 못 한다.
    cur.execute("""
        CREATE TABLE IF NOT EXISTS supplier_reports (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            supplier    TEXT NOT NULL,
            period      TEXT NOT NULL,          -- 'YYYY-MM'
            start_date  TEXT NOT NULL,
            end_date    TEXT NOT NULL,
            payload     TEXT NOT NULL,          -- 집계 결과 스냅샷(JSON)
            status      TEXT NOT NULL DEFAULT 'draft',   -- draft / approved / sent
            created_by  TEXT,
            created_at  TEXT DEFAULT (datetime('now', 'localtime')),
            approved_by TEXT,
            approved_at TEXT,
            approve_signature TEXT,
            sent_to     TEXT,
            sent_at     TEXT,
            UNIQUE(supplier, period)
        )
    """)

    # 5. 활동 로그 — 등록/수정/삭제 등 주요 액션 기록 (admin만 열람)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS activity_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER,
            username TEXT,
            role TEXT,
            action TEXT NOT NULL,        -- 예: '성적서 등록', '규격 항목 수정' 등
            target_type TEXT,            -- 예: 'inspection', 'spec_item', 'intake', 'user'
            target_id TEXT,
            detail TEXT,
            created_at TEXT DEFAULT (datetime('now', 'localtime'))
        )
    """)

    # 6. 자재 그룹(조립품) — 입고 시엔 조립품 자재번호 하나지만, 분해해서 부품별로 따로 검사하는 경우
    cur.execute("""
        CREATE TABLE IF NOT EXISTS material_groups (
            group_no TEXT PRIMARY KEY,   -- 조립품 자재번호(입고 리스트/검사입력에서 이 번호로 취급)
            group_name TEXT,
            created_at TEXT DEFAULT (datetime('now', 'localtime'))
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS material_group_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            group_no TEXT NOT NULL,
            material_no TEXT NOT NULL,   -- 분해된 부품의 실제 자재번호 (specs.material_no와 매칭)
            item_order INTEGER DEFAULT 0,
            FOREIGN KEY (group_no) REFERENCES material_groups(group_no)
        )
    """)

    # 7. 규격 일괄등록에서 문제(확인필요/실패) 있었던 자재 — 나중에 다시 보고 해결 처리할 수 있게 기록
    cur.execute("""
        CREATE TABLE IF NOT EXISTS spec_review_flags (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            material_no TEXT,
            source_filename TEXT,
            reason TEXT NOT NULL,
            resolved INTEGER NOT NULL DEFAULT 0,
            created_at TEXT DEFAULT (datetime('now', 'localtime'))
        )
    """)

    # 8. 업체 정보 — 이름, 이메일, 연락처
    cur.execute("""
        CREATE TABLE IF NOT EXISTS suppliers (
            name TEXT PRIMARY KEY,
            email TEXT,
            contact TEXT,
            notes TEXT,
            created_at TEXT DEFAULT (datetime('now', 'localtime'))
        )
    """)
    existing_supplier_cols = [row[1] for row in cur.execute("PRAGMA table_info(suppliers)").fetchall()]
    if "address" not in existing_supplier_cols:
        cur.execute("ALTER TABLE suppliers ADD COLUMN address TEXT")
    if "biz_no" not in existing_supplier_cols:
        cur.execute("ALTER TABLE suppliers ADD COLUMN biz_no TEXT")
    if "contact_name" not in existing_supplier_cols:
        cur.execute("ALTER TABLE suppliers ADD COLUMN contact_name TEXT")
    if "contact2" not in existing_supplier_cols:
        cur.execute("ALTER TABLE suppliers ADD COLUMN contact2 TEXT")
    if "items" not in existing_supplier_cols:
        cur.execute("ALTER TABLE suppliers ADD COLUMN items TEXT")

    # 8-1. 업체 담당자 — 역할(영업/품질/구매/기타)별로 여러 명 등록 가능(2026-09-07).
    #      suppliers의 contact_name/contact/contact2/email은 옛 단일 담당자 필드로,
    #      호환을 위해 그대로 남겨두고 폴백용으로 쓴다(get_default_contact 참고).
    cur.execute("""
        CREATE TABLE IF NOT EXISTS supplier_contacts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            supplier_name TEXT NOT NULL,
            role TEXT NOT NULL DEFAULT '기타',
            contact_name TEXT,
            phone TEXT,
            email TEXT,
            notes TEXT,
            sort_order INTEGER DEFAULT 0,
            created_at TEXT DEFAULT (datetime('now', 'localtime'))
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_supplier_contacts_name ON supplier_contacts(supplier_name)")

    # 9. 부적합 통보서 (NCR)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS ncr (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            ncr_no TEXT UNIQUE,
            inspection_id INTEGER,
            material_no TEXT,
            material_name TEXT,
            supplier TEXT,
            defect_description TEXT,
            action_required TEXT,
            due_date TEXT,
            issued_by TEXT,
            issued_date TEXT,
            photos TEXT DEFAULT '[]',
            email_sent_at TEXT,
            status TEXT DEFAULT 'draft',
            confirmed_by TEXT,
            confirmed_at TEXT,
            sent_to TEXT,
            created_at TEXT DEFAULT (datetime('now', 'localtime'))
        )
    """)
    # 기존 DB에 컬럼 없으면 추가 (마이그레이션)
    existing_cols = [r[1] for r in conn.execute("PRAGMA table_info(ncr)").fetchall()]
    for col, definition in [
        ("confirmed_by", "TEXT"),
        ("confirmed_at", "TEXT"),
        ("sent_to",      "TEXT"),
        # 최종결정권자 승인 서명 이미지 경로 — 협력사로 나가는 문서라 서명 근거가 남아야 함
        ("confirm_signature", "TEXT"),
        # 수기입력(성적서 미연결) 통보서용 — 연결된 성적서가 있으면 insp_ 조인값을 쓰고,
        # 없으면(inspection_id NULL) 이 값을 대신 보여준다
        ("lot_number", "TEXT"),
        ("receive_date", "TEXT"),
        # 새 양식 필드 (2026-09-04)
        ("cc_recipient",  "TEXT"),
        ("sample_qty",    "INTEGER"),
        ("defect_qty",    "INTEGER"),
        ("special_note",  "TEXT"),
        ("lot_qty",       "TEXT"),
    ]:
        if col not in existing_cols:
            conn.execute(f"ALTER TABLE ncr ADD COLUMN {col} {definition}")

    # 9-1. 개선요청서 — NCR보다 가벼운 사전조치 문서 (2026-09-15, 서명 없음)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS improvement_requests (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            request_no TEXT UNIQUE,
            inspection_id INTEGER,
            material_no TEXT,
            material_name TEXT,
            supplier TEXT,
            lot_number TEXT,
            po_number TEXT,
            lot_qty TEXT,
            defect_categories TEXT,
            defect_category_etc TEXT,
            request_detail TEXT,
            confirmed_name TEXT,
            issued_by TEXT,
            issued_date TEXT,
            status TEXT DEFAULT 'draft',
            email_sent_at TEXT,
            sent_to TEXT,
            created_at TEXT DEFAULT (datetime('now', 'localtime'))
        )
    """)
    existing_imp_cols = [r[1] for r in conn.execute(
        "PRAGMA table_info(improvement_requests)"
    ).fetchall()]
    for col, definition in [
        ("sample_qty", "INTEGER"),
        ("defect_qty", "INTEGER"),
    ]:
        if col not in existing_imp_cols:
            _safe_add_column(cur, f"ALTER TABLE improvement_requests ADD COLUMN {col} {definition}")
    cur.execute("""
        CREATE TABLE IF NOT EXISTS improvement_request_photos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            improvement_request_id INTEGER NOT NULL,
            kind TEXT NOT NULL DEFAULT 'defect',
            photo_path TEXT NOT NULL,
            sort_order INTEGER DEFAULT 0,
            created_at TEXT DEFAULT (datetime('now', 'localtime'))
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_imp_photos_req ON improvement_request_photos(improvement_request_id)")

    # 10. 앱 설정 (SMTP 등)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS app_settings (
            key TEXT PRIMARY KEY,
            value TEXT
        )
    """)

    # 11. 계측기 마스터
    cur.execute("""
        CREATE TABLE IF NOT EXISTS gauge_master (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            gauge_no TEXT UNIQUE,
            name TEXT NOT NULL,
            model TEXT,
            location TEXT,
            last_calibrated TEXT,
            expiry_date TEXT,
            notes TEXT,
            created_at TEXT DEFAULT (datetime('now','localtime'))
        )
    """)

    # 12. 검사 진행 현황 (검사 입력폼 열고 있는 사람 추적)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS inspection_progress (
            intake_id INTEGER PRIMARY KEY,
            inspectors TEXT DEFAULT '[]',
            updated_at TEXT DEFAULT (datetime('now','localtime'))
        )
    """)

    # 13. 반품 처리
    cur.execute("""
        CREATE TABLE IF NOT EXISTS return_requests (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            inspection_id INTEGER NOT NULL,
            material_no TEXT,
            material_name TEXT,
            supplier TEXT,
            return_date TEXT,
            reason TEXT,
            quantity INTEGER,
            status TEXT DEFAULT '반품요청',
            resolved_inspection_id INTEGER,
            created_by TEXT,
            created_at TEXT DEFAULT (datetime('now','localtime'))
        )
    """)

    # 전수검사 헤더 — inspection_id 1:1 대응. 상태: open / complete
    cur.execute("""
        CREATE TABLE IF NOT EXISTS full_inspections (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            inspection_id INTEGER NOT NULL UNIQUE,
            inspect_date TEXT,
            complete_date TEXT,
            status TEXT DEFAULT 'open',
            created_at TEXT DEFAULT (datetime('now','localtime')),
            FOREIGN KEY (inspection_id) REFERENCES inspections(id)
        )
    """)
    # 전수검사 유닛별 데이터 — 유닛 1개(시리얼번호) = 1행
    cur.execute("""
        CREATE TABLE IF NOT EXISTS full_inspection_units (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            full_inspection_id INTEGER NOT NULL,
            unit_no INTEGER NOT NULL,
            serial_no TEXT DEFAULT '',
            values_json TEXT DEFAULT '{}',
            result TEXT DEFAULT '',
            remark TEXT DEFAULT '',
            gauge_name TEXT DEFAULT '',
            FOREIGN KEY (full_inspection_id) REFERENCES full_inspections(id)
        )
    """)
    existing_fi_cols = [row[1] for row in cur.execute("PRAGMA table_info(full_inspection_units)").fetchall()]
    if "gauge_name" not in existing_fi_cols:
        cur.execute("ALTER TABLE full_inspection_units ADD COLUMN gauge_name TEXT DEFAULT ''")
    # 유닛 단위 특채(개별 예외처리) — 2026-09-28. 자동판정이 NG인 유닛이라도 검사자·
    # 최종결정권자가 "이 유닛만은 예외로 쓴다"고 표시할 수 있어야 한다(로트 전체 검사가
    # 끝나기 전에 이미 개별로 확인된 유닛을 먼저 생산 투입할 수 있게 하기 위함, S/N 조회
    # 화면과 연동). _safe_add_column 사용(CLAUDE.md 26-1절, 멀티워커 ALTER TABLE 경합 방지).
    _safe_add_column(cur, "ALTER TABLE full_inspection_units ADD COLUMN special_marked INTEGER DEFAULT 0")

    # ---- 여기부터 신규 삽입 ----
    # 출고(완제품 S/N·QR·사진) 관리 — 입고검사(IQC)와 완전히 별개의 신규 하위시스템.
    # finished_goods_serials/outbound_items는 FK 없이 문자열(serial_no)로만 느슨하게 연결한다
    # — 미등록 S/N, 다른 배치에 이미 쓰인 S/N도 경고만 하고 출고 항목으로 받아들여야 하기
    # 때문에 FK로 강제하면 안 된다(material_bom_links가 materials에 FK 없는 것과 같은 이유).
    cur.execute("""
        CREATE TABLE IF NOT EXISTS finished_goods_serials (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            serial_no   TEXT NOT NULL UNIQUE,
            issued_by   TEXT,
            issued_at   TEXT DEFAULT (datetime('now','localtime'))
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS outbound_batches (
            id            INTEGER PRIMARY KEY AUTOINCREMENT,
            customer      TEXT,
            ship_date     TEXT,
            handler       TEXT,
            created_by    TEXT,
            created_at    TEXT DEFAULT (datetime('now','localtime')),
            updated_at    TEXT
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS outbound_items (
            id              INTEGER PRIMARY KEY AUTOINCREMENT,
            batch_id        INTEGER NOT NULL REFERENCES outbound_batches(id),
            serial_no       TEXT NOT NULL,
            product_name    TEXT,
            quantity        INTEGER,
            created_at      TEXT DEFAULT (datetime('now','localtime'))
        )
    """)

    existing_oi_cols = [row[1] for row in cur.execute("PRAGMA table_info(outbound_items)").fetchall()]
    # 2026-09-15 확장: 출고 전 품질확인항목(PASS/FAIL/SPECIAL/해당없음) + 전체판정 수동 오버라이드.
    # 2026-09-20: 5→9개로 확장(check_sticker/paint/cable/indicator 추가).
    # 자동판정 값 자체는 저장하지 않는다 — compute_outbound_item_auto_result()가 9개
    # check_* 컬럼에서 매번 계산한다(CLAUDE.md 8-1절, 저장값-계산값 불일치 사고 방지).
    for _ob_col in OUTBOUND_CHECK_FIELDS + ["result_override"]:
        if _ob_col not in existing_oi_cols:
            cur.execute(f"ALTER TABLE outbound_items ADD COLUMN {_ob_col} TEXT")

    if "inspected_by" not in existing_oi_cols:
        # 2026-09-21 확장: 항목(S/N)을 등록한 사용자 표시이름. app.py의 outbound_item_add()가
        # 생성 시점에 채운다. 저장된 항목의 수정/삭제 권한을 "이 항목의 검사자 본인" 또는
        # "이 차수의 등록자"(outbound_batches.created_by)로 제한하는 데 쓴다(사용자 확정,
        # admin 예외·outbound_delete 권한 예외 없음). 마이그레이션 이전에 생성된 기존 항목은
        # NULL로 남는다 — 그 항목은 차수 등록자만 수정/삭제 가능해진다(의도된 동작, 버그 아님).
        cur.execute("ALTER TABLE outbound_items ADD COLUMN inspected_by TEXT")

    cur.execute("""
        CREATE TABLE IF NOT EXISTS outbound_item_photos (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            item_id     INTEGER NOT NULL REFERENCES outbound_items(id),
            file_path   TEXT NOT NULL,
            uploaded_at TEXT DEFAULT (datetime('now','localtime'))
        )
    """)
    existing_obp_cols = [row[1] for row in cur.execute("PRAGMA table_info(outbound_item_photos)").fetchall()]
    if "kind" not in existing_obp_cols:
        # 'indicator'(인디케이터 사진) / 'body'(본체사진). 기존 사진은 전부 'indicator'로
        # 간주한다 — 지금까지는 사진 종류 구분이 없어 전부 한 종류였고, 사용자가 준
        # 실제 서식에서도 첫 번째 사진 칸이 인디케이터 사진이기 때문(합리적 기본값).
        cur.execute("ALTER TABLE outbound_item_photos ADD COLUMN kind TEXT NOT NULL DEFAULT 'indicator'")

    # ---- 2026-09-15 확장: 모델명 자동분류 규칙 3개 매핑표 ----
    # 정규식 패턴 구조(CKMR.../CKCB...)는 코드에 고정하고, 여기 값(코드->라벨)만
    # 화면에서 CRUD 가능하게 한다(YAGNI — 패턴 구조 자체를 바꿀 요구가 생기면 재설계).
    cur.execute("""
        CREATE TABLE IF NOT EXISTS outbound_rule_voltage (
            code  TEXT PRIMARY KEY,
            label TEXT NOT NULL
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS outbound_rule_suffix (
            code  TEXT PRIMARY KEY,
            label TEXT NOT NULL
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS outbound_rule_pcode (
            code  TEXT PRIMARY KEY,
            label TEXT NOT NULL
        )
    """)
    # ---- 2026-09-15 확장(모델명 분류 규칙) 끝 ----

    # ---- 2026-09-21 신규: 스티커 부착 기준 참고(출고 스캔 "기준 보기" 팝업용) ----
    # assembly_masters(5-0절)/material_bom_links(18절)와 완전히 별개 — CKMR 모델별
    # 스티커 4종 부착면/수량 표 + 참고이미지를 보여주기 위한 순수 조회용 참고 데이터.
    cur.execute("""
        CREATE TABLE IF NOT EXISTS sticker_reference_models (
            model_no      TEXT PRIMARY KEY,
            voltage_code  TEXT,
            suffix_code   TEXT,
            pcode         TEXT,
            image_path    TEXT,
            updated_at    TEXT DEFAULT (datetime('now','localtime'))
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS sticker_reference_items (
            id            INTEGER PRIMARY KEY AUTOINCREMENT,
            model_no      TEXT NOT NULL REFERENCES sticker_reference_models(model_no),
            no            TEXT,
            label         TEXT NOT NULL,
            attach_face   TEXT,
            qty           TEXT,
            sort_order    INTEGER NOT NULL
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_sticker_ref_items_model "
                "ON sticker_reference_items(model_no)")
    # ---- 2026-09-21 신규 끝 ----

    # ---- 2026-10-10 신규: 완제품 모델 마스터(TASK #66 패키지 1B) ----
    # product_model_history는 append-only — 수정·삭제 함수/UI를 만들지 말 것.
    cur.execute("""
        CREATE TABLE IF NOT EXISTS product_models (
            id                   INTEGER PRIMARY KEY AUTOINCREMENT,
            model_no             TEXT NOT NULL,
            model_key            TEXT NOT NULL UNIQUE,
            display_name         TEXT,
            voltage_code         TEXT,
            voltage_kv           INTEGER,
            suffix_code          TEXT,
            pcode                TEXT,
            phase_type           TEXT NOT NULL,
            housing_required_qty INTEGER NOT NULL,
            is_active            INTEGER NOT NULL DEFAULT 1,
            revision             INTEGER NOT NULL DEFAULT 1,
            note                 TEXT,
            created_by           TEXT,
            created_at           TEXT DEFAULT (datetime('now','localtime')),
            updated_by           TEXT,
            updated_at           TEXT
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS product_model_history (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            model_id    INTEGER NOT NULL REFERENCES product_models(id),
            revision    INTEGER NOT NULL,
            change_type TEXT NOT NULL,
            before_json TEXT,
            after_json  TEXT,
            reason      TEXT,
            changed_by  TEXT,
            changed_at  TEXT DEFAULT (datetime('now','localtime'))
        )
    """)
    # ---- 2026-10-10 신규 끝 ----

    # ---- 2026-10-10 신규: 생산 LOT(TASK #66 패키지 1C-a) ----
    # production_lot_history는 append-only — 수정·삭제 함수/UI를 만들지 말 것.
    cur.execute("""
        CREATE TABLE IF NOT EXISTS production_lots (
            id                  INTEGER PRIMARY KEY AUTOINCREMENT,
            lot_no              TEXT NOT NULL UNIQUE,
            lot_date            TEXT NOT NULL,
            seq                 INTEGER NOT NULL,
            model_id            INTEGER NOT NULL REFERENCES product_models(id),
            model_snapshot_json TEXT,
            status              TEXT NOT NULL DEFAULT 'draft',
            is_test             INTEGER NOT NULL DEFAULT 0,
            lot_revision        INTEGER NOT NULL DEFAULT 0,
            note                TEXT,
            cancel_reason       TEXT,
            created_by          TEXT,
            created_at          TEXT DEFAULT (datetime('now','localtime')),
            confirmed_by        TEXT,
            confirmed_at        TEXT,
            cancelled_by        TEXT,
            cancelled_at        TEXT,
            updated_by          TEXT,
            updated_at          TEXT,
            UNIQUE(lot_date, seq)
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS production_lot_history (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            lot_id      INTEGER NOT NULL REFERENCES production_lots(id),
            change_type TEXT NOT NULL,
            before_json TEXT,
            after_json  TEXT,
            reason      TEXT,
            changed_by  TEXT,
            changed_at  TEXT DEFAULT (datetime('now','localtime'))
        )
    """)
    # ---- 2026-10-10 신규 끝 ----

    # ---- 2026-10-10 신규: 완제품 내부 ID와 S/N(TASK #66 패키지 1C-b) ----
    # finished_unit_history는 append-only — 수정·삭제 함수/UI를 만들지 말 것.
    cur.execute("""
        CREATE TABLE IF NOT EXISTS finished_units (
            id              INTEGER PRIMARY KEY AUTOINCREMENT,
            lot_id          INTEGER NOT NULL REFERENCES production_lots(id),
            origin_state    TEXT NOT NULL DEFAULT 'production_lot',
            status          TEXT NOT NULL DEFAULT 'active',
            excluded_reason TEXT,
            excluded_by     TEXT,
            excluded_at     TEXT,
            model_mismatch  INTEGER NOT NULL DEFAULT 0,
            created_by      TEXT,
            created_at      TEXT DEFAULT (datetime('now','localtime'))
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS unit_serial_aliases (
            id              INTEGER PRIMARY KEY AUTOINCREMENT,
            unit_id         INTEGER NOT NULL REFERENCES finished_units(id),
            serial_raw      TEXT NOT NULL,
            serial_key      TEXT NOT NULL UNIQUE,
            kind            TEXT NOT NULL DEFAULT 'current',
            replaced_reason TEXT,
            created_by      TEXT,
            created_at      TEXT DEFAULT (datetime('now','localtime'))
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS finished_unit_history (
            id           INTEGER PRIMARY KEY AUTOINCREMENT,
            unit_id      INTEGER NOT NULL REFERENCES finished_units(id),
            lot_id       INTEGER NOT NULL,
            change_type  TEXT NOT NULL,
            before_json  TEXT,
            after_json   TEXT,
            reason       TEXT,
            lot_revision INTEGER NOT NULL,
            changed_by   TEXT,
            changed_at   TEXT DEFAULT (datetime('now','localtime'))
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_finished_units_lot ON finished_units(lot_id)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_unit_alias_unit ON unit_serial_aliases(unit_id)")
    # ---- 2026-10-10 신규 끝 ----

    # ---- 2026-10-10 신규: 제품검사 항목 버전과 LOT 검사 시작 스냅샷(TASK #136 패키지 1D-a) ----
    # 두 테이블 모두 append-only — 수정·삭제 함수/UI를 만들지 말 것.
    cur.execute("""
        CREATE TABLE IF NOT EXISTS product_model_item_versions (
            id           INTEGER PRIMARY KEY AUTOINCREMENT,
            model_id     INTEGER NOT NULL REFERENCES product_models(id),
            version      INTEGER NOT NULL,
            items_json   TEXT NOT NULL,
            content_hash TEXT NOT NULL,
            reason       TEXT,
            created_by   TEXT,
            created_at   TEXT NOT NULL,
            UNIQUE(model_id, version)
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS production_lot_inspections (
            id                  INTEGER PRIMARY KEY AUTOINCREMENT,
            lot_id              INTEGER NOT NULL UNIQUE REFERENCES production_lots(id),
            started_by          TEXT,
            started_at          TEXT NOT NULL,
            model_snapshot_json TEXT NOT NULL,
            lot_snapshot_json   TEXT NOT NULL,
            item_version_id     INTEGER NOT NULL REFERENCES product_model_item_versions(id),
            item_version        INTEGER NOT NULL,
            items_json          TEXT NOT NULL,
            snapshot_hash       TEXT NOT NULL,
            units_json          TEXT NOT NULL DEFAULT '[]'
        )
    """)
    _ins_cols = [r[1] for r in cur.execute("PRAGMA table_info(production_lot_inspections)").fetchall()]
    if "units_json" not in _ins_cols:
        _safe_add_column(cur, "ALTER TABLE production_lot_inspections ADD COLUMN units_json TEXT NOT NULL DEFAULT '[]'")
    # ---- 2026-10-10 신규 끝 ----

    # ---- 2026-10-10 신규: 제품검사 측정값(TASK #139 패키지 1D-b1) ----
    # lot_inspection_results는 append-only — UPDATE/DELETE 구문·함수·UI를 만들지 말 것(정정은 rev+1 새 행).
    cur.execute("""
        CREATE TABLE IF NOT EXISTS lot_inspection_results (
            id            INTEGER PRIMARY KEY AUTOINCREMENT,
            inspection_id INTEGER NOT NULL REFERENCES production_lot_inspections(id),
            unit_id       INTEGER NOT NULL REFERENCES finished_units(id),
            item_key      TEXT NOT NULL,
            idx           INTEGER NOT NULL,
            rev           INTEGER NOT NULL,
            entry_type    TEXT NOT NULL,
            value_input   TEXT,
            value         TEXT,
            is_missing    INTEGER NOT NULL DEFAULT 0,
            verdict       TEXT NOT NULL,
            reason        TEXT,
            roster_member INTEGER NOT NULL,
            serial_at_entry TEXT,
            lot_revision  INTEGER NOT NULL,
            entered_by    TEXT,
            entered_by_user_id INTEGER,
            entered_at    TEXT NOT NULL,
            file_sha256   TEXT,
            UNIQUE(inspection_id, unit_id, item_key, idx, rev)
        )
    """)
    _lir_cols = [r[1] for r in cur.execute("PRAGMA table_info(lot_inspection_results)").fetchall()]
    if "entered_by_user_id" not in _lir_cols:
        _safe_add_column(cur, "ALTER TABLE lot_inspection_results ADD COLUMN entered_by_user_id INTEGER")
    if "file_sha256" not in _lir_cols:
        _safe_add_column(cur, "ALTER TABLE lot_inspection_results ADD COLUMN file_sha256 TEXT")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_lir_slot ON lot_inspection_results"
                "(inspection_id, unit_id, item_key, idx)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_lir_unit ON lot_inspection_results(inspection_id, unit_id)")
    # ---- 2026-10-10 신규(1D-b1) 끝 ----

    # ---- 2026-10-11 신규: 제품검사 완료·재개 (TASK #145 패키지 1D-c1) ----
    # lot_inspection_unit_events는 append-only — UPDATE/DELETE 구문·함수·UI를 만들지 말 것.
    # 현재 상태 = unit별 MAX(seq) 행의 event_type.
    cur.execute("""
        CREATE TABLE IF NOT EXISTS lot_inspection_unit_events (
            id            INTEGER PRIMARY KEY AUTOINCREMENT,
            inspection_id INTEGER NOT NULL REFERENCES production_lot_inspections(id),
            unit_id       INTEGER NOT NULL REFERENCES finished_units(id),
            seq           INTEGER NOT NULL,
            event_type    TEXT NOT NULL,
            source        TEXT NOT NULL,
            reason        TEXT,
            results_digest TEXT NOT NULL,
            summary_json  TEXT,
            roster_member INTEGER NOT NULL,
            serial_at_event TEXT,
            lot_revision  INTEGER NOT NULL,
            by_name       TEXT,
            by_user_id    INTEGER,
            at            TEXT NOT NULL,
            UNIQUE(inspection_id, unit_id, seq)
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_liue_unit ON lot_inspection_unit_events(inspection_id, unit_id)")
    # ---- 2026-10-11 신규(1D-c1) 끝 ----

    # ---- 2026-10-11 신규: 제품 성적서 승인 (TASK #148 패키지 4a) ----
    # 두 테이블 모두 UPDATE/DELETE 금지(append-only). 상태는 저장하지 않고 revision별 마지막 이벤트로 도출한다.
    # event_type은 CHECK 제약 없음(4e가 review 계열 추가). CREATE IF NOT EXISTS만 — 기존 테이블 ALTER 없음.
    cur.execute("""
        CREATE TABLE IF NOT EXISTS product_report_revisions (
            id            INTEGER PRIMARY KEY AUTOINCREMENT,
            inspection_id INTEGER NOT NULL REFERENCES production_lot_inspections(id),
            unit_id       INTEGER NOT NULL REFERENCES finished_units(id),
            submission_no INTEGER NOT NULL,
            results_digest TEXT NOT NULL,
            completion_seq INTEGER NOT NULL,
            snapshot_json TEXT NOT NULL,
            content_hash  TEXT NOT NULL,
            serial_at_submit TEXT,
            lot_revision  INTEGER,
            writer_user_id INTEGER,
            writer_name   TEXT,
            approver_user_id INTEGER,
            approver_name TEXT,
            created_at    TEXT NOT NULL,
            UNIQUE(inspection_id, unit_id, submission_no)
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_prr_unit ON product_report_revisions(inspection_id, unit_id)")
    cur.execute("""
        CREATE TABLE IF NOT EXISTS product_report_events (
            id            INTEGER PRIMARY KEY AUTOINCREMENT,
            revision_id   INTEGER NOT NULL REFERENCES product_report_revisions(id),
            seq           INTEGER NOT NULL,
            event_type    TEXT NOT NULL,
            reason        TEXT,
            reject_kind   TEXT,
            doc_rev       INTEGER,
            actor_name    TEXT,
            actor_user_id INTEGER,
            payload_json  TEXT,
            prev_event_hash TEXT,
            event_hash    TEXT NOT NULL,
            at            TEXT NOT NULL,
            UNIQUE(revision_id, seq)
        )
    """)
    # ---- 2026-10-11 신규(4a) 끝 ----

    # ---- 2026-10-11 신규: NG 사건 (TASK #150 패키지 4b-1) ----
    # 두 테이블 모두 UPDATE/DELETE 금지(append-only). 사건 상태는 저장하지 않고 ng_event_log에서 파생한다.
    # log_type/action_kind/close_kind에 CHECK 없음(4b-2가 cancelled·unusable 추가). CREATE IF NOT EXISTS만.
    cur.execute("""
        CREATE TABLE IF NOT EXISTS ng_events (
            id            INTEGER PRIMARY KEY AUTOINCREMENT,
            inspection_id INTEGER NOT NULL REFERENCES production_lot_inspections(id),
            unit_id       INTEGER NOT NULL REFERENCES finished_units(id),
            item_key      TEXT NOT NULL,
            event_no      INTEGER NOT NULL,
            origin_event_id INTEGER,
            origin_round_no INTEGER,
            opened_trigger TEXT NOT NULL,
            fail_snapshot_json TEXT NOT NULL,
            header_hash   TEXT NOT NULL,
            serial_at_open TEXT,
            lot_revision  INTEGER,
            opened_by     TEXT,
            opened_by_user_id INTEGER,
            opened_at     TEXT NOT NULL,
            UNIQUE(inspection_id, unit_id, event_no)
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_nge_unit ON ng_events(inspection_id, unit_id)")
    cur.execute("""
        CREATE TABLE IF NOT EXISTS ng_event_log (
            id            INTEGER PRIMARY KEY AUTOINCREMENT,
            ng_event_id   INTEGER NOT NULL REFERENCES ng_events(id),
            seq           INTEGER NOT NULL,
            log_type      TEXT NOT NULL,
            action_kind   TEXT,
            round_no      INTEGER,
            outcome       TEXT,
            close_kind    TEXT,
            reason        TEXT,
            memo          TEXT,
            revision_id   INTEGER,
            report_event_seq INTEGER,
            payload_json  TEXT,
            actor_name    TEXT,
            actor_user_id INTEGER,
            prev_hash     TEXT,
            event_hash    TEXT NOT NULL,
            at            TEXT NOT NULL,
            UNIQUE(ng_event_id, seq)
        )
    """)
    # ---- 2026-10-11 신규(4b-1) 끝 ----

    # ---- 2026-10-11 신규: S/N 사용 금지 결정 이력 (TASK #154 패키지 4b-2) ----
    # S/N 단위 append-only 해시 체인. event_type 'unusable'|'unusable_revoked'(CHECK 없음). 트리거 설치보다 먼저 만들 것.
    cur.execute("""
        CREATE TABLE IF NOT EXISTS unit_disposal_events (
            id            INTEGER PRIMARY KEY AUTOINCREMENT,
            inspection_id INTEGER NOT NULL REFERENCES production_lot_inspections(id),
            unit_id       INTEGER NOT NULL REFERENCES finished_units(id),
            seq           INTEGER NOT NULL,
            event_type    TEXT NOT NULL,
            reason        TEXT NOT NULL,
            payload_json  TEXT,
            serial_at_event TEXT,
            lot_revision  INTEGER,
            actor_name    TEXT,
            actor_user_id INTEGER,
            prev_hash     TEXT,
            event_hash    TEXT NOT NULL,
            at            TEXT NOT NULL,
            UNIQUE(inspection_id, unit_id, seq)
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_ude_unit ON unit_disposal_events(unit_id, seq)")
    # ---- 2026-10-11 신규(4b-2) 끝 ----

    # ---- 2026-10-11 신규: 증거성 테이블 7개 append-only DB 트리거 (TASK #153) ----
    _install_append_only_triggers(cur, APPEND_ONLY_TABLES)
    # ---- 2026-10-11 신규(트리거) 끝 ----

    # ---- 2026-09-15 확장: 차수(=배치) 계획 ----
    existing_ob_cols = [row[1] for row in cur.execute("PRAGMA table_info(outbound_batches)").fetchall()]
    for col in ("round_no", "confirmed_by", "confirmed_at", "confirm_signature"):
        if col not in existing_ob_cols:
            cur.execute(f"ALTER TABLE outbound_batches ADD COLUMN {col} TEXT")

    cur.execute("""
        CREATE TABLE IF NOT EXISTS outbound_planned_items (
            id         INTEGER PRIMARY KEY AUTOINCREMENT,
            batch_id   INTEGER NOT NULL REFERENCES outbound_batches(id),
            serial_no  TEXT NOT NULL,
            created_at TEXT DEFAULT (datetime('now','localtime'))
        )
    """)
    # ---- 2026-09-15 확장(차수 계획) 끝 ----

    # ---- 2026-09-15 확장: QR 라벨 출력 이력 ----
    cur.execute("""
        CREATE TABLE IF NOT EXISTS outbound_qr_exports (
            id            INTEGER PRIMARY KEY AUTOINCREMENT,
            batch_id      INTEGER NOT NULL REFERENCES outbound_batches(id),
            generated_by  TEXT,
            generated_at  TEXT DEFAULT (datetime('now','localtime')),
            item_count    INTEGER
        )
    """)
    # ---- 2026-09-15 확장(QR 라벨 출력 이력) 끝 ----

    # ---- 2026-09-26 확장: 하우징 성적서 프로젝트 1단계 (차수별 비고·시각, 검사자 집합,
    #      심화검사) — PROGRESS.md 2026-09-25 설계 확정분 참고 ----
    existing_fi_hdr_cols = [row[1] for row in cur.execute("PRAGMA table_info(full_inspections)").fetchall()]
    for col in ("stage1_remark", "stage2_remark", "stage3_remark"):
        if col not in existing_fi_hdr_cols:
            _safe_add_column(cur, f"ALTER TABLE full_inspections ADD COLUMN {col} TEXT DEFAULT ''")
    for col in ("stage1_started_at", "stage1_completed_at",
                "stage2_started_at", "stage2_completed_at",
                "stage3_started_at", "stage3_completed_at"):
        if col not in existing_fi_hdr_cols:
            _safe_add_column(cur, f"ALTER TABLE full_inspections ADD COLUMN {col} TEXT")

    # 칸을 저장할 때마다 그 사용자를 이 성적서를 건드린 사람 집합에 자동 추가(누적, 삭제 없음)
    # — inspection_progress(일회성 프레즌스 추적, 제출 시 삭제됨)와 목적이 다르다.
    cur.execute("""
        CREATE TABLE IF NOT EXISTS full_inspection_editors (
            id                  INTEGER PRIMARY KEY AUTOINCREMENT,
            full_inspection_id  INTEGER NOT NULL,
            user_id             INTEGER,
            display_name        TEXT,
            created_at          TEXT DEFAULT (datetime('now','localtime')),
            UNIQUE(full_inspection_id, user_id)
        )
    """)

    # 심화검사(PD/충격 불량 시 추가 시험) — 이번 단계는 테이블만 생성(채우는 화면은 다음 단계)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS housing_deep_inspections (
            id                       INTEGER PRIMARY KEY AUTOINCREMENT,
            full_inspection_unit_id  INTEGER NOT NULL,
            test_type                TEXT NOT NULL,
            vi_state                 TEXT,
            chamber_stage            TEXT,
            readings_json            TEXT DEFAULT '{}',
            result                   TEXT,
            remark                   TEXT DEFAULT '',
            created_by               TEXT,
            created_at               TEXT DEFAULT (datetime('now','localtime')),
            FOREIGN KEY (full_inspection_unit_id) REFERENCES full_inspection_units(id)
        )
    """)
    # ---- 2026-09-26 확장 끝 ----

    # ---- 신규 삽입 끝 ----

    conn.commit()
    conn.close()


# ---------- 사용자 계정 ----------

def get_user_by_username(username):
    conn = get_conn()
    row = conn.execute("SELECT * FROM users WHERE username = ?", (username,)).fetchone()
    conn.close()
    return row


def get_user(user_id):
    conn = get_conn()
    row = conn.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
    conn.close()
    return row


def list_users():
    conn = get_conn()
    rows = conn.execute("SELECT * FROM users ORDER BY username").fetchall()
    conn.close()
    return rows


def count_users():
    conn = get_conn()
    n = conn.execute("SELECT COUNT(*) AS n FROM users").fetchone()["n"]
    conn.close()
    return n


def create_user(username, password, display_name, permissions=""):
    conn = get_conn()
    conn.execute("""
        INSERT INTO users (username, password, display_name, permissions)
        VALUES (?, ?, ?, ?)
    """, (username, password, display_name, permissions))
    conn.commit()
    conn.close()


def update_user_permissions(user_id, permissions):
    conn = get_conn()
    conn.execute("UPDATE users SET permissions = ? WHERE id = ?", (permissions, user_id))
    conn.commit()
    conn.close()


def user_has_permission(user_row, perm):
    if user_row is None:
        return False
    perms = (user_row["permissions"] or "").split(",")
    return perm in perms


# ---------- 최종결정권자 ----------

MAX_FINAL_APPROVERS = 2   # 최종결정권자는 최대 2명까지만 지정할 수 있다


def list_final_approvers():
    """최종결정권자로 지정된 계정 목록."""
    conn = get_conn()
    rows = conn.execute(
        "SELECT * FROM users WHERE is_final_approver = 1 ORDER BY id"
    ).fetchall()
    conn.close()
    return rows


def set_final_approver(user_id, enabled):
    """최종결정권자 지정/해제. 최대 인원을 넘기면 (False, 에러메시지)."""
    conn = get_conn()
    try:
        user = conn.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
        if user is None:
            return False, "존재하지 않는 계정이야."

        if enabled:
            if not user_has_permission(user, "approve"):
                return False, "최종결정권자로 지정하려면 먼저 '승인' 권한을 줘야 해."
            current = conn.execute(
                "SELECT COUNT(*) FROM users WHERE is_final_approver = 1 AND id != ?", (user_id,)
            ).fetchone()[0]
            if current >= MAX_FINAL_APPROVERS:
                return False, (f"최종결정권자는 최대 {MAX_FINAL_APPROVERS}명까지야. "
                               f"다른 사람 체크를 먼저 풀어줘.")

        conn.execute("UPDATE users SET is_final_approver = ? WHERE id = ?",
                     (1 if enabled else 0, user_id))
        conn.commit()
        return True, None
    finally:
        conn.close()


def delete_user(user_id):
    conn = get_conn()
    conn.execute("DELETE FROM users WHERE id = ?", (user_id,))
    conn.commit()
    conn.close()


def update_user_password(user_id, password):
    conn = get_conn()
    conn.execute("UPDATE users SET password = ? WHERE id = ?", (password, user_id))
    conn.commit()
    conn.close()


def update_user_profile(user_id, username, display_name):
    """계정 아이디/표시이름 수정. username 중복이면 (False, 에러메시지) 반환."""
    conn = get_conn()
    dup = conn.execute("SELECT id FROM users WHERE username = ? AND id != ?", (username, user_id)).fetchone()
    if dup:
        conn.close()
        return False, f"아이디 '{username}'는 이미 사용 중이야."
    conn.execute("UPDATE users SET username = ?, display_name = ? WHERE id = ?",
                (username, display_name, user_id))
    conn.commit()
    conn.close()
    return True, None


# ---------- 활동 로그 ----------

def log_activity(user_id, username, role, action, target_type=None, target_id=None, detail=None):
    conn = get_conn()
    cur = conn.cursor()
    cur.execute("""
        INSERT INTO activity_log (user_id, username, role, action, target_type, target_id, detail)
        VALUES (?, ?, ?, ?, ?, ?, ?)
    """, (user_id, username, role, action, target_type, str(target_id) if target_id is not None else None, detail))
    log_id = cur.lastrowid
    conn.commit()
    conn.close()
    return log_id


def list_activity_logs(limit=300):
    conn = get_conn()
    rows = conn.execute(
        "SELECT * FROM activity_log ORDER BY id DESC LIMIT ?", (limit,)
    ).fetchall()
    conn.close()
    return rows


# ---------- 규격표 조회/등록 ----------

def get_materials():
    """등록된 모든 자재번호 목록 (항목이 아직 없는 자재도 포함)"""
    conn = get_conn()
    rows = conn.execute(
        """SELECT m.*, mc.major_category, mc.middle_category
             FROM materials m
             LEFT JOIN material_classification_map mc ON mc.material_no = m.material_no
            ORDER BY m.material_no"""
    ).fetchall()
    conn.close()
    return rows


def get_material(material_no):
    conn = get_conn()
    row = conn.execute("""SELECT m.*, mc.major_category, mc.middle_category
                            FROM materials m
                            LEFT JOIN material_classification_map mc ON mc.material_no = m.material_no
                           WHERE m.material_no = ?""", (material_no,)).fetchone()
    conn.close()
    return row


def upsert_material(material_no, material_name=None):
    conn = get_conn()
    existing = conn.execute("SELECT material_no FROM materials WHERE material_no = ?", (material_no,)).fetchone()
    if existing:
        if material_name:
            conn.execute("UPDATE materials SET material_name = ? WHERE material_no = ?", (material_name, material_no))
    else:
        today = _dt.now().strftime("%Y-%m-%d")
        conn.execute("""
            INSERT INTO materials (material_no, material_name, revision_date)
            VALUES (?, ?, ?)
        """, (material_no, material_name, today))
    conn.commit()
    conn.close()


def update_drawing_file(material_no, drawing_file):
    """수동 도면 파일 지정 — drawing_file은 도면 폴더 내 파일명(예: 'CKMR4610.pdf')"""
    conn = get_conn()
    conn.execute("UPDATE materials SET drawing_file = ? WHERE material_no = ?",
                 (drawing_file or None, material_no))
    conn.commit()
    conn.close()


def update_material_standard_info(material_no, drawing_version, revision_date, edition, unit):
    """기준서(SAM 양식)에 들어가는 자재별 정보 — 도면버전/개정일자/판수/단위."""
    conn = get_conn()
    conn.execute("""
        UPDATE materials SET drawing_version = ?, revision_date = ?, edition = ?, unit = ?
        WHERE material_no = ?
    """, (drawing_version, revision_date, edition, unit, material_no))
    conn.commit()
    conn.close()


def rename_material(old_no, new_no, new_name):
    """
    자재번호/자재명 변경 — 지금부터만 적용(과거 입고/검사/성적서 기록은 옛 번호 그대로 둠).
    materials·specs·material_group_items(부품으로 소속된 경우)와 새 분류 매핑을 새 번호로 갱신.
    반환: (성공여부, 에러메시지)
    """
    conn = get_conn()
    try:
        # 중복/분류 충돌 확인과 실제 번호 변경 사이에 다른 쓰기가 끼어들지 않도록 한 트랜잭션에서
        # 처리한다. 오류가 나면 자재·규격·분류가 서로 다른 번호를 가리키지 않게 전체 롤백한다.
        conn.execute("BEGIN IMMEDIATE")
        old_mapping = None
        new_mapping = None
        if old_no != new_no:
            dup = conn.execute("SELECT 1 FROM materials WHERE material_no = ?", (new_no,)).fetchone()
            if dup:
                conn.rollback()
                return False, f"자재번호 '{new_no}'는 이미 사용 중이야."
            dup_group = conn.execute("SELECT 1 FROM material_groups WHERE group_no = ?", (new_no,)).fetchone()
            if dup_group:
                conn.rollback()
                return False, f"'{new_no}'는 이미 조립품 그룹 번호로 쓰이고 있어."

            old_mapping = conn.execute("""
                SELECT major_category, middle_category
                  FROM material_classification_map
                 WHERE material_no = ?
            """, (old_no,)).fetchone()
            new_mapping = conn.execute("""
                SELECT major_category, middle_category
                  FROM material_classification_map
                 WHERE material_no = ?
            """, (new_no,)).fetchone()
            if (old_mapping and new_mapping and
                    (old_mapping["major_category"], old_mapping["middle_category"]) !=
                    (new_mapping["major_category"], new_mapping["middle_category"])):
                conn.rollback()
                return False, (f"자재번호 '{new_no}'에 서로 다른 대분류/중분류 매핑이 있어 "
                               "번호를 변경할 수 없어.")

        existing = conn.execute("SELECT 1 FROM materials WHERE material_no = ?", (old_no,)).fetchone()
        if existing:
            conn.execute("UPDATE materials SET material_no = ?, material_name = ? WHERE material_no = ?",
                        (new_no, new_name, old_no))
        else:
            # old_no가 아직 자재로 등록 안 된 상태(예: 검사대기 목록에서 "규격 미등록" 링크로 바로
            # 들어와서 자재명만 처음 입력하는 경우) — UPDATE는 매칭되는 행이 없어 조용히 아무 일도
            # 안 하므로, 이 경우엔 새로 INSERT해야 한다.
            today = _dt.now().strftime("%Y-%m-%d")
            conn.execute("INSERT INTO materials (material_no, material_name, revision_date) VALUES (?, ?, ?)",
                        (new_no, new_name, today))
        conn.execute("UPDATE specs SET material_no = ?, material_name = ? WHERE material_no = ?",
                    (new_no, new_name, old_no))
        conn.execute("UPDATE material_group_items SET material_no = ? WHERE material_no = ?",
                    (new_no, old_no))

        if old_no != new_no and old_mapping:
            if new_mapping:
                # 동일한 매핑이 새 번호에 미리 준비된 경우 PK 충돌을 피하고 낡은 키만 제거한다.
                conn.execute("DELETE FROM material_classification_map WHERE material_no = ?", (old_no,))
            else:
                conn.execute("""
                    UPDATE material_classification_map
                       SET material_no = ?, updated_at = CURRENT_TIMESTAMP
                     WHERE material_no = ?
                """, (new_no, old_no))

        conn.commit()
        return True, None
    except sqlite3.Error as exc:
        conn.rollback()
        return False, f"자재번호 변경 중 데이터베이스 오류가 발생했어: {exc}"
    finally:
        conn.close()


def upsert_materials_bulk(rows):
    """rows: [{"material_no":.., "material_name":..}, ...]"""
    conn = get_conn()
    today = _dt.now().strftime("%Y-%m-%d")
    for r in rows:
        existing = conn.execute("SELECT material_no FROM materials WHERE material_no = ?", (r["material_no"],)).fetchone()
        if existing:
            if r.get("material_name"):
                conn.execute("UPDATE materials SET material_name = ? WHERE material_no = ?",
                            (r["material_name"], r["material_no"]))
        else:
            conn.execute("""
                INSERT INTO materials (material_no, material_name, revision_date)
                VALUES (?, ?, ?)
            """, (r["material_no"], r.get("material_name"), today))
    conn.commit()
    conn.close()


def list_material_categories():
    """등록된 자재 분류명 전체 (가나다순)."""
    conn = get_conn()
    rows = conn.execute("SELECT name FROM material_categories ORDER BY name").fetchall()
    conn.close()
    return [r["name"] for r in rows]


def add_material_category(name):
    """분류명을 마스터에 등록. 이미 있으면(공백/대소문자 무시) 기존 정본 표기를 반환.
    name이 빈 값이면 None."""
    name = (name or "").strip()
    if not name:
        return None
    conn = get_conn()
    for row in conn.execute("SELECT name FROM material_categories").fetchall():
        if row["name"].strip().casefold() == name.casefold():
            conn.close()
            return row["name"]
    conn.execute("INSERT INTO material_categories (name) VALUES (?)", (name,))
    conn.commit()
    conn.close()
    return name


def update_material_category(material_no, category):
    """자재 하나의 분류 저장. category가 비어있지 않으면 마스터에 자동 등록 후 정본
    표기로 저장. 빈 값이면 NULL로 해제. 반환값: 실제 저장된 분류명(또는 None)."""
    canonical = add_material_category(category) if (category or "").strip() else None
    conn = get_conn()
    conn.execute("UPDATE materials SET category = ? WHERE material_no = ?", (canonical, material_no))
    conn.commit()
    conn.close()
    return canonical


UNASSIGNED_CLASSIFICATION = "__unassigned__"


def list_material_classification_pairs(major_category=None):
    """대분류/중분류 조합 목록. 화면의 종속 드롭다운이 공통으로 사용한다."""
    conn = get_conn()
    params = []
    sql = "SELECT major_category, middle_category FROM material_classification_pairs"
    major_category = (major_category or "").strip()
    if major_category:
        sql += " WHERE major_category = ?"
        params.append(major_category)
    sql += " ORDER BY major_category, middle_category"
    rows = [dict(r) for r in conn.execute(sql, params).fetchall()]
    conn.close()
    return rows


def get_material_classification(material_no):
    """자재번호의 확정 대분류·중분류 매핑. 미지정이면 None."""
    conn = get_conn()
    row = conn.execute("""
        SELECT major_category, middle_category
          FROM material_classification_map
         WHERE material_no = ?
    """, ((material_no or "").strip(),)).fetchone()
    conn.close()
    return row


def list_major_categories():
    return list(dict.fromkeys(
        r["major_category"] for r in list_material_classification_pairs()
    ))


def add_material_classification_pair(major_category, middle_category, conn=None, source="manual"):
    """대/중분류 조합을 등록하고 공백을 제거한 정본 tuple을 반환한다."""
    major = (major_category or "").strip()
    middle = (middle_category or "").strip()
    if not major or not middle:
        raise ValueError("대분류와 중분류를 모두 입력해줘.")
    own_conn = conn is None
    conn = conn or get_conn()
    try:
        conn.execute("""INSERT OR IGNORE INTO material_classification_pairs
                        (major_category, middle_category, source) VALUES (?, ?, ?)""",
                     (major, middle, source or "manual"))
        if own_conn:
            conn.commit()
        return major, middle
    finally:
        if own_conn:
            conn.close()


def update_material_classification(material_no, major_category, middle_category, source="manual"):
    """자재번호의 대/중분류 저장. 둘 다 비우면 미지정(매핑 삭제)으로 되돌린다."""
    material_no = (material_no or "").strip()
    major = (major_category or "").strip()
    middle = (middle_category or "").strip()
    if not material_no:
        raise ValueError("자재번호가 없어 분류를 저장할 수 없어.")
    if bool(major) != bool(middle):
        raise ValueError("대분류와 중분류를 함께 선택해줘.")
    conn = get_conn()
    try:
        if not major:
            conn.execute("DELETE FROM material_classification_map WHERE material_no = ?", (material_no,))
            conn.commit()
            return None
        add_material_classification_pair(major, middle, conn=conn)
        conn.execute("""
            INSERT INTO material_classification_map
                (material_no, major_category, middle_category, source, updated_at)
            VALUES (?, ?, ?, ?, datetime('now', 'localtime'))
            ON CONFLICT(material_no) DO UPDATE SET
                major_category=excluded.major_category,
                middle_category=excluded.middle_category,
                source=excluded.source,
                updated_at=datetime('now', 'localtime')
        """, (material_no, major, middle, source or "manual"))
        conn.commit()
        return major, middle
    finally:
        conn.close()


def ensure_material_classification_asset(asset_path=None):
    """저장소의 확정 분류 JSON을 파일 해시 기준으로 멱등 반영한다.

    필수값 공백, 같은 자재번호의 상충 분류, 메타데이터 건수 불일치는 DB 쓰기 전에
    거부한다. 새 파일에서는 이전 asset 매핑을 현재 파일 기준으로 교체하지만, 파일 밖에서
    담당자가 수동 추가한 매핑은 보존한다. materials의 이름과 기존 category는 건드리지 않는다.
    """
    if asset_path is None:
        asset_path = os.path.join(os.path.dirname(__file__), "data", "material_classifications.json")
    with open(asset_path, "rb") as f:
        raw = f.read()
    asset_hash = hashlib.sha256(raw).hexdigest()
    payload = json.loads(raw.decode("utf-8"))
    raw_mappings = payload.get("mappings")
    if not isinstance(raw_mappings, list):
        raise ValueError("자재 분류 데이터의 mappings 배열이 없어.")

    normalized = {}
    for idx, item in enumerate(raw_mappings, start=1):
        if not isinstance(item, dict):
            raise ValueError(f"자재 분류 데이터 {idx}번째 행 형식이 올바르지 않아.")
        material_no = str(item.get("material_no") or "").strip()
        major = str(item.get("major_category") or "").strip()
        middle = str(item.get("middle_category") or "").strip()
        if not material_no or not major or not middle:
            raise ValueError(f"자재 분류 데이터 {idx}번째 행에 필수값 공백이 있어.")
        pair = (major, middle)
        if material_no in normalized and normalized[material_no] != pair:
            raise ValueError(f"자재번호 '{material_no}'의 대/중분류가 서로 충돌해.")
        normalized[material_no] = pair

    expected_count = payload.get("mapping_count")
    if expected_count is not None and int(expected_count) != len(normalized):
        raise ValueError("자재 분류 데이터의 mapping_count와 실제 고유 자재 수가 달라.")

    conn = get_conn()
    try:
        # Render의 여러 gunicorn 워커가 동시에 시작해도 한 워커씩 해시를 확인·반영한다.
        # SELECT 뒤에 쓰기 트랜잭션으로 승격하면 stale snapshot/locked 경합이 날 수 있으므로
        # 처음부터 IMMEDIATE로 직렬화한다(데이터 914건이라 잠금 구간은 매우 짧다).
        conn.execute("BEGIN IMMEDIATE")
        current = conn.execute(
            "SELECT value FROM app_settings WHERE key = 'material_classification_asset_sha256'"
        ).fetchone()
        if current and current["value"] == asset_hash:
            conn.rollback()
            return {"applied": False, "count": len(normalized), "sha256": asset_hash}

        conn.execute("DELETE FROM material_classification_map WHERE source = 'asset'")
        conn.execute("""DELETE FROM material_classification_pairs
                         WHERE source = 'asset'
                           AND NOT EXISTS (
                               SELECT 1 FROM material_classification_map mc
                                WHERE mc.major_category = material_classification_pairs.major_category
                                  AND mc.middle_category = material_classification_pairs.middle_category
                           )""")
        for major, middle in sorted(set(normalized.values())):
            add_material_classification_pair(major, middle, conn=conn, source="asset")
        conn.executemany("""
            INSERT INTO material_classification_map
                (material_no, major_category, middle_category, source, updated_at)
            VALUES (?, ?, ?, 'asset', datetime('now', 'localtime'))
            ON CONFLICT(material_no) DO UPDATE SET
                major_category=excluded.major_category,
                middle_category=excluded.middle_category,
                source='asset',
                updated_at=datetime('now', 'localtime')
        """, [(material_no, *normalized[material_no]) for material_no in sorted(normalized)])
        conn.execute("""INSERT INTO app_settings(key, value) VALUES
                      ('material_classification_asset_sha256', ?)
                      ON CONFLICT(key) DO UPDATE SET value=excluded.value""", (asset_hash,))
        conn.commit()
        return {"applied": True, "count": len(normalized), "sha256": asset_hash}
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


CHECK_INSPECTION_AQL_MIGRATION_KEY = "check_inspection_aql_from_classification_v1"


def ensure_check_inspection_aql_migration(asset_path=None):
    """확정 분류 자산의 241개 대상 자재 기존 규격을 체크검사로 한 번만 전환한다.

    완료 플래그와 specs 갱신을 같은 IMMEDIATE 트랜잭션에서 처리한다. 플래그가 기록된
    뒤에는 분류나 사용자가 바꾼 AQL을 다시 덮어쓰지 않는다. 자산에 대상 번호가 정확히
    241개가 아니면 운영 데이터를 건드리기 전에 중단한다.
    """
    # 완료된 DB는 이후 배포에서 분류 자산이 바뀌거나 사라져도 영향을 받지 않아야 한다.
    # 따라서 파일 경로 계산·open·JSON 파싱보다 먼저 플래그를 확인한다.
    conn = get_conn()
    try:
        done = conn.execute(
            "SELECT value FROM app_settings WHERE key = ?",
            (CHECK_INSPECTION_AQL_MIGRATION_KEY,)
        ).fetchone()
    finally:
        conn.close()
    if done:
        return {"applied": False, "target_count": 241, "updated_spec_count": 0}

    if asset_path is None:
        asset_path = os.path.join(os.path.dirname(__file__), "data", "material_classifications.json")
    with open(asset_path, "rb") as f:
        payload = json.loads(f.read().decode("utf-8"))
    mappings = payload.get("mappings")
    if not isinstance(mappings, list):
        raise ValueError("자재 분류 데이터의 mappings 배열이 없어.")

    target_material_nos = set()
    seen = {}
    for idx, item in enumerate(mappings, start=1):
        if not isinstance(item, dict):
            raise ValueError(f"자재 분류 데이터 {idx}번째 행 형식이 올바르지 않아.")
        material_no = str(item.get("material_no") or "").strip()
        major = str(item.get("major_category") or "").strip()
        middle = str(item.get("middle_category") or "").strip()
        if not material_no or not major or not middle:
            raise ValueError(f"자재 분류 데이터 {idx}번째 행에 필수값 공백이 있어.")
        pair = (major, middle)
        if material_no in seen and seen[material_no] != pair:
            raise ValueError(f"자재번호 '{material_no}'의 대/중분류가 서로 충돌해.")
        seen[material_no] = pair
        if (major in ("기타 소자류", "조립부품류 (소모품)")
                or (major == "상품" and middle == "엘보/부싱 액세서리")):
            target_material_nos.add(material_no)

    if len(target_material_nos) != 241:
        raise ValueError(
            f"체크검사 전환 대상은 정확히 241개여야 하는데 {len(target_material_nos)}개야.")

    conn = get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        # 선조회와 트랜잭션 시작 사이에 다른 worker가 먼저 끝냈을 수 있으므로 재확인한다.
        done = conn.execute(
            "SELECT value FROM app_settings WHERE key = ?",
            (CHECK_INSPECTION_AQL_MIGRATION_KEY,)
        ).fetchone()
        if done:
            conn.rollback()
            return {"applied": False, "target_count": 241, "updated_spec_count": 0}

        material_nos = sorted(target_material_nos)
        placeholders = ",".join("?" for _ in material_nos)
        cursor = conn.execute(
            f"UPDATE specs SET aql = '체크검사' WHERE material_no IN ({placeholders})",
            material_nos,
        )
        updated_count = cursor.rowcount
        conn.execute(
            "INSERT INTO app_settings(key, value) VALUES (?, '1')",
            (CHECK_INSPECTION_AQL_MIGRATION_KEY,),
        )
        conn.commit()
        return {"applied": True, "target_count": 241,
                "updated_spec_count": updated_count}
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def _include_exclude_clause(columns, include_words=None, exclude_words=None):
    """포함/제외 다중 단어 검색 조건 생성 — 자재 관리(search_materials)/자재 찾기
    (search_bom_materials) 공용 헬퍼(8-1절 원칙, 복붙 금지). 포함 단어 여러 개 = OR
    (하나라도 매칭 컬럼에 있으면 통과), 제외 단어 여러 개 = OR(하나라도 있으면 제외 —
    드모르간 법칙으로 NOT (OR 그룹) 하나로 표현).

    columns: SQL 컬럼 표현식 리스트(예: ["m.material_no", "m.material_name"]).
    반환: (조건식 문자열 — 괄호 포함, 앞에 "AND" 없음, params 리스트). 조건이 없으면 ("", []).
    호출부에서 기존 SQL 조립 방식(리스트에 append 하거나 " AND {식}"으로 이어붙이는 등)에
    맞춰 붙이면 된다."""
    include_words = [w.strip() for w in (include_words or []) if w and w.strip()]
    exclude_words = [w.strip() for w in (exclude_words or []) if w and w.strip()]
    parts, params = [], []

    def _or_group(words):
        # COALESCE 필수 — search_bom_materials처럼 LEFT JOIN이라 컬럼이 NULL일 수 있으면
        # "NOT (... OR NULL ...)" 이 SQL 3치논리상 NULL이 되어 조건 없는 exclude에서도
        # 그 행이 통째로 빠져버린다(실측으로 확인한 버그, 2026-09-14).
        ors = []
        for w in words:
            like = f"%{w}%"
            for col in columns:
                ors.append(f"COALESCE({col}, '') LIKE ?")
                params.append(like)
        return "(" + " OR ".join(ors) + ")"

    if include_words:
        parts.append(_or_group(include_words))
    if exclude_words:
        parts.append("NOT " + _or_group(exclude_words))

    if not parts:
        return "", []
    return " AND ".join(parts), params


def search_materials(query=None, search_by="all", category=None, include=None, exclude=None,
                     major_category=None, middle_category=None):
    """
    query: 검색어. search_by: 'material_no' / 'material_name' / 'method' / 'spec' / 'all'
    'method'(검사방식)는 specs.inspect_method에서, 'spec'(규격 표기)는 specs.spec_display에서 매칭.
    'all'(전체)은 자재번호·자재명·규격 표기·검사방식을 모두 훑는다.
    category는 이전 단일분류 URL/호출 호환용이다. major_category/middle_category가 새
    대·중분류 필터이며 둘 다 있으면 AND 정확일치한다. '__unassigned__'는 미지정이다.
    include/exclude: 포함/제외 다중 단어(자재번호·자재명만 대상, 규격표기/검사방식은
    제외 — 확정된 결정). _include_exclude_clause() 참고.
    """
    conn = get_conn()
    category = (category or "").strip() or None
    cat_sql = " AND m.category = ?" if category else ""
    cat_params = [category] if category else []
    major_category = (major_category or "").strip() or None
    middle_category = (middle_category or "").strip() or None
    class_parts, class_params = [], []
    if major_category == UNASSIGNED_CLASSIFICATION or middle_category == UNASSIGNED_CLASSIFICATION:
        class_parts.append("mc.material_no IS NULL")
    else:
        if major_category:
            class_parts.append("mc.major_category = ?")
            class_params.append(major_category)
        if middle_category:
            class_parts.append("mc.middle_category = ?")
            class_params.append(middle_category)
    class_sql = "".join(f" AND {part}" for part in class_parts)
    ie_expr, ie_params = _include_exclude_clause(["m.material_no", "m.material_name"], include, exclude)
    ie_sql = f" AND {ie_expr}" if ie_expr else ""
    select_cols = "m.material_no, m.material_name, m.category, mc.major_category, mc.middle_category"
    class_join = " LEFT JOIN material_classification_map mc ON mc.material_no = m.material_no"

    if search_by == "method_empty":
        # 검사방식(inspect_method)이 비어 있는 항목을 가진 자재. 검색어와 무관하게 동작한다.
        rows = conn.execute(f"""
            SELECT DISTINCT {select_cols} FROM materials m{class_join}
            JOIN specs s ON s.material_no = m.material_no
            WHERE (s.inspect_method IS NULL OR TRIM(s.inspect_method) = ''){cat_sql}{class_sql}{ie_sql}
            ORDER BY m.material_no
        """, (*cat_params, *class_params, *ie_params)).fetchall()
        conn.close()
        return rows
    if not query:
        rows = conn.execute(f"""
            SELECT {select_cols} FROM materials m{class_join}
            WHERE 1=1{cat_sql}{class_sql}{ie_sql}
            ORDER BY m.material_no
        """, (*cat_params, *class_params, *ie_params)).fetchall()
        conn.close()
        return rows

    like = f"%{query}%"
    if search_by == "material_no":
        rows = conn.execute(f"""
            SELECT {select_cols} FROM materials m{class_join}
            WHERE m.material_no LIKE ?{cat_sql}{class_sql}{ie_sql}
            ORDER BY m.material_no
        """, (like, *cat_params, *class_params, *ie_params)).fetchall()
    elif search_by == "material_name":
        rows = conn.execute(f"""
            SELECT {select_cols} FROM materials m{class_join}
            WHERE m.material_name LIKE ?{cat_sql}{class_sql}{ie_sql}
            ORDER BY m.material_no
        """, (like, *cat_params, *class_params, *ie_params)).fetchall()
    elif search_by == "method":
        rows = conn.execute(f"""
            SELECT DISTINCT {select_cols} FROM materials m{class_join}
            JOIN specs s ON s.material_no = m.material_no
            WHERE s.inspect_method LIKE ?{cat_sql}{class_sql}{ie_sql}
            ORDER BY m.material_no
        """, (like, *cat_params, *class_params, *ie_params)).fetchall()
    elif search_by == "spec":
        rows = conn.execute(f"""
            SELECT DISTINCT {select_cols} FROM materials m{class_join}
            JOIN specs s ON s.material_no = m.material_no
            WHERE s.spec_display LIKE ?{cat_sql}{class_sql}{ie_sql}
            ORDER BY m.material_no
        """, (like, *cat_params, *class_params, *ie_params)).fetchall()
    else:
        # 전체: 자재번호·자재명·규격 표기·검사방식 어디에 있든 잡는다.
        # (규격 표기 spec_display를 빠뜨려서 규격에만 있는 검색어가 안 걸리던 버그 수정)
        rows = conn.execute(f"""
            SELECT DISTINCT {select_cols} FROM materials m{class_join}
            LEFT JOIN specs s ON s.material_no = m.material_no
            WHERE (m.material_no LIKE ? OR m.material_name LIKE ?
               OR s.spec_display LIKE ? OR s.inspect_method LIKE ?){cat_sql}{class_sql}{ie_sql}
            ORDER BY m.material_no
        """, (like, like, like, like, *cat_params, *class_params, *ie_params)).fetchall()
    conn.close()
    return rows


def get_specs_by_material(material_no):
    conn = get_conn()
    rows = conn.execute(
        "SELECT * FROM specs WHERE material_no = ? ORDER BY item_order, id",
        (material_no,)
    ).fetchall()
    conn.close()
    return rows


# ---------- 자재 그룹(조립품) ----------

def list_material_groups():
    conn = get_conn()
    rows = conn.execute("SELECT * FROM material_groups ORDER BY group_no").fetchall()
    conn.close()
    return rows


def get_material_group(group_no):
    conn = get_conn()
    row = conn.execute("SELECT * FROM material_groups WHERE group_no = ?", (group_no,)).fetchone()
    conn.close()
    return row


def get_group_members(group_no):
    """그룹에 속한 부품 자재번호 목록 (등록 순서대로), 각 부품의 material_name도 같이 반환."""
    conn = get_conn()
    rows = conn.execute("""
        SELECT gi.id, gi.group_no, gi.material_no, gi.item_order,
               (SELECT material_name FROM specs WHERE specs.material_no = gi.material_no LIMIT 1) AS material_name
        FROM material_group_items gi
        WHERE gi.group_no = ?
        ORDER BY gi.item_order, gi.id
    """, (group_no,)).fetchall()
    conn.close()
    return rows


def create_material_group(group_no, group_name):
    conn = get_conn()
    conn.execute("INSERT OR REPLACE INTO material_groups (group_no, group_name) VALUES (?, ?)",
                (group_no, group_name))
    conn.commit()
    conn.close()


def delete_material_group(group_no):
    conn = get_conn()
    conn.execute("DELETE FROM material_group_items WHERE group_no = ?", (group_no,))
    conn.execute("DELETE FROM material_groups WHERE group_no = ?", (group_no,))
    conn.commit()
    conn.close()


def add_group_member(group_no, material_no, item_order=0):
    conn = get_conn()
    conn.execute("""
        INSERT INTO material_group_items (group_no, material_no, item_order)
        VALUES (?, ?, ?)
    """, (group_no, material_no, item_order))
    conn.commit()
    conn.close()


def remove_group_member(member_id):
    conn = get_conn()
    conn.execute("DELETE FROM material_group_items WHERE id = ?", (member_id,))
    conn.commit()
    conn.close()


# ---------- 규격 일괄등록 확인필요 자재 ----------

def add_review_flag(material_no, source_filename, reason):
    conn = get_conn()
    conn.execute("""
        INSERT INTO spec_review_flags (material_no, source_filename, reason)
        VALUES (?, ?, ?)
    """, (material_no, source_filename, reason))
    conn.commit()
    conn.close()


def list_unresolved_review_flags():
    conn = get_conn()
    rows = conn.execute("""
        SELECT * FROM spec_review_flags WHERE resolved = 0 ORDER BY id DESC
    """).fetchall()
    conn.close()
    return rows


def resolve_review_flag(flag_id):
    conn = get_conn()
    conn.execute("UPDATE spec_review_flags SET resolved = 1 WHERE id = ?", (flag_id,))
    conn.commit()
    conn.close()


def add_spec(material_no, material_name, item_name, spec_display,
             judge_type, lower_limit, upper_limit, inspect_method, aql, item_order=0):
    conn = get_conn()
    cur = conn.execute("""
        INSERT INTO specs (material_no, material_name, item_name, spec_display,
                            judge_type, lower_limit, upper_limit, inspect_method, aql, item_order)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (material_no, material_name, item_name, spec_display,
          judge_type, lower_limit, upper_limit, inspect_method, aql, item_order))
    new_id = cur.lastrowid
    if not conn.execute("SELECT 1 FROM materials WHERE material_no = ?", (material_no,)).fetchone():
        conn.execute("INSERT INTO materials (material_no, material_name) VALUES (?, ?)",
                    (material_no, material_name))
    elif material_name:
        conn.execute("UPDATE materials SET material_name = ? WHERE material_no = ? AND (material_name IS NULL OR material_name = '')",
                    (material_name, material_no))
    conn.commit()
    conn.close()
    return new_id


def replace_specs_for_material(material_no, material_name, items):
    """
    한 자재의 규격표를 통째로 교체 등록 (파일 재업로드 시 중복 없이 덮어쓰기).
    items: [{"item_name","spec_display","judge_type","lower_limit","upper_limit",
             "inspect_method","aql","item_order"}]
    """
    conn = get_conn()
    cur = conn.cursor()
    cur.execute("DELETE FROM specs WHERE material_no = ?", (material_no,))
    for it in items:
        cur.execute("""
            INSERT INTO specs (material_no, material_name, item_name, spec_display,
                                judge_type, lower_limit, upper_limit, inspect_method, aql, item_order)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (material_no, material_name, it["item_name"], it["spec_display"],
              it["judge_type"], it["lower_limit"], it["upper_limit"],
              it["inspect_method"], it["aql"], it["item_order"]))
    if not cur.execute("SELECT 1 FROM materials WHERE material_no = ?", (material_no,)).fetchone():
        cur.execute("INSERT INTO materials (material_no, material_name) VALUES (?, ?)",
                    (material_no, material_name))
    elif material_name:
        cur.execute("UPDATE materials SET material_name = ? WHERE material_no = ?", (material_name, material_no))
    conn.commit()
    conn.close()


def get_spec_item(spec_id):
    conn = get_conn()
    row = conn.execute("SELECT * FROM specs WHERE id = ?", (spec_id,)).fetchone()
    conn.close()
    return row


def update_spec_item(spec_id, item_name, spec_display, judge_type,
                      lower_limit, upper_limit, inspect_method, aql, item_order):
    conn = get_conn()
    conn.execute("""
        UPDATE specs SET item_name=?, spec_display=?, judge_type=?, lower_limit=?,
                          upper_limit=?, inspect_method=?, aql=?, item_order=?
        WHERE id=?
    """, (item_name, spec_display, judge_type, lower_limit, upper_limit,
          inspect_method, aql, item_order, spec_id))
    conn.commit()
    conn.close()


def delete_spec_item(spec_id):
    conn = get_conn()
    conn.execute("DELETE FROM specs WHERE id = ?", (spec_id,))
    conn.commit()
    conn.close()


def delete_spec_items_bulk(spec_ids):
    if not spec_ids:
        return
    conn = get_conn()
    placeholders = ",".join("?" for _ in spec_ids)
    conn.execute(f"DELETE FROM specs WHERE id IN ({placeholders})", spec_ids)
    conn.commit()
    conn.close()


def delete_materials_bulk(material_nos):
    """자재 통째로(규격+마스터) 삭제 — 규격 목록 화면에서 자재 단위 선택삭제용."""
    if not material_nos:
        return
    conn = get_conn()
    placeholders = ",".join("?" for _ in material_nos)
    conn.execute(f"DELETE FROM specs WHERE material_no IN ({placeholders})", material_nos)
    conn.execute(f"DELETE FROM materials WHERE material_no IN ({placeholders})", material_nos)
    conn.commit()
    conn.close()


# ---------- 입고 리스트 (붙여넣기 등록) ----------

def find_duplicate_intakes(rows):
    """이미 intake_list에 등록된 동일 항목 찾기. 둘 중 하나라도 걸리면 중복:
    (1) po_number 있으면 material_no+po_number 일치
    (2) material_no+receive_date+supplier 일치 (2026-09-16 사용자 요청 추가 —
        발주번호가 달라도 같은 업체·같은 날짜에 같은 자재번호면 중복으로 본다.
        예전엔 po_number가 있는 행에서 (1)만 보고 (2)는 아예 안 봤음).
    매칭된 기존 레코드의 id/status/quantity/po_number를 'existing' 키에 담아
    같이 돌려준다(2026-09-16 합산 기능 추가) — 합산할 때 어느 기존 행에 합칠지,
    이미 검사완료돼서 합산이 불가능한지(그럴 땐 별도 등록으로 폴백) 판단하는 데 쓴다."""
    if not rows:
        return []
    conn = get_conn()
    dups = []
    for r in rows:
        mn  = r["material_no"]
        po  = (r.get("po_number") or "").strip()
        rd  = (r.get("receive_date") or "").strip()
        sup = (r.get("supplier") or "").strip()
        hit = None
        if po:
            hit = conn.execute(
                "SELECT id, status, quantity, po_number FROM intake_list WHERE material_no=? AND po_number=?",
                (mn, po)).fetchone()
        if not hit:
            hit = conn.execute(
                "SELECT id, status, quantity, po_number FROM intake_list WHERE material_no=? AND receive_date=? AND supplier=?",
                (mn, rd, sup)).fetchone()
        if hit:
            d = dict(r)
            d["existing"] = {"id": hit["id"], "status": hit["status"],
                              "quantity": hit["quantity"], "po_number": hit["po_number"] or ""}
            dups.append(d)
    conn.close()
    return dups


def merge_intake_duplicate(existing_id, add_quantity, add_po_number):
    """중복 입고 합산 처리(2026-09-16 사용자 요청) — 기존 intake_list 행의 수량에
    더하고, 발주번호를 '/'로 이어붙인다(이미 포함된 발주번호면 또 안 붙임 — 사용자
    확정). 숫자로 못 바꾸는 수량이 섞여 있으면 값 유실 방지를 위해 원문을 이어붙인다
    (_merge_duplicate_intake_rows()의 같은 원칙 재사용, 8-1절)."""
    conn = get_conn()
    row = conn.execute("SELECT quantity, po_number FROM intake_list WHERE id=?", (existing_id,)).fetchone()
    if row is None:
        conn.close()
        return
    try:
        existing_qty = int(row["quantity"]) if row["quantity"] not in (None, "") else 0
    except (TypeError, ValueError):
        existing_qty = None
    try:
        new_qty = int(add_quantity) if add_quantity not in (None, "") else 0
    except (TypeError, ValueError):
        new_qty = None
    if existing_qty is None or new_qty is None:
        parts = [str(p) for p in (row["quantity"], add_quantity) if p not in (None, "")]
        merged_qty = ", ".join(parts) if parts else None
    else:
        merged_qty = existing_qty + new_qty

    po_parts = [p for p in (row["po_number"] or "").split("/") if p]
    new_po = (add_po_number or "").strip()
    if new_po and new_po not in po_parts:
        po_parts.append(new_po)
    merged_po = "/".join(po_parts)

    conn.execute("UPDATE intake_list SET quantity=?, po_number=? WHERE id=?",
                (merged_qty, merged_po, existing_id))
    conn.commit()
    conn.close()


def add_intake_bulk(rows):
    """
    rows: list of dict (material_no, quantity, supplier, receive_date, po_number, product_name, assembly_no, is_priority)
    붙여넣기로 여러 건을 한 번에 등록.
    반환: 새로 생긴 intake_list.id 목록(rows와 같은 순서) — 호출부가 전수검사 자재
    자동시작 같은 후처리를 하려면 이 id로 db.get_intake()를 다시 불러야 한다.
    """
    conn = get_conn()
    cur = conn.cursor()
    new_ids = []
    for r in rows:
        cur.execute("""
            INSERT INTO intake_list (material_no, quantity, supplier, receive_date, po_number, product_name, assembly_no, is_priority)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """, (r["material_no"], r.get("quantity"), r.get("supplier"),
              r.get("receive_date"), r.get("po_number"), r.get("product_name"), r.get("assembly_no"),
              1 if r.get("is_priority") else 0))
        new_ids.append(cur.lastrowid)
    conn.commit()
    conn.close()
    return new_ids


def set_intake_status(intake_id, status):
    conn = get_conn()
    conn.execute("UPDATE intake_list SET status = ? WHERE id = ?", (status, intake_id))
    conn.commit()
    conn.close()


def set_intake_priority(intake_id, is_priority):
    conn = get_conn()
    conn.execute("UPDATE intake_list SET is_priority = ? WHERE id = ?",
                (1 if is_priority else 0, intake_id))
    conn.commit()
    conn.close()


def list_intake(status=None):
    conn = get_conn()
    if status:
        rows = conn.execute(
            "SELECT * FROM intake_list WHERE status = ? ORDER BY id DESC", (status,)
        ).fetchall()
    else:
        rows = conn.execute("SELECT * FROM intake_list ORDER BY id DESC").fetchall()
    conn.close()
    return rows


def search_intake(query, status=None, include=None, exclude=None):
    """자재번호/제품명/납품업체/발주번호 기준으로 입고 리스트 검색.
    include/exclude(포함/제외 다중 단어)도 같은 4개 필드를 대상으로 본다(_include_exclude_clause 재사용,
    8-1절 원칙, 검사대기목록 화면의 포함/제외 검색 — 2026-09-14)."""
    cols = ["material_no", "product_name", "supplier", "po_number"]
    where = ["1=1"]
    params = []
    if status:
        where.append("status = ?")
        params.append(status)
    if query:
        like = f"%{query}%"
        where.append("(material_no LIKE ? OR product_name LIKE ? OR supplier LIKE ? OR po_number LIKE ?)")
        params += [like, like, like, like]
    ie_sql, ie_params = _include_exclude_clause(cols, include, exclude)
    if ie_sql:
        where.append(ie_sql)
        params += ie_params

    conn = get_conn()
    rows = conn.execute(f"""
        SELECT * FROM intake_list
        WHERE {' AND '.join(where)}
        ORDER BY id DESC
    """, params).fetchall()
    conn.close()
    return rows


def get_intake(intake_id):
    conn = get_conn()
    row = conn.execute("SELECT * FROM intake_list WHERE id = ?", (intake_id,)).fetchone()
    conn.close()
    return row


def delete_intake_bulk(intake_ids):
    """입고 리스트(검사 대기 목록)에서 선택한 건들을 삭제. 이미 검사완료(상태='검사완료')된 건은
    성적서와 연결돼있을 수 있으니 여기서는 지우지 않고 건너뜀 — 대기 중인 건만 삭제."""
    if not intake_ids:
        return 0
    conn = get_conn()
    placeholders = ",".join("?" for _ in intake_ids)
    cur = conn.execute(
        f"DELETE FROM intake_list WHERE id IN ({placeholders}) AND status = '대기'",
        intake_ids
    )
    deleted = cur.rowcount
    conn.commit()
    conn.close()
    return deleted


# ---------- 과거 입고 이력 (검사/승인과 무관한 순수 조회용 — intake_list와 별개 테이블) ----------

def add_intake_history_bulk(rows, source, imported_by):
    """rows: list of dict(material_no, product_name, supplier, receive_date, po_number, quantity)
    검사/승인 큐에 절대 올라가지 않는다 — 그냥 intake_history 테이블에 그대로 적재만 한다."""
    conn = get_conn()
    cur = conn.cursor()
    for r in rows:
        cur.execute("""
            INSERT INTO intake_history
                (material_no, product_name, supplier, receive_date, po_number, quantity, source, imported_by)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """, (r["material_no"], r.get("product_name"), r.get("supplier"),
              r.get("receive_date"), r.get("po_number"), r.get("quantity"),
              source, imported_by))
    conn.commit()
    conn.close()


def find_duplicate_intake_history(rows):
    """material_no+supplier+receive_date+po_number+quantity가 전부 같은 행이 이미
    intake_history에 있으면 그 rows만 골라 돌려준다. 같은 엑셀 파일을 실수로 두 번
    올렸을 때 감지하기 위함. 값이 같은 게 실제로 반복 입고일 수도 있어서 자동으로
    막지는 않고(등록 여부는 호출부가 결정), 감지만 한다."""
    if not rows:
        return []
    conn = get_conn()
    dups = []
    for r in rows:
        hit = conn.execute("""
            SELECT id FROM intake_history
            WHERE material_no = ?
              AND IFNULL(supplier, '') = IFNULL(?, '')
              AND IFNULL(receive_date, '') = IFNULL(?, '')
              AND IFNULL(po_number, '') = IFNULL(?, '')
              AND IFNULL(quantity, -1) = IFNULL(?, -1)
        """, (r["material_no"], r.get("supplier"), r.get("receive_date"),
              r.get("po_number"), r.get("quantity"))).fetchone()
        if hit:
            dups.append(r)
    conn.close()
    return dups


def list_intake_history():
    """과거 입고 이력 전체 조회 — 검색/필터는 app.py의 공용 헬퍼
    (_list_search_params/_row_passes_search)가 화면 쪽에서 Python으로 거른다.
    검사이력(history.html) 등 4개 화면과 같은 방식으로 통일했다(2026-09-08 —
    예전엔 이 테이블만 SQL WHERE로 직접 검색해서 화면마다 검색 UI/방식이 달랐음)."""
    conn = get_conn()
    rows = conn.execute("SELECT * FROM intake_history ORDER BY receive_date DESC, id DESC").fetchall()
    conn.close()
    return rows


def delete_intake_history_bulk(ids):
    """목록 화면에서 선택 삭제. 이 테이블엔 상태 개념이 없으니 그냥 지운다."""
    if not ids:
        return 0
    conn = get_conn()
    placeholders = ",".join("?" for _ in ids)
    cur = conn.execute(f"DELETE FROM intake_history WHERE id IN ({placeholders})", ids)
    deleted = cur.rowcount
    conn.commit()
    conn.close()
    return deleted


def delete_intake_history_by_source(source):
    """대량 등록 직후 "방금 등록한 것 전체 되돌리기" 용 — source 문자열은 그 등록
    액션(파일명+시각)마다 고유하게 만들어지므로 이걸로 배치 단위 삭제가 된다."""
    if not source:
        return 0
    conn = get_conn()
    cur = conn.execute("DELETE FROM intake_history WHERE source = ?", (source,))
    deleted = cur.rowcount
    conn.commit()
    conn.close()
    return deleted


# ---------- 검사(성적서) 생성/조회 ----------

def create_inspection(header, items_with_results, overall_result, intake_id=None,
                       est_time_label=None, actual_time_sec=None, created_by_user_id=None,
                       total_time_sec=None):
    """
    header: dict (material_no, material_name, supplier, po_number,
                   receive_date, inspect_date, inspector, quantity)
    items_with_results: list of dict (item_name, measured_value, max_value, min_value, result)

    intake_id가 있으면 BEGIN IMMEDIATE로 잠근 뒤 활성 성적서가 없는지 다시 한 번
    확인하고 INSERT한다 — 호출부(라우트)가 미리 active_inspection_for_intake()로
    걸러내지만, 그 확인과 이 INSERT 사이에 동시 요청이 끼어들 수 있는 경합
    (TOCTOU)이 이론적으로 있었다(2026-09-10 감사에서 발견). 이미 있으면
    ValueError를 낸다 — 호출부가 이를 "이미 등록됨"으로 처리해야 한다.
    """
    conn = get_conn()
    if intake_id:
        conn.execute("BEGIN IMMEDIATE")
        existing = conn.execute(
            "SELECT id FROM inspections WHERE intake_id = ? AND status != 'superseded'",
            (intake_id,),
        ).fetchone()
        if existing is not None:
            conn.rollback()
            conn.close()
            raise ValueError(f"intake_id={intake_id}에 이미 활성 성적서(#{existing['id']})가 있어.")
    cur = conn.cursor()
    cur.execute("""
        INSERT INTO inspections (intake_id, material_no, material_name, supplier, po_number,
                                  receive_date, inspect_date, inspector, quantity,
                                  overall_result, status, est_time_label, actual_time_sec,
                                  total_time_sec, created_by_user_id)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'pending', ?, ?, ?, ?)
    """, (intake_id, header["material_no"], header.get("material_name"), header.get("supplier"),
          header.get("po_number"), header.get("receive_date"), header.get("inspect_date"),
          header.get("inspector"), header.get("quantity"), overall_result,
          est_time_label, actual_time_sec, total_time_sec, created_by_user_id))
    inspection_id = cur.lastrowid

    if intake_id:
        cur.execute("UPDATE intake_list SET status = '검사완료' WHERE id = ?", (intake_id,))

    for it in items_with_results:
        cur.execute("""
            INSERT INTO inspection_items (inspection_id, item_name, measured_value,
                                           max_value, min_value, result, gauge_expiry, gauge_name, part_material_no,
                                           aql_snapshot, sample_qty_snapshot, ac_allowance_snapshot, aql_snapshot_set)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (inspection_id, it["item_name"], it.get("measured_value"),
              it.get("max_value"), it.get("min_value"), it.get("result"),
              it.get("gauge_expiry"), it.get("gauge_name"), it.get("part_material_no") or header["material_no"],
              it.get("aql_snapshot"), it.get("sample_qty_snapshot"),
              it.get("ac_allowance_snapshot"), 1 if "aql_snapshot" in it else 0))

    conn.commit()
    conn.close()
    return inspection_id


def rename_inspection_item(item_id, new_item_name):
    """inspection_items 한 행의 item_name만 고친다.

    자재 규격에 항목기호(item_name)가 중복 등록됐다가 하나를 다른 기호로 고친 뒤,
    이미 저장된 성적서 쪽의 item_name도 맞춰줘야 할 때 쓰는 정정용 함수
    (2026-08-30, 실제 데이터 사고 수습 — 정상 플로우에서는 쓰이지 않음)."""
    conn = get_conn()
    conn.execute("UPDATE inspection_items SET item_name=? WHERE id=?", (new_item_name, item_id))
    conn.commit()
    conn.close()


def set_item_override(inspection_id, part_material_no, item_name, override_result, reason,
                       user_id, user_name):
    """검사 항목의 자동판정 결과를 사람이 사후에 합격/불합격으로 덮어쓴다.
    inspection_items.result 원본은 건드리지 않는다."""
    conn = get_conn()
    existing = conn.execute(
        "SELECT id FROM inspection_item_overrides WHERE inspection_id=? AND part_material_no=? AND item_name=?",
        (inspection_id, part_material_no, item_name)).fetchone()
    if existing:
        conn.execute("""
            UPDATE inspection_item_overrides
               SET override_result=?, reason=?, created_by_user_id=?, created_by_name=?,
                   created_at=datetime('now','localtime')
             WHERE id=?
        """, (override_result, reason, user_id, user_name, existing["id"]))
    else:
        conn.execute("""
            INSERT INTO inspection_item_overrides
                (inspection_id, part_material_no, item_name, override_result, reason,
                 created_by_user_id, created_by_name)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (inspection_id, part_material_no, item_name, override_result, reason, user_id, user_name))
    conn.commit()
    conn.close()


def clear_item_override(inspection_id, part_material_no, item_name):
    """오버라이드를 지워서 자동판정 결과로 되돌린다."""
    conn = get_conn()
    conn.execute(
        "DELETE FROM inspection_item_overrides WHERE inspection_id=? AND part_material_no=? AND item_name=?",
        (inspection_id, part_material_no, item_name))
    conn.commit()
    conn.close()


def update_inspection_overall_result(inspection_id, overall_result):
    """항목 오버라이드 반영 후 성적서 전체 판정(표시용)을 갱신한다.
    status/approval_type/서명/content_hash/pdf_hash는 절대 건드리지 않는다 —
    최종 결정 자체를 바꾸려면 여전히 승인회수 후 재결정이 필요하다."""
    conn = get_conn()
    conn.execute("UPDATE inspections SET overall_result=? WHERE id=?", (overall_result, inspection_id))
    conn.commit()
    conn.close()


def sync_material_names_from_master():
    """specs.material_name(등록 당시 복사본)이 materials.material_name(정본)과 어긋난
    자재들을 바로잡고, 그 여파로 material_name이 빈 채 저장된 성적서(inspections)도
    같이 채운다. 원인: 성적서 생성 코드가 예전엔 정본(materials) 대신 이 복사본을
    읽어서, 복사본이 비어있거나 오래된 자재는 검사이력·승인 화면에 자재명이
    안 보였다(2026-09-01 실사용자 리포트로 발견). 코드는 이미 정본을 보게 고쳤고,
    이건 기존에 어긋나 있던 데이터 자체를 정리하는 일회성 정리 함수.
    반환: {"specs_fixed": N, "inspections_fixed": N}"""
    conn = get_conn()
    materials_rows = conn.execute(
        "SELECT material_no, material_name FROM materials WHERE material_name IS NOT NULL AND material_name != ''"
    ).fetchall()
    specs_fixed = 0
    insp_fixed = 0
    for m in materials_rows:
        cur = conn.execute(
            "UPDATE specs SET material_name = ? WHERE material_no = ? AND (material_name IS NULL OR material_name != ?)",
            (m["material_name"], m["material_no"], m["material_name"]),
        )
        specs_fixed += cur.rowcount
        cur = conn.execute(
            "UPDATE inspections SET material_name = ? WHERE material_no = ? AND (material_name IS NULL OR material_name = '')",
            (m["material_name"], m["material_no"]),
        )
        insp_fixed += cur.rowcount
    conn.commit()
    conn.close()
    return {"specs_fixed": specs_fixed, "inspections_fixed": insp_fixed}


def update_inspection_items(inspection_id, inspect_date, inspector, items_with_results, overall_result,
                             est_time_label=None, actual_time_sec=None, total_time_sec=None):
    """pending 상태 성적서 측정값·판정 전체 갱신"""
    conn = get_conn()
    header_row = conn.execute("SELECT material_no FROM inspections WHERE id=?", (inspection_id,)).fetchone()
    default_material_no = header_row["material_no"] if header_row else None
    if est_time_label is not None or actual_time_sec is not None:
        conn.execute("""
            UPDATE inspections SET inspect_date=?, inspector=?, overall_result=?,
                                    est_time_label=?, actual_time_sec=?, total_time_sec=? WHERE id=?
        """, (inspect_date, inspector, overall_result, est_time_label, actual_time_sec, total_time_sec, inspection_id))
    else:
        conn.execute("""
            UPDATE inspections SET inspect_date=?, inspector=?, overall_result=? WHERE id=?
        """, (inspect_date, inspector, overall_result, inspection_id))
    conn.execute("DELETE FROM inspection_items WHERE inspection_id=?", (inspection_id,))
    for it in items_with_results:
        conn.execute("""
            INSERT INTO inspection_items (inspection_id, item_name, measured_value, max_value, min_value, result,
                                          gauge_expiry, gauge_name, part_material_no,
                                          aql_snapshot, sample_qty_snapshot, ac_allowance_snapshot, aql_snapshot_set)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
        """, (inspection_id, it["item_name"], it.get("measured_value"),
              it.get("max_value"), it.get("min_value"), it.get("result"),
              it.get("gauge_expiry"), it.get("gauge_name"), it.get("part_material_no") or default_material_no,
              it.get("aql_snapshot"), it.get("sample_qty_snapshot"),
              it.get("ac_allowance_snapshot"), 1 if "aql_snapshot" in it else 0))
    conn.commit()
    conn.close()


def get_inspection(inspection_id):
    conn = get_conn()
    header = conn.execute(
        "SELECT * FROM inspections WHERE id = ?", (inspection_id,)
    ).fetchone()
    # specs 테이블에서 spec_display·AQL·검사방법 조인 — part_material_no(조립품이면 부품별 실제 자재,
    # 일반 검사면 헤더 자재와 동일)와 항목명이 같으면 매핑
    items = conn.execute("""
        SELECT ii.*,
               COALESCE(s.spec_display, ii.item_name) AS spec_display,
               CASE WHEN ii.aql_snapshot_set = 1 THEN ii.aql_snapshot ELSE s.aql END AS aql,
               CASE WHEN ii.aql_snapshot_set = 1 THEN ii.sample_qty_snapshot ELSE NULL END AS sample_qty,
               CASE WHEN ii.aql_snapshot_set = 1 THEN ii.ac_allowance_snapshot ELSE NULL END AS ac_allowance,
               s.inspect_method AS inspect_method,
               s.lower_limit AS lower_limit,
               s.upper_limit AS upper_limit,
               s.judge_type AS judge_type,
               s.material_name AS part_material_name,
               io.override_result AS override_result,
               io.reason          AS override_reason,
               io.created_by_name AS override_by,
               io.created_at      AS override_at
        FROM inspection_items ii
        LEFT JOIN specs s
               ON s.material_no = COALESCE(ii.part_material_no,
                                            (SELECT material_no FROM inspections WHERE id = ?))
              AND s.item_name   = ii.item_name
        LEFT JOIN inspection_item_overrides io
               ON io.inspection_id = ii.inspection_id
              AND io.part_material_no = COALESCE(ii.part_material_no,
                                            (SELECT material_no FROM inspections WHERE id = ?))
              AND io.item_name = ii.item_name
        WHERE ii.inspection_id = ?
        ORDER BY ii.id
    """, (inspection_id, inspection_id, inspection_id)).fetchall()
    conn.close()
    return header, items


def list_unset_inspection_aql_snapshots():
    """신규 스냅샷 컬럼 도입 전 검사 항목을 이관하기 위한 읽기 전용 목록.

    현재 규격의 AQL은 이관 시점에 한 번만 참고하고, 이후에는 set 플래그가 1인 행을
    다시 건드리지 않는다. NULL AQL도 그대로 고정할 수 있도록 spec_aql 별칭을 쓴다.
    """
    conn = get_conn()
    rows = conn.execute("""
        SELECT ii.id, i.quantity,
               (SELECT s.aql
                  FROM specs s
                 WHERE s.material_no = COALESCE(ii.part_material_no, i.material_no)
                   AND s.item_name = ii.item_name
                 ORDER BY s.id
                 LIMIT 1) AS spec_aql
          FROM inspection_items ii
          JOIN inspections i ON i.id = ii.inspection_id
         WHERE COALESCE(ii.aql_snapshot_set, 0) = 0
         ORDER BY ii.id
    """).fetchall()
    conn.close()
    return rows


def set_inspection_aql_snapshots(rows):
    """[(item_id, aql, sample_qty, ac), ...]를 한 트랜잭션으로 최초 1회 고정."""
    if not rows:
        return 0
    conn = get_conn()
    cur = conn.cursor()
    changed = 0
    try:
        cur.execute("BEGIN IMMEDIATE")
        for item_id, aql, sample_qty, ac_allowance in rows:
            result = cur.execute("""
                UPDATE inspection_items
                   SET aql_snapshot=?, sample_qty_snapshot=?, ac_allowance_snapshot=?,
                       aql_snapshot_set=1
                 WHERE id=? AND COALESCE(aql_snapshot_set, 0)=0
            """, (aql, sample_qty, ac_allowance, item_id))
            changed += result.rowcount
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return changed


def list_inspections(status=None):
    conn = get_conn()
    if status:
        rows = conn.execute(
            "SELECT * FROM inspections WHERE status = ? ORDER BY id DESC", (status,)
        ).fetchall()
    else:
        rows = conn.execute(
            "SELECT * FROM inspections ORDER BY id DESC"
        ).fetchall()
    conn.close()
    return rows


def list_output_history(q="", date_from="", date_to=""):
    """이미 PDF/xlsx가 생성된 성적서 목록 (출력 이력).
    q로 자재번호·자재명·업체·승인자 검색, date_from/date_to로 승인일 범위 필터."""
    conn = get_conn()
    sql = "SELECT * FROM inspections WHERE pdf_path IS NOT NULL"
    params = []
    if q:
        sql += " AND (material_no LIKE ? OR material_name LIKE ? OR supplier LIKE ? OR approver LIKE ?)"
        like = f"%{q}%"
        params += [like, like, like, like]
    if date_from:
        sql += " AND DATE(approved_at) >= ?"
        params.append(date_from)
    if date_to:
        sql += " AND DATE(approved_at) <= ?"
        params.append(date_to)
    sql += " ORDER BY approved_at DESC, id DESC"
    rows = conn.execute(sql, params).fetchall()
    conn.close()
    return rows


def list_pending_output_inspections():
    """승인은 됐지만 아직 PDF/xlsx 출력을 안 한 성적서들."""
    conn = get_conn()
    rows = conn.execute("""
        SELECT * FROM inspections
        WHERE status = 'approved' AND pdf_path IS NULL
        ORDER BY id DESC
    """).fetchall()
    conn.close()
    return rows


def get_today_stats():
    """금일 현황 — 입고/검사/불량/소요시간(인원별)."""
    from datetime import date
    today = date.today().isoformat()
    conn = get_conn()

    # 오늘 입고 건수/품목수
    intake_rows = conn.execute(
        "SELECT COUNT(*) as cnt, COUNT(DISTINCT material_no) as items FROM intake_list WHERE date(created_at) = ?",
        (today,)
    ).fetchone()

    # 오늘 검사 완료 건수 — 제출(created_at)이 아니라 실제 검사한 날짜(inspect_date) 기준.
    # 검사를 오늘 했어도 제출은 나중에(또는 그 반대) 할 수 있어서 created_at으로 집계하면 어긋난다.
    inspected = conn.execute(
        "SELECT COUNT(*) FROM inspections WHERE date(inspect_date) = ?", (today,)
    ).fetchone()[0]

    # 오늘 불량/검토필요 건수
    defects = conn.execute(
        "SELECT COUNT(*) FROM inspections WHERE date(inspect_date) = ? AND overall_result NOT IN ('합격', '')",
        (today,)
    ).fetchone()[0]

    # 인원별 검사 소요시간 (오늘) — total_time_sec(성적서 총 측정시간) 우선, 없으면 actual_time_sec fallback
    time_rows = conn.execute("""
        SELECT inspector,
               SUM(COALESCE(total_time_sec, actual_time_sec, 0)) as total_sec
        FROM inspections
        WHERE date(inspect_date) = ?
          AND (total_time_sec > 0 OR actual_time_sec > 0)
        GROUP BY inspector
    """, (today,)).fetchall()

    conn.close()

    return {
        "intake_count": intake_rows["cnt"] if intake_rows else 0,
        "intake_items": intake_rows["items"] if intake_rows else 0,
        "inspected": inspected,
        "defects": defects,
        "time_by_user": [{"inspector": r["inspector"], "total_sec": r["total_sec"]} for r in time_rows],
    }


def get_defect_history(start_date=None, end_date=None):
    """기간 내 불합격/검토필요 성적서 목록 + 업체별/자재별 통계."""
    conn = get_conn()

    # 기간 조건은 두 쿼리가 똑같이 쓰므로 한 번만 만든다
    period_sql = ""
    period_params = []
    if start_date:
        period_sql += " AND i.inspect_date >= ?"
        period_params.append(start_date)
    if end_date:
        period_sql += " AND i.inspect_date <= ?"
        period_params.append(end_date)

    # ncr_count — 이 성적서로 부적합 통보서가 이미 작성됐는지 (불량 이력 화면의 후속조치 안내용)
    rows = conn.execute(
        f"""SELECT i.*,
                   (SELECT COUNT(*) FROM ncr n WHERE n.inspection_id = i.id) AS ncr_count
              FROM inspections i
             WHERE i.overall_result IS NOT NULL
               AND i.overall_result NOT IN ('합격', '')
                   {period_sql}
             ORDER BY i.inspect_date DESC, i.id DESC""",
        period_params
    ).fetchall()

    # 전체 검사 건수 (같은 기간) — 불량률 계산용
    total_rows = conn.execute(
        f"""SELECT i.supplier, COUNT(*) AS cnt
              FROM inspections i
             WHERE 1=1 {period_sql}
             GROUP BY i.supplier""",
        period_params
    ).fetchall()
    total_by_supplier = {r["supplier"]: r["cnt"] for r in total_rows}

    conn.close()

    # 업체별 집계
    supplier_stats = {}
    for r in rows:
        s = r["supplier"] or "(미입력)"
        if s not in supplier_stats:
            supplier_stats[s] = {"supplier": s, "defects": 0,
                                  "total": total_by_supplier.get(r["supplier"], 0)}
        supplier_stats[s]["defects"] += 1

    by_supplier = sorted(supplier_stats.values(), key=lambda x: x["defects"], reverse=True)
    for s in by_supplier:
        s["rate"] = round(s["defects"] / s["total"] * 100, 1) if s["total"] else 0

    return {"rows": rows, "by_supplier": by_supplier}


def get_defect_followup(completed_start=None, completed_end=None):
    """후속조치 추적 전용 — 각 단계별로 분류된 목록을 반환."""
    conn = get_conn()

    def q(sql, params=()):
        return [dict(r) for r in conn.execute(sql, params).fetchall()]

    # 1. 재검사 대기 — 반려됨, 아직 재검사 안 됨
    #    재검사됐으면 원본이 superseded로 바뀌므로 rejected만 보면 됨
    recheck = q("""
        SELECT i.id, i.material_no, m.material_name, i.supplier, i.inspector,
               i.inspect_date, i.receive_date, i.reject_reason, i.status
          FROM inspections i
          LEFT JOIN materials m ON m.material_no = i.material_no
         WHERE i.status = 'rejected'
         ORDER BY i.inspect_date DESC, i.id DESC
    """)

    # 2. 통보서 작성 필요 — 불합격 항목 하나라도 있는 성적서(상태 무관) + NCR 없음 + 대체 아님
    #    B안: 승인 여부와 무관하게 fail 항목이 있으면 통보서 대상(pending 포함)
    #    ncr_waived=1(통보서 불필요 처리)인 건은 제외 — 대신 아래 6번 버킷으로 따로 뺀다.
    ncr_write = q("""
        SELECT DISTINCT i.id, i.material_no, m.material_name, i.supplier, i.inspector,
               i.inspect_date, i.receive_date, i.remark_approver, i.status, i.overall_result
          FROM inspections i
          LEFT JOIN materials m ON m.material_no = i.material_no
         INNER JOIN inspection_items ii ON ii.inspection_id = i.id
         WHERE ii.result NOT IN ('합격', '미측정', '', '규격미입력')
           AND ii.result IS NOT NULL
           AND NOT EXISTS (SELECT 1 FROM ncr n WHERE n.inspection_id = i.id)
           AND i.status != 'superseded'
           AND (i.ncr_waived IS NULL OR i.ncr_waived = 0)
         ORDER BY i.inspect_date DESC, i.id DESC
    """)

    # 6. 통보서 불필요 처리됨 — 위 2번과 조건은 같지만 ncr_waived=1인 건만. 작성 대기열에서
    #    빠졌다고 사라진 게 아니라 여기로 옮겨져서 계속 추적 가능하게 한다.
    ncr_waived = q("""
        SELECT DISTINCT i.id, i.material_no, m.material_name, i.supplier, i.inspector,
               i.inspect_date, i.receive_date, i.ncr_waived_reason, i.status, i.overall_result
          FROM inspections i
          LEFT JOIN materials m ON m.material_no = i.material_no
         INNER JOIN inspection_items ii ON ii.inspection_id = i.id
         WHERE ii.result NOT IN ('합격', '미측정', '', '규격미입력')
           AND ii.result IS NOT NULL
           AND NOT EXISTS (SELECT 1 FROM ncr n WHERE n.inspection_id = i.id)
           AND i.status != 'superseded'
           AND i.ncr_waived = 1
         ORDER BY i.inspect_date DESC, i.id DESC
    """)

    # 3. 통보서 확인 필요 — NCR draft 상태
    ncr_review = q("""
        SELECT n.id AS ncr_id, n.ncr_no, n.material_no, m.material_name, n.supplier,
               n.issued_by, n.created_at, n.due_date, n.defect_description,
               i.inspect_date, i.receive_date, i.inspector, i.id AS inspection_id
          FROM ncr n
          LEFT JOIN materials m ON m.material_no = n.material_no
          LEFT JOIN inspections i ON i.id = n.inspection_id
         WHERE n.status = 'draft'
         ORDER BY n.created_at DESC, n.id DESC
    """)

    # 4. 통보서 발송 필요 — NCR confirmed 상태
    ncr_send = q("""
        SELECT n.id AS ncr_id, n.ncr_no, n.material_no, m.material_name, n.supplier,
               n.issued_by, n.confirmed_by, n.confirmed_at, n.due_date, n.defect_description,
               i.inspect_date, i.receive_date, i.inspector, i.id AS inspection_id
          FROM ncr n
          LEFT JOIN materials m ON m.material_no = n.material_no
          LEFT JOIN inspections i ON i.id = n.inspection_id
         WHERE n.status = 'confirmed'
         ORDER BY n.confirmed_at DESC, n.id DESC
    """)

    # 5. 완료 — NCR sent (기간 필터로 archive 조회)
    period_sql = ""
    period_params = []
    if completed_start:
        period_sql += " AND date(n.email_sent_at) >= ?"
        period_params.append(completed_start)
    if completed_end:
        period_sql += " AND date(n.email_sent_at) <= ?"
        period_params.append(completed_end)

    completed = q(f"""
        SELECT n.id AS ncr_id, n.ncr_no, n.material_no, m.material_name, n.supplier,
               n.sent_to, n.email_sent_at, n.issued_by, n.confirmed_by, n.defect_description,
               i.inspect_date, i.receive_date, i.inspector, i.id AS inspection_id
          FROM ncr n
          LEFT JOIN materials m ON m.material_no = n.material_no
          LEFT JOIN inspections i ON i.id = n.inspection_id
         WHERE n.status = 'sent' {period_sql}
         ORDER BY n.email_sent_at DESC, n.id DESC
    """, period_params)

    conn.close()
    return {
        "recheck": recheck,
        "ncr_write": ncr_write,
        "ncr_review": ncr_review,
        "ncr_send": ncr_send,
        "completed": completed,
        "ncr_waived": ncr_waived,
    }


def get_material_inspection_history(material_no):
    """자재번호로 전체 검사 이력 + 항목별 측정값을 반환."""
    conn = get_conn()
    inspections = conn.execute("""
        SELECT id, inspect_date, supplier, inspector, overall_result, status, approval_type
        FROM inspections WHERE material_no = ?
        ORDER BY inspect_date ASC, id ASC
    """, (material_no,)).fetchall()

    items_by_inspection = {}
    for insp in inspections:
        rows = conn.execute("""
            SELECT item_name, measured_value, result, max_value, min_value
            FROM inspection_items WHERE inspection_id = ?
            ORDER BY id
        """, (insp["id"],)).fetchall()
        items_by_inspection[insp["id"]] = rows

    conn.close()
    return inspections, items_by_inspection


def get_repeat_defects(min_count=3):
    """업체+자재 조합별 불량 건수, min_count 이상인 것만 반환.
    '진행중'(2026-09-28 신설 — 하우징 STAGE 검사가 아직 안 끝난 상태)은 불량이
    아니라 '아직 측정 중'일 뿐이므로 제외한다(사용자가 대시보드에 실제로는 불량이
    없는데 반복불량 경고가 뜬다고 지적해서 발견 — 이 함수가 그 값이 생기기 전에
    짜여서 '합격/공백 아니면 전부 불량'으로 취급하고 있었다)."""
    conn = get_conn()
    rows = conn.execute("""
        SELECT supplier, material_no, material_name, COUNT(*) as cnt,
               MAX(inspect_date) as last_date
        FROM inspections
        WHERE overall_result NOT IN ('합격', '', '진행중') AND overall_result IS NOT NULL
        GROUP BY supplier, material_no
        HAVING COUNT(*) >= ?
        ORDER BY cnt DESC, last_date DESC
    """, (min_count,)).fetchall()
    conn.close()
    return rows


def get_defect_count_for(supplier, material_no):
    """특정 업체+자재의 전체 불량 건수 반환. get_repeat_defects()와 같은 기준
    ('진행중' 제외, 8-1절 — 새로 정의하지 않음)."""
    if not supplier or not material_no:
        return 0
    conn = get_conn()
    count = conn.execute("""
        SELECT COUNT(*) FROM inspections
        WHERE supplier = ? AND material_no = ?
          AND overall_result NOT IN ('합격', '', '진행중') AND overall_result IS NOT NULL
    """, (supplier, material_no)).fetchone()[0]
    conn.close()
    return count


# ---------- 계측기 마스터 ----------

# 측정 방식(specs.inspect_method)에 적히지만 계측기가 아닌 값들.
# (육안/외관=시각검사, 전수=검사 범위) — 계측기 종류 자동등록·자동매칭 대상에서 뺀다.
NON_GAUGE_METHODS = {"육안", "외관", "전수"}

def distinct_inspect_methods():
    """규격에 실제로 쓰인 측정 방식들을 (이름, 사용건수)로 돌려준다. 빈칸 제외.
    계측기 종류 자동등록·드롭다운 옵션의 원본이 된다."""
    conn = get_conn()
    rows = conn.execute("""
        SELECT TRIM(inspect_method) AS m, COUNT(*) AS c
          FROM specs
         WHERE inspect_method IS NOT NULL AND TRIM(inspect_method) <> ''
      GROUP BY TRIM(inspect_method)
      ORDER BY c DESC
    """).fetchall()
    conn.close()
    return [(r["m"], r["c"]) for r in rows]

def list_gauges():
    conn = get_conn()
    rows = conn.execute(
        "SELECT * FROM gauge_master"
    ).fetchall()
    conn.close()
    return rows

def search_gauges(query=None):
    """계측기 검색 — 관리번호/이름/모델/위치를 단일 자유텍스트로 훑는다
    (suppliers.html과 동일한 방식, 필드가 적어 드롭다운형 검색은 과함)."""
    conn = get_conn()
    if not query:
        rows = conn.execute("SELECT * FROM gauge_master").fetchall()
        conn.close()
        return rows
    like = f"%{query}%"
    rows = conn.execute("""
        SELECT * FROM gauge_master
         WHERE gauge_no LIKE ? OR name LIKE ? OR model LIKE ? OR location LIKE ?
    """, (like, like, like, like)).fetchall()
    conn.close()
    return rows

def get_gauge(gauge_id):
    conn = get_conn()
    row = conn.execute("SELECT * FROM gauge_master WHERE id = ?", (gauge_id,)).fetchone()
    conn.close()
    return row

def upsert_gauge(gauge_id, gauge_no, name, model, location, last_calibrated, expiry_date, notes):
    conn = get_conn()
    if gauge_id:
        conn.execute("""
            UPDATE gauge_master SET gauge_no=?, name=?, model=?, location=?,
              last_calibrated=?, expiry_date=?, notes=? WHERE id=?
        """, (gauge_no, name, model, location, last_calibrated, expiry_date, notes, gauge_id))
    else:
        conn.execute("""
            INSERT INTO gauge_master (gauge_no, name, model, location, last_calibrated, expiry_date, notes)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (gauge_no, name, model, location, last_calibrated, expiry_date, notes))
    conn.commit()
    conn.close()

def delete_gauge(gauge_id):
    conn = get_conn()
    conn.execute("DELETE FROM gauge_master WHERE id = ?", (gauge_id,))
    conn.commit()
    conn.close()

def get_gauge_master_warnings(days=30):
    """만료일이 days일 이내(이미 만료 포함)인 계측기 목록."""
    from datetime import date, timedelta
    today = date.today().isoformat()
    limit = (date.today() + timedelta(days=days)).isoformat()
    conn = get_conn()
    rows = conn.execute("""
        SELECT *, julianday(expiry_date) - julianday('now') as days_left
        FROM gauge_master
        WHERE expiry_date IS NOT NULL AND expiry_date != ''
          AND expiry_date <= ?
        ORDER BY expiry_date ASC
    """, (limit,)).fetchall()
    conn.close()
    return rows


def get_gauge_expiry_warnings(days=15):
    """검교정 유효기간이 days일 이내인 항목 목록."""
    from datetime import date, timedelta
    today = date.today().isoformat()
    limit = (date.today() + timedelta(days=days)).isoformat()
    conn = get_conn()
    rows = conn.execute("""
        SELECT ii.gauge_expiry, ii.item_name, i.material_no, i.supplier
        FROM inspection_items ii
        JOIN inspections i ON ii.inspection_id = i.id
        WHERE ii.gauge_expiry IS NOT NULL
          AND ii.gauge_expiry != ''
          AND ii.gauge_expiry <= ?
          AND ii.gauge_expiry >= ?
        ORDER BY ii.gauge_expiry ASC
    """, (limit, today)).fetchall()
    conn.close()
    return rows


# ---------- 업체 정보 ----------

def list_suppliers():
    conn = get_conn()
    rows = conn.execute("SELECT * FROM suppliers ORDER BY name").fetchall()
    conn.close()
    return rows

def get_supplier(name):
    conn = get_conn()
    row = conn.execute("SELECT * FROM suppliers WHERE name = ?", (name,)).fetchone()
    conn.close()
    return row

def upsert_supplier(name, email, contact, notes, address="", biz_no="", contact_name="", contact2="", items=""):
    conn = get_conn()
    conn.execute("""
        INSERT INTO suppliers (name, email, contact, notes, address, biz_no, contact_name, contact2, items)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(name) DO UPDATE SET email=excluded.email,
            contact=excluded.contact, notes=excluded.notes,
            address=excluded.address, biz_no=excluded.biz_no,
            contact_name=excluded.contact_name, contact2=excluded.contact2,
            items=excluded.items
    """, (name, email, contact, notes, address, biz_no, contact_name, contact2, items))
    conn.commit()
    conn.close()

def delete_supplier(name):
    conn = get_conn()
    conn.execute("DELETE FROM supplier_contacts WHERE supplier_name = ?", (name,))
    conn.execute("DELETE FROM suppliers WHERE name = ?", (name,))
    conn.commit()
    conn.close()


def search_suppliers(query=None):
    """업체 검색 — 단일 자유텍스트로 업체명·주소·취급품목·(레거시)연락처 필드·
    담당자(역할/이름/전화/이메일)까지 전부 훑는다. 필드 지정 없이 아무 정보로나
    찾는 게 비개발자 사용자에게 더 직관적이라 spec.html의 by= 드롭다운과는 다르게
    필드 선택 없이 만들었다."""
    conn = get_conn()
    if not query:
        rows = conn.execute("SELECT * FROM suppliers ORDER BY name").fetchall()
        conn.close()
        return rows
    like = f"%{query}%"
    rows = conn.execute("""
        SELECT DISTINCT s.* FROM suppliers s
        LEFT JOIN supplier_contacts c ON c.supplier_name = s.name
        WHERE s.name LIKE ? OR s.items LIKE ? OR s.address LIKE ?
           OR s.contact_name LIKE ? OR s.contact LIKE ? OR s.contact2 LIKE ? OR s.email LIKE ?
           OR c.role LIKE ? OR c.contact_name LIKE ? OR c.phone LIKE ? OR c.email LIKE ?
        ORDER BY s.name
    """, (like,) * 11).fetchall()
    conn.close()
    return rows


# ---------- 업체 담당자 (역할별 다중 등록) ----------

def list_supplier_contacts(supplier_name):
    conn = get_conn()
    rows = conn.execute(
        "SELECT * FROM supplier_contacts WHERE supplier_name = ? ORDER BY sort_order, id",
        (supplier_name,)
    ).fetchall()
    conn.close()
    return rows


def add_supplier_contact(supplier_name, role, contact_name, phone, email, notes):
    conn = get_conn()
    max_order = conn.execute(
        "SELECT MAX(sort_order) FROM supplier_contacts WHERE supplier_name = ?", (supplier_name,)
    ).fetchone()[0]
    sort_order = (max_order or 0) + 1
    cur = conn.execute("""
        INSERT INTO supplier_contacts (supplier_name, role, contact_name, phone, email, notes, sort_order)
        VALUES (?, ?, ?, ?, ?, ?, ?)
    """, (supplier_name, role or "기타", contact_name, phone, email, notes, sort_order))
    conn.commit()
    contact_id = cur.lastrowid
    conn.close()
    return contact_id


def update_supplier_contact(contact_id, role, contact_name, phone, email, notes):
    conn = get_conn()
    conn.execute("""
        UPDATE supplier_contacts SET role=?, contact_name=?, phone=?, email=?, notes=?
        WHERE id=?
    """, (role or "기타", contact_name, phone, email, notes, contact_id))
    conn.commit()
    conn.close()


def delete_supplier_contact(contact_id):
    conn = get_conn()
    conn.execute("DELETE FROM supplier_contacts WHERE id = ?", (contact_id,))
    conn.commit()
    conn.close()


_CONTACT_ROLE_PRIORITY = {
    "ncr": ["품질", "영업", "구매"],
    "report": ["영업", "품질", "구매"],
    "improvement": ["품질", "영업", "구매"],
}


def get_default_contact(supplier_name, purpose):
    """문서 종류(purpose: 'ncr' 또는 'report')에 맞는 기본 발송 대상 담당자를 고른다.

    role 우선순위에 맞는 담당자 중 sort_order가 가장 빠른 사람 → 없으면 등록된
    담당자 중 첫 번째 → 그마저 없으면 suppliers의 옛 단일 필드(email/contact_name)로
    폴백한다. 반환 형식은 기존 supplier_info["email"] 사용부와 호환되는 dict.
    """
    contacts = list_supplier_contacts(supplier_name)
    if contacts:
        priority = _CONTACT_ROLE_PRIORITY.get(purpose, [])
        for role in priority:
            for c in contacts:
                if c["role"] == role and c["email"]:
                    return {"email": c["email"], "contact_name": c["contact_name"], "role": c["role"]}
        for c in contacts:
            if c["email"]:
                return {"email": c["email"], "contact_name": c["contact_name"], "role": c["role"]}
    supplier = get_supplier(supplier_name)
    if supplier:
        return {"email": supplier["email"] or "", "contact_name": supplier["contact_name"] or "", "role": None}
    return {"email": "", "contact_name": "", "role": None}


# ---------- 부적합 통보서 (NCR) ----------

def _next_ncr_no():
    from datetime import date
    today = date.today().strftime("%Y%m%d")
    conn = get_conn()
    count = conn.execute(
        "SELECT COUNT(*) FROM ncr WHERE ncr_no LIKE ?", (f"NCR-{today}-%",)
    ).fetchone()[0]
    conn.close()
    return f"NCR-{today}-{count + 1:03d}"

def create_ncr(inspection_id, material_no, material_name, supplier, defect_description,
               action_required, due_date, issued_by, issued_date, lot_number=None, receive_date=None,
               cc_recipient=None, sample_qty=None, defect_qty=None, special_note=None, lot_qty=None,
               occurrence_type='입고검사', defect_type=None):
    ncr_no = _next_ncr_no()
    conn = get_conn()
    cur = conn.execute("""
        INSERT INTO ncr (ncr_no, inspection_id, material_no, material_name, supplier,
            defect_description, action_required, due_date, issued_by, issued_date, status,
            lot_number, receive_date, cc_recipient, sample_qty, defect_qty, special_note, lot_qty,
            occurrence_type, defect_type)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'draft', ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (ncr_no, inspection_id, material_no, material_name, supplier,
          defect_description, action_required, due_date, issued_by, issued_date,
          lot_number, receive_date, cc_recipient, sample_qty, defect_qty, special_note, lot_qty,
          occurrence_type, defect_type))
    conn.commit()
    ncr_id = cur.lastrowid
    conn.close()
    return ncr_id, ncr_no


def update_ncr(ncr_id, material_no, material_name, supplier, defect_description,
               due_date, issued_date, lot_number=None, receive_date=None,
               cc_recipient=None, sample_qty=None, defect_qty=None, special_note=None,
               lot_qty=None, occurrence_type='입고검사', defect_type=None):
    """부적합 통보서 내용 수정. 발송/확인 여부와 무관하게 항상 저장하되,
    내용이 바뀐 문서에 옛 확인/발송 흔적이 남지 않도록 status를 'draft'로
    되돌리고 confirmed_by/confirmed_at/confirm_signature/email_sent_at/sent_to를
    같이 지운다(8-2-10절의 '승인 회수 시 서명·해시 초기화'와 같은 원칙)."""
    conn = get_conn()
    conn.execute("""
        UPDATE ncr SET
            material_no = ?, material_name = ?, supplier = ?,
            defect_description = ?, due_date = ?, issued_date = ?,
            lot_number = ?, receive_date = ?, cc_recipient = ?,
            sample_qty = ?, defect_qty = ?, special_note = ?, lot_qty = ?,
            occurrence_type = ?, defect_type = ?,
            status = 'draft', confirmed_by = NULL, confirmed_at = NULL,
            confirm_signature = NULL, email_sent_at = NULL, sent_to = NULL
        WHERE id = ?
    """, (material_no, material_name, supplier, defect_description, due_date, issued_date,
          lot_number, receive_date, cc_recipient, sample_qty, defect_qty, special_note, lot_qty,
          occurrence_type, defect_type, ncr_id))
    conn.commit()
    conn.close()


def confirm_ncr(ncr_id, confirmed_by, signature_path=None):
    """부적합 통보서 확인 완료. 최종결정권자의 승인 서명 경로를 같이 저장한다."""
    conn = get_conn()
    conn.execute("""
        UPDATE ncr SET status = 'confirmed',
            confirmed_by = ?, confirmed_at = datetime('now','localtime'),
            confirm_signature = ?
        WHERE id = ?
    """, (confirmed_by, signature_path, ncr_id))
    conn.commit()
    conn.close()

def get_ncr(ncr_id):
    conn = get_conn()
    row = conn.execute("SELECT * FROM ncr WHERE id = ?", (ncr_id,)).fetchone()
    conn.close()
    return row

_NCR_LIST_SELECT = """
    SELECT n.*,
           i.inspector       AS insp_inspector,
           i.receive_date    AS insp_receive_date,
           i.inspect_date    AS insp_inspect_date,
           i.overall_result  AS insp_overall_result,
           i.status          AS insp_status,
           i.approval_type   AS insp_approval_type
    FROM ncr n
    LEFT JOIN inspections i ON i.id = n.inspection_id
"""


def list_ncr(inspection_id=None):
    """검사자/입고일/검사일/자동판정/승인상태로도 검색할 수 있게 원본 성적서(inspections)
    정보를 같이 조인해서 넘긴다(insp_ 접두어 — ncr.status와 헷갈리지 않게 구분)."""
    conn = get_conn()
    if inspection_id:
        rows = conn.execute(
            _NCR_LIST_SELECT + " WHERE n.inspection_id = ? ORDER BY n.id DESC", (inspection_id,)
        ).fetchall()
    else:
        rows = conn.execute(_NCR_LIST_SELECT + " ORDER BY n.id DESC").fetchall()
    conn.close()
    return rows

def add_ncr_photo(ncr_id, photo_path):
    import json
    conn = get_conn()
    row = conn.execute("SELECT photos FROM ncr WHERE id = ?", (ncr_id,)).fetchone()
    photos = json.loads(row["photos"] or "[]")
    photos.append(photo_path)
    conn.execute("UPDATE ncr SET photos = ? WHERE id = ?", (json.dumps(photos), ncr_id))
    conn.commit()
    conn.close()

def mark_ncr_email_sent(ncr_id, sent_to):
    conn = get_conn()
    conn.execute("""
        UPDATE ncr SET status = 'sent',
            email_sent_at = datetime('now','localtime'),
            sent_to = ?
        WHERE id = ?
    """, (sent_to, ncr_id))
    conn.commit()
    conn.close()


# ---------- 개선요청서 (부적합 통보서보다 가벼운 사전 조치 문서, 서명 없음) ----------

def _next_improvement_no():
    from datetime import date
    today = date.today().strftime("%Y%m%d")
    conn = get_conn()
    count = conn.execute(
        "SELECT COUNT(*) FROM improvement_requests WHERE request_no LIKE ?", (f"IMP-{today}-%",)
    ).fetchone()[0]
    conn.close()
    return f"IMP-{today}-{count + 1:03d}"


def create_improvement_request(inspection_id, material_no, material_name, supplier,
                                lot_number, po_number, lot_qty, defect_categories,
                                defect_category_etc, request_detail, confirmed_name,
                                issued_by, issued_date, sample_qty=None, defect_qty=None):
    request_no = _next_improvement_no()
    conn = get_conn()
    cur = conn.execute("""
        INSERT INTO improvement_requests (request_no, inspection_id, material_no, material_name,
            supplier, lot_number, po_number, lot_qty, defect_categories, defect_category_etc,
            request_detail, confirmed_name, issued_by, issued_date, sample_qty, defect_qty, status)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'draft')
    """, (request_no, inspection_id, material_no, material_name, supplier, lot_number,
          po_number, lot_qty, defect_categories, defect_category_etc, request_detail,
          confirmed_name, issued_by, issued_date, sample_qty, defect_qty))
    conn.commit()
    req_id = cur.lastrowid
    conn.close()
    return req_id, request_no


def update_improvement_request(req_id, material_no, material_name, supplier,
                                lot_number, po_number, lot_qty, defect_categories,
                                defect_category_etc, request_detail, confirmed_name,
                                issued_date, sample_qty=None, defect_qty=None):
    """개선요청서 내용 수정. NCR과 같은 원칙 — 발송 여부 무관하게 수정 허용하되
    저장 시 status를 'draft'로 되돌리고 발송 흔적을 지운다(서명 개념은 없음)."""
    conn = get_conn()
    conn.execute("""
        UPDATE improvement_requests SET
            material_no = ?, material_name = ?, supplier = ?, lot_number = ?, po_number = ?,
            lot_qty = ?, defect_categories = ?, defect_category_etc = ?, request_detail = ?,
            confirmed_name = ?, issued_date = ?, sample_qty = ?, defect_qty = ?,
            status = 'draft', email_sent_at = NULL, sent_to = NULL
        WHERE id = ?
    """, (material_no, material_name, supplier, lot_number, po_number, lot_qty,
          defect_categories, defect_category_etc, request_detail, confirmed_name,
          issued_date, sample_qty, defect_qty, req_id))
    conn.commit()
    conn.close()


def get_improvement_request(req_id):
    conn = get_conn()
    row = conn.execute("SELECT * FROM improvement_requests WHERE id = ?", (req_id,)).fetchone()
    conn.close()
    return row


_IMPROVEMENT_LIST_SELECT = """
    SELECT r.*,
           i.inspector       AS insp_inspector,
           i.receive_date    AS insp_receive_date,
           i.inspect_date    AS insp_inspect_date,
           i.overall_result  AS insp_overall_result,
           i.status          AS insp_status,
           i.approval_type   AS insp_approval_type
    FROM improvement_requests r
    LEFT JOIN inspections i ON i.id = r.inspection_id
"""


def list_improvement_requests(inspection_id=None):
    """성적서 정보를 LEFT JOIN해서 같이 넘긴다 — inspection_id가 NULL인 독립작성 건도
    정상 조회된다(insp_ 컬럼은 전부 NULL로 나올 뿐)."""
    conn = get_conn()
    if inspection_id:
        rows = conn.execute(
            _IMPROVEMENT_LIST_SELECT + " WHERE r.inspection_id = ? ORDER BY r.id DESC", (inspection_id,)
        ).fetchall()
    else:
        rows = conn.execute(_IMPROVEMENT_LIST_SELECT + " ORDER BY r.id DESC").fetchall()
    conn.close()
    return rows


def add_improvement_photo(req_id, kind, photo_path):
    conn = get_conn()
    cur = conn.execute(
        "INSERT INTO improvement_request_photos (improvement_request_id, kind, photo_path) VALUES (?, ?, ?)",
        (req_id, kind, photo_path))
    conn.commit()
    photo_id = cur.lastrowid
    conn.close()
    return photo_id


def list_improvement_photos(req_id, kind=None):
    conn = get_conn()
    if kind:
        rows = conn.execute(
            "SELECT * FROM improvement_request_photos WHERE improvement_request_id = ? AND kind = ? ORDER BY id",
            (req_id, kind)).fetchall()
    else:
        rows = conn.execute(
            "SELECT * FROM improvement_request_photos WHERE improvement_request_id = ? ORDER BY id",
            (req_id,)).fetchall()
    conn.close()
    return rows


def delete_improvement_photo(photo_id):
    conn = get_conn()
    row = conn.execute("SELECT photo_path FROM improvement_request_photos WHERE id = ?", (photo_id,)).fetchone()
    conn.execute("DELETE FROM improvement_request_photos WHERE id = ?", (photo_id,))
    conn.commit()
    conn.close()
    return row["photo_path"] if row else None


def mark_improvement_sent(req_id, sent_to):
    conn = get_conn()
    conn.execute("""
        UPDATE improvement_requests SET status = 'sent',
            email_sent_at = datetime('now','localtime'),
            sent_to = ?
        WHERE id = ?
    """, (sent_to, req_id))
    conn.commit()
    conn.close()


def delete_improvement_requests(req_ids):
    """개선요청서 일괄 삭제. 첨부 사진의 실제 파일명 목록을 반환하니 호출부가 디스크에서도
    지울 것(delete_outbound_item()과 같은 계약)."""
    if not req_ids:
        return []
    conn = get_conn()
    placeholders = ",".join("?" * len(req_ids))
    photo_paths = [r["photo_path"] for r in conn.execute(
        f"SELECT photo_path FROM improvement_request_photos WHERE improvement_request_id IN ({placeholders})",
        req_ids).fetchall()]
    conn.execute(f"DELETE FROM improvement_request_photos WHERE improvement_request_id IN ({placeholders})", req_ids)
    conn.execute(f"DELETE FROM improvement_requests WHERE id IN ({placeholders})", req_ids)
    conn.commit()
    conn.close()
    return photo_paths


# ---------- 커스텀 성적서 템플릿 ----------

def list_custom_templates():
    """모든 커스텀 템플릿 + 각 템플릿에 지정된 자재 수."""
    conn = get_conn()
    rows = conn.execute("""
        SELECT t.*,
               (SELECT COUNT(*) FROM materials m WHERE m.custom_template_id = t.id) AS material_count
          FROM custom_templates t
         ORDER BY t.updated_at DESC, t.id DESC
    """).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def get_custom_template(template_id):
    conn = get_conn()
    row = conn.execute("SELECT * FROM custom_templates WHERE id = ?", (template_id,)).fetchone()
    conn.close()
    return dict(row) if row else None


def create_custom_template(name, layout_json="[]", created_by=None,
                           canvas_w=495, canvas_h=700,
                           page_size="A4", orientation="portrait"):
    conn = get_conn()
    cur = conn.execute("""
        INSERT INTO custom_templates
            (name, layout_json, created_by, canvas_w, canvas_h, page_size, orientation)
        VALUES (?, ?, ?, ?, ?, ?, ?)
    """, (name, layout_json, created_by, canvas_w, canvas_h, page_size, orientation))
    tid = cur.lastrowid
    conn.commit()
    conn.close()
    return tid


def update_custom_template(template_id, name=None, layout_json=None,
                           canvas_w=None, canvas_h=None,
                           page_size=None, orientation=None):
    fields, params = [], []
    for col, val in [("name", name), ("layout_json", layout_json),
                     ("canvas_w", canvas_w), ("canvas_h", canvas_h),
                     ("page_size", page_size), ("orientation", orientation)]:
        if val is not None:
            fields.append(f"{col} = ?")
            params.append(val)
    if not fields:
        return
    fields.append("updated_at = datetime('now', 'localtime')")
    params.append(template_id)
    conn = get_conn()
    conn.execute(f"UPDATE custom_templates SET {', '.join(fields)} WHERE id = ?", params)
    conn.commit()
    conn.close()


def delete_custom_template(template_id):
    """템플릿 삭제 — 이 템플릿을 쓰던 자재는 기본 양식으로 되돌린다."""
    conn = get_conn()
    conn.execute("UPDATE materials SET custom_template_id = NULL WHERE custom_template_id = ?",
                 (template_id,))
    conn.execute("DELETE FROM custom_templates WHERE id = ?", (template_id,))
    conn.commit()
    conn.close()


def materials_for_template(template_id):
    """이 템플릿이 지정된 자재 목록."""
    conn = get_conn()
    rows = conn.execute("""
        SELECT material_no, material_name FROM materials
         WHERE custom_template_id = ? ORDER BY material_no
    """, (template_id,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def set_material_template(material_no, template_id):
    """자재에 커스텀 템플릿 지정(None이면 기본 양식으로 되돌림)."""
    conn = get_conn()
    conn.execute("UPDATE materials SET custom_template_id = ? WHERE material_no = ?",
                 (template_id, material_no))
    conn.commit()
    conn.close()


def get_material_template_id(material_no):
    """자재에 지정된 커스텀 템플릿 id(없으면 None)."""
    conn = get_conn()
    row = conn.execute("SELECT custom_template_id FROM materials WHERE material_no = ?",
                       (material_no,)).fetchone()
    conn.close()
    return row["custom_template_id"] if row and row["custom_template_id"] else None


def latest_inspection_for_material(material_no):
    """이 자재의 가장 최근 승인 성적서(미리보기용). 없으면 None."""
    conn = get_conn()
    row = conn.execute("""
        SELECT * FROM inspections
         WHERE material_no = ? AND status = 'approved'
         ORDER BY id DESC LIMIT 1
    """, (material_no,)).fetchone()
    conn.close()
    return dict(row) if row else None


def get_prev_inspection_values(material_no):
    """이 자재의 가장 최근 성적서(superseded 제외)의 측정값을 반환. 없으면 None."""
    conn = get_conn()
    row = conn.execute("""
        SELECT id, inspector, inspect_date FROM inspections
         WHERE material_no=? AND status != 'superseded'
         ORDER BY id DESC LIMIT 1
    """, (material_no,)).fetchone()
    if not row:
        conn.close()
        return None
    items = conn.execute("""
        SELECT item_name, measured_value FROM inspection_items
         WHERE inspection_id=? ORDER BY id
    """, (row["id"],)).fetchall()
    conn.close()
    return {
        "inspection_id": row["id"],
        "inspector": row["inspector"] or "",
        "inspect_date": row["inspect_date"] or "",
        "values": {r["item_name"]: r["measured_value"] for r in items},
    }


# ---------- 전수검사 ----------

def get_full_inspect_config(material_no):
    """자재의 전수검사 설정 dict. 없으면 None."""
    import json as _j
    conn = get_conn()
    row = conn.execute("SELECT full_inspect_config FROM materials WHERE material_no = ?",
                       (material_no,)).fetchone()
    conn.close()
    if not row or not row["full_inspect_config"]:
        return None
    try:
        return _j.loads(row["full_inspect_config"])
    except Exception:
        return None


def set_full_inspect_config(material_no, config_dict):
    """전수검사 열 설정 저장. config_dict=None이면 해제."""
    import json as _j
    conn = get_conn()
    val = _j.dumps(config_dict, ensure_ascii=False) if config_dict is not None else None
    conn.execute("UPDATE materials SET full_inspect_config = ? WHERE material_no = ?",
                 (val, material_no))
    conn.commit()
    conn.close()


def set_housing_voltage(material_no, voltage):
    """하우징 전압분류를 full_inspect_config JSON의 'voltage' 키에 저장(기존 template/note
    키는 보존, 병합 방식). voltage가 빈 문자열/None이면 키를 지워서 명시값을 해제한다
    (해제되면 app.py/report_builder.py의 정규식·문자열 기반 자동판별로 되돌아간다)."""
    cfg = get_full_inspect_config(material_no) or {}
    if voltage:
        cfg["voltage"] = voltage
    else:
        cfg.pop("voltage", None)
    set_full_inspect_config(material_no, cfg)


def clone_housing_material(material_no, material_name, template_material_no, voltage):
    """하우징 자재 신규 등록 — template_material_no(기존 하우징 자재)의 materials/specs
    설정을 그대로 복제해서 새 자재번호로 만든다. /housing/voltage-groups 화면의
    "새 하우징 자재 등록" 폼이 쓴다(register_housing_605013P300.py 1회성 스크립트를
    일반화한 것, 2026-10-03).

    **반드시 full_inspect_config에 template='housing'을 명시적으로 강제한다** —
    복제 원본은 list_housing_materials()에서 고른 자재라 이미 이 키를 갖고 있지만,
    방어적으로 한 번 더 덮어쓴다(이 키가 빠지면 list_housing_materials()에 영원히
    안 걸려서 /housing/inspect에 안 뜨는 실제 사고가 있었다 — 2026-10-03).

    반환: (성공여부, 에러메시지 또는 None).
    """
    conn = get_conn()
    if conn.execute("SELECT 1 FROM materials WHERE material_no=?", (material_no,)).fetchone():
        conn.close()
        return False, f"자재번호 '{material_no}'는 이미 등록돼 있어."
    src_mat = conn.execute(
        "SELECT * FROM materials WHERE material_no=?", (template_material_no,)
    ).fetchone()
    if src_mat is None:
        conn.close()
        return False, f"복제 원본 자재 '{template_material_no}'를 찾을 수 없어."

    cols = [c for c in src_mat.keys() if c != "created_at"]
    values = []
    for c in cols:
        if c == "material_no":
            values.append(material_no)
        elif c == "material_name":
            values.append(material_name)
        else:
            values.append(src_mat[c])
    placeholders = ",".join("?" * len(cols))
    try:
        conn.execute(
            f"INSERT INTO materials ({','.join(cols)}) VALUES ({placeholders})", values
        )
    except sqlite3.IntegrityError:
        conn.close()
        return False, f"자재번호 '{material_no}'는 이미 등록돼 있어."

    src_specs = conn.execute(
        "SELECT * FROM specs WHERE material_no=? ORDER BY item_order, id",
        (template_material_no,),
    ).fetchall()
    try:
        if src_specs:
            spec_cols = [c for c in src_specs[0].keys() if c != "id"]
            for sp in src_specs:
                sp_values = [
                    material_no if c == "material_no" else
                    material_name if c == "material_name" else
                    sp[c]
                    for c in spec_cols
                ]
                sp_placeholders = ",".join("?" * len(spec_cols))
                conn.execute(
                    f"INSERT INTO specs ({','.join(spec_cols)}) VALUES ({sp_placeholders})",
                    sp_values,
                )
        conn.commit()
    except Exception:
        # 규격 복제 도중 실패하면 자재(materials)만 반쪽으로 남아 template 키도 없이
        # 떠버리는 사고(605013P300과 같은 구조)로 이어진다 — 통째로 롤백한다.
        conn.rollback()
        conn.close()
        return False, "규격 복제 중 오류가 나서 등록을 취소했어. 다시 시도해줘."
    conn.close()

    cfg = get_full_inspect_config(material_no) or {}
    cfg["template"] = "housing"
    if voltage:
        cfg["voltage"] = voltage
    set_full_inspect_config(material_no, cfg)
    return True, None


def get_or_create_full_inspection(inspection_id):
    """전수검사 헤더 조회 또는 신규 생성. dict 반환."""
    conn = get_conn()
    row = conn.execute("SELECT * FROM full_inspections WHERE inspection_id = ?",
                       (inspection_id,)).fetchone()
    if row is None:
        conn.execute("INSERT INTO full_inspections (inspection_id) VALUES (?)", (inspection_id,))
        conn.commit()
        row = conn.execute("SELECT * FROM full_inspections WHERE inspection_id = ?",
                           (inspection_id,)).fetchone()
    result = dict(row)
    conn.close()
    return result


def get_full_inspection(inspection_id):
    """전수검사 헤더 조회. 없으면 None."""
    conn = get_conn()
    row = conn.execute("SELECT * FROM full_inspections WHERE inspection_id = ?",
                       (inspection_id,)).fetchone()
    conn.close()
    return dict(row) if row else None


def update_full_inspection(inspection_id, inspect_date=None, complete_date=None, status=None):
    conn = get_conn()
    fi = conn.execute("SELECT id FROM full_inspections WHERE inspection_id = ?",
                      (inspection_id,)).fetchone()
    if fi is None:
        conn.close()
        return
    fields, params = [], []
    if inspect_date is not None:
        fields.append("inspect_date = ?"); params.append(inspect_date)
    if complete_date is not None:
        fields.append("complete_date = ?"); params.append(complete_date)
    if status is not None:
        fields.append("status = ?"); params.append(status)
    if fields:
        params.append(fi["id"])
        conn.execute(f"UPDATE full_inspections SET {', '.join(fields)} WHERE id = ?", params)
        conn.commit()
    conn.close()


def list_full_inspection_units(inspection_id):
    """unit_no 오름차순으로 유닛 목록 반환."""
    import json as _j
    conn = get_conn()
    fi = conn.execute("SELECT id FROM full_inspections WHERE inspection_id = ?",
                      (inspection_id,)).fetchone()
    if fi is None:
        conn.close()
        return []
    rows = conn.execute("""
        SELECT * FROM full_inspection_units
         WHERE full_inspection_id = ?
         ORDER BY unit_no
    """, (fi["id"],)).fetchall()
    conn.close()
    result = []
    for r in rows:
        d = dict(r)
        try:
            d["values"] = _j.loads(d["values_json"] or "{}")
        except Exception:
            d["values"] = {}
        result.append(d)
    return result


def save_full_inspection_units(inspection_id, units):
    """units = [{"unit_no":1,"serial_no":"...","values":{...},"result":"OK","remark":""}]
    기존 전부 삭제 후 재삽입(upsert 대신 단순 교체)."""
    import json as _j
    conn = get_conn()
    fi = conn.execute("SELECT id FROM full_inspections WHERE inspection_id = ?",
                      (inspection_id,)).fetchone()
    if fi is None:
        conn.execute("INSERT INTO full_inspections (inspection_id) VALUES (?)", (inspection_id,))
        conn.commit()
        fi = conn.execute("SELECT id FROM full_inspections WHERE inspection_id = ?",
                          (inspection_id,)).fetchone()
    fid = fi["id"]
    conn.execute("DELETE FROM full_inspection_units WHERE full_inspection_id = ?", (fid,))
    for u in units:
        conn.execute("""
            INSERT INTO full_inspection_units
                (full_inspection_id, unit_no, serial_no, values_json, result, remark, gauge_name)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (fid, u.get("unit_no", 0), u.get("serial_no", ""),
              _j.dumps(u.get("values", {}), ensure_ascii=False),
              u.get("result", ""), u.get("remark", ""), u.get("gauge_name", "")))
    conn.commit()
    conn.close()


# ---- 2026-09-26 신규: 칸 단위 upsert+병합 (delete+reinsert 레이스컨디션 버그 수정) ----
# save_full_inspection_units()는 성적서 "최초 등록" 시 1회성 벌크삽입 전용으로 계속 쓴다
# (동시편집 위험 없음). 그 이후 자동저장/수동저장은 아래 함수로 "바뀐 칸만" 반영한다.

def upsert_full_inspection_cells(inspection_id, cells, compute_result=None):
    """cells = [{"unit_no":1,"key":"N","value":"12.3"}, ...] 를 unit_no별로 그룹핑해서
    DB에 있는 기존 값과 "병합"한 뒤 UPSERT한다(통째로 교체하지 않음 — 2026-09-26 이전
    delete+reinsert 방식은 두 사람이 같은 성적서의 다른 열을 동시에 입력하면 한쪽이
    저장한 값을 다른쪽 저장이 지워버리는 실제 레이스컨디션 버그가 있었다).

    key가 'serial_no'/'remark'/'gauge_name'이면 해당 컬럼에 직접 반영, 그 외 key는
    기존 values_json(dict)에 그 key만 병합(나머지 기존 key는 그대로 유지).

    compute_result(unit_no, values_dict) -> "OK"/"NG"/"" 콜백을 병합된 유닛 "전체" values
    기준으로 호출해 result를 계산한다 — 판정 로직(스펙 범위 비교, 차수별 필수항목 완료
    여부 등)은 app.py의 _fi_auto_result가 갖고 있어 database.py는 몰라도 되게 콜백으로
    분리했다(8-1절 관례: DB 계층에 업무판정 로직을 심지 않음). unit_no를 넘기는 이유는
    하우징처럼 유닛마다 "필수 항목"이 다를 수 있어서다(1차 AQL 샘플 외 유닛은 1차 항목이
    필요 없음, 2026-09-28).

    같은 unit_no를 두 요청이 동시에 건드릴 때 "읽고(병합)-쓰는" 사이에 다른 요청이 끼어들어
    한쪽 병합 결과가 유실되는 걸 막기 위해, 유닛별 읽기+쓰기 전체를 BEGIN IMMEDIATE
    트랜잭션 하나로 묶는다(create_inspection()의 TOCTOU 방지 패턴과 동일, CLAUDE.md 20절 참고).

    반환: [{"unit_no","serial_no","remark","gauge_name","values","result"}, ...]
    """
    import json as _j
    conn = get_conn()
    conn.execute("BEGIN IMMEDIATE")
    try:
        fi = conn.execute("SELECT id FROM full_inspections WHERE inspection_id = ?",
                          (inspection_id,)).fetchone()
        if fi is None:
            conn.execute("INSERT INTO full_inspections (inspection_id) VALUES (?)", (inspection_id,))
            fi = conn.execute("SELECT id FROM full_inspections WHERE inspection_id = ?",
                              (inspection_id,)).fetchone()
        fid = fi["id"]

        by_unit = {}
        for c in cells:
            try:
                un = int(c.get("unit_no", 0))
            except (TypeError, ValueError):
                continue
            if un <= 0:
                continue
            by_unit.setdefault(un, []).append(c)

        saved = []
        for unit_no, unit_cells in by_unit.items():
            row = conn.execute("""
                SELECT * FROM full_inspection_units
                 WHERE full_inspection_id = ? AND unit_no = ?
            """, (fid, unit_no)).fetchone()
            if row:
                serial_no = row["serial_no"] or ""
                remark = row["remark"] or ""
                gauge_name = row["gauge_name"] or ""
                try:
                    values = _j.loads(row["values_json"] or "{}")
                except Exception:
                    values = {}
            else:
                serial_no, remark, gauge_name, values = "", "", "", {}

            for c in unit_cells:
                key = c.get("key")
                if not key or not isinstance(key, str):
                    continue
                val = str(c.get("value", "")).strip()
                if key == "serial_no":
                    serial_no = val
                elif key == "remark":
                    remark = val
                elif key == "gauge_name":
                    gauge_name = val
                else:
                    values[key] = val

            result = compute_result(unit_no, values) if compute_result else ""
            values_json = _j.dumps(values, ensure_ascii=False)

            if row:
                conn.execute("""
                    UPDATE full_inspection_units
                       SET serial_no = ?, values_json = ?, result = ?, remark = ?, gauge_name = ?
                     WHERE id = ?
                """, (serial_no, values_json, result, remark, gauge_name, row["id"]))
            else:
                conn.execute("""
                    INSERT INTO full_inspection_units
                        (full_inspection_id, unit_no, serial_no, values_json, result, remark, gauge_name)
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                """, (fid, unit_no, serial_no, values_json, result, remark, gauge_name))

            saved.append({"unit_no": unit_no, "serial_no": serial_no, "remark": remark,
                          "gauge_name": gauge_name, "values": values, "result": result})

        conn.commit()
        return saved
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def record_full_inspection_editor(full_inspection_id, user_id, display_name):
    """칸을 저장할 때마다 그 사용자를 이 성적서를 건드린 사람 집합에 추가(중복 자동 무시).
    성적서 '검사자'란에 누적된 이름 전부를 보여주는 용도(화면 반영은 다음 단계)."""
    conn = get_conn()
    conn.execute("""
        INSERT OR IGNORE INTO full_inspection_editors (full_inspection_id, user_id, display_name)
        VALUES (?, ?, ?)
    """, (full_inspection_id, user_id, display_name))
    conn.commit()
    conn.close()


def list_full_inspection_editor_names(inspection_id):
    """이 성적서(전수검사)를 건드린 사람 이름 목록(입력순). 하우징 승인목록 '검사자' 칸용."""
    conn = get_conn()
    rows = conn.execute("""
        SELECT e.display_name FROM full_inspection_editors e
        JOIN full_inspections fi ON fi.id = e.full_inspection_id
        WHERE fi.inspection_id = ?
        ORDER BY e.id
    """, (inspection_id,)).fetchall()
    conn.close()
    return [r["display_name"] for r in rows if r["display_name"]]


def get_full_inspection_unit_id(inspection_id, unit_no):
    """(inspection_id, unit_no) -> full_inspection_units.id. 아직 그 유닛이 한 번도
    저장된 적 없으면 None(심화검사 저장 전 반드시 일반 저장이 먼저 있어야 함)."""
    conn = get_conn()
    row = conn.execute("""
        SELECT fiu.id FROM full_inspection_units fiu
        JOIN full_inspections fi ON fi.id = fiu.full_inspection_id
        WHERE fi.inspection_id = ? AND fiu.unit_no = ?
    """, (inspection_id, unit_no)).fetchone()
    conn.close()
    return row["id"] if row else None


def set_full_inspection_unit_special(inspection_id, unit_no, special):
    """유닛 단위 특채 표시를 켜고 끈다(2026-09-28). NG로 자동판정된 유닛을 검사자·
    최종결정권자가 개별 예외처리할 때 쓴다 — 이 유닛의 result 자체는 안 건드리고
    (다음 저장 때 _fi_auto_result가 또 덮어쓸 수 있으므로) 별도 플래그만 세운다.
    아직 저장된 적 없는 유닛이면 아무 일도 안 하고 False를 돌려준다."""
    unit_id = get_full_inspection_unit_id(inspection_id, unit_no)
    if unit_id is None:
        return False
    conn = get_conn()
    conn.execute("UPDATE full_inspection_units SET special_marked=? WHERE id=?",
                 (1 if special else 0, unit_id))
    conn.commit()
    conn.close()
    return True


def add_housing_deep_inspection(full_inspection_unit_id, test_type, vi_state, chamber_stage,
                                 readings, result, remark, created_by):
    """PD/충격 심화검사 1건 기록. readings는 {"u_pre":..,"u_pd":..,"u_i":..,"u_e":..} 형태의
    dict(PD/충격 공용 4개 전압필드, 사용자 확정 — PROGRESS.md 2026-09-25 참고) -> JSON 저장.
    반환: 새로 생성된 행 id."""
    import json as _j
    conn = get_conn()
    cur = conn.execute("""
        INSERT INTO housing_deep_inspections
            (full_inspection_unit_id, test_type, vi_state, chamber_stage, readings_json, result, remark, created_by)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    """, (full_inspection_unit_id, test_type, vi_state, chamber_stage,
          _j.dumps(readings or {}, ensure_ascii=False), result or "", remark or "", created_by))
    conn.commit()
    new_id = cur.lastrowid
    conn.close()
    return new_id


def count_housing_deep_inspections_by_unit(full_inspection_id):
    """이 전수검사(성적서 1건) 안의 유닛별 심화검사 건수. {unit_no: 건수} 형태 — 화면(입력/
    승인상세) 렌더링 시 유닛마다 별도 쿼리를 안 날리게 한 번에 계산한다(CLAUDE.md 24절
    N+1 방지 패턴과 동일한 취지)."""
    conn = get_conn()
    rows = conn.execute("""
        SELECT fiu.unit_no AS unit_no, COUNT(*) AS cnt
        FROM housing_deep_inspections hdi
        JOIN full_inspection_units fiu ON fiu.id = hdi.full_inspection_unit_id
        WHERE fiu.full_inspection_id = ?
        GROUP BY fiu.unit_no
    """, (full_inspection_id,)).fetchall()
    conn.close()
    return {r["unit_no"]: r["cnt"] for r in rows}


def list_inspection_ids_with_deep_inspections(inspection_ids):
    """주어진 성적서 id 목록 중 심화검사가 하나라도 있는 것들의 id 집합. 하우징 검사이력
    목록에 "심화검사 있음" 컬럼을 붙일 때 성적서마다 따로 쿼리하지 않으려고 한 번에 계산."""
    ids = [i for i in (inspection_ids or []) if i]
    if not ids:
        return set()
    conn = get_conn()
    ph = ",".join("?" * len(ids))
    rows = conn.execute(f"""
        SELECT DISTINCT fi.inspection_id AS iid
        FROM housing_deep_inspections hdi
        JOIN full_inspection_units fiu ON fiu.id = hdi.full_inspection_unit_id
        JOIN full_inspections fi ON fi.id = fiu.full_inspection_id
        WHERE fi.inspection_id IN ({ph})
    """, ids).fetchall()
    conn.close()
    return {r["iid"] for r in rows}


def list_housing_materials():
    """template=='housing'로 설정된 자재 전체. {material_no: material_name} 형태.
    하우징 전용 목록 3화면(검사입력/승인/이력)이 공유하는 자재 판별 기준(get_full_inspect_config
    의 JSON 계약을 그대로 따름, CLAUDE.md 8-1절)."""
    import json as _j
    conn = get_conn()
    rows = conn.execute(
        "SELECT material_no, material_name, full_inspect_config FROM materials "
        "WHERE full_inspect_config IS NOT NULL"
    ).fetchall()
    conn.close()
    out = {}
    for r in rows:
        try:
            cfg = _j.loads(r["full_inspect_config"])
        except Exception:
            continue
        if isinstance(cfg, dict) and cfg.get("template") == "housing":
            out[r["material_no"]] = r["material_name"]
    return out


_FI_STAGE_TIME_COLS = {
    "stage1_started_at", "stage1_completed_at",
    "stage2_started_at", "stage2_completed_at",
    "stage3_started_at", "stage3_completed_at",
}


def update_full_inspection_stage_times(inspection_id, **fields):
    """차수(1/2/3차)별 시작·완료 시각 기록. fields의 키는 _FI_STAGE_TIME_COLS 안에 있는
    것만 반영한다(동적 컬럼명 SQL 인젝션 방지용 허용목록). '이미 값이 있으면 갱신하지
    않는다'는 최초 완료시점만 남기는 규칙은 호출부(app.py)가 판단해서 이 함수엔 필요한
    필드만 넘긴다."""
    conn = get_conn()
    fi = conn.execute("SELECT id FROM full_inspections WHERE inspection_id = ?",
                      (inspection_id,)).fetchone()
    if fi is None:
        conn.close()
        return
    set_parts, params = [], []
    for k, v in fields.items():
        if k not in _FI_STAGE_TIME_COLS:
            continue
        set_parts.append(f"{k} = ?")
        params.append(v)
    if set_parts:
        params.append(fi["id"])
        conn.execute(f"UPDATE full_inspections SET {', '.join(set_parts)} WHERE id = ?", params)
        conn.commit()
    conn.close()
# ---- 2026-09-26 신규 끝 ----


def delete_full_inspection(inspection_id):
    conn = get_conn()
    fi = conn.execute("SELECT id FROM full_inspections WHERE inspection_id = ?",
                      (inspection_id,)).fetchone()
    if fi:
        conn.execute("DELETE FROM full_inspection_units WHERE full_inspection_id = ?", (fi["id"],))
        conn.execute("DELETE FROM full_inspections WHERE id = ?", (fi["id"],))
        conn.commit()
    conn.close()


# ---------- 앱 설정 ----------

def get_setting(key, default=None):
    conn = get_conn()
    row = conn.execute("SELECT value FROM app_settings WHERE key = ?", (key,)).fetchone()
    conn.close()
    return row["value"] if row else default

def set_setting(key, value):
    conn = get_conn()
    conn.execute(
        "INSERT INTO app_settings (key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (key, value)
    )
    conn.commit()
    conn.close()

def traceability_enabled():
    """제품검사 이력추적 기능 스위치 — 기본 OFF, 켜기 전까지 신규 기능이 동작하지 않게 하는 용도."""
    return get_setting("traceability_enabled", "0") == "1"

def set_traceability_enabled(on):
    """추적 기능 스위치 저장('1'/'0') — 같은 키를 traceability_enabled()가 읽는다."""
    set_setting("traceability_enabled", "1" if on else "0")


# ---------- 완제품 모델 마스터 (TASK #66 패키지 1B) ----------

_MODEL_NO_RE = re.compile(r"^[A-Za-z0-9_-]+$")
_PM_IDENTITY_FIELDS = ("voltage_code", "voltage_kv", "suffix_code", "pcode",
                       "phase_type", "housing_required_qty")


def normalize_model_key(s):
    """모델번호 정규화 키 — 모든 공백 제거 후 대문자."""
    return re.sub(r"\s+", "", s or "").upper()


def _pm_clean(data):
    """폼 입력을 검증·정리해 dict로 돌려준다(하우징 수량은 형식에서 서버가 강제)."""
    def txt(k):
        v = data.get(k)
        return ("" if v is None else str(v).strip()) or None
    phase = txt("phase_type")
    if phase not in ("single", "three"):
        raise ValueError("제품 형식은 단상 또는 3상이어야 해.")
    kv = txt("voltage_kv")
    if kv is not None:
        try:
            kv = int(kv)
        except ValueError:
            raise ValueError("전압(kV)은 숫자로 입력해줘.")
    return {"display_name": txt("display_name"), "voltage_code": txt("voltage_code"),
            "voltage_kv": kv, "suffix_code": txt("suffix_code"), "pcode": txt("pcode"),
            "phase_type": phase, "housing_required_qty": 1 if phase == "single" else 3,
            "note": txt("note")}


def _pm_snapshot(row):
    return {k: row[k] for k in ("display_name", "note") + _PM_IDENTITY_FIELDS}


_PM_FULL_KEYS = ("model_no", "model_key", "display_name", "voltage_code", "voltage_kv",
                 "suffix_code", "pcode", "phase_type", "housing_required_qty", "is_active",
                 "revision", "note")


def _pm_full(conn, model_id):
    """이력용 전체 상태 스냅샷 — 같은 트랜잭션 연결에서 한 행을 읽어 고정 키 순서 dict로."""
    row = conn.execute("SELECT * FROM product_models WHERE id=?", (model_id,)).fetchone()
    return {k: row[k] for k in _PM_FULL_KEYS}


def _pm_history(conn, model_id, revision, change_type, before, after, reason, user):
    conn.execute(
        "INSERT INTO product_model_history (model_id, revision, change_type, before_json, "
        "after_json, reason, changed_by) VALUES (?,?,?,?,?,?,?)",
        (model_id, revision, change_type,
         json.dumps(before, ensure_ascii=False) if before is not None else None,
         json.dumps(after, ensure_ascii=False) if after is not None else None,
         reason, user))


def list_product_models(search=None, active=None):
    """모델번호·표시명 부분일치 검색. active가 True/False면 사용 여부로 거른다."""
    sql, args = "SELECT * FROM product_models WHERE 1=1", []
    if search:
        sql += " AND (model_no LIKE ? OR display_name LIKE ?)"
        args += [f"%{search}%"] * 2
    if active is not None:
        sql += " AND is_active = ?"
        args.append(1 if active else 0)
    conn = get_conn()
    rows = conn.execute(sql + " ORDER BY model_no", args).fetchall()
    conn.close()
    return rows


def get_product_model(model_id):
    conn = get_conn()
    row = conn.execute("SELECT * FROM product_models WHERE id=?", (model_id,)).fetchone()
    conn.close()
    return row


def get_product_model_history(model_id):
    """변경 이력(최신순)."""
    conn = get_conn()
    rows = conn.execute("SELECT * FROM product_model_history WHERE model_id=? "
                        "ORDER BY id DESC", (model_id,)).fetchall()
    conn.close()
    return rows


@_write_retry
def create_product_model(data, user):
    """모델 등록 + create 이력을 한 트랜잭션으로. 모델번호 중복·형식 오류는 ValueError."""
    model_no = (data.get("model_no") or "").strip()
    if not model_no or len(model_no) > 40 or not _MODEL_NO_RE.match(model_no):
        raise ValueError("모델번호는 40자 이하 영문·숫자·하이픈·밑줄만 쓸 수 있어.")
    c = _pm_clean(data)
    key = normalize_model_key(model_no)
    conn = get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        if conn.execute("SELECT 1 FROM product_models WHERE model_key=?", (key,)).fetchone():
            raise ValueError("이미 등록된 모델번호야.")
        cur = conn.execute(
            "INSERT INTO product_models (model_no, model_key, display_name, voltage_code, "
            "voltage_kv, suffix_code, pcode, phase_type, housing_required_qty, note, created_by) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?)",
            (model_no, key, c["display_name"], c["voltage_code"], c["voltage_kv"],
             c["suffix_code"], c["pcode"], c["phase_type"], c["housing_required_qty"],
             c["note"], user))
        mid = cur.lastrowid
        _pm_history(conn, mid, 1, "create", None, _pm_full(conn, mid), None, user)
        conn.commit()
        return mid
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@_write_retry
def update_product_model(model_id, data, user, reason):
    """표시명·비고는 개정 없이, 식별 정보 변경은 개정+1(사유 필수). model_no는 못 바꾼다."""
    c = _pm_clean(data)
    reason = (reason or "").strip() or None
    conn = get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        row = conn.execute("SELECT * FROM product_models WHERE id=?", (model_id,)).fetchone()
        if not row:
            raise ValueError("모델을 찾을 수 없어.")
        before, after = _pm_snapshot(row), c
        if before == after:
            conn.rollback()
            return row["revision"]
        ident_changed = any(before[k] != after[k] for k in _PM_IDENTITY_FIELDS)
        if ident_changed and not reason:
            raise ValueError("식별 정보를 바꾸려면 사유를 입력해줘.")
        if ident_changed and conn.execute("SELECT 1 FROM production_lots WHERE model_id=? LIMIT 1",
                                          (model_id,)).fetchone():
            raise ValueError("이미 LOT에서 쓰는 모델은 식별 정보를 바꿀 수 없어. 새 모델로 등록해줘.")
        rev = row["revision"] + (1 if ident_changed else 0)
        full_before = _pm_full(conn, model_id)
        conn.execute(
            "UPDATE product_models SET display_name=?, voltage_code=?, voltage_kv=?, "
            "suffix_code=?, pcode=?, phase_type=?, housing_required_qty=?, note=?, revision=?, "
            "updated_by=?, updated_at=datetime('now','localtime') WHERE id=?",
            (c["display_name"], c["voltage_code"], c["voltage_kv"], c["suffix_code"], c["pcode"],
             c["phase_type"], c["housing_required_qty"], c["note"], rev, user, model_id))
        _pm_history(conn, model_id, rev, "update", full_before, _pm_full(conn, model_id),
                    reason, user)
        conn.commit()
        return rev
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@_write_retry
def set_product_model_active(model_id, is_active, user, reason):
    """사용 중지/재사용 전환 + 이력(deactivate/reactivate). 개정번호는 유지."""
    conn = get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        row = conn.execute("SELECT * FROM product_models WHERE id=?", (model_id,)).fetchone()
        if not row:
            raise ValueError("모델을 찾을 수 없어.")
        flag = 1 if is_active else 0
        if row["is_active"] == flag:
            conn.rollback()
            return
        before = _pm_full(conn, model_id)
        conn.execute("UPDATE product_models SET is_active=?, updated_by=?, "
                     "updated_at=datetime('now','localtime') WHERE id=?", (flag, user, model_id))
        _pm_history(conn, model_id, row["revision"], "reactivate" if flag else "deactivate",
                    before, _pm_full(conn, model_id),
                    (reason or "").strip() or None, user)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


# ---------- 생산 LOT (TASK #66 패키지 1C-a) ----------

def kst_date_str():
    """오늘 날짜(한국, YYYYMMDD) — 서버가 UTC여도 UTC+9를 직접 계산한다."""
    return _dt.now(_tz(_td(hours=9))).strftime("%Y%m%d")


def product_model_lot_reference_count(model_id):
    """이 모델을 쓰는 생산 LOT 수(상태 불문)."""
    conn = get_conn()
    n = conn.execute("SELECT COUNT(*) FROM production_lots WHERE model_id=?", (model_id,)).fetchone()[0]
    conn.close()
    return n


def _pl_full(conn, lot_id):
    """이력용 전체 상태 스냅샷 — 같은 연결에서 읽어 고정 키 순서 dict로."""
    r = conn.execute("SELECT * FROM production_lots WHERE id=?", (lot_id,)).fetchone()
    snap = json.loads(r["model_snapshot_json"]) if r["model_snapshot_json"] else None
    return {"lot_no": r["lot_no"], "lot_date": r["lot_date"], "seq": r["seq"],
            "model_id": r["model_id"], "model_snapshot": snap, "status": r["status"],
            "is_test": r["is_test"], "lot_revision": r["lot_revision"], "note": r["note"],
            "cancel_reason": r["cancel_reason"]}


def _pl_history(conn, lot_id, change_type, before, after, reason, user):
    conn.execute(
        "INSERT INTO production_lot_history (lot_id, change_type, before_json, after_json, "
        "reason, changed_by) VALUES (?,?,?,?,?,?)",
        (lot_id, change_type,
         json.dumps(before, ensure_ascii=False) if before is not None else None,
         json.dumps(after, ensure_ascii=False), reason, user))


def _pl_lock(conn, lot_id, need_draft=False):
    """BEGIN IMMEDIATE 후 LOT 행을 읽는다. 없으면/상태가 안 맞으면 ValueError."""
    conn.execute("BEGIN IMMEDIATE")
    row = conn.execute("SELECT * FROM production_lots WHERE id=?", (lot_id,)).fetchone()
    if not row:
        raise ValueError("LOT을 찾을 수 없어.")
    if need_draft and row["status"] != "draft":
        raise ValueError("작성 중인 LOT만 바꿀 수 있어.")
    return row


@_write_retry
def create_production_lot(model_id, is_test, note, user):
    """생산 LOT 생성(작성중) + create 이력. 번호는 한국 날짜 기준 일련번호, 취소 번호도 재사용 안 함."""
    note = (note or "").strip() or None
    conn = get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        m = conn.execute("SELECT is_active FROM product_models WHERE id=?", (model_id,)).fetchone()
        if not m:
            raise ValueError("모델을 찾을 수 없어.")
        if not m["is_active"]:
            raise ValueError("사용 중지된 모델이야.")
        lot_date = kst_date_str()
        seq = conn.execute("SELECT COALESCE(MAX(seq),0)+1 FROM production_lots WHERE lot_date=?",
                           (lot_date,)).fetchone()[0]
        cur = conn.execute(
            "INSERT INTO production_lots (lot_no, lot_date, seq, model_id, status, is_test, note, "
            "created_by) VALUES (?,?,?,?, 'draft', ?,?,?)",
            (f"PROD-{lot_date}-{seq:03d}", lot_date, seq, model_id, 1 if is_test else 0, note, user))
        lid = cur.lastrowid
        _pl_history(conn, lid, "create", None, _pl_full(conn, lid), None, user)
        conn.commit()
        return lid
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@_write_retry
def update_production_lot(lot_id, note, is_test, user):
    """작성중 LOT의 비고·테스트 표시 수정 + update 이력. 값이 같으면 이력 없음."""
    note = (note or "").strip() or None
    flag = 1 if is_test else 0
    conn = get_conn()
    try:
        row = _pl_lock(conn, lot_id, need_draft=True)
        if row["note"] == note and row["is_test"] == flag:
            conn.rollback()
            return
        before = _pl_full(conn, lot_id)
        conn.execute("UPDATE production_lots SET note=?, is_test=?, updated_by=?, "
                     "updated_at=datetime('now','localtime') WHERE id=?", (note, flag, user, lot_id))
        _pl_history(conn, lot_id, "update", before, _pl_full(conn, lot_id), None, user)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@_write_retry
def confirm_production_lot(lot_id, user):
    """작성중 LOT 확정 — 정상 S/N이 1개 이상일 때만, 그 시점 모델 값을 스냅샷으로 굳히고 confirm 이력."""
    conn = get_conn()
    try:
        row = _pl_lock(conn, lot_id, need_draft=True)
        m = conn.execute("SELECT * FROM product_models WHERE id=?", (row["model_id"],)).fetchone()
        if not m or not m["is_active"]:
            raise ValueError("사용 중지된 모델이야.")
        if not conn.execute("SELECT 1 FROM finished_units WHERE lot_id=? AND status='active' LIMIT 1",
                            (lot_id,)).fetchone():
            raise ValueError("정상 S/N이 1개 이상 있어야 확정할 수 있어.")
        snap = {"model_id": m["id"], "model_no": m["model_no"], "model_key": m["model_key"],
                "model_revision": m["revision"], "display_name": m["display_name"],
                "voltage_code": m["voltage_code"], "voltage_kv": m["voltage_kv"],
                "suffix_code": m["suffix_code"], "pcode": m["pcode"], "phase_type": m["phase_type"],
                "housing_required_qty": m["housing_required_qty"]}
        before = _pl_full(conn, lot_id)
        conn.execute("UPDATE production_lots SET status='confirmed', model_snapshot_json=?, "
                     "confirmed_by=?, confirmed_at=datetime('now','localtime'), updated_by=?, "
                     "updated_at=datetime('now','localtime') WHERE id=?",
                     (json.dumps(snap, ensure_ascii=False), user, user, lot_id))
        _pl_history(conn, lot_id, "confirm", before, _pl_full(conn, lot_id), None, user)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def lot_has_started_inspection(lot_id):
    """이 LOT의 제품검사 시작 기록(production_lot_inspections)이 있는지."""
    conn = get_conn()
    try:
        return conn.execute("SELECT 1 FROM production_lot_inspections WHERE lot_id=?",
                            (lot_id,)).fetchone() is not None
    finally:
        conn.close()


_FINAL_DOC_SQL = ("SELECT 1 FROM product_report_revisions r JOIN production_lot_inspections i "
                  "ON i.id=r.inspection_id JOIN product_report_events e ON e.revision_id=r.id "
                  "AND e.seq=(SELECT MAX(x.seq) FROM product_report_events x WHERE x.revision_id=r.id) "
                  "WHERE e.event_type='approved' AND ")


def has_final_documents_for_lot(lot_id, conn=None):
    """LOT에 최종 승인된 제품 성적서(마지막 이벤트=approved)가 있는지. 쓰기 트랜잭션 안에서는 자기 conn을 넘길 것."""
    own = conn is None
    c = get_conn() if own else conn
    try:
        return c.execute(_FINAL_DOC_SQL + "i.lot_id=? LIMIT 1", (lot_id,)).fetchone() is not None
    finally:
        if own:
            c.close()


def unit_has_final_document(unit_id, conn=None):
    """완제품의 제품 성적서가 최종 승인됐는지(어느 revision이든 마지막 이벤트=approved). 쓰기 트랜잭션 안에서는 자기 conn을 넘길 것."""
    own = conn is None
    c = get_conn() if own else conn
    try:
        return c.execute(_FINAL_DOC_SQL + "r.unit_id=? LIMIT 1", (unit_id,)).fetchone() is not None
    finally:
        if own:
            c.close()


_FINAL_DOC_UNIT_MSG = "이 S/N의 성적서가 이미 최종 승인 또는 출력되어 제외/정정할 수 없어요. 승인 회수가 먼저 필요해요."
_UNUSABLE_UNIT_MSG = "이 S/N은 사용 금지로 결정돼서 입력·정정·승인 요청을 할 수 없어요. 사용 금지를 되돌린 뒤에 할 수 있어요."


def unit_is_unusable(unit_id, conn=None):
    """S/N의 마지막 사용 금지 이벤트가 unusable이면 참(회수되면 거짓). 쓰기 트랜잭션 안에서는 자기 conn을 넘길 것."""
    own = conn is None
    c = get_conn() if own else conn
    try:
        r = c.execute("SELECT event_type FROM unit_disposal_events WHERE unit_id=? ORDER BY seq DESC LIMIT 1",
                      (unit_id,)).fetchone()
        return bool(r and r["event_type"] == "unusable")
    finally:
        if own:
            c.close()


@_write_retry
def cancel_production_lot(lot_id, reason, user):
    """LOT 전체 취소(사유 필수) + cancel 이력. 행·번호는 남긴다.
    검사 시작 전이면 정상 S/N을 모두 'LOT 취소'로 제외(같은 트랜잭션), 시작 후면 LOT 상태만 취소하고
    S/N·검사 기록은 보존한다. 최종 승인·출력된 성적서가 있으면 취소 불가.
    lot_revision은 올리지 않는다 — 취소는 LOT 상태 전이 1건이고, S/N별 제외 이력에는 취소 시점의 lot_revision이 기록된다."""
    reason = (reason or "").strip()
    if not reason:
        raise ValueError("취소 사유를 입력해줘.")
    conn = get_conn()
    try:
        row = _pl_lock(conn, lot_id)
        if row["status"] == "cancelled":
            raise ValueError("이미 취소된 LOT이야.")
        if has_final_documents_for_lot(lot_id, conn):
            raise ValueError("최종 승인 또는 출력된 성적서가 있어서 취소할 수 없어요. 승인 회수가 먼저 필요해요.")
        before = _pl_full(conn, lot_id)
        if not lot_has_started_inspection(lot_id):
            for u in conn.execute("SELECT id FROM finished_units WHERE lot_id=? AND status='active' "
                                  "ORDER BY id", (lot_id,)).fetchall():
                ub = _fu_full(conn, u["id"])
                conn.execute("UPDATE finished_units SET status='excluded', excluded_reason='LOT 취소', "
                             "excluded_by=?, excluded_at=datetime('now','localtime') WHERE id=?",
                             (user, u["id"]))
                _fu_history(conn, u["id"], lot_id, "exclude", ub, _fu_full(conn, u["id"]),
                            "LOT 취소", row["lot_revision"], user)
        conn.execute("UPDATE production_lots SET status='cancelled', cancel_reason=?, "
                     "cancelled_by=?, cancelled_at=datetime('now','localtime'), updated_by=?, "
                     "updated_at=datetime('now','localtime') WHERE id=?", (reason, user, user, lot_id))
        _pl_history(conn, lot_id, "cancel", before, _pl_full(conn, lot_id), reason, user)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


_PL_SELECT = ("SELECT l.*, m.model_no AS model_no, m.display_name AS model_display_name "
              "FROM production_lots l JOIN product_models m ON m.id = l.model_id")


def list_production_lots(search=None, status=None):
    """LOT 번호·모델번호·비고 부분일치 + 상태 필터, 최신순."""
    sql, args = (_PL_SELECT.replace(
        " FROM production_lots l",
        ", (SELECT COUNT(*) FROM finished_units u WHERE u.lot_id=l.id AND u.status='active') "
        "AS active_count, (SELECT COUNT(*) FROM finished_units u WHERE u.lot_id=l.id "
        "AND u.status='excluded') AS excluded_count FROM production_lots l", 1)
        + " WHERE 1=1"), []
    if search:
        sql += " AND (l.lot_no LIKE ? OR m.model_no LIKE ? OR l.note LIKE ?)"
        args += [f"%{search}%"] * 3
    if status:
        sql += " AND l.status = ?"
        args.append(status)
    conn = get_conn()
    rows = conn.execute(sql + " ORDER BY l.id DESC", args).fetchall()
    conn.close()
    return rows


def get_production_lot(lot_id):
    conn = get_conn()
    row = conn.execute(_PL_SELECT + " WHERE l.id=?", (lot_id,)).fetchone()
    conn.close()
    return row


def get_production_lot_history(lot_id):
    """LOT 변경 이력(최신순)."""
    conn = get_conn()
    rows = conn.execute("SELECT * FROM production_lot_history WHERE lot_id=? ORDER BY id DESC",
                        (lot_id,)).fetchall()
    conn.close()
    return rows


# ---------- 제품검사 항목 버전 · LOT 검사 시작 (TASK #136 패키지 1D-a) ----------

_ITEM_KEY_RE = re.compile(r"^[A-Za-z0-9_]{1,40}$")
_ITEM_TYPES = ("numeric", "ok_ng", "photo")
_ITEM_KEYS = ("key", "name", "input_type", "lower", "upper", "display_digits",
              "required_n", "allow_missing", "sort_order")


def _canon(obj):
    """해시용 정규화 직렬화 — 키 정렬·비ASCII 유지·공백 없음."""
    return json.dumps(obj, sort_keys=True, ensure_ascii=False, separators=(",", ":"))


def _sha(obj):
    return hashlib.sha256(_canon(obj).encode("utf-8")).hexdigest()


def _utc_now():
    """기존 테이블의 datetime('now','localtime')과 같은 형식·기준(서버 로컬 시각) — 한 화면에서 시각이 어긋나지 않게 맞춘다."""
    return _dt.now().strftime("%Y-%m-%d %H:%M:%S")


# 전각 숫자(０-９)·전각 부호/점·유니코드 마이너스만 반각으로 바꾼다. ¹ ① 같은 호환 숫자는 바꾸지 않아 거부된다.
_ITEM_NUM_MAP = {**{0xFF10 + i: 0x30 + i for i in range(10)},
                 0xFF0B: ord("+"), 0xFF0D: ord("-"), 0x2212: ord("-"), 0xFF0E: ord(".")}


def _item_num(v, label, n):
    """빈 값은 None, 숫자 아니면 ValueError. Decimal로 읽어 float 오차 없이 '10.5'처럼 표준 십진 문자열로 돌려준다
    (5와 5.0과 5.00은 모두 '5', -0은 0). 1D-b 판정이 Decimal로 비교하도록 문자열로 보관한다.
    전각 숫자는 반각으로, 유니코드 마이너스(−)는 '-'로 바꾼 뒤 [부호]숫자[.숫자] 형태만 받는다
    (지수 표기·밑줄·쉼표·NaN·무한대는 거부 — 거대한 지수로 긴 문자열이 만들어지는 것을 막는다)."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    s = str(v).translate(_ITEM_NUM_MAP).strip()
    if not re.fullmatch(r"[+-]?(\d+(\.\d*)?|\.\d+)", s, re.ASCII) or len(s) > 24:
        raise ValueError(f"{n}번째 항목: {label}은(는) 숫자로 입력해줘.")
    x = _Dec(s)
    if x == 0:
        return "0"
    return format(x.normalize(), "f")


def _item_int(v, label, n, lo, hi, default):
    if v is None or (isinstance(v, str) and not v.strip()):
        return default
    try:
        i = int(str(v).strip())
    except ValueError:
        raise ValueError(f"{n}번째 항목: {label}은(는) 정수로 입력해줘.")
    if not lo <= i <= hi:
        raise ValueError(f"{n}번째 항목: {label}은(는) {lo}~{hi} 사이여야 해.")
    return i


def _norm_items(items):
    """항목 목록 검증·정규화 → 키 순서 고정 dict 리스트. 오류는 한국어 ValueError.
    key가 비어 있으면 'item_N'으로 자동 부여한다(새 행 입력용)."""
    if not isinstance(items, (list, tuple)) or not items:
        raise ValueError("검사 항목을 1개 이상 입력해줘.")
    out, used = [], set()
    for raw in items:
        k = (raw.get("key") or "").strip() if isinstance(raw, dict) else ""
        if k:
            used.add(k)
    auto = 0
    for pos, raw in enumerate(items, 1):
        if not isinstance(raw, dict):
            raise ValueError(f"{pos}번째 항목 형식이 올바르지 않아.")
        name = (raw.get("name") or "").strip()
        if not name:
            raise ValueError(f"{pos}번째 항목: 이름을 입력해줘.")
        if len(name) > 60:
            raise ValueError(f"{pos}번째 항목: 이름은 60자 이하로 입력해줘.")
        key = (raw.get("key") or "").strip()
        if not key:
            while True:
                auto += 1
                key = f"item_{auto}"
                if key not in used:
                    break
        if not _ITEM_KEY_RE.match(key):
            raise ValueError(f"{pos}번째 항목: 항목 키는 40자 이하 영문·숫자·밑줄만 쓸 수 있어.")
        if any(o["key"] == key for o in out):
            raise ValueError(f"{pos}번째 항목: 항목 키 '{key}'가 중복이야.")
        used.add(key)
        itype = (raw.get("input_type") or "").strip()
        if itype not in _ITEM_TYPES:
            raise ValueError(f"{pos}번째 항목: 입력 방식은 숫자·적합/부적합·사진 중 하나여야 해.")
        lower = upper = digits = None
        if itype == "numeric":
            lower = _item_num(raw.get("lower"), "하한", pos)
            upper = _item_num(raw.get("upper"), "상한", pos)
            if lower is None and upper is None:
                raise ValueError(f"{pos}번째 항목: 숫자 항목은 하한이나 상한 중 하나는 필요해.")
            if lower is not None and upper is not None and _Dec(lower) > _Dec(upper):
                raise ValueError(f"{pos}번째 항목: 하한이 상한보다 클 수 없어.")
            digits = _item_int(raw.get("display_digits"), "표시 소수 자릿수", pos, 0, 6, 2)
        am = raw.get("allow_missing")
        out.append({"key": key, "name": name, "input_type": itype, "lower": lower, "upper": upper,
                    "display_digits": digits,
                    "required_n": _item_int(raw.get("required_n"), "필요 개수", pos, 1, 50, 1),
                    "allow_missing": 1 if am in (1, True, "1", "on", "true") else 0,
                    "sort_order": pos})
    return out


def _miv_dict(row):
    d = dict(row)
    d["items"] = json.loads(d["items_json"])
    return d


def list_model_item_versions(model_id):
    """모델의 항목 버전 목록(최신순, items 파싱 포함)."""
    conn = get_conn()
    rows = conn.execute("SELECT * FROM product_model_item_versions WHERE model_id=? "
                        "ORDER BY version DESC", (model_id,)).fetchall()
    conn.close()
    return [_miv_dict(r) for r in rows]


def get_current_model_items(model_id):
    """모델의 최신 항목 버전(dict, items 파싱 포함). 없으면 None."""
    conn = get_conn()
    row = conn.execute("SELECT * FROM product_model_item_versions WHERE model_id=? "
                       "ORDER BY version DESC LIMIT 1", (model_id,)).fetchone()
    conn.close()
    return _miv_dict(row) if row else None


def copy_items_from_model(src_model_id):
    """다른 모델의 현재 항목 목록을 돌려준다(저장하지 않음 — 화면에서 고친 뒤 저장)."""
    cur = get_current_model_items(src_model_id)
    if not cur:
        raise ValueError("복사할 모델에 저장된 검사 항목이 없어.")
    return cur["items"]


@_write_retry
def save_model_items(model_id, items, reason, actor):
    """항목 목록을 새 버전으로 저장 → 버전 번호. 현재와 같으면 거부, 2번째 버전부터 사유 필수."""
    norm = _norm_items(items)
    reason = (reason or "").strip() or None
    conn = get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        if not conn.execute("SELECT 1 FROM product_models WHERE id=?", (model_id,)).fetchone():
            raise ValueError("모델을 찾을 수 없어.")
        cur = conn.execute("SELECT version, content_hash FROM product_model_item_versions "
                           "WHERE model_id=? ORDER BY version DESC LIMIT 1", (model_id,)).fetchone()
        h = _sha(norm)
        if cur and cur["content_hash"] == h:
            raise ValueError("변경 사항이 없어요.")
        if cur and not reason:
            raise ValueError("항목을 바꾸려면 사유를 입력해줘.")
        ver = (cur["version"] if cur else 0) + 1
        conn.execute("INSERT INTO product_model_item_versions (model_id, version, items_json, "
                     "content_hash, reason, created_by, created_at) VALUES (?,?,?,?,?,?,?)",
                     (model_id, ver, json.dumps(norm, ensure_ascii=False), h, reason, actor,
                      _utc_now()))
        conn.commit()
        return ver
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@_write_retry
def start_lot_inspection(lot_id, actor):
    """확정된 LOT의 제품검사 시작 — 그 시점의 모델·LOT·항목 버전을 스냅샷으로 굳힌다(시작 후 불변)."""
    conn = get_conn()
    try:
        row = _pl_lock(conn, lot_id)
        if row["status"] != "confirmed":
            raise ValueError("확정된 LOT만 제품검사를 시작할 수 있어요.")
        if conn.execute("SELECT 1 FROM production_lot_inspections WHERE lot_id=?",
                        (lot_id,)).fetchone():
            raise ValueError("이미 제품검사가 시작된 LOT이에요.")
        if not conn.execute("SELECT 1 FROM finished_units WHERE lot_id=? AND status='active' LIMIT 1",
                            (lot_id,)).fetchone():
            raise ValueError("정상 S/N이 1개 이상 있어야 시작할 수 있어요.")
        m = conn.execute("SELECT is_active FROM product_models WHERE id=?", (row["model_id"],)).fetchone()
        if not m or not m["is_active"]:
            raise ValueError("사용 중지된 모델이라 시작할 수 없어요.")
        v = conn.execute("SELECT * FROM product_model_item_versions WHERE model_id=? "
                         "ORDER BY version DESC LIMIT 1", (row["model_id"],)).fetchone()
        if not v:
            raise ValueError("모델의 검사 항목이 설정되어 있지 않아요.")
        model_snap = _pm_full(conn, row["model_id"])
        lot_snap = _pl_full(conn, lot_id)
        # 시작 당시 정상 S/N 목록 — unit_id 오름차순으로 동결(이후 추가·제외·정정과 분리된 검사 대상 기준)
        units_snap = []
        for ur in conn.execute("SELECT id FROM finished_units WHERE lot_id=? AND status='active' "
                               "ORDER BY id", (lot_id,)).fetchall():
            f = _fu_full(conn, ur["id"])
            units_snap.append({"unit_id": f["unit_id"], "serial_raw": f["serial_raw"],
                               "serial_key": f["serial_key"], "status": f["status"],
                               "model_mismatch": f["model_mismatch"]})
        started_at = _utc_now()
        h = _sha({"lot_id": lot_id, "started_by": actor, "started_at": started_at,
                  "model": model_snap, "lot": lot_snap, "item_version_id": v["id"],
                  "item_version": v["version"], "items": json.loads(v["items_json"]),
                  "units": units_snap})
        conn.execute("INSERT INTO production_lot_inspections (lot_id, started_by, started_at, "
                     "model_snapshot_json, lot_snapshot_json, item_version_id, item_version, "
                     "items_json, snapshot_hash, units_json) VALUES (?,?,?,?,?,?,?,?,?,?)",
                     (lot_id, actor, started_at, json.dumps(model_snap, ensure_ascii=False),
                      json.dumps(lot_snap, ensure_ascii=False), v["id"], v["version"],
                      v["items_json"], h, json.dumps(units_snap, ensure_ascii=False)))
        _pl_history(conn, lot_id, "inspection_start", lot_snap, lot_snap, None, actor)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def get_lot_inspection(lot_id):
    """LOT의 제품검사 시작 기록(스냅샷 JSON 파싱 포함). 시작 전이면 None."""
    conn = get_conn()
    row = conn.execute("SELECT * FROM production_lot_inspections WHERE lot_id=?",
                       (lot_id,)).fetchone()
    conn.close()
    if not row:
        return None
    d = dict(row)
    d["items"] = json.loads(d["items_json"])
    d["model_snapshot"] = json.loads(d["model_snapshot_json"])
    d["lot_snapshot"] = json.loads(d["lot_snapshot_json"])
    d["units"] = json.loads(d["units_json"] or "[]")
    return d


# ---------- 제품검사 측정값 입력·정정 (TASK #139 패키지 1D-b1) ----------
# 계약(승인·성적서 출력 등 이후 단계도 반드시 지킬 것)
#  - lot_inspection_results는 append-only: 이 테이블을 고치거나 지우는 구문·함수·UI를 만들지 않는다.
#    정정은 같은 칸(item_key, idx)에 rev+1 행을 추가하고, 현재값 = 칸의 MAX(rev).
#  - production_lot_inspections(시작 스냅샷·units_json·snapshot_hash)는 시작 후 절대 바꾸지 않는다.
#  - 검사 결과를 쓰는 모든 함수(입력·정정, 이후 승인·출력)는 _insp_begin()을 거쳐야 한다
#    (잠금 + 취소 LOT 거부 + 검사 시작 여부 확인). 읽기는 취소 후에도 허용.
#  - 판정은 정규화한 입력 원값(Decimal) 기준이고 display_digits 반올림은 쓰지 않는다.

class SlotError(ValueError):
    """칸 하나에 대한 거부 — 화면이 그 칸에 메시지를 띄울 수 있게 item_key·idx를 싣는다."""
    def __init__(self, msg, item_key=None, idx=None):
        super().__init__(msg)
        self.item_key, self.idx = item_key, idx


_MEASURE_WS = {0xA0: " ", 0x202F: " ", 0x2009: " "}
_THOUSANDS_RE = re.compile(r"[+-]?\d{1,3}(,\d{3})+(\.\d+)?", re.ASCII)
_NUM_RE = re.compile(r"[+-]?(\d+(\.\d*)?|\.\d+)", re.ASCII)


def parse_measure_number(raw):
    """측정값 문자열 → 표준 십진 문자열. 전각 숫자·U+2212는 반각으로(NFKC는 쓰지 않아 ¹ ① 등은 거부),
    NBSP류 공백은 일반 공백으로 바꿔 앞뒤만 지우고 숫자 중간 공백은 거부. 쉼표는 1,234.5 같은 천단위일 때만 허용."""
    s = str(raw if raw is not None else "").translate(_ITEM_NUM_MAP).translate(_MEASURE_WS).strip()
    if not s:
        raise ValueError("값을 입력해줘.")
    if re.search(r"\s", s):
        raise ValueError("숫자 중간에 공백이 있어요. 공백 없이 입력해줘.")
    if "," in s:
        if not _THOUSANDS_RE.fullmatch(s):
            raise ValueError("모호한 입력이야. 소수점은 점(.)으로 다시 입력해줘.")
        s = s.replace(",", "")
    if len(s) > 24 or not _NUM_RE.fullmatch(s):
        raise ValueError("숫자로 입력해줘. (24자 이하, 지수 표기·문자는 쓸 수 없어요)")
    x = _Dec(s)
    return "0" if x == 0 else format(x.normalize(), "f")


def judge_measure(item, value):
    """'pass'/'fail'/'nospec'. 숫자: 한계는 None 허용(단측), 경계값은 합격. 적합/부적합: OK→pass, NG→fail."""
    if item["input_type"] == "ok_ng":
        return "pass" if value == "OK" else "fail"
    lo, up = item.get("lower"), item.get("upper")
    if lo is None and up is None:
        return "nospec"
    v = _Dec(value)
    if lo is not None and v < _Dec(lo):
        return "fail"
    if up is not None and v > _Dec(up):
        return "fail"
    return "pass"


def summarize_unit_item(item, cur):
    """한 완제품·한 항목의 현재 상태. cur = {idx: 현재 행}. 저장하지 않고 조회 때 계산한다.
    fail / incomplete(빈 칸 있음) / missing_only(전부 결측 — 합격 아님) / nospec / pass."""
    rows = [cur[i] for i in range(1, item["required_n"] + 1) if i in cur]
    if item["input_type"] == "photo":  # 합격이 아니라 "필요 장수 모두 등록됨"일 뿐
        return "recorded" if len(rows) >= item["required_n"] else "incomplete"
    if any(r["verdict"] == "fail" for r in rows):
        return "fail"
    if len(rows) < item["required_n"]:
        return "incomplete"
    if all(r["is_missing"] for r in rows):
        return "missing_only"
    if any(r["verdict"] == "nospec" for r in rows):
        return "nospec"
    return "pass"


def product_photo_base(inspection_id, unit_id, item_key, idx):
    """제품검사 사진 저장 파일명 base(확장자 없음) — 파일명 규칙의 단일 출처."""
    return f"pi{inspection_id}_{unit_id}_{item_key}_{idx}_{uuid.uuid4().hex[:8]}"


def is_product_photo_name(name):
    return bool(re.fullmatch(r"pi\d+_\d+_[A-Za-z0-9_]{1,40}_\d+_[0-9a-f]{8}\.jpg", name or ""))


def _slot_current(conn, inspection_id, unit_id=None):
    """칸별 현재(최신 rev) 행 목록."""
    sql = ("SELECT r.* FROM lot_inspection_results r WHERE r.inspection_id=? "
           + ("AND r.unit_id=? " if unit_id is not None else "")
           + "AND r.rev = (SELECT MAX(x.rev) FROM lot_inspection_results x WHERE "
             "x.inspection_id=r.inspection_id AND x.unit_id=r.unit_id AND x.item_key=r.item_key "
             "AND x.idx=r.idx)")
    args = (inspection_id,) if unit_id is None else (inspection_id, unit_id)
    return conn.execute(sql, args).fetchall()


def _insp_begin(conn, lot_id):
    """검사 결과를 쓰는 모든 함수의 시작점 — 잠금 후 취소 LOT·검사 미시작 거부. (LOT행, 검사시작행)."""
    lot = _pl_lock(conn, lot_id)
    if lot["status"] == "cancelled":
        raise ValueError("취소된 LOT은 열람만 할 수 있어요.")
    insp = conn.execute("SELECT * FROM production_lot_inspections WHERE lot_id=?", (lot_id,)).fetchone()
    if not insp:
        raise ValueError("제품검사가 아직 시작되지 않았어요.")
    return lot, insp


def lot_inspection_writable(lot_id):
    """(쓸 수 있나, 안 되는 이유) — 화면·라우트의 사전 점검용(읽기 전용, 진짜 방어선은 _insp_begin)."""
    lot = get_production_lot(lot_id)
    if not lot:
        return False, "LOT을 찾을 수 없어요."
    if lot["status"] == "cancelled":
        return False, "취소된 LOT은 열람만 할 수 있어요."
    if not lot_has_started_inspection(lot_id):
        return False, "제품검사가 아직 시작되지 않았어요."
    return True, ""


def resolve_unit_for_inspection(lot_id, serial_text):
    """S/N → 완제품. result: ok / past_alias(과거 S/N — 조용히 연결하지 않고 현재 S/N 안내) /
    other_lot / not_found / excluded. 읽기 전용."""
    key = normalize_serial_key(serial_text)
    out = {"result": "not_found", "unit_id": None, "current_serial": None, "matched_serial": None}
    conn = get_conn()
    try:
        a = conn.execute("SELECT a.unit_id, a.kind, a.serial_raw, u.lot_id, u.status FROM unit_serial_aliases a "
                         "JOIN finished_units u ON u.id=a.unit_id WHERE a.serial_key=?", (key,)).fetchone() if key else None
        if not a:
            return out
        cs = conn.execute("SELECT serial_raw FROM unit_serial_aliases WHERE unit_id=? AND kind='current'",
                          (a["unit_id"],)).fetchone()
        out.update(unit_id=a["unit_id"], matched_serial=a["serial_raw"],
                   current_serial=cs["serial_raw"] if cs else None)
        if a["lot_id"] != lot_id:
            out["result"] = "other_lot"
        elif a["kind"] == "past":
            out["result"] = "past_alias"
        elif a["status"] != "active":
            out["result"] = "excluded"
        else:
            out["result"] = "ok"
        return out
    finally:
        conn.close()


def _roster_ids(insp):
    return {u["unit_id"] for u in json.loads(insp["units_json"] or "[]")}


def list_inspection_units(lot_id):
    """LOT 완제품마다 시작 시 포함 여부와 항목별 진행·전체 판정을 계산해서 돌려준다(검사 시작 전이면 [])."""
    insp = get_lot_inspection(lot_id)
    if not insp:
        return []
    items = insp["items"]
    roster = {u["unit_id"] for u in insp["units"]}
    conn = get_conn()
    try:
        cur = {}
        for r in _slot_current(conn, insp["id"]):
            cur.setdefault(r["unit_id"], {}).setdefault(r["item_key"], {})[r["idx"]] = r
        latest = _unit_events_latest(conn, insp["id"])
        disposal = _unit_disposal_latest(conn, insp["id"])
    finally:
        conn.close()
    out = []
    for u in get_production_lot_units(lot_id):
        ucur = cur.get(u["id"], {})
        prog = []
        for it in items:
            c = ucur.get(it["key"], {})
            prog.append({"key": it["key"], "name": it["name"], "filled": len(c),
                         "required": it["required_n"], "state": summarize_unit_item(it, c)})
        states = [p["state"] for p in prog]
        overall = ("fail" if "fail" in states else
                   "incomplete" if (not states or "incomplete" in states) else "all_filled")
        ev = latest.get(u["id"])
        # ready = 모든 항목의 칸이 다 찼음(fail이 있어도 칸이 차 있으면 입력 끝)
        all_filled = bool(items) and all(p["filled"] >= p["required"] for p in prog)
        if ev and ev["event_type"] == "complete":
            stage = "completed"
        elif all_filled:
            stage = "ready"
        elif ev or ucur:
            stage = "in_progress"
        else:
            stage = "not_started"
        has_fail = "fail" in states
        completion = None
        if ev:
            completion = {"event_type": ev["event_type"], "source": ev["source"], "reason": ev["reason"],
                          "at": ev["at"], "by_name": ev["by_name"], "seq": ev["seq"],
                          "results_digest": ev["results_digest"], "reopened": ev["event_type"] == "reopen"}
        out.append({"unit_id": u["id"], "serial_raw": u["serial_raw"], "past_serials": u["past_serials"],
                    "status": u["status"], "model_mismatch": bool(u["model_mismatch"]),
                    "roster_member": u["id"] in roster, "progress": prog, "overall": overall,
                    "stage": stage, "completion": completion, "has_fail": has_fail,
                    "unusable": bool(disposal.get(u["id"]) and disposal[u["id"]]["event_type"] == "unusable")})
    return out


def get_unit_results(inspection_id, unit_id):
    """{item_key: {idx: {"current": 행, "history": [rev 오름차순 전체 행]}}}"""
    conn = get_conn()
    try:
        rows = conn.execute("SELECT * FROM lot_inspection_results WHERE inspection_id=? AND unit_id=? "
                            "ORDER BY item_key, idx, rev", (inspection_id, unit_id)).fetchall()
    finally:
        conn.close()
    out = {}
    for r in rows:
        slot = out.setdefault(r["item_key"], {}).setdefault(r["idx"], {"current": None, "history": []})
        slot["history"].append(dict(r))
        slot["current"] = dict(r)
    return out


@_write_retry
def save_unit_results(lot_id, unit_id, entries, actor, can_correct_others=False, actor_id=None,
                      photo_files=None, can_correct_completed=False):
    """측정값 최초 입력·정정 저장 → 저장 행 수. 한 칸이라도 거부되면 전체 rollback.
    actor=표시명(감사용 스냅샷), actor_id=계정 id("본인" 판정 기준; None이면 본인 불인정).
    entries=[{item_key, idx, raw, base_rev(None=최초 입력), reason, missing}].
    photo_files={(item_key, idx): (저장파일명, 원본파일명, sha256)} — 사진 업로드 라우트만 넘긴다
    (일괄 저장 라우트는 안 넘기므로 photo 칸은 거부 → 클라이언트가 임의 파일명을 끼워 넣을 수 없다).
    판정·저장 본문은 _save_results_conn 하나(NG 재검사도 같은 본문을 쓴다)."""
    if not entries:
        raise ValueError("저장할 내용이 없어요.")
    conn = get_conn()
    try:
        lot, insp = _insp_begin(conn, lot_id)
        n = _save_results_conn(conn, lot, insp, unit_id, entries, actor, actor_id, can_correct_others,
                               photo_files, can_correct_completed)
        conn.commit()
        return n
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def _save_results_conn(conn, lot, insp, unit_id, entries, actor, actor_id, can_correct_others=False,
                       photo_files=None, can_correct_completed=False, retest_ctx=None):
    """save_unit_results의 본문 — 이미 _insp_begin으로 잠긴 트랜잭션 안에서만 부를 것(commit은 호출한 쪽).
    retest_ctx={"reason": ...}면 NG 재검사 저장: entry_type='retest', reason 자동, 본인/타인 정정 권한·변경없음 검사 생략,
    사진·base_rev 없는 입력 거부, 완료 S/N 재개는 호출자가 한다(여기서는 하지 않는다)."""
    photo_files = photo_files or {}
    lot_id = insp["lot_id"]
    u = conn.execute("SELECT lot_id, status FROM finished_units WHERE id=?", (unit_id,)).fetchone()
    if not u or u["lot_id"] != lot_id:
        raise ValueError("이 LOT의 완제품이 아니에요.")
    if u["status"] != "active":
        raise ValueError("제외된 완제품은 입력/정정할 수 없어요.")
    if unit_is_unusable(unit_id, conn):
        raise ValueError(_UNUSABLE_UNIT_MSG)
    last_ev =conn.execute("SELECT event_type FROM lot_inspection_unit_events WHERE inspection_id=? "
                           "AND unit_id=? ORDER BY seq DESC LIMIT 1", (insp["id"], unit_id)).fetchone()
    was_done = bool(last_ev and last_ev["event_type"] == "complete")
    if retest_ctx:  # 재검사는 완료 S/N도 받는다(재개는 호출자) — 단, 최종 승인된 S/N은 불가
        if unit_has_final_document(unit_id, conn):
            raise ValueError(_FINAL_DOC_UNIT_MSG)
        was_done = False
    if was_done:  # 완료된 S/N은 정정 권한자만, 정정하면 같은 트랜잭션에서 자동 재개
        if unit_has_final_document(unit_id, conn):
            raise ValueError(_FINAL_DOC_UNIT_MSG)
        if not can_correct_completed:
            raise ValueError("완료된 S/N이에요. 값을 고치려면 정정 권한이 필요해요.")
    items = {i["key"]: i for i in json.loads(insp["items_json"])}
    roster = 1 if unit_id in _roster_ids(insp) else 0
    sa = conn.execute("SELECT serial_raw FROM unit_serial_aliases WHERE unit_id=? AND kind='current'",
                      (unit_id,)).fetchone()
    serial = sa["serial_raw"] if sa else None
    seen = set()
    for e in entries:
        key = e.get("item_key")
        try:
            idx = int(e.get("idx"))
        except (TypeError, ValueError):
            raise SlotError("칸 번호가 올바르지 않아요.", key)
        it = items.get(key)
        if not it:
            raise SlotError("알 수 없는 검사 항목이에요.", key, idx)
        is_photo = it["input_type"] == "photo"
        if is_photo and (key, idx) not in photo_files:
            raise SlotError("사진은 사진 업로드 버튼으로만 등록할 수 있어요.", key, idx)
        if not is_photo and (key, idx) in photo_files:
            raise SlotError("사진 항목이 아니에요.", key, idx)
        if retest_ctx and is_photo:
            raise SlotError("사진 항목은 재검사로 입력할 수 없어요.", key, idx)
        if not 1 <= idx <= it["required_n"]:
            raise SlotError("칸 번호가 범위를 벗어났어요.", key, idx)
        if (key, idx) in seen:
            raise SlotError("같은 칸이 한 번에 두 번 들어왔어요.", key, idx)
        seen.add((key, idx))
        base = e.get("base_rev")
        base = None if base in (None, "") else base
        if base is not None:
            try:
                base = int(base)
            except (TypeError, ValueError):
                raise SlotError("회차 정보가 올바르지 않아요.", key, idx)
        if retest_ctx and base is None:
            raise SlotError("재검사는 기존 값이 있는 칸만 입력할 수 있어요.", key, idx)
        if was_done and base is None:
            raise SlotError("완료된 S/N은 새로 입력할 수 없어요. 먼저 완료를 취소(재개)해줘.", key, idx)
        cur = conn.execute("SELECT * FROM lot_inspection_results WHERE inspection_id=? AND unit_id=? "
                           "AND item_key=? AND idx=? ORDER BY rev DESC LIMIT 1",
                           (insp["id"], unit_id, key, idx)).fetchone()
        reason = (e.get("reason") or "").strip() or None
        if base is None:
            if cur:
                raise SlotError("다른 사람이 이미 입력했어요. 정정으로 다시 입력해줘.", key, idx)
            rev, etype, reason = 1, "initial", None
        else:
            if not cur:
                raise SlotError("정정할 기존 값이 없어요.", key, idx)
            if base != cur["rev"]:
                raise SlotError(f"그 사이에 정정됐어요(현재 {cur['rev']}회차). 새로고침해줘.", key, idx)
            if retest_ctx:
                reason, rev, etype = retest_ctx["reason"], cur["rev"] + 1, "retest"
            else:
                if not reason:
                    raise SlotError("정정 사유를 입력해줘.", key, idx)
                is_own = actor_id is not None and cur["entered_by_user_id"] == actor_id
                if not is_own and not can_correct_others:
                    raise SlotError("다른 사람이 입력한 값은 정정 권한이 있어야 고칠 수 있어요.", key, idx)
                rev, etype = cur["rev"] + 1, "correction"
        sha = None
        try:
            if is_photo:
                if e.get("missing"):
                    raise ValueError("사진은 결측으로 입력할 수 없어요.")
                val, vin, sha = photo_files[(key, idx)]
                if not re.fullmatch(rf"pi{insp['id']}_{unit_id}_{re.escape(key)}_{idx}_[0-9a-f]{{8}}\.jpg", val or ""):
                    raise ValueError("사진 파일 정보가 올바르지 않아요.")
                miss, verdict = 0, "recorded"
            elif e.get("missing"):
                if not it["allow_missing"]:
                    raise ValueError("이 항목은 결측으로 입력할 수 없어요.")
                vin = val = None
                miss, verdict = 1, "missing"
            else:
                vin = str(e.get("raw") if e.get("raw") is not None else "").strip()
                if it["input_type"] == "ok_ng":
                    val = vin.upper()
                    if val not in ("OK", "NG"):
                        raise ValueError("적합 또는 부적합을 선택해줘.")
                else:
                    val = parse_measure_number(vin)
                miss, verdict = 0, judge_measure(it, val)
        except SlotError:
            raise
        except ValueError as ve:
            raise SlotError(str(ve), key, idx)
        if not retest_ctx and not is_photo and cur and cur["is_missing"] == miss and cur["value"] == val:
            raise SlotError("변경 없음 — 현재 값과 같아요.", key, idx)
        try:
            conn.execute(
                "INSERT INTO lot_inspection_results (inspection_id, unit_id, item_key, idx, rev, "
                "entry_type, value_input, value, is_missing, verdict, reason, roster_member, "
                "serial_at_entry, lot_revision, entered_by, entered_by_user_id, entered_at, file_sha256) "
                "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (insp["id"], unit_id, key, idx, rev, etype, vin, val, miss, verdict, reason,
                 roster, serial, lot["lot_revision"], actor, actor_id, _utc_now(), sha))
        except sqlite3.IntegrityError as ie:
            if "lot_inspection_results" not in str(ie) or "UNIQUE" not in str(ie):
                raise
            raise SlotError("다른 사람이 먼저 입력했어요. 새로고침해줘.", key, idx)
    if was_done:
        reasons = list(dict.fromkeys((e.get("reason") or "").strip() for e in entries))
        why = "; ".join(r for r in reasons if r)[:200] or "값 정정"
        _append_unit_event(conn, insp, unit_id, "reopen", "correction", why, actor, actor_id)
    return len(entries)


def verify_lot_inspection_snapshot(lot_id):
    """저장된 컬럼에서 시작 스냅샷 해시를 다시 계산해 snapshot_hash와 비교(시작 후 불변 확인)."""
    insp = get_lot_inspection(lot_id)
    if not insp:
        return {"ok": False, "error": "검사가 시작되지 않았어요."}
    h = _sha({"lot_id": lot_id, "started_by": insp["started_by"], "started_at": insp["started_at"],
              "model": insp["model_snapshot"], "lot": insp["lot_snapshot"],
              "item_version_id": insp["item_version_id"], "item_version": insp["item_version"],
              "items": insp["items"], "units": insp["units"]})
    return {"ok": h == insp["snapshot_hash"], "expected": insp["snapshot_hash"], "actual": h}


def list_post_start_units(lot_id):
    """시작 후 합류한 완제품(시작 명단에 없는 것)과 시작 이후의 추가·복구 이력."""
    insp = get_lot_inspection(lot_id)
    if not insp:
        return {"units": [], "events": []}
    roster = {u["unit_id"] for u in insp["units"]}
    units = [u for u in get_production_lot_units(lot_id) if u["id"] not in roster]
    conn = get_conn()
    try:
        ev = [dict(r) for r in conn.execute(
            "SELECT * FROM finished_unit_history WHERE lot_id=? AND change_type IN ('add','restore') "
            "AND changed_at >= ? ORDER BY id", (lot_id, insp["started_at"]))]
    finally:
        conn.close()
    return {"units": units, "events": ev}


# ---------- 제품검사 S/N 완료·재개 (TASK #145 패키지 1D-c1) ----------
# 계약: lot_inspection_unit_events는 append-only(UPDATE/DELETE 금지). 완료 = "검사 입력 끝"이지 합격이 아니다.
# 완료 요청은 '대기'를 저장하지 않고 한 트랜잭션에서 전 조건을 다시 검증해 성공하면 complete 이벤트만 만든다.

class CompletionBlocked(ValueError):
    """완료 거부 — blockers=[{code, item_key, name, msg}, ...]."""
    def __init__(self, blockers):
        self.blockers = blockers
        super().__init__(blockers[0]["msg"] if blockers else "완료할 수 없어요.")


def _blk(code, msg, item_key=None, name=None):
    return {"code": code, "item_key": item_key, "name": name, "msg": msg}


def _digest_rows(rows):
    """현재 결과 행 목록 → digest (슬롯 정렬 후 sha256). 완료 이벤트·4번 단계가 같은 값을 쓴다."""
    lst = sorted(([r["item_key"], r["idx"], r["rev"], r["value"], r["is_missing"], r["file_sha256"]]
                  for r in rows), key=lambda x: (x[0], x[1]))
    return _sha(lst)


def _results_digest(conn, inspection_id, unit_id):
    return _digest_rows(_slot_current(conn, inspection_id, unit_id))


def digest_from_results(results):
    """get_unit_results()가 돌려준 값(화면에 그린 것)에서 같은 digest를 계산 — 화면과 digest가 어긋나지 않게."""
    return _digest_rows([s["current"] for slots in results.values() for s in slots.values()])


def unit_results_digest(lot_id, unit_id):
    insp = get_lot_inspection(lot_id)
    if not insp:
        return None
    conn = get_conn()
    try:
        return _results_digest(conn, insp["id"], unit_id)
    finally:
        conn.close()


def _unit_events_latest(conn, inspection_id):
    """{unit_id: 최신(MAX seq) 이벤트 행}"""
    rows = conn.execute(
        "SELECT e.* FROM lot_inspection_unit_events e WHERE e.inspection_id=? AND e.seq = "
        "(SELECT MAX(x.seq) FROM lot_inspection_unit_events x WHERE x.inspection_id=e.inspection_id "
        "AND x.unit_id=e.unit_id)", (inspection_id,)).fetchall()
    return {r["unit_id"]: r for r in rows}


def unit_completion_state(lot_id, unit_id):
    """list_inspection_units 한 건(stage·completion·has_fail 포함). 없으면 None."""
    return next((u for u in list_inspection_units(lot_id) if u["unit_id"] == unit_id), None)


def list_unit_events(inspection_id, unit_id):
    conn = get_conn()
    try:
        return [dict(r) for r in conn.execute(
            "SELECT * FROM lot_inspection_unit_events WHERE inspection_id=? AND unit_id=? ORDER BY seq",
            (inspection_id, unit_id))]
    finally:
        conn.close()


def unit_photo_rows(lot_id, unit_id):
    """현재 사진 행 [{item_key, name, idx, value, file_sha256}] — 완료 요청 전 파일 확인용(읽기)."""
    insp = get_lot_inspection(lot_id)
    if not insp:
        return []
    photo = {i["key"]: i["name"] for i in insp["items"] if i["input_type"] == "photo"}
    conn = get_conn()
    try:
        rows = _slot_current(conn, insp["id"], unit_id)
    finally:
        conn.close()
    return [{"item_key": r["item_key"], "name": photo[r["item_key"]], "idx": r["idx"], "value": r["value"],
             "file_sha256": r["file_sha256"]} for r in rows if r["item_key"] in photo]


def _unit_progress(items, cur):
    """항목별 {key,name,filled,required,state,fail_count,missing_count}. cur={item_key:{idx:행}}."""
    out = []
    for it in items:
        c = cur.get(it["key"], {})
        rows = [c[i] for i in range(1, it["required_n"] + 1) if i in c]
        out.append({"key": it["key"], "name": it["name"], "filled": len(rows), "required": it["required_n"],
                    "state": summarize_unit_item(it, c),
                    "fail_count": sum(1 for r in rows if r["verdict"] == "fail"),
                    "missing_count": sum(1 for r in rows if r["is_missing"])})
    return out


def _append_unit_event(conn, insp, unit_id, event_type, source, reason, actor, actor_id, summary=None):
    """이벤트 한 줄 추가. 이미 _insp_begin으로 잠긴 트랜잭션 안에서만 부를 것(commit은 호출한 쪽)."""
    seq = conn.execute("SELECT COALESCE(MAX(seq),0)+1 FROM lot_inspection_unit_events WHERE inspection_id=? "
                       "AND unit_id=?", (insp["id"], unit_id)).fetchone()[0]
    lot = conn.execute("SELECT lot_revision FROM production_lots WHERE id=?", (insp["lot_id"],)).fetchone()
    sa = conn.execute("SELECT serial_raw FROM unit_serial_aliases WHERE unit_id=? AND kind='current'",
                      (unit_id,)).fetchone()
    conn.execute(
        "INSERT INTO lot_inspection_unit_events (inspection_id, unit_id, seq, event_type, source, reason, "
        "results_digest, summary_json, roster_member, serial_at_event, lot_revision, by_name, by_user_id, at) "
        "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (insp["id"], unit_id, seq, event_type, source, reason, _results_digest(conn, insp["id"], unit_id),
         _canon(summary) if summary is not None else None, 1 if unit_id in _roster_ids(insp) else 0,
         sa["serial_raw"] if sa else None, lot["lot_revision"], actor, actor_id, _utc_now()))
    if event_type == "reopen":  # 승인 대기 중인 제품 성적서는 재개(사용자·정정·시스템 모두) 때 자동 철회 — 한 지점
        _withdraw_pending_report(conn, insp, unit_id, reason, source, actor, actor_id)
    return seq


def _unit_in_lot(conn, lot_id, unit_id):
    u = conn.execute("SELECT lot_id, status, model_mismatch FROM finished_units WHERE id=?", (unit_id,)).fetchone()
    if not u or u["lot_id"] != lot_id:
        raise ValueError("이 LOT의 완제품이 아니에요.")
    return u


@_write_retry
def complete_unit_inspection(lot_id, unit_id, actor, actor_id, expected_digest, photo_blockers=None):
    """S/N 검사 완료(=입력 끝) 기록. 차단 사유가 하나라도 있으면 CompletionBlocked(.blockers) — 입력은 그대로.
    NG(fail)·전부 결측·모델 불일치는 차단하지 않고 summary에만 남긴다. photo_blockers=라우트가 락 밖에서 확인한 파일 문제."""
    conn = get_conn()
    try:
        lot, insp = _insp_begin(conn, lot_id)
        u = _unit_in_lot(conn, lot_id, unit_id)
        if u["status"] != "active":
            raise CompletionBlocked([_blk("unit_excluded", "제외된 완제품은 완료할 수 없어요.")])
        if unit_is_unusable(unit_id, conn):
            raise CompletionBlocked([_blk("unit_unusable", _UNUSABLE_UNIT_MSG)])
        last =conn.execute("SELECT event_type FROM lot_inspection_unit_events WHERE inspection_id=? AND unit_id=? "
                            "ORDER BY seq DESC LIMIT 1", (insp["id"], unit_id)).fetchone()
        if last and last["event_type"] == "complete":
            raise CompletionBlocked([_blk("already_completed", "이미 완료됐어요.")])
        cur = {}
        for r in _slot_current(conn, insp["id"], unit_id):
            cur.setdefault(r["item_key"], {})[r["idx"]] = r
        items = json.loads(insp["items_json"])
        prog = _unit_progress(items, cur)
        blockers = []
        for it, p in zip(items, prog):
            if p["filled"] < p["required"]:
                unit_word = "장" if it["input_type"] == "photo" else "개"
                blockers.append(_blk("incomplete", f"{p['name']}: {p['required']}{unit_word} 중 {p['filled']}{unit_word}만 "
                                     "입력됐어요.", p["key"], p["name"]))
            elif p["state"] == "nospec":
                blockers.append(_blk("nospec", f"{p['name']}: 규격이 없어서 완료할 수 없어요.", p["key"], p["name"]))
        digest = _results_digest(conn, insp["id"], unit_id)
        if expected_digest != digest:
            blockers.insert(0, _blk("stale", "화면을 본 뒤에 값이 바뀌었어요. 새로고침해서 다시 확인해줘."))
        blockers.extend(photo_blockers or [])
        # locked_by_other: 항목 잠금(1D-c2)에서 추가 예정
        if blockers:
            raise CompletionBlocked(blockers)
        summary = {"items": prog,
                   "has_fail": any(p["state"] == "fail" for p in prog),
                   "fail_items": [p["key"] for p in prog if p["fail_count"]],
                   "missing_only_items": [p["key"] for p in prog if p["state"] == "missing_only"],
                   "model_mismatch": bool(u["model_mismatch"])}
        ng_ids = _ng_open_missing(conn, insp, unit_id, actor, actor_id, "completion")  # NG 항목별 사건(멱등)
        try:
            seq = _append_unit_event(conn, insp, unit_id, "complete", "user", None, actor, actor_id, summary)
        except sqlite3.IntegrityError as ie:
            if "lot_inspection_unit_events" not in str(ie) or "UNIQUE" not in str(ie):
                raise
            raise ValueError("다른 사람이 먼저 처리했어요. 새로고침해줘.")
        conn.commit()
        return {"seq": seq, "digest": digest, "has_fail": summary["has_fail"], "summary": summary,
                "ng_event_ids": ng_ids}
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@_write_retry
def reopen_unit_completion(lot_id, unit_id, reason, actor, actor_id, source="user"):
    """완료 취소(재개) — 사유 필수. 값 정정에 따른 자동 재개(source='correction')는 save_unit_results가 한다."""
    reason = (reason or "").strip()
    if not reason:
        raise ValueError("완료를 취소하는 사유를 입력해줘.")
    conn = get_conn()
    try:
        lot, insp = _insp_begin(conn, lot_id)
        u = _unit_in_lot(conn, lot_id, unit_id)
        if u["status"] != "active":
            raise ValueError("제외된 완제품은 완료를 취소할 수 없어요.")
        last = conn.execute("SELECT event_type FROM lot_inspection_unit_events WHERE inspection_id=? AND unit_id=? "
                            "ORDER BY seq DESC LIMIT 1", (insp["id"], unit_id)).fetchone()
        if not last or last["event_type"] != "complete":
            raise ValueError("완료된 S/N이 아니에요.")
        if unit_has_final_document(unit_id, conn):
            raise ValueError(_FINAL_DOC_UNIT_MSG)
        if unit_is_unusable(unit_id, conn):
            raise ValueError(_UNUSABLE_UNIT_MSG)
        seq = _append_unit_event(conn, insp, unit_id, "reopen", source, reason[:200], actor, actor_id)
        conn.commit()
        return seq
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def unit_completion_valid(lot_id, unit_id):
    """(유효한가, 이유) — LOT이 취소되지 않았고, S/N이 active이며, 최신 이벤트가 complete이고 현재 digest가 그 이벤트의 digest와 같을 때만 참(4번 단계 안전망).
    제외·LOT 취소는 이벤트·digest를 바꾸지 않으므로 여기서 따로 확인한다(제외 중엔 무효, 복구하면 기존 완료가 다시 유효)."""
    insp = get_lot_inspection(lot_id)
    if not insp:
        return False, "제품검사가 시작되지 않았어요."
    conn = get_conn()
    try:
        return _unit_completion_valid_conn(conn, insp, unit_id)
    finally:
        conn.close()


def _unit_completion_valid_conn(conn, insp, unit_id):
    """unit_completion_valid의 본문 — 호출자가 연 conn(트랜잭션 안)에서 같은 판정을 한다. insp는 id·lot_id를 가진 행/dict."""
    lot_id = insp["lot_id"]
    lot = conn.execute("SELECT status FROM production_lots WHERE id=?", (lot_id,)).fetchone()
    if not lot or lot["status"] == "cancelled":
        return False, "LOT이 취소됐어요."
    u = conn.execute("SELECT status FROM finished_units WHERE id=? AND lot_id=?", (unit_id, lot_id)).fetchone()
    if not u or u["status"] != "active":
        return False, "제외된 S/N이에요."
    if unit_is_unusable(unit_id, conn):
        return False, "사용 금지로 결정된 S/N이에요."
    ev = _unit_events_latest(conn, insp["id"]).get(unit_id)
    if not ev or ev["event_type"] != "complete":
        return False, "완료되지 않은 S/N이에요."
    if _results_digest(conn, insp["id"], unit_id) != ev["results_digest"]:
        return False, "완료 이후 값이 바뀌었어요."
    return True, ""


# ---------- NG 사건·후속조치·재검사 (TASK #150 패키지 4b-1) ----------
# 계약: ng_events(불변 헤더)·ng_event_log(해시 체인)는 append-only — UPDATE/DELETE 금지.
# 사건 단위 = (S/N, 검사 항목). 미종결 사건은 같은 (S/N, 항목)에 동시에 1개(_insp_begin 락 안 코드로 강제).
# 상태는 저장하지 않고 로그에서 파생: open / special_requested / closed_retest / closed_special.
# 모든 쓰기는 _insp_begin(BEGIN IMMEDIATE) 안에서 소속·상태·최종문서 여부를 다시 확인한다.

class NgBlocked(ValueError):
    """NG 사건 관련 거부 — blockers=[{code, item_key, name, msg}, ...]."""
    def __init__(self, blockers):
        self.blockers = blockers
        super().__init__(blockers[0]["msg"] if blockers else "처리할 수 없어요.")


_NG_ACTIONS = ("retest", "repair_retest", "reassemble", "special_request", "other")
_NG_LOG_FIELDS = ("action_kind", "round_no", "outcome", "close_kind", "reason", "memo", "revision_id",
                  "report_event_seq")


def _ng_h(s):
    return hashlib.sha256(s.encode("utf-8")).hexdigest() if s else None


def _ng_header_hash(inspection_id, unit_id, item_key, event_no, origin_event_id, origin_round_no, opened_trigger,
                    fail_snapshot_json, serial_at_open, lot_revision, opened_by, opened_by_user_id, opened_at):
    """사건 헤더 해시 — 열 때 한 번 계산해 저장하고 verify가 같은 식으로 다시 계산한다."""
    return _sha({"inspection_id": inspection_id, "unit_id": unit_id, "item_key": item_key, "event_no": event_no,
                 "origin_event_id": origin_event_id, "origin_round_no": origin_round_no,
                 "opened_trigger": opened_trigger, "snapshot_hash": _ng_h(fail_snapshot_json),
                 "serial_at_open": serial_at_open, "lot_revision": lot_revision, "opened_by": opened_by,
                 "opened_by_user_id": opened_by_user_id, "opened_at": opened_at})


def _ng_log_hash(header_hash, ng_event_id, seq, log_type, action_kind, round_no, outcome, close_kind, reason, memo,
                 revision_id, report_event_seq, payload_json, actor_name, actor_user_id, at, prev_hash):
    """로그 한 줄 해시의 단일 출처 — 쓰기와 verify가 같은 식을 쓴다. header_hash를 매 줄에 넣어 헤더 변조도 잡는다."""
    return _sha({"header_hash": header_hash, "ng_event_id": ng_event_id, "seq": seq, "log_type": log_type,
                 "action_kind": action_kind, "round_no": round_no, "outcome": outcome, "close_kind": close_kind,
                 "reason_hash": _ng_h(reason), "memo_hash": _ng_h(memo), "revision_id": revision_id,
                 "report_event_seq": report_event_seq, "payload_hash": _ng_h(payload_json),
                 "actor_name": actor_name, "actor_user_id": actor_user_id, "at": at, "prev_event_hash": prev_hash})


def _ng_log_append(conn, ng_event_id, log_type, actor, actor_id, payload=None, **fields):
    """사건 로그 한 줄 추가(seq·prev_hash·event_hash의 단일 출처). _insp_begin 잠금 안에서만 부를 것."""
    bad = set(fields) - set(_NG_LOG_FIELDS)
    if bad:
        raise ValueError(f"알 수 없는 로그 필드: {sorted(bad)}")
    hd = conn.execute("SELECT header_hash FROM ng_events WHERE id=?", (ng_event_id,)).fetchone()
    if not hd:
        raise ValueError("NG 사건을 찾을 수 없어요.")
    last = conn.execute("SELECT seq, event_hash FROM ng_event_log WHERE ng_event_id=? ORDER BY seq DESC LIMIT 1",
                        (ng_event_id,)).fetchone()
    seq = (last["seq"] if last else 0) + 1
    prev = last["event_hash"] if last else None
    f = {k: fields.get(k) for k in _NG_LOG_FIELDS}
    pj = _canon(payload) if payload is not None else None
    at = _utc_now()
    eh = _ng_log_hash(hd["header_hash"], ng_event_id, seq, log_type, f["action_kind"], f["round_no"], f["outcome"],
                      f["close_kind"], f["reason"], f["memo"], f["revision_id"], f["report_event_seq"], pj,
                      actor, actor_id, at, prev)
    conn.execute(
        "INSERT INTO ng_event_log (ng_event_id, seq, log_type, action_kind, round_no, outcome, close_kind, reason, "
        "memo, revision_id, report_event_seq, payload_json, actor_name, actor_user_id, prev_hash, event_hash, at) "
        "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (ng_event_id, seq, log_type, f["action_kind"], f["round_no"], f["outcome"], f["close_kind"], f["reason"],
         f["memo"], f["revision_id"], f["report_event_seq"], pj, actor, actor_id, prev, eh, at))
    return seq


def _ng_state(logs, rev_last):
    """로그 목록(seq 순) → 상태. 특채 종결은 그 revision의 마지막 이벤트가 approved인 동안만 유효.
    취소(cancelled)는 마지막 줄일 때만 — 취소 뒤엔 쓰기가 막혀 항상 마지막 줄이다."""
    if logs and logs[-1]["log_type"] == "cancelled":
        return "cancelled"
    closed = [l for l in logs if l["log_type"] == "closed"]
    if closed:
        c = closed[-1]
        if c["close_kind"] == "retest_pass":
            return "closed_retest"
        if c["close_kind"] == "special" and rev_last.get(c["revision_id"]) == "approved":
            return "closed_special"
    acts = [l for l in logs if l["log_type"] in ("action", "retest_round")]
    if acts and acts[-1]["log_type"] == "action" and acts[-1]["action_kind"] == "special_request":
        return "special_requested"
    return "open"


def _ng_events_derived(conn, inspection_id, unit_id=None):
    """사건 목록 + 로그 + 파생 상태(쿼리 3개 고정 — 건수와 무관, N+1 없음)."""
    w, a = "e.inspection_id=?", [inspection_id]
    if unit_id is not None:
        w += " AND e.unit_id=?"
        a.append(unit_id)
    evs = [dict(r) for r in conn.execute(f"SELECT e.* FROM ng_events e WHERE {w} ORDER BY e.unit_id, e.event_no", a)]
    if not evs:
        return []
    logs = {}
    for r in conn.execute(f"SELECT l.* FROM ng_event_log l JOIN ng_events e ON e.id=l.ng_event_id WHERE {w} "
                          "ORDER BY l.ng_event_id, l.seq", a):
        logs.setdefault(r["ng_event_id"], []).append(dict(r))
    rev_ids = sorted({l["revision_id"] for ls in logs.values() for l in ls
                      if l["log_type"] == "closed" and l["close_kind"] == "special" and l["revision_id"]})
    rev_last = {}
    if rev_ids:
        ph = ",".join("?" * len(rev_ids))
        for r in conn.execute(
                f"SELECT e.revision_id, e.event_type FROM product_report_events e WHERE e.revision_id IN ({ph}) "
                "AND e.seq=(SELECT MAX(x.seq) FROM product_report_events x WHERE x.revision_id=e.revision_id)",
                rev_ids):
            rev_last[r["revision_id"]] = r["event_type"]
    for e in evs:
        ls = logs.get(e["id"], [])
        e["fail_snapshot"] = json.loads(e["fail_snapshot_json"])
        e["logs"] = ls
        e["round_no"] = sum(1 for l in ls if l["log_type"] == "retest_round")
        e["state"] = _ng_state(ls, rev_last)
    return evs


_NG_OPEN_STATES = ("open", "special_requested")


def _ng_cur(conn, inspection_id, unit_id):
    cur = {}
    for r in _slot_current(conn, inspection_id, unit_id):
        cur.setdefault(r["item_key"], {})[r["idx"]] = r
    return cur


def _ng_unit_gate(conn, insp, unit_id):
    """S/N의 NG 게이트(읽기). blockers=제출·승인을 막는 사유(ng_event_missing/ng_open/ng_unresolved),
    pending=아직 종결 안 된 사건들(특채 요청 포함). 완제품 성적서 제출·승인(T4)과 unit_ng_gate_clear가 쓴다."""
    items = json.loads(insp["items_json"])
    cur = _ng_cur(conn, insp["id"], unit_id)
    fail_keys = [it["key"] for it in items if summarize_unit_item(it, cur.get(it["key"], {})) == "fail"]
    names = {it["key"]: it["name"] for it in items}
    evs = _ng_events_derived(conn, insp["id"], unit_id)
    pending = {e["item_key"]: e for e in evs if e["state"] in _NG_OPEN_STATES}
    covered = {e["item_key"] for e in evs if e["state"] == "closed_special"}
    blockers = []
    for k in fail_keys:
        e = pending.get(k)
        if not e:
            if k not in covered:
                blockers.append(_blk("ng_event_missing", f"{names[k]}: NG 사건이 없어요. 먼저 NG 사건을 열어줘.", k, names[k]))
        elif e["state"] == "open":
            blockers.append(_blk("ng_open", f"{names[k]}: NG 후속조치(특채 요청)가 아직 없어요.", k, names[k]))
    for k, e in pending.items():
        if k not in fail_keys:
            blockers.append(_blk("ng_unresolved", f"{names.get(k, k)}: 값이 정정됐지만 NG 사건이 종결되지 않았어요.", k,
                                 names.get(k)))
    return {"fail_items": fail_keys, "events": evs, "pending": list(pending.values()), "blockers": blockers,
            "has_ng": bool(fail_keys or pending)}


def unit_ng_gate_clear(conn_or_none, unit_id):
    """(NG 때문에 막힘 없음, 사유 목록) — 5번 단계(출하 등)용 읽기 계약. 미종결 사건·사건 없는 NG가 없어야 참.
    특채 종결은 그 성적서가 승인된 동안만 유효(회수되면 자동 해제)."""
    own = conn_or_none is None
    conn = get_conn() if own else conn_or_none
    try:
        insp = conn.execute("SELECT i.* FROM production_lot_inspections i JOIN finished_units u ON u.lot_id=i.lot_id "
                            "WHERE u.id=?", (unit_id,)).fetchone()
        if not insp:
            return True, []
        if unit_is_unusable(unit_id, conn):
            return False, [_blk("unit_unusable", "사용 금지로 결정된 S/N이에요.")]
        g = _ng_unit_gate(conn, insp, unit_id)
        bl = list(g["blockers"])
        for e in g["pending"]:
            if e["state"] == "special_requested":
                bl.append(_blk("ng_special_requested", "특채 승인 전이에요.", e["item_key"]))
        return (not bl), bl
    finally:
        if own:
            conn.close()


def _ng_open_missing(conn, insp, unit_id, actor, actor_id, trigger, only_keys=None, origin=None):
    """현재 fail인 항목 중 미종결 사건이 없는 것마다 사건을 연다(멱등 — 있으면 건너뜀) → 새 사건 id 목록.
    _insp_begin 잠금 안에서만 부를 것. origin=(원 사건 id, 회차)."""
    items = json.loads(insp["items_json"])
    cur = _ng_cur(conn, insp["id"], unit_id)
    evs = _ng_events_derived(conn, insp["id"], unit_id)
    busy = {e["item_key"] for e in evs if e["state"] in _NG_OPEN_STATES}
    closed_sp = {e["item_key"] for e in evs if e["state"] == "closed_special"}
    last = conn.execute("SELECT seq, event_type FROM lot_inspection_unit_events WHERE inspection_id=? AND unit_id=? "
                        "ORDER BY seq DESC LIMIT 1", (insp["id"], unit_id)).fetchone()
    if trigger == "completion":
        comp_seq = (last["seq"] if last else 0) + 1  # 이어서 추가될 complete 이벤트의 seq
    else:
        comp_seq = last["seq"] if last and last["event_type"] == "complete" else None
    digest = _results_digest(conn, insp["id"], unit_id)
    sa = conn.execute("SELECT serial_raw FROM unit_serial_aliases WHERE unit_id=? AND kind='current'",
                      (unit_id,)).fetchone()
    serial = sa["serial_raw"] if sa else None
    lot_rev = conn.execute("SELECT lot_revision FROM production_lots WHERE id=?", (insp["lot_id"],)).fetchone()[0]
    ids = []
    for it in items:
        k = it["key"]
        c = cur.get(k, {})
        if (only_keys is not None and k not in only_keys) or k in busy or k in closed_sp:
            continue
        if summarize_unit_item(it, c) != "fail":
            continue
        cells = [{"idx": i, "rev": c[i]["rev"], "value": c[i]["value"], "verdict": c[i]["verdict"]}
                 for i in sorted(c) if c[i]["verdict"] == "fail"]
        snap = _canon({"cells": cells, "digest": digest, "completion_seq": comp_seq})
        no = conn.execute("SELECT COALESCE(MAX(event_no),0)+1 FROM ng_events WHERE inspection_id=? AND unit_id=?",
                          (insp["id"], unit_id)).fetchone()[0]
        at = _utc_now()
        oe, orn = origin if origin else (None, None)
        hh = _ng_header_hash(insp["id"], unit_id, k, no, oe, orn, trigger, snap, serial, lot_rev, actor, actor_id, at)
        cu = conn.execute(
            "INSERT INTO ng_events (inspection_id, unit_id, item_key, event_no, origin_event_id, origin_round_no, "
            "opened_trigger, fail_snapshot_json, header_hash, serial_at_open, lot_revision, opened_by, "
            "opened_by_user_id, opened_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (insp["id"], unit_id, k, no, oe, orn, trigger, snap, hh, serial, lot_rev, actor, actor_id, at))
        ids.append(cu.lastrowid)
        busy.add(k)
    return ids


def _ng_writable_event(conn, insp, unit_id, ng_event_id):
    """NG 쓰기 공통 점검(잠금 안) → 파생된 사건 dict. 소속 LOT·active S/N·최종문서 없음·사건 소속·미종결."""
    u = _unit_in_lot(conn, insp["lot_id"], unit_id)
    if u["status"] != "active":
        raise ValueError("제외된 완제품은 NG 처리를 할 수 없어요.")
    if unit_has_final_document(unit_id, conn):
        raise ValueError(_FINAL_DOC_UNIT_MSG)
    if unit_is_unusable(unit_id, conn):
        raise ValueError(_UNUSABLE_UNIT_MSG)
    ev = next((e for e in _ng_events_derived(conn, insp["id"], unit_id) if e["id"] == ng_event_id), None)
    if not ev:
        raise ValueError("이 S/N의 NG 사건이 아니에요.")
    if ev["state"].startswith("closed") or ev["state"] == "cancelled":
        raise ValueError("이미 종결되거나 취소된 NG 사건이에요.")
    return ev


@_write_retry
def open_ng_events(lot_id, unit_id, actor, actor_id):
    """완료된 S/N에서 사건 없는 NG 항목의 사건을 수동으로 연다(레거시·예외용) → 새 사건 id 목록."""
    conn = get_conn()
    try:
        lot, insp = _insp_begin(conn, lot_id)
        u = _unit_in_lot(conn, lot_id, unit_id)
        if u["status"] != "active":
            raise ValueError("제외된 완제품은 NG 처리를 할 수 없어요.")
        if unit_has_final_document(unit_id, conn):
            raise ValueError(_FINAL_DOC_UNIT_MSG)
        if unit_is_unusable(unit_id, conn):
            raise ValueError(_UNUSABLE_UNIT_MSG)
        last = conn.execute("SELECT event_type FROM lot_inspection_unit_events WHERE inspection_id=? AND unit_id=? "
                            "ORDER BY seq DESC LIMIT 1", (insp["id"], unit_id)).fetchone()
        if not last or last["event_type"] != "complete":
            raise ValueError("완료된 S/N에서만 NG 사건을 열 수 있어요.")
        ids = _ng_open_missing(conn, insp, unit_id, actor, actor_id, "manual_open")
        if not ids:
            raise ValueError("새로 열 NG 사건이 없어요.")
        _ng_withdraw_pending(conn, insp, unit_id, "NG 사건 수동 열기", actor, actor_id)
        conn.commit()
        return ids
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@_write_retry
def add_ng_action(lot_id, unit_id, ng_event_id, action_kind, reason, memo, actor, actor_id):
    """후속조치 기록(retest|repair_retest|reassemble|special_request|other) → 로그 seq. 특채 요청은 사유 필수."""
    if action_kind not in _NG_ACTIONS:
        raise ValueError("알 수 없는 후속조치예요.")
    reason = (reason or "").strip()[:500]
    memo = (memo or "").strip()[:500] or None
    if action_kind == "special_request" and not reason:
        raise ValueError("특채 요청 사유를 입력해줘.")
    conn = get_conn()
    try:
        lot, insp = _insp_begin(conn, lot_id)
        _ng_writable_event(conn, insp, unit_id, ng_event_id)
        seq = _ng_log_append(conn, ng_event_id, "action", actor, actor_id, action_kind=action_kind,
                             reason=reason or None, memo=memo)
        # 승인 대기 성적서는 특채 요청이 바뀌든 다른 조치가 붙든 같은 트랜잭션에서 철회(대기가 아니면 무동작)
        _ng_withdraw_pending(conn, insp, unit_id, f"NG 후속조치 변경({action_kind}) · 사건 #{ng_event_id}", actor, actor_id)
        conn.commit()
        return seq
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@_write_retry
def submit_ng_retest(lot_id, unit_id, ng_event_id, entries, extra_item_keys, seen_digest, actor, actor_id):
    """NG 재검사 입력. entries=[{item_key, idx, raw|missing}] — 사건 항목은 '현재 fail인 칸'만, extra_item_keys에 든
    영향 항목은 이미 값이 있는 칸 어느 것이나. 합격(재검사 칸 모두 pass AND 그 항목 모든 칸 현재 pass)이면 같은
    트랜잭션에서 사건을 closed(retest_pass)로 닫는다. 영향 항목이 새로 fail이면 별도 사건을 연다."""
    entries = list(entries or [])
    extra = list(dict.fromkeys(extra_item_keys or []))
    if not entries:
        raise ValueError("재검사 값을 입력해줘.")
    conn = get_conn()
    try:
        lot, insp = _insp_begin(conn, lot_id)
        ev = _ng_writable_event(conn, insp, unit_id, ng_event_id)
        items = {i["key"]: i for i in json.loads(insp["items_json"])}
        key0 = ev["item_key"]
        it0 = items.get(key0)
        if not it0 or it0["input_type"] not in ("numeric", "ok_ng"):
            raise ValueError("숫자·적합/부적합 항목만 재검사할 수 있어요.")
        busy = {e["item_key"] for e in _ng_events_derived(conn, insp["id"], unit_id) if e["state"] in _NG_OPEN_STATES}
        for k in extra:
            if k == key0 or k not in items or items[k]["input_type"] not in ("numeric", "ok_ng"):
                raise ValueError("영향 항목이 올바르지 않아요.")
            if k in busy:
                raise NgBlocked([_blk("ng_extra_busy", f"{items[k]['name']}: 이미 진행 중인 NG 사건이 있어요.", k,
                                      items[k]["name"])])
        if seen_digest != _results_digest(conn, insp["id"], unit_id):
            raise ValueError("화면을 본 뒤에 값이 바뀌었어요. 새로고침해서 다시 확인해줘.")
        cur = _ng_cur(conn, insp["id"], unit_id)
        fail_idx = {i for i, r in cur.get(key0, {}).items() if r["verdict"] == "fail"}
        prepared, own_idx = [], []
        for e in entries:
            k = e.get("item_key")
            try:
                idx = int(e.get("idx"))
            except (TypeError, ValueError):
                raise SlotError("칸 번호가 올바르지 않아요.", k)
            row = cur.get(k, {}).get(idx)
            if k == key0:
                if idx not in fail_idx:
                    raise SlotError("불량이 난 칸만 재검사할 수 있어요.", k, idx)
                own_idx.append(idx)
            elif k in extra:
                if not row:
                    raise SlotError("값이 없는 칸은 재검사할 수 없어요.", k, idx)
            else:
                raise SlotError("이 사건의 재검사 대상이 아니에요.", k, idx)
            prepared.append({"item_key": k, "idx": idx, "raw": e.get("raw"), "missing": bool(e.get("missing")),
                             "base_rev": row["rev"]})
        if not own_idx:
            raise ValueError("불량이 난 칸의 재검사 값을 입력해줘.")
        round_no = ev["round_no"] + 1
        why = f"NG 재검사 {round_no}차 · 사건 #{ng_event_id}"
        _save_results_conn(conn, lot, insp, unit_id, prepared, actor, actor_id, retest_ctx={"reason": why})
        last = conn.execute("SELECT event_type FROM lot_inspection_unit_events WHERE inspection_id=? AND unit_id=? "
                            "ORDER BY seq DESC LIMIT 1", (insp["id"], unit_id)).fetchone()
        reopened = bool(last and last["event_type"] == "complete")
        _ng_withdraw_pending(conn, insp, unit_id, why, actor, actor_id)  # 이미 재개된 S/N의 대기 성적서도 포함
        if reopened:  # 값이 바뀌었으니 완료를 풀고 다시 완료하게 한다(대기 성적서 철회는 _append_unit_event 훅)
            _append_unit_event(conn, insp, unit_id, "reopen", "system", why, actor, actor_id)
        cur = _ng_cur(conn, insp["id"], unit_id)
        cells0 = cur.get(key0, {})
        passed = (all(cells0[i]["verdict"] == "pass" for i in own_idx)
                  and summarize_unit_item(it0, cells0) == "pass")
        new_fail = {k for k in extra if summarize_unit_item(items[k], cur.get(k, {})) == "fail"}
        new_ids = (_ng_open_missing(conn, insp, unit_id, actor, actor_id, "retest_extra", only_keys=new_fail,
                                    origin=(ng_event_id, round_no)) if new_fail else [])
        rows = [{"item_key": p["item_key"], "idx": p["idx"], "rev": cur[p["item_key"]][p["idx"]]["rev"],
                 "value": cur[p["item_key"]][p["idx"]]["value"],
                 "verdict": cur[p["item_key"]][p["idx"]]["verdict"], "extra": p["item_key"] != key0}
                for p in prepared]
        digest = _results_digest(conn, insp["id"], unit_id)
        _ng_log_append(conn, ng_event_id, "retest_round", actor, actor_id, round_no=round_no,
                       outcome="pass" if passed else "fail",
                       payload={"round_no": round_no, "rows": rows, "digest_after": digest, "new_event_ids": new_ids})
        if passed:
            _ng_log_append(conn, ng_event_id, "closed", actor, actor_id, round_no=round_no, close_kind="retest_pass",
                           outcome="pass")
        conn.commit()
        return {"round_no": round_no, "passed": passed, "closed": passed, "reopened": reopened,
                "new_event_ids": new_ids, "digest": digest}
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def list_ng_events(inspection_id, unit_id):
    """S/N의 사건 목록(헤더+로그+파생 상태+항목명) — 읽기 전용."""
    conn = get_conn()
    try:
        insp = conn.execute("SELECT * FROM production_lot_inspections WHERE id=?", (inspection_id,)).fetchone()
        names = {i["key"]: i["name"] for i in json.loads(insp["items_json"])} if insp else {}
        evs = _ng_events_derived(conn, inspection_id, unit_id)
        for e in evs:
            e["item_name"] = names.get(e["item_key"], e["item_key"])
            e["cancel_blocked"] = _ng_cancel_blockers(conn, insp, e["unit_id"], e) if insp else []
        return evs
    finally:
        conn.close()


def unit_ng_overview(lot_id, unit_id):
    """S/N의 NG 현황(읽기 전용) — {fail_items, blockers, events(항목명 포함), pending}. 검사가 시작 전이면 None."""
    insp = get_lot_inspection(lot_id)
    if not insp:
        return None
    conn = get_conn()
    try:
        row = conn.execute("SELECT * FROM production_lot_inspections WHERE id=?", (insp["id"],)).fetchone()
        g = _ng_unit_gate(conn, row, unit_id)
    finally:
        conn.close()
    names = {i["key"]: i["name"] for i in insp["items"]}
    for e in g["events"]:
        e["item_name"] = names.get(e["item_key"], e["item_key"])
    return g


def ng_states_for_inspection(inspection_id):
    """목록 화면용 — {unit_id: [{id, item_key, state, round_no}]}. 쿼리 3개 고정(N+1 없음)."""
    conn = get_conn()
    try:
        out = {}
        for e in _ng_events_derived(conn, inspection_id):
            out.setdefault(e["unit_id"], []).append({"id": e["id"], "item_key": e["item_key"], "state": e["state"],
                                                     "round_no": e["round_no"]})
        return out
    finally:
        conn.close()


def _ng_check_special_close(conn, e, l):
    """closed(special) 로그가 가리키는 성적서 승인 이벤트가 실제로 있었는지 교차 확인 → 오류 문구 목록.
    승인이 나중에 회수(revoked)됐는지는 보지 않는다(파생 상태의 몫) — 그 승인 이벤트가 존재했는지만 본다."""
    tag = f"사건 #{e['id']} seq {l['seq']}(특채 종결)"
    rev = conn.execute("SELECT inspection_id, unit_id FROM product_report_revisions WHERE id=?",
                       (l["revision_id"],)).fetchone() if l["revision_id"] is not None else None
    if not rev:
        return [f"{tag}: 연결된 성적서 revision이 없어요"]
    if rev["inspection_id"] != e["inspection_id"] or rev["unit_id"] != e["unit_id"]:
        return [f"{tag}: 연결된 성적서가 이 사건의 검사·S/N과 달라요"]
    ae = conn.execute("SELECT event_type, reason, payload_json FROM product_report_events WHERE revision_id=? AND seq=?",
                      (l["revision_id"], l["report_event_seq"])).fetchone() if l["report_event_seq"] is not None else None
    if not ae or ae["event_type"] != "approved":
        return [f"{tag}: 연결된 승인 이벤트가 없어요"]
    errs = []
    try:
        ids = ((json.loads(ae["payload_json"]) if ae["payload_json"] else {}) or {}).get("special", {}).get("ng_event_ids") or []
    except (ValueError, AttributeError):
        ids = []
    if e["id"] not in ids:
        errs.append(f"{tag}: 승인 이벤트의 특채 대상 사건 목록에 이 사건이 없어요")
    if (ae["reason"] or None) != (l["reason"] or None):
        errs.append(f"{tag}: 승인 이벤트의 특채 사유와 종결 로그 사유가 달라요")
    return errs


def verify_ng_events(inspection_id, unit_id=None):
    """저장된 컬럼에서 헤더 해시와 로그 체인을 다시 계산해 비교 → {ok, events, errors}."""
    conn = get_conn()
    try:
        errors = []
        evs = _ng_events_derived(conn, inspection_id, unit_id)
        for e in evs:
            hh = _ng_header_hash(e["inspection_id"], e["unit_id"], e["item_key"], e["event_no"],
                                 e["origin_event_id"], e["origin_round_no"], e["opened_trigger"],
                                 e["fail_snapshot_json"], e["serial_at_open"], e["lot_revision"], e["opened_by"],
                                 e["opened_by_user_id"], e["opened_at"])
            if hh != e["header_hash"]:
                errors.append(f"사건 #{e['id']}: 헤더 해시 불일치")
            prev, n = None, 0
            for l in e["logs"]:
                n += 1
                eh = _ng_log_hash(e["header_hash"], e["id"], l["seq"], l["log_type"], l["action_kind"],
                                  l["round_no"], l["outcome"], l["close_kind"], l["reason"], l["memo"],
                                  l["revision_id"], l["report_event_seq"], l["payload_json"], l["actor_name"],
                                  l["actor_user_id"], l["at"], l["prev_hash"])
                if l["seq"] != n:
                    errors.append(f"사건 #{e['id']}: seq {l['seq']} 순서 이상")
                if l["prev_hash"] != prev:
                    errors.append(f"사건 #{e['id']} seq {l['seq']}: 이전 해시 연결 끊김")
                if eh != l["event_hash"]:
                    errors.append(f"사건 #{e['id']} seq {l['seq']}: 해시 불일치")
                prev = l["event_hash"]
                if l["log_type"] == "closed" and l["close_kind"] == "special":
                    errors.extend(_ng_check_special_close(conn, e, l))
        return {"ok": not errors, "events": len(evs), "errors": errors}
    finally:
        conn.close()


# ---------- 제품 성적서 제출·승인 (TASK #148 패키지 4a) ----------
# 계약: product_report_revisions/events는 append-only(UPDATE/DELETE 금지). 승인 단위=S/N 1건.
# 상태는 revision별 마지막 이벤트(submitted=승인 대기, approved=승인됨, withdrawn/rejected/revoked=비활성).
# 모든 쓰기는 _insp_begin(LOT 단위 BEGIN IMMEDIATE) 안에서 상태·자격·digest를 다시 확인한다.
# 수입검사 approve/approve_revoke/compute_content_hash는 inspections 테이블 전제라 재사용하지 않는다.

class ReportBlocked(ValueError):
    """제품 성적서 제출·승인·반려·회수 거부 — blockers=[{code, item_key, name, msg}, ...]."""
    def __init__(self, blockers):
        self.blockers = blockers
        super().__init__(blockers[0]["msg"] if blockers else "처리할 수 없어요.")


_REJECT_KINDS = ("document_fix", "quality_impact")


def _explicit_perms(conn, user_id):
    """users.permissions에 직접 체크된 권한만(콤마 구분) — admin 자동 전권을 인정하지 않는다(사용자 확정 Q4)."""
    if user_id is None:
        return set()
    r = conn.execute("SELECT permissions FROM users WHERE id=?", (user_id,)).fetchone()
    return {p.strip() for p in (r["permissions"] or "").split(",") if p.strip()} if r else set()


def unit_revoke_blockers(conn, unit_id):
    """승인 회수를 막는 사유 목록 — 출하 완료 후 회수 차단 훅(지금은 항상 빈 목록)."""
    return []


def _report_event_hash(revision_id, seq, event_type, reason, actor_name, actor_user_id, reject_kind,
                       doc_rev, at, prev_hash, payload_json, content_hash):
    """이벤트 해시의 단일 출처 — 쓰기와 verify가 같은 식을 쓴다."""
    h = lambda s: hashlib.sha256(s.encode("utf-8")).hexdigest() if s else None
    return _sha({"revision_id": revision_id, "seq": seq, "event_type": event_type, "reason_hash": h(reason),
                 "actor_name": actor_name, "actor_user_id": actor_user_id, "reject_kind": reject_kind,
                 "doc_rev": doc_rev, "at": at, "prev_event_hash": prev_hash, "payload_hash": h(payload_json),
                 "content_hash": content_hash})


def _report_event_append(conn, revision_id, event_type, actor, actor_id, reason=None, reject_kind=None,
                         doc_rev=None, payload=None):
    """제품 성적서 이벤트 한 줄 추가(seq·prev_event_hash·event_hash의 단일 출처). _insp_begin 잠금 안에서만 부를 것."""
    last = conn.execute("SELECT seq, event_hash FROM product_report_events WHERE revision_id=? "
                        "ORDER BY seq DESC LIMIT 1", (revision_id,)).fetchone()
    seq = (last["seq"] if last else 0) + 1
    prev = last["event_hash"] if last else None
    ch = conn.execute("SELECT content_hash FROM product_report_revisions WHERE id=?", (revision_id,)).fetchone()[0]
    pj = _canon(payload) if payload is not None else None
    at = _utc_now()
    eh = _report_event_hash(revision_id, seq, event_type, reason, actor, actor_id, reject_kind, doc_rev, at,
                            prev, pj, ch)
    conn.execute(
        "INSERT INTO product_report_events (revision_id, seq, event_type, reason, reject_kind, doc_rev, "
        "actor_name, actor_user_id, payload_json, prev_event_hash, event_hash, at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
        (revision_id, seq, event_type, reason, reject_kind, doc_rev, actor, actor_id, pj, prev, eh, at))
    return seq


def _report_state(conn, inspection_id, unit_id):
    """unit의 가장 최근 revision 상태 {state, revision_id, submission_no, last_event}. 없으면 state=None."""
    r = conn.execute("SELECT id, submission_no FROM product_report_revisions WHERE inspection_id=? AND unit_id=? "
                     "ORDER BY submission_no DESC LIMIT 1", (inspection_id, unit_id)).fetchone()
    if not r:
        return {"state": None, "revision_id": None, "submission_no": 0, "last_event": None}
    ev = conn.execute("SELECT * FROM product_report_events WHERE revision_id=? ORDER BY seq DESC LIMIT 1",
                      (r["id"],)).fetchone()
    return {"state": ev["event_type"] if ev else None, "revision_id": r["id"],
            "submission_no": r["submission_no"], "last_event": ev}


def _withdraw_pending_report(conn, insp, unit_id, reason, source, actor, actor_id):
    """마지막 이벤트가 submitted인 revision만 철회. 이미 반려·철회·승인된 revision에는 아무것도 안 한다(중복 방지)."""
    st = _report_state(conn, insp["id"], unit_id)
    if st["state"] != "submitted":
        return None
    return _report_event_append(conn, st["revision_id"], "withdrawn", actor, actor_id,
                                reason=(reason or "완료 취소"), payload={"source": source, "cause": reason})


def _sig_ok(sig):
    return isinstance(sig, dict) and bool(sig.get("file")) and bool(sig.get("sha256"))


def _sig_payload(sig, role):
    return {"role": role, "file": sig["file"], "sha256": sig["sha256"], "stamp": sig.get("stamp")}


def _snapshot_rows(rows):
    return [{"item_key": r["item_key"], "idx": r["idx"], "rev": r["rev"], "value_input": r["value_input"],
             "value": r["value"], "verdict": r["verdict"], "is_missing": r["is_missing"],
             "file_sha256": r["file_sha256"], "entered_by": r["entered_by"],
             "entered_by_user_id": r["entered_by_user_id"], "entered_at": r["entered_at"],
             "entry_type": r["entry_type"], "reason": r["reason"]} for r in rows]


def _ng_snapshot_events(evs):
    """제출 스냅샷(schema 2)의 NG 섹션 — 사건 헤더·로그 전체(승인자가 요청 사유와 재검사 이력을 보게)."""
    keep = ("seq", "log_type", "action_kind", "round_no", "outcome", "close_kind", "reason", "memo",
            "actor_name", "actor_user_id", "at", "event_hash")
    return [{"id": e["id"], "item_key": e["item_key"], "event_no": e["event_no"], "state": e["state"],
             "opened_trigger": e["opened_trigger"], "origin_event_id": e["origin_event_id"],
             "header_hash": e["header_hash"], "fail_snapshot": e["fail_snapshot"],
             "logs": [{k: l[k] for k in keep} for l in e["logs"]]} for e in evs]


def _ng_fingerprint(evs):
    """사건 구성이 제출 뒤 바뀌었는지 비교용 — [(사건 id, 로그 수)] (스냅샷 dict·파생 dict 둘 다 받는다)."""
    return sorted((e["id"], len(e["logs"])) for e in evs)


def _ng_withdraw_pending(conn, insp, unit_id, why, actor, actor_id):
    """NG 사건에 변경이 생기면 승인 대기 성적서를 같은 트랜잭션에서 철회(대기가 아니면 아무것도 안 함)."""
    return _withdraw_pending_report(conn, insp, unit_id, why, "ng", actor, actor_id)


# ---------- NG 사건 취소 · S/N 사용 금지 (TASK #154 패키지 4b-2) ----------
# 사건 취소는 ng_event_log에 log_type='cancelled' 한 줄만 추가한다(컬럼 변경 없음, 취소된 사건은 재오픈 안 함).
# 사용 금지는 S/N 단위라 별도 해시 체인 unit_disposal_events(append-only)에 쓴다 — UPDATE/DELETE 금지.

_NG_CANCEL_STILL_FAIL_MSG = ("이 항목에 아직 불합격 칸이 있어요. 값이 합격으로 정정된 사건만 취소할 수 있어요. "
                             "불량이 맞다면 재검사나 특채 요청을 진행해줘.")


def _ng_cancel_blockers(conn, insp, unit_id, ev):
    """사건 취소를 막는 사유 목록(빈 목록=취소 가능) — 쓰기 함수와 화면이 같이 쓰는 단일 출처.
    취소 가능 = 미종결 + 승인본 없음 + 사용 금지 아님 + 그 사건 항목의 현재 값에 불합격 칸이 없음(fail 아님이면 허용)."""
    bl = []
    if ev["state"] not in _NG_OPEN_STATES:
        bl.append(_blk("ng_closed", "이미 종결되거나 취소된 NG 사건이에요.", ev["item_key"]))
    if unit_has_final_document(unit_id, conn):
        bl.append(_blk("final_document", _FINAL_DOC_UNIT_MSG, ev["item_key"]))
    if unit_is_unusable(unit_id, conn):
        bl.append(_blk("unit_unusable", _UNUSABLE_UNIT_MSG, ev["item_key"]))
    it = next((i for i in json.loads(insp["items_json"]) if i["key"] == ev["item_key"]), None)
    if it and summarize_unit_item(it, _ng_cur(conn, insp["id"], unit_id).get(it["key"], {})) == "fail":
        bl.append(_blk("ng_cancel_still_fail", _NG_CANCEL_STILL_FAIL_MSG, ev["item_key"]))
    return bl


@_write_retry
def cancel_ng_event(lot_id, unit_id, ng_event_id, reason, actor, actor_id):
    """NG 사건 취소(값이 합격 쪽으로 정정돼 불량이 사라진 사건) → 로그 seq. 권한은 process_ng_cancel 명시 보유자만."""
    conn = get_conn()
    try:
        lot, insp = _insp_begin(conn, lot_id)
        ev = _ng_writable_event(conn, insp, unit_id, ng_event_id)
        if "process_ng_cancel" not in _explicit_perms(conn, actor_id):
            raise NgBlocked([_blk("no_cancel_perm", "NG 사건 취소 권한(process_ng_cancel)이 이 계정에 직접 체크되어 있지 않아요. "
                                  "관리자 계정도 계정 상세에서 직접 체크해야 해요.", ev["item_key"])])
        reason = (reason or "").strip()[:500]
        if not reason:
            raise NgBlocked([_blk("reason_required", "취소 사유를 입력해줘.", ev["item_key"])])
        bl = _ng_cancel_blockers(conn, insp, unit_id, ev)
        if bl:
            raise NgBlocked(bl)
        c = _ng_cur(conn, insp["id"], unit_id).get(ev["item_key"], {})
        it = next(i for i in json.loads(insp["items_json"]) if i["key"] == ev["item_key"])
        payload = {"item_key": ev["item_key"], "item_verdict": summarize_unit_item(it, c),
                   "snapshot_cells": ev["fail_snapshot"].get("cells", []),
                   "cells_now": [{"idx": i, "rev": c[i]["rev"], "value": c[i]["value"], "verdict": c[i]["verdict"]}
                                 for i in sorted(c)],
                   "digest": _results_digest(conn, insp["id"], unit_id)}
        seq = _ng_log_append(conn, ng_event_id, "cancelled", actor, actor_id, payload=payload, outcome="cancelled",
                             reason=reason)
        _ng_withdraw_pending(conn, insp, unit_id, f"NG 사건 취소 · 사건 #{ng_event_id}", actor, actor_id)
        conn.commit()
        return seq
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def _disposal_event_hash(inspection_id, unit_id, seq, event_type, reason, payload_json, serial_at_event,
                         lot_revision, actor_name, actor_user_id, at, prev_hash):
    """사용 금지 이벤트 해시의 단일 출처 — 쓰기와 verify가 같은 식을 쓴다."""
    return _sha({"inspection_id": inspection_id, "unit_id": unit_id, "seq": seq, "event_type": event_type,
                 "reason_hash": _ng_h(reason), "payload_hash": _ng_h(payload_json),
                 "serial_at_event": serial_at_event, "lot_revision": lot_revision, "actor_name": actor_name,
                 "actor_user_id": actor_user_id, "at": at, "prev_event_hash": prev_hash})


def _disposal_append(conn, insp, unit_id, event_type, reason, payload, actor, actor_id):
    """사용 금지 이벤트 한 줄 추가 → seq. _insp_begin 잠금 안에서만 부를 것."""
    last = conn.execute("SELECT seq, event_hash FROM unit_disposal_events WHERE inspection_id=? AND unit_id=? "
                        "ORDER BY seq DESC LIMIT 1", (insp["id"], unit_id)).fetchone()
    seq = (last["seq"] if last else 0) + 1
    prev = last["event_hash"] if last else None
    sa = conn.execute("SELECT serial_raw FROM unit_serial_aliases WHERE unit_id=? AND kind='current'",
                      (unit_id,)).fetchone()
    serial = sa["serial_raw"] if sa else None
    lot_rev = conn.execute("SELECT lot_revision FROM production_lots WHERE id=?", (insp["lot_id"],)).fetchone()[0]
    pj = _canon(payload) if payload is not None else None
    at = _utc_now()
    eh = _disposal_event_hash(insp["id"], unit_id, seq, event_type, reason, pj, serial, lot_rev, actor, actor_id,
                              at, prev)
    conn.execute(
        "INSERT INTO unit_disposal_events (inspection_id, unit_id, seq, event_type, reason, payload_json, "
        "serial_at_event, lot_revision, actor_name, actor_user_id, prev_hash, event_hash, at) "
        "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (insp["id"], unit_id, seq, event_type, reason, pj, serial, lot_rev, actor, actor_id, prev, eh, at))
    return seq


def _unit_disposal_latest(conn, inspection_id):
    """{unit_id: 마지막(MAX seq) 사용 금지 이벤트 행} — 쿼리 1회(N+1 없음)."""
    rows = conn.execute(
        "SELECT e.* FROM unit_disposal_events e WHERE e.inspection_id=? AND e.seq = "
        "(SELECT MAX(x.seq) FROM unit_disposal_events x WHERE x.inspection_id=e.inspection_id "
        "AND x.unit_id=e.unit_id)", (inspection_id,)).fetchall()
    return {r["unit_id"]: r for r in rows}


@_write_retry
def decide_unit_unusable(lot_id, unit_id, reason, actor, actor_id, signature, expected_digest):
    """S/N 사용 금지 결정 → {seq, withdrew_report}. process_final 명시 보유자만, 사유+서명 필수,
    미종결 NG 사건이 1건 이상 있어야 하며 승인본이 있으면 거부. 대기 성적서는 같은 트랜잭션에서 철회."""
    conn = get_conn()
    try:
        lot, insp = _insp_begin(conn, lot_id)
        u = _unit_in_lot(conn, lot_id, unit_id)
        if u["status"] != "active":
            raise ValueError("제외된 완제품은 사용 금지를 결정할 수 없어요.")
        if "process_final" not in _explicit_perms(conn, actor_id):
            raise NgBlocked([_blk("no_final_perm", "제품검사 최종결정 권한(process_final)이 필요해요.")])
        if unit_is_unusable(unit_id, conn):
            raise ValueError("이미 사용 금지된 S/N이에요.")
        if unit_has_final_document(unit_id, conn):
            raise ValueError("이 S/N의 성적서가 최종 승인돼 있어서 사용 금지를 결정할 수 없어요. 승인 회수가 먼저 필요해요.")
        reason = (reason or "").strip()[:500]
        if not reason:
            raise NgBlocked([_blk("reason_required", "사용 금지 사유를 입력해줘.")])
        if not _sig_ok(signature):
            raise NgBlocked([_blk("no_signature", "등록된 내 서명이 없어요. 서명을 먼저 등록해줘.")])
        digest = _results_digest(conn, insp["id"], unit_id)
        if expected_digest != digest:
            raise ValueError("화면을 본 뒤에 값이 바뀌었어요. 새로고침해서 다시 확인해줘.")
        pend = [e for e in _ng_events_derived(conn, insp["id"], unit_id) if e["state"] in _NG_OPEN_STATES]
        if not pend:
            raise NgBlocked([_blk("no_open_ng", "열려 있는 NG 사건이 있어야 사용 금지로 결정할 수 있어요. 먼저 NG 사건을 열어줘.")])
        withdrew = _report_state(conn, insp["id"], unit_id)["state"] == "submitted"
        seq = _disposal_append(
            conn, insp, unit_id, "unusable", reason,
            {"signature": _sig_payload(signature, "decider"),
             "ng_events": [{"id": e["id"], "item_key": e["item_key"], "state": e["state"], "round_no": e["round_no"],
                            "log_count": len(e["logs"])} for e in pend],
             "digest": digest, "withdrew_report": withdrew}, actor, actor_id)
        if withdrew:
            _withdraw_pending_report(conn, insp, unit_id, "S/N 사용 금지 결정", "unusable", actor, actor_id)
        conn.commit()
        return {"seq": seq, "withdrew_report": withdrew}
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@_write_retry
def revoke_unit_unusable(lot_id, unit_id, reason, actor, actor_id):
    """사용 금지 회수 → seq. process_final 명시 보유자 누구나, 사유 필수, 서명 없음. 제외된 S/N도 허용.
    회수해도 NG 사건·완료 상태는 금지 전 그대로다(건드리지 않는다)."""
    conn = get_conn()
    try:
        lot, insp = _insp_begin(conn, lot_id)
        _unit_in_lot(conn, lot_id, unit_id)
        if "process_final" not in _explicit_perms(conn, actor_id):
            raise NgBlocked([_blk("no_final_perm", "제품검사 최종결정 권한(process_final)이 필요해요.")])
        reason = (reason or "").strip()[:500]
        if not reason:
            raise NgBlocked([_blk("reason_required", "회수 사유를 입력해줘.")])
        last = conn.execute("SELECT seq, event_type FROM unit_disposal_events WHERE inspection_id=? AND unit_id=? "
                            "ORDER BY seq DESC LIMIT 1", (insp["id"], unit_id)).fetchone()
        if not last or last["event_type"] != "unusable":
            raise ValueError("사용 금지 상태가 아니에요. 새로고침해줘.")
        seq = _disposal_append(conn, insp, unit_id, "unusable_revoked", reason, {"ref_seq": last["seq"]},
                               actor, actor_id)
        conn.commit()
        return seq
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def list_unit_disposal(inspection_id, unit_id):
    """S/N의 사용 금지 이력(seq 순, payload 파싱) — 읽기 전용."""
    conn = get_conn()
    try:
        out = []
        for r in conn.execute("SELECT * FROM unit_disposal_events WHERE inspection_id=? AND unit_id=? ORDER BY seq",
                              (inspection_id, unit_id)):
            d = dict(r)
            d["payload"] = json.loads(d["payload_json"]) if d["payload_json"] else {}
            out.append(d)
        return out
    finally:
        conn.close()


def verify_unit_disposal(inspection_id, unit_id=None):
    """저장된 컬럼에서 사용 금지 체인을 다시 계산해 비교 → {ok, events, errors}.
    seq 연속·prev_hash·해시·유형 순서(unusable ↔ unusable_revoked 번갈아)·ref_seq·payload.ng_events 소속을 확인."""
    conn = get_conn()
    try:
        w, a = "inspection_id=?", [inspection_id]
        if unit_id is not None:
            w += " AND unit_id=?"
            a.append(unit_id)
        rows = [dict(r) for r in conn.execute(f"SELECT * FROM unit_disposal_events WHERE {w} ORDER BY unit_id, seq", a)]
        errors, by_unit = [], {}
        for r in rows:
            by_unit.setdefault(r["unit_id"], []).append(r)
        for uid, evs in by_unit.items():
            prev, prev_type, prev_seq = None, None, None
            for n, r in enumerate(evs, 1):
                tag = f"S/N unit {uid} seq {r['seq']}"
                if r["seq"] != n:
                    errors.append(f"{tag}: seq 순서 이상")
                if r["prev_hash"] != prev:
                    errors.append(f"{tag}: 이전 해시 연결 끊김")
                eh = _disposal_event_hash(r["inspection_id"], uid, r["seq"], r["event_type"], r["reason"],
                                          r["payload_json"], r["serial_at_event"], r["lot_revision"], r["actor_name"],
                                          r["actor_user_id"], r["at"], r["prev_hash"])
                if eh != r["event_hash"]:
                    errors.append(f"{tag}: 해시 불일치")
                want = "unusable" if prev_type in (None, "unusable_revoked") else "unusable_revoked"
                if r["event_type"] != want:
                    errors.append(f"{tag}: 이벤트 순서 이상({r['event_type']})")
                try:
                    pl = json.loads(r["payload_json"]) if r["payload_json"] else {}
                except ValueError:
                    pl, _ = {}, errors.append(f"{tag}: payload를 읽을 수 없어요")
                if r["event_type"] == "unusable_revoked" and pl.get("ref_seq") != prev_seq:
                    errors.append(f"{tag}: 회수가 직전 사용 금지를 가리키지 않아요")
                for ne in pl.get("ng_events") or []:
                    ok = conn.execute("SELECT 1 FROM ng_events WHERE id=? AND inspection_id=? AND unit_id=?",
                                      (ne.get("id"), inspection_id, uid)).fetchone()
                    if not ok:
                        errors.append(f"{tag}: NG 사건 #{ne.get('id')}이 이 S/N 소속이 아니에요")
                prev, prev_type, prev_seq = r["event_hash"], r["event_type"], r["seq"]
        return {"ok": not errors, "events": len(rows), "errors": errors}
    finally:
        conn.close()


@_write_retry
def submit_product_report(lot_id, unit_id, actor, actor_id, expected_digest, approver_user_id,
                          writer_signature, photo_blockers=None):
    """S/N 1건의 성적서 승인 요청 → {revision_id, submission_no, content_hash, seq}. 위반은 ReportBlocked(.blockers).
    writer_signature={file, sha256, stamp} — 파일 처리는 라우트 몫, 여기서는 기록만."""
    conn = get_conn()
    try:
        lot, insp = _insp_begin(conn, lot_id)
        u = _unit_in_lot(conn, lot_id, unit_id)
        if u["status"] != "active":
            raise ReportBlocked([_blk("unit_excluded", "제외된 완제품은 승인 요청할 수 없어요.")])
        st = _report_state(conn, insp["id"], unit_id)
        if st["state"] == "submitted":
            raise ReportBlocked([_blk("already_submitted", "이미 승인 요청 중이에요.")])
        if st["state"] == "approved":
            raise ReportBlocked([_blk("already_approved", "이미 최종 승인됐어요. 승인 회수가 먼저 필요해요.")])
        valid, why = _unit_completion_valid_conn(conn, insp, unit_id)
        if not valid:
            raise ReportBlocked([_blk("not_completed", why)])
        digest = _results_digest(conn, insp["id"], unit_id)
        blockers = []
        if expected_digest != digest:
            blockers.append(_blk("stale", "화면을 본 뒤에 값이 바뀌었어요. 새로고침해서 다시 확인해줘."))
        blockers.extend(photo_blockers or [])
        items = json.loads(insp["items_json"])
        names = {i["key"]: i["name"] for i in items}
        rows = _slot_current(conn, insp["id"], unit_id)
        gate = _ng_unit_gate(conn, insp, unit_id)  # NG는 사건이 있고 전부 특채 요청된 때만 제출 허용
        blockers.extend(gate["blockers"])
        if not _sig_ok(writer_signature):
            blockers.append(_blk("no_signature", "등록된 내 서명이 없어요. 서명을 먼저 등록해줘."))
        wp = _explicit_perms(conn, actor_id)
        participants = {r[0] for r in conn.execute(
            "SELECT entered_by_user_id FROM lot_inspection_results WHERE inspection_id=? AND unit_id=?",
            (insp["id"], unit_id)) if r[0] is not None}
        participants |= {r[0] for r in conn.execute(
            "SELECT by_user_id FROM lot_inspection_unit_events WHERE inspection_id=? AND unit_id=? "
            "AND event_type='complete'", (insp["id"], unit_id)) if r[0] is not None}
        if actor_id is None or "process_submit" not in wp:
            blockers.append(_blk("no_submit_perm", "승인 요청 권한(process_submit)이 없어요."))
        elif actor_id not in participants:
            blockers.append(_blk("not_participant", "이 S/N 검사에 참여한 사람만 승인 요청할 수 있어요."))
        ap = conn.execute("SELECT id, username, display_name FROM users WHERE id=?", (approver_user_id,)).fetchone() \
            if approver_user_id is not None else None
        if not ap or "process_final" not in _explicit_perms(conn, ap["id"]):
            blockers.append(_blk("no_approver", "최종결정권자 지정이 필요해요. 관리자에게 제품검사 최종결정 권한 지정을 요청해줘."))
        elif ap["id"] == actor_id:
            blockers.append(_blk("same_person", "작성자와 최종 승인자는 다른 사람이어야 해요."))
        if blockers:
            raise ReportBlocked(blockers)
        cev = conn.execute("SELECT * FROM lot_inspection_unit_events WHERE inspection_id=? AND unit_id=? "
                           "AND event_type='complete' ORDER BY seq DESC LIMIT 1", (insp["id"], unit_id)).fetchone()
        sa = conn.execute("SELECT serial_raw FROM unit_serial_aliases WHERE unit_id=? AND kind='current'",
                          (unit_id,)).fetchone()
        serial = sa["serial_raw"] if sa else None
        hist = conn.execute("SELECT * FROM lot_inspection_results WHERE inspection_id=? AND unit_id=? "
                            "ORDER BY item_key, idx, rev", (insp["id"], unit_id)).fetchall()
        snapshot = {
            "schema": 2,
            "inspection": {"id": insp["id"], "snapshot_hash": insp["snapshot_hash"],
                           "item_version": insp["item_version"], "started_at": insp["started_at"], "items": items},
            "lot": json.loads(insp["lot_snapshot_json"]), "model": json.loads(insp["model_snapshot_json"]),
            "unit": {"unit_id": unit_id, "serial_at_submit": serial, "model_mismatch": bool(u["model_mismatch"]),
                     "roster_member": unit_id in _roster_ids(insp)},
            "slots": sorted(_snapshot_rows(rows), key=lambda x: (x["item_key"], x["idx"])),
            "history": _snapshot_rows(hist),
            "completion": {"seq": cev["seq"], "at": cev["at"], "by_name": cev["by_name"],
                           "by_user_id": cev["by_user_id"], "source": cev["source"], "results_digest": cev["results_digest"]},
            "missing_count": sum(1 for r in rows if r["is_missing"]),
            "ng": {"events": _ng_snapshot_events(gate["events"])},
        }
        sj = _canon(snapshot)
        try:
            cur = conn.execute(
                "INSERT INTO product_report_revisions (inspection_id, unit_id, submission_no, results_digest, "
                "completion_seq, snapshot_json, content_hash, serial_at_submit, lot_revision, writer_user_id, "
                "writer_name, approver_user_id, approver_name, created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (insp["id"], unit_id, st["submission_no"] + 1, digest, cev["seq"], sj, _sha(snapshot), serial,
                 lot["lot_revision"], actor_id, actor, ap["id"], ap["display_name"] or ap["username"], _utc_now()))
            rid = cur.lastrowid
        except sqlite3.IntegrityError as ie:
            if "product_report_revisions" not in str(ie) or "UNIQUE" not in str(ie):
                raise
            raise ValueError("다른 사람이 먼저 처리했어요. 새로고침해줘.")
        seq = _report_event_append(conn, rid, "submitted", actor, actor_id,
                                   payload={"signature": _sig_payload(writer_signature, "writer")})
        conn.commit()
        return {"revision_id": rid, "submission_no": st["submission_no"] + 1, "content_hash": _sha(snapshot), "seq": seq}
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def _report_lot_id(revision_id):
    """revision → lot_id(잠금 전 사전 조회). 없으면 ValueError."""
    conn = get_conn()
    try:
        r = conn.execute("SELECT i.lot_id FROM product_report_revisions r JOIN production_lot_inspections i "
                         "ON i.id=r.inspection_id WHERE r.id=?", (revision_id,)).fetchone()
    finally:
        conn.close()
    if not r:
        raise ValueError("성적서를 찾을 수 없어요.")
    return r["lot_id"]


def _report_locked(conn, revision_id):
    """잠금 안에서 revision·마지막 이벤트를 다시 읽는다 → (lot, insp, revision행, 마지막 이벤트행)."""
    lot_id = _report_lot_id(revision_id)
    lot, insp = _insp_begin(conn, lot_id)
    rev = conn.execute("SELECT * FROM product_report_revisions WHERE id=?", (revision_id,)).fetchone()
    last = conn.execute("SELECT * FROM product_report_events WHERE revision_id=? ORDER BY seq DESC LIMIT 1",
                        (revision_id,)).fetchone()
    return lot, insp, rev, last


def _need_designated_approver(conn, rev, actor_id):
    """지정된 승인자 본인 + process_final 명시 보유(admin 폴백 없음) + 작성자와 다른 계정."""
    if actor_id is None or actor_id != rev["approver_user_id"]:
        raise ReportBlocked([_blk("not_approver", "이 성적서에 지정된 최종결정권자만 처리할 수 있어요.")])
    if "process_final" not in _explicit_perms(conn, actor_id):
        raise ReportBlocked([_blk("no_final_perm", "제품검사 최종결정 권한(process_final)이 필요해요.")])
    if actor_id == rev["writer_user_id"]:
        raise ReportBlocked([_blk("same_person", "작성자와 최종 승인자는 다른 사람이어야 해요.")])


@_write_retry
def approve_product_report(revision_id, actor, actor_id, approver_signature, expected_digest, photo_blockers=None,
                           special_reason=None):
    """최종 승인 → {seq, doc_rev, special}. doc_rev=이전 approved 이벤트 수(최초 승인본 0, 정정본부터 1).
    NG(특채 요청된 사건)가 있으면 special_reason(특채 사유) 필수 — 승인 이벤트 reason·payload.special에 남고,
    같은 트랜잭션에서 그 사건들에 closed(special) 로그가 붙는다(하나라도 실패하면 전체 롤백)."""
    conn = get_conn()
    try:
        lot, insp, rev, last = _report_locked(conn, revision_id)
        if last["event_type"] != "submitted":
            raise ReportBlocked([_blk("not_pending", "승인 대기 상태가 아니에요. 새로고침해줘.")])
        _need_designated_approver(conn, rev, actor_id)
        valid, why = _unit_completion_valid_conn(conn, insp, rev["unit_id"])
        if not valid:
            raise ReportBlocked([_blk("not_completed", why)])
        blockers = []
        digest = _results_digest(conn, insp["id"], rev["unit_id"])
        if digest != rev["results_digest"] or (expected_digest is not None and expected_digest != rev["results_digest"]):
            blockers.append(_blk("stale", "제출 뒤에 값이 바뀌었어요. 새로고침해줘."))
        blockers.extend(photo_blockers or [])
        gate = _ng_unit_gate(conn, insp, rev["unit_id"])  # 제출 뒤 NG 구성이 달라졌는지 잠금 안에서 다시 계산
        blockers.extend(gate["blockers"])
        snap_ng = json.loads(rev["snapshot_json"]).get("ng")
        if snap_ng is not None and _ng_fingerprint(snap_ng["events"]) != _ng_fingerprint(gate["events"]):
            blockers.append(_blk("stale", "제출 뒤에 NG 사건이 바뀌었어요. 새로고침해줘."))
        special = (special_reason or "").strip()[:500]
        special_events = [e for e in gate["pending"]] if not gate["blockers"] else []
        if special_events and not special:
            blockers.append(_blk("special_reason_required", "NG 특채 승인 사유를 입력해줘."))
        if not _sig_ok(approver_signature):
            blockers.append(_blk("no_signature", "등록된 내 서명이 없어요. 서명을 먼저 등록해줘."))
        if blockers:
            raise ReportBlocked(blockers)
        doc_rev = conn.execute(
            "SELECT COUNT(*) FROM product_report_events e JOIN product_report_revisions r ON r.id=e.revision_id "
            "WHERE r.inspection_id=? AND r.unit_id=? AND e.event_type='approved'",
            (insp["id"], rev["unit_id"])).fetchone()[0]
        pl = {"signature": _sig_payload(approver_signature, "approver")}
        if special_events:
            pl["special"] = {"reason": special, "ng_event_ids": sorted(e["id"] for e in special_events)}
        seq = _report_event_append(conn, revision_id, "approved", actor, actor_id, reason=special or None,
                                   doc_rev=doc_rev, payload=pl)
        for e in special_events:  # 이미 닫힌 사건은 pending이 아니라 다시 닫지 않는다(재승인은 새 revision의 새 closed 줄)
            _ng_log_append(conn, e["id"], "closed", actor, actor_id, round_no=e["round_no"], outcome="special",
                           close_kind="special", reason=special, revision_id=revision_id, report_event_seq=seq)
        conn.commit()
        return {"seq": seq, "doc_rev": doc_rev, "special": bool(special_events)}
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@_write_retry
def reject_product_report(revision_id, actor, actor_id, reason, reject_kind):
    """반려(사유·종류 필수) → seq. quality_impact면 같은 트랜잭션에서 완료를 시스템 재개한다(document_fix는 완료 유지)."""
    reason = (reason or "").strip()
    if not reason:
        raise ReportBlocked([_blk("reason_required", "반려 사유를 입력해줘.")])
    if reject_kind not in _REJECT_KINDS:
        raise ReportBlocked([_blk("kind_required", "반려 종류를 선택해줘.")])
    conn = get_conn()
    try:
        lot, insp, rev, last = _report_locked(conn, revision_id)
        if last["event_type"] != "submitted":
            raise ReportBlocked([_blk("not_pending", "승인 대기 상태가 아니에요. 새로고침해줘.")])
        _need_designated_approver(conn, rev, actor_id)
        seq = _report_event_append(conn, revision_id, "rejected", actor, actor_id, reason=reason,
                                   reject_kind=reject_kind)
        if reject_kind == "quality_impact":
            ev = _unit_events_latest(conn, insp["id"]).get(rev["unit_id"])
            if ev and ev["event_type"] == "complete":  # rejected가 먼저 기록돼서 재개 훅의 철회는 중복되지 않는다
                _append_unit_event(conn, insp, rev["unit_id"], "reopen", "system",
                                   ("성적서 반려(품질 영향): " + reason)[:200], actor, actor_id)
        conn.commit()
        return seq
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@_write_retry
def revoke_product_report(revision_id, actor, actor_id, reason):
    """승인 회수(사유 필수) → seq. 서명·해시는 그대로 남고 '회수됨'으로만 바뀐다. process_final 명시 보유자."""
    reason = (reason or "").strip()
    if not reason:
        raise ReportBlocked([_blk("reason_required", "회수 사유를 입력해줘.")])
    conn = get_conn()
    try:
        lot, insp, rev, last = _report_locked(conn, revision_id)
        if last["event_type"] != "approved":
            raise ReportBlocked([_blk("not_approved", "승인된 성적서가 아니에요. 새로고침해줘.")])
        if actor_id is None or "process_final" not in _explicit_perms(conn, actor_id):
            raise ReportBlocked([_blk("no_final_perm", "제품검사 최종결정 권한(process_final)이 필요해요.")])
        blk = unit_revoke_blockers(conn, rev["unit_id"])
        if blk:
            raise ReportBlocked(blk)
        seq = _report_event_append(conn, revision_id, "revoked", actor, actor_id, reason=reason,
                                   payload={"content_hash_before": rev["content_hash"]})
        conn.commit()
        return seq
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def _event_dict(r):
    d = dict(r)
    d["payload"] = json.loads(d["payload_json"]) if d.get("payload_json") else None
    return d


def list_report_events(revision_id):
    conn = get_conn()
    try:
        return [_event_dict(r) for r in conn.execute(
            "SELECT * FROM product_report_events WHERE revision_id=? ORDER BY seq", (revision_id,))]
    finally:
        conn.close()


_REPORT_LIST_SQL = (
    "SELECT r.id, r.inspection_id, r.unit_id, r.submission_no, r.results_digest, r.content_hash, "
    "r.serial_at_submit, r.writer_user_id, r.writer_name, r.approver_user_id, r.approver_name, r.created_at, "
    "i.lot_id, l.lot_no, e.event_type AS state, e.at AS state_at, e.doc_rev AS doc_rev, e.reject_kind AS reject_kind, "
    "(SELECT COUNT(*) FROM product_report_events x WHERE x.revision_id=r.id AND x.event_type='approved') AS was_approved "
    "FROM product_report_revisions r JOIN production_lot_inspections i ON i.id=r.inspection_id "
    "JOIN production_lots l ON l.id=i.lot_id JOIN product_report_events e ON e.revision_id=r.id "
    "AND e.seq=(SELECT MAX(x.seq) FROM product_report_events x WHERE x.revision_id=r.id) WHERE 1=1")


def list_product_reports(state=None, lot_id=None, approver_user_id=None, writer_user_id=None):
    """revision 목록(최신순, 한 쿼리). state=마지막 이벤트 종류('submitted'|'approved'|...)."""
    sql, args = _REPORT_LIST_SQL, []
    for col, val in (("e.event_type", state), ("i.lot_id", lot_id), ("r.approver_user_id", approver_user_id),
                     ("r.writer_user_id", writer_user_id)):
        if val is not None:
            sql += f" AND {col}=?"
            args.append(val)
    conn = get_conn()
    try:
        return [dict(r) for r in conn.execute(sql + " ORDER BY r.id DESC", args)]
    finally:
        conn.close()


def list_unit_reports(lot_id, unit_id):
    """한 S/N의 revision 이력(최신순, 상태 포함)."""
    conn = get_conn()
    try:
        return [dict(r) for r in conn.execute(
            _REPORT_LIST_SQL + " AND i.lot_id=? AND r.unit_id=? ORDER BY r.submission_no DESC", (lot_id, unit_id))]
    finally:
        conn.close()


def get_product_report(revision_id):
    """revision 한 건(snapshot 파싱·상태·이벤트 포함). 없으면 None."""
    conn = get_conn()
    try:
        rows = [dict(r) for r in conn.execute(_REPORT_LIST_SQL + " AND r.id=?", (revision_id,))]
        if not rows:
            return None
        full =conn.execute("SELECT snapshot_json, completion_seq, lot_revision FROM product_report_revisions "
                            "WHERE id=?", (revision_id,)).fetchone()
    finally:
        conn.close()
    d = rows[0]
    d.update(snapshot=json.loads(full["snapshot_json"]), completion_seq=full["completion_seq"],
             lot_revision=full["lot_revision"], events=list_report_events(revision_id))
    return d


def verify_product_report(revision_id, photo_dir=None):
    """저장된 값에서 해시 체인·스냅샷·현재 결과를 다시 계산해 대조(읽기 전용).
    photo_dir를 주면 스냅샷 사진 파일의 존재·sha256도 확인한다(None이면 photo_ok=None).
    live_digest_ok는 승인 대기·승인 상태일 때만 의미가 있어 그 외엔 None."""
    out = {"ok": False, "chain_ok": False, "content_ok": False, "live_digest_ok": None, "photo_ok": None,
           "problems": []}
    conn = get_conn()
    try:
        rev = conn.execute("SELECT r.*, i.lot_id FROM product_report_revisions r JOIN production_lot_inspections i "
                           "ON i.id=r.inspection_id WHERE r.id=?", (revision_id,)).fetchone()
        if not rev:
            out["problems"].append("성적서를 찾을 수 없어요.")
            return out
        evs = conn.execute("SELECT * FROM product_report_events WHERE revision_id=? ORDER BY seq", (revision_id,)).fetchall()
        chain, prev = True, None
        for n, e in enumerate(evs, 1):
            want = _report_event_hash(revision_id, e["seq"], e["event_type"], e["reason"], e["actor_name"],
                                      e["actor_user_id"], e["reject_kind"], e["doc_rev"], e["at"],
                                      e["prev_event_hash"], e["payload_json"], rev["content_hash"])
            if e["seq"] != n or e["prev_event_hash"] != prev or e["event_hash"] != want:
                chain = False
                out["problems"].append(f"이벤트 #{n} 해시 체인이 맞지 않아요.")
                break
            prev = e["event_hash"]
        if not evs:
            chain = False
            out["problems"].append("이벤트가 없어요.")
        out["chain_ok"] = chain
        snap = json.loads(rev["snapshot_json"])
        content = _sha(snap) == rev["content_hash"] and _digest_rows(snap["slots"]) == rev["results_digest"]
        out["content_ok"] = content
        if not content:
            out["problems"].append("제출 당시 스냅샷이 저장된 해시와 달라요.")
        state = evs[-1]["event_type"] if evs else None
        if state in ("submitted", "approved"):
            live = _results_digest(conn, rev["inspection_id"], rev["unit_id"]) == rev["results_digest"]
            out["live_digest_ok"] = live
            if not live:
                out["problems"].append("현재 결과 값이 제출 당시와 달라요.")
        if photo_dir is not None:
            photo_ok = True
            for s in snap["slots"]:
                if s["verdict"] != "recorded":
                    continue
                p = os.path.join(photo_dir, os.path.basename(s["value"] or ""))
                try:
                    with open(p, "rb") as f:
                        okh = hashlib.sha256(f.read()).hexdigest() == s["file_sha256"]
                except OSError:
                    okh = False
                if not okh:
                    photo_ok = False
                    out["problems"].append(f"사진 파일 확인 실패: {s['item_key']}-{s['idx']}")
            out["photo_ok"] = photo_ok
        out["ok"] = chain and content and out["live_digest_ok"] is not False and out["photo_ok"] is not False
        return out
    finally:
        conn.close()


# ---------- 완제품 내부 ID와 S/N (TASK #66 패키지 1C-b) ----------

class SerialBatchError(ValueError):
    """S/N 일괄 추가 거부 — errors=[(줄번호, 원문, 사유), ...]. 하나라도 있으면 전부 취소."""
    def __init__(self, errors):
        self.errors = errors
        super().__init__(f"S/N {len(errors)}건에 문제가 있어서 아무것도 추가하지 않았어.")


def normalize_serial_key(s):
    """모든 공백(탭·NBSP 포함) 제거 + 대문자. 하이픈은 유지."""
    return re.sub(r"\s+", "", s or "").upper()


def parse_serial_lines(text):
    """줄바꿈·탭·세미콜론·쉼표로 분리, 앞뒤 공백 제거, 빈 항목 무시(공백으로는 분리하지 않는다)."""
    return [p.strip() for p in re.split(r"[\r\n\t;,]+", text or "") if p.strip()]


def serial_key_conflicts(conn, keys):
    """{키: 사유} — 완제품 별칭(현재·과거) 및 기존 출하 자료 3종(읽기 전용)과 겹치는 키만."""
    keys, out = set(keys), {}
    for r in conn.execute("SELECT a.serial_key, l.lot_no FROM unit_serial_aliases a "
                          "JOIN finished_units u ON u.id=a.unit_id "
                          "JOIN production_lots l ON l.id=u.lot_id"):
        if r["serial_key"] in keys:
            out[r["serial_key"]] = f"이미 등록된 S/N(LOT {r['lot_no']})"
    sources = (
        ("SELECT serial_no, NULL AS rno, NULL AS bid FROM finished_goods_serials",
         lambda r: "기존 출하 자료(S/N 발급 이력)에 이미 있는 S/N"),
        ("SELECT p.serial_no, b.round_no AS rno, p.batch_id AS bid FROM outbound_planned_items p "
         "LEFT JOIN outbound_batches b ON b.id=p.batch_id",
         lambda r: f"기존 출하 자료(차수 {r['rno'] or '#' + str(r['bid'])} 계획)에 이미 있는 S/N"),
        ("SELECT i.serial_no, b.round_no AS rno, i.batch_id AS bid FROM outbound_items i "
         "LEFT JOIN outbound_batches b ON b.id=i.batch_id",
         lambda r: f"기존 출하 자료(차수 {r['rno'] or '#' + str(r['bid'])} 출하 항목)에 이미 있는 S/N"),
    )
    for sql, why in sources:
        for r in conn.execute(sql):
            k = normalize_serial_key(r["serial_no"])
            if k in keys and k not in out:
                out[k] = why(r)
    return out


def _lot_model_codes(conn, lot):
    """LOT 모델의 (전압코드, 접미사, P코드) — 확정이면 스냅샷, 작성중이면 현재 모델 값."""
    if lot["model_snapshot_json"]:
        m = json.loads(lot["model_snapshot_json"])
    else:
        m = conn.execute("SELECT * FROM product_models WHERE id=?", (lot["model_id"],)).fetchone()
    return (m["voltage_code"] or "", m["suffix_code"] or "", m["pcode"] or "")


def _serial_parse(raw, rules, codes):
    """(키, 형식 오류 문구 또는 None, 모델 불일치 여부) — 형식은 기존 CKMR 본체 패턴 재사용."""
    key = normalize_serial_key(raw)
    mt = _match_ckmr(key, rules)
    if not mt:
        return key, "S/N 형식이 CKMR 본체 형식이 아니야.", 0
    mismatch = (mt["volt_code"] != codes[0] or mt["suffix_code"] is None
                or mt["suffix_code"] != codes[1] or (mt["pcode"] or "") != codes[2])
    return key, None, 1 if mismatch else 0


def _fu_full(conn, unit_id):
    """완제품 전체 상태 스냅샷(고정 키 순서) — 현재 별칭 포함."""
    u = conn.execute("SELECT * FROM finished_units WHERE id=?", (unit_id,)).fetchone()
    a = conn.execute("SELECT serial_raw, serial_key FROM unit_serial_aliases "
                     "WHERE unit_id=? AND kind='current'", (unit_id,)).fetchone()
    return {"unit_id": u["id"], "lot_id": u["lot_id"], "origin_state": u["origin_state"],
            "status": u["status"], "excluded_reason": u["excluded_reason"],
            "excluded_by": u["excluded_by"], "excluded_at": u["excluded_at"],
            "model_mismatch": u["model_mismatch"],
            "serial_raw": a["serial_raw"] if a else None,
            "serial_key": a["serial_key"] if a else None}


def _fu_history(conn, unit_id, lot_id, change_type, before, after, reason, rev, user):
    conn.execute(
        "INSERT INTO finished_unit_history (unit_id, lot_id, change_type, before_json, after_json, "
        "reason, lot_revision, changed_by) VALUES (?,?,?,?,?,?,?,?)",
        (unit_id, lot_id, change_type,
         json.dumps(before, ensure_ascii=False) if before is not None else None,
         json.dumps(after, ensure_ascii=False), reason, rev, user))


def _unit_lot_begin(conn, lot_id, reason):
    """잠금 후 LOT 검사 — 취소 거부, 확정이면 사유 필수. (LOT행, 사유, 변경 전 스냅샷, 변경 후 개정번호)."""
    row = _pl_lock(conn, lot_id)
    if row["status"] == "cancelled":
        raise ValueError("취소된 LOT은 바꿀 수 없어.")
    reason = (reason or "").strip() or None
    if row["status"] == "confirmed" and not reason:
        raise ValueError("확정된 LOT은 사유를 입력해야 바꿀 수 있어.")
    rev = row["lot_revision"] + (1 if row["status"] == "confirmed" else 0)
    return row, reason, _pl_full(conn, lot_id), rev


def _unit_lot_finish(conn, lot_id, rev, change_type, before, reason, user):
    """LOT 개정번호·갱신자 반영 + LOT 이력 1행."""
    conn.execute("UPDATE production_lots SET lot_revision=?, updated_by=?, "
                 "updated_at=datetime('now','localtime') WHERE id=?", (rev, user, lot_id))
    _pl_history(conn, lot_id, change_type, before, _pl_full(conn, lot_id), reason, user)


def _unit_lot_id(conn, unit_id):
    r = conn.execute("SELECT lot_id FROM finished_units WHERE id=?", (unit_id,)).fetchone()
    if not r:
        raise ValueError("완제품을 찾을 수 없어.")
    return r["lot_id"]


def get_finished_unit(unit_id):
    conn = get_conn()
    row = conn.execute("SELECT * FROM finished_units WHERE id=?", (unit_id,)).fetchone()
    conn.close()
    return row


@_write_retry
def add_units(lot_id, serial_text, user, reason=None):
    """S/N 일괄 추가 — 하나라도 오류면 전부 취소(SerialBatchError). (추가 건수, 경고[(번호, S/N, 사유)])."""
    items = parse_serial_lines(serial_text)
    if not items:
        raise ValueError("추가할 S/N을 입력해줘.")
    rules = get_classify_rules()
    conn = get_conn()
    try:
        lot, reason, before, rev = _unit_lot_begin(conn, lot_id, reason)
        codes = _lot_model_codes(conn, lot)
        parsed = [(i + 1, raw) + _serial_parse(raw, rules, codes) for i, raw in enumerate(items)]
        conflicts = serial_key_conflicts(conn, [p[2] for p in parsed if p[3] is None])
        errors, seen = [], set()
        for no, raw, key, err, mm in parsed:
            if err is None and key in seen:
                err = "같은 입력에서 중복"
            elif err is None:
                err = conflicts.get(key)
            seen.add(key)
            if err:
                errors.append((no, raw, err))
        if errors:
            raise SerialBatchError(errors)
        warnings = []
        for no, raw, key, err, mm in parsed:
            uid = conn.execute("INSERT INTO finished_units (lot_id, model_mismatch, created_by) "
                               "VALUES (?,?,?)", (lot_id, mm, user)).lastrowid
            try:
                conn.execute("INSERT INTO unit_serial_aliases (unit_id, serial_raw, serial_key, kind, "
                             "created_by) VALUES (?,?,?, 'current', ?)", (uid, raw, key, user))
            except sqlite3.IntegrityError as e:
                # 사전 검사와 같은 트랜잭션이라 정상 경로에선 못 오지만, S/N 키의 UNIQUE 위반만 안내 문구로
                # 바꾸고 다른 무결성 오류는 가리지 않고 그대로 전달한다
                if "unit_serial_aliases.serial_key" not in str(e):
                    raise
                raise ValueError("이미 등록된 S/N이야. 다시 확인해줘.")
            _fu_history(conn, uid, lot_id, "add", None, _fu_full(conn, uid), reason, rev, user)
            if mm:
                warnings.append((no, raw, "LOT 모델과 S/N의 모델 코드가 달라."))
        _unit_lot_finish(conn, lot_id, rev, "unit_add", before, reason, user)
        conn.commit()
        return len(parsed), warnings
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@_write_retry
def exclude_unit(unit_id, reason, user):
    """완제품 제외(사유 필수). S/N 번호는 그대로 예약 유지."""
    return _unit_change(unit_id, reason, user, "exclude", "active", "정상 상태인 완제품만 제외할 수 있어.")


@_write_retry
def restore_unit(unit_id, reason, user):
    """제외된 완제품 복구(사유 필수)."""
    return _unit_change(unit_id, reason, user, "restore", "excluded", "제외된 완제품만 복구할 수 있어.")


def _unit_change(unit_id, reason, user, kind, need_status, state_msg):
    reason = (reason or "").strip()
    if not reason:
        raise ValueError("사유를 입력해줘.")
    conn = get_conn()
    try:
        lot_id = _unit_lot_id(conn, unit_id)
        lot, reason, lot_before, rev = _unit_lot_begin(conn, lot_id, reason)
        if conn.execute("SELECT status FROM finished_units WHERE id=?", (unit_id,)).fetchone()[0] != need_status:
            raise ValueError(state_msg)
        if kind == "exclude" and unit_has_final_document(unit_id, conn):
            raise ValueError(_FINAL_DOC_UNIT_MSG)
        before = _fu_full(conn, unit_id)
        if kind == "exclude":
            conn.execute("UPDATE finished_units SET status='excluded', excluded_reason=?, "
                         "excluded_by=?, excluded_at=datetime('now','localtime') WHERE id=?",
                         (reason, user, unit_id))
        else:
            conn.execute("UPDATE finished_units SET status='active', excluded_reason=NULL, "
                         "excluded_by=NULL, excluded_at=NULL WHERE id=?", (unit_id,))
        _fu_history(conn, unit_id, lot_id, kind, before, _fu_full(conn, unit_id), reason, rev, user)
        _unit_lot_finish(conn, lot_id, rev, "unit_" + kind, lot_before, reason, user)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@_write_retry
def correct_serial(unit_id, new_serial, reason, user):
    """S/N 정정(사유 필수) — 기존 번호는 과거 별칭으로 남아 계속 예약된다."""
    reason = (reason or "").strip()
    if not reason:
        raise ValueError("사유를 입력해줘.")
    rules = get_classify_rules()
    conn = get_conn()
    try:
        lot_id = _unit_lot_id(conn, unit_id)
        lot, reason, lot_before, rev = _unit_lot_begin(conn, lot_id, reason)
        if unit_has_final_document(unit_id, conn):
            raise ValueError(_FINAL_DOC_UNIT_MSG)
        raw = (new_serial or "").strip()
        key, err, mm = _serial_parse(raw, rules, _lot_model_codes(conn, lot))
        if err:
            raise ValueError(err)
        if conn.execute("SELECT 1 FROM unit_serial_aliases WHERE unit_id=? AND serial_key=?",
                        (unit_id, key)).fetchone():
            raise ValueError("이미 이 완제품의 S/N이야.")
        why = serial_key_conflicts(conn, [key]).get(key)
        if why:
            raise ValueError(why)
        before = _fu_full(conn, unit_id)
        conn.execute("UPDATE unit_serial_aliases SET kind='past', replaced_reason=? "
                     "WHERE unit_id=? AND kind='current'", (reason, unit_id))
        try:
            conn.execute("INSERT INTO unit_serial_aliases (unit_id, serial_raw, serial_key, kind, "
                         "created_by) VALUES (?,?,?, 'current', ?)", (unit_id, raw, key, user))
        except sqlite3.IntegrityError as e:
            if "unit_serial_aliases.serial_key" not in str(e):
                raise
            raise ValueError("이미 등록된 S/N이야. 다시 확인해줘.")
        conn.execute("UPDATE finished_units SET model_mismatch=? WHERE id=?", (mm, unit_id))
        _fu_history(conn, unit_id, lot_id, "correct", before, _fu_full(conn, unit_id), reason, rev, user)
        _unit_lot_finish(conn, lot_id, rev, "serial_correct", lot_before, reason, user)
        conn.commit()
        return mm
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def get_production_lot_units(lot_id):
    """LOT의 완제품 목록(내부 ID 오름차순) — 현재 S/N, 과거 S/N 목록 포함."""
    conn = get_conn()
    units = [dict(r) for r in conn.execute(
        "SELECT u.*, a.serial_raw AS serial_raw FROM finished_units u "
        "LEFT JOIN unit_serial_aliases a ON a.unit_id=u.id AND a.kind='current' "
        "WHERE u.lot_id=? ORDER BY u.id", (lot_id,))]
    past = {}
    for r in conn.execute("SELECT a.unit_id, a.serial_raw FROM unit_serial_aliases a "
                          "JOIN finished_units u ON u.id=a.unit_id "
                          "WHERE u.lot_id=? AND a.kind='past' ORDER BY a.id", (lot_id,)):
        past.setdefault(r["unit_id"], []).append(r["serial_raw"])
    conn.close()
    for u in units:
        u["past_serials"] = past.get(u["id"], [])
    return units


def get_unit_history_for_lot(lot_id):
    """완제품 변경 이력(최신순) — 변경 시점의 S/N(serial)과 정정 전 S/N(prev_serial) 포함."""
    conn = get_conn()
    rows = [dict(r) for r in conn.execute(
        "SELECT * FROM finished_unit_history WHERE lot_id=? ORDER BY id DESC", (lot_id,))]
    conn.close()
    for h in rows:
        after = json.loads(h["after_json"]) if h["after_json"] else {}
        before = json.loads(h["before_json"]) if h["before_json"] else {}
        h["serial"] = after.get("serial_raw")
        h["prev_serial"] = before.get("serial_raw") if before.get("serial_raw") != h["serial"] else None
    return rows


def _pm_seed_rows(conn):
    """시드 계산 공용부 — STICKER_MODEL_MATCH_KEYS 21종 × 분류 규칙으로 후보 행 생성."""
    rules = get_classify_rules()
    suffix_map = dict(rules["suffix"])
    existing = {r["model_key"] for r in conn.execute("SELECT model_key FROM product_models")}
    out = []
    for model_no, (volt, suf, pc) in STICKER_MODEL_MATCH_KEYS.items():
        volt_label = rules["voltage"].get(volt)
        suf_label = suffix_map.get(suf)
        missing = []
        if not volt_label:
            missing.append(f"전압코드 {volt}")
        if not suf_label:
            missing.append(f"접미사 {suf}")
        label = " ".join(p for p in (volt_label, suf_label, f"({pc}P)" if pc else None) if p)
        m = re.search(r"\d+", volt_label or "")
        phase = "single" if suf == "S" else "three"
        out.append({"model_no": model_no, "voltage_code": volt, "suffix_code": suf, "pcode": pc,
                    "label": label, "voltage_kv": int(m.group()) if m else None,
                    "phase_type": phase, "housing_required_qty": 1 if phase == "single" else 3,
                    "exists": normalize_model_key(model_no) in existing,
                    "missing_rules": missing})
    return out


def product_model_seed_preview():
    """초기 모델 21종 시드 미리보기(쓰기 없음)."""
    conn = get_conn()
    try:
        return _pm_seed_rows(conn)
    finally:
        conn.close()


@_write_retry
def apply_product_model_seed(user):
    """없는 model_key만 INSERT(덮어쓰기 금지) + seed 이력. (삽입 건수, 건너뜀 건수) 반환."""
    conn = get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        inserted = skipped = 0
        for r in _pm_seed_rows(conn):
            if r["exists"]:
                skipped += 1
                continue
            name = (r["model_no"] + " " + r["label"]).strip()
            cur = conn.execute(
                "INSERT INTO product_models (model_no, model_key, display_name, voltage_code, "
                "voltage_kv, suffix_code, pcode, phase_type, housing_required_qty, created_by) "
                "VALUES (?,?,?,?,?,?,?,?,?,?)",
                (r["model_no"], normalize_model_key(r["model_no"]), name, r["voltage_code"],
                 r["voltage_kv"], r["suffix_code"], r["pcode"], r["phase_type"],
                 r["housing_required_qty"], user))
            _pm_history(conn, cur.lastrowid, 1, "seed", None,
                        _pm_full(conn, cur.lastrowid), None, user)
            inserted += 1
        conn.commit()
        return inserted, skipped
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def update_inspection_status(inspection_id, status, approver=None, reject_reason=None,
                             approval_type=None, approved_at=None):
    """approved_at을 넘기면 그 값을 쓰고, 안 넘기면 현재시각(now, localtime)을 쓴다.
    (복구 등록 — 원래 승인됐던 날짜를 그대로 기록하고 싶을 때 씀.)
    """
    conn = get_conn()
    approved_at_sql = "?" if approved_at else "datetime('now', 'localtime')"
    params = [status, approver]
    if approved_at:
        params.append(approved_at)
    if approval_type is not None:
        conn.execute(f"""
            UPDATE inspections
            SET status = ?, approver = ?, approved_at = {approved_at_sql},
                reject_reason = ?, approval_type = ?
            WHERE id = ?
        """, (*params, reject_reason, approval_type, inspection_id))
    else:
        conn.execute(f"""
            UPDATE inspections
            SET status = ?, approver = ?, approved_at = {approved_at_sql},
                reject_reason = ?
            WHERE id = ?
        """, (*params, reject_reason, inspection_id))
    conn.commit()
    conn.close()


def set_report_files(inspection_id, signature_path=None, pdf_path=None, xlsx_path=None):
    conn = get_conn()
    conn.execute("""
        UPDATE inspections SET signature_path = ?, pdf_path = ? WHERE id = ?
    """, (signature_path, pdf_path, inspection_id))
    conn.commit()
    conn.close()


REMARK_FIELDS = {"inspector": "remark_inspector", "manager": "remark_manager", "approver": "remark_approver"}


def update_inspection_remark(inspection_id, role_key, text):
    """role_key: 'inspector' / 'manager' / 'approver' — 각자 자기 비고란만 갱신."""
    column = REMARK_FIELDS.get(role_key)
    if column is None:
        return
    conn = get_conn()
    conn.execute(f"UPDATE inspections SET {column} = ? WHERE id = ?", (text, inspection_id))
    conn.commit()
    conn.close()


# ---------- 검사 진행 현황 ----------

def _ip_get(conn, intake_id):
    import json
    row = conn.execute(
        "SELECT inspectors FROM inspection_progress WHERE intake_id = ?", (intake_id,)
    ).fetchone()
    return json.loads(row["inspectors"]) if row else []


def register_inspector(intake_id, name):
    """검사 입력폼 열 때 본인 이름을 진행 현황에 등록."""
    import json
    conn = get_conn()
    names = _ip_get(conn, intake_id)
    if name not in names:
        names.append(name)
    conn.execute("""
        INSERT INTO inspection_progress (intake_id, inspectors, updated_at)
        VALUES (?, ?, datetime('now','localtime'))
        ON CONFLICT(intake_id) DO UPDATE SET
            inspectors = excluded.inspectors,
            updated_at = excluded.updated_at
    """, (intake_id, json.dumps(names, ensure_ascii=False)))
    conn.commit()
    conn.close()


def withdraw_inspector(intake_id, name):
    """검사원 제외 — 본인 이름만 제거. 마지막 1명은 제외 불가(호출 전 확인 필요)."""
    import json
    conn = get_conn()
    names = _ip_get(conn, intake_id)
    names = [n for n in names if n != name]
    if names:
        conn.execute("""
            INSERT INTO inspection_progress (intake_id, inspectors, updated_at)
            VALUES (?, ?, datetime('now','localtime'))
            ON CONFLICT(intake_id) DO UPDATE SET
                inspectors = excluded.inspectors,
                updated_at = excluded.updated_at
        """, (intake_id, json.dumps(names, ensure_ascii=False)))
    else:
        conn.execute("DELETE FROM inspection_progress WHERE intake_id = ?", (intake_id,))
    conn.commit()
    conn.close()
    return names  # 남은 검사원 목록 반환


def clear_inspection_progress(intake_id):
    """검사 제출 완료 시 진행 현황 삭제."""
    conn = get_conn()
    conn.execute("DELETE FROM inspection_progress WHERE intake_id = ?", (intake_id,))
    conn.commit()
    conn.close()


def get_progress_by_intake_ids(intake_ids):
    """intake_id 목록에 대한 진행 현황 일괄 조회. 반환: {intake_id: [name, ...]}"""
    import json
    if not intake_ids:
        return {}
    conn = get_conn()
    placeholders = ",".join("?" for _ in intake_ids)
    rows = conn.execute(
        f"SELECT intake_id, inspectors FROM inspection_progress WHERE intake_id IN ({placeholders})",
        intake_ids
    ).fetchall()
    conn.close()
    return {r["intake_id"]: json.loads(r["inspectors"] or "[]") for r in rows}


# ---------- 반품 처리 ----------

def create_return_request(inspection_id, material_no, material_name, supplier,
                          return_date, reason, quantity, created_by):
    conn = get_conn()
    cur = conn.execute("""
        INSERT INTO return_requests
            (inspection_id, material_no, material_name, supplier,
             return_date, reason, quantity, created_by)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    """, (inspection_id, material_no, material_name, supplier,
          return_date, reason, quantity, created_by))
    conn.commit()
    rid = cur.lastrowid
    conn.close()
    return rid


def update_return_request(return_id, return_date, reason, quantity):
    """반품 처리 내용 수정(2026-09-16 신설). 반품 사유/날짜/수량이 바뀌면 이미 진행된
    처리(재납품대기/재검사완료 등)의 근거가 달라지므로, 상태를 '반품요청'으로 되돌리고
    resolved_inspection_id를 지운다(8-2-10절 승인회수 원칙과 같은 정신 — NCR/개선요청서
    수정 시 발송·확인 흔적을 지우는 것과 동일한 이유)."""
    conn = get_conn()
    conn.execute("""
        UPDATE return_requests SET
            return_date = ?, reason = ?, quantity = ?,
            status = '반품요청', resolved_inspection_id = NULL
        WHERE id = ?
    """, (return_date, reason, quantity, return_id))
    conn.commit()
    conn.close()


_RETURN_LIST_SELECT = """
    SELECT r.*,
           i.inspector       AS insp_inspector,
           i.receive_date    AS insp_receive_date,
           i.inspect_date    AS insp_inspect_date,
           i.overall_result  AS insp_overall_result,
           i.status          AS insp_status,
           i.approval_type   AS insp_approval_type
    FROM return_requests r
    LEFT JOIN inspections i ON i.id = r.inspection_id
"""


def list_return_requests(status=None):
    """검사자/입고일/검사일/자동판정/승인상태로도 검색할 수 있게 원본 성적서(inspections)
    정보를 같이 조인해서 넘긴다(insp_ 접두어 — return_requests.status와 헷갈리지 않게 구분)."""
    conn = get_conn()
    if status:
        rows = conn.execute(
            _RETURN_LIST_SELECT + " WHERE r.status = ? ORDER BY r.id DESC", (status,)
        ).fetchall()
    else:
        rows = conn.execute(_RETURN_LIST_SELECT + " ORDER BY r.id DESC").fetchall()
    conn.close()
    return rows


def get_return_request(return_id):
    conn = get_conn()
    row = conn.execute("SELECT * FROM return_requests WHERE id = ?", (return_id,)).fetchone()
    conn.close()
    return row


def update_return_status(return_id, status, resolved_inspection_id=None):
    conn = get_conn()
    conn.execute("""
        UPDATE return_requests SET status = ?, resolved_inspection_id = ?
        WHERE id = ?
    """, (status, resolved_inspection_id, return_id))
    conn.commit()
    conn.close()


def get_return_requests_by_inspection(inspection_id):
    conn = get_conn()
    rows = conn.execute(
        "SELECT * FROM return_requests WHERE inspection_id = ? ORDER BY id DESC",
        (inspection_id,)
    ).fetchall()
    conn.close()
    return rows


def delete_return_requests(return_ids):
    """반품 요청 일괄 삭제 (admin 전용)."""
    if not return_ids:
        return
    conn = get_conn()
    placeholders = ",".join("?" * len(return_ids))
    conn.execute(f"DELETE FROM return_requests WHERE id IN ({placeholders})", return_ids)
    conn.commit()
    conn.close()


def delete_inspections(inspection_ids):
    """성적서 복수 삭제 (inspection_items 포함). pending/rejected 건은 연결 intake를 대기로 되돌림."""
    if not inspection_ids:
        return
    ids = list(inspection_ids)
    ph = ",".join("?" * len(ids))
    conn = get_conn()
    # pending/rejected 건의 intake_id 수집 → 나중에 대기로 복구
    rows = conn.execute(
        f"SELECT intake_id FROM inspections WHERE id IN ({ph}) AND status IN ('pending','rejected')",
        ids
    ).fetchall()
    intake_ids_to_reset = [r["intake_id"] for r in rows if r["intake_id"]]

    # 전수검사 자식 테이블 먼저 삭제 (FK 걸려 있음)
    fi_ids = [r["id"] for r in conn.execute(
        f"SELECT id FROM full_inspections WHERE inspection_id IN ({ph})", ids
    ).fetchall()]
    if fi_ids:
        fph = ",".join("?" * len(fi_ids))
        conn.execute(f"DELETE FROM full_inspection_units WHERE full_inspection_id IN ({fph})", fi_ids)
        conn.execute(f"DELETE FROM full_inspections WHERE id IN ({fph})", fi_ids)

    conn.execute(f"DELETE FROM inspection_items WHERE inspection_id IN ({ph})", ids)
    conn.execute(f"DELETE FROM inspections WHERE id IN ({ph})", ids)

    # 남은 inspection이 없는 intake만 대기로 되돌림
    for iid in intake_ids_to_reset:
        remaining = conn.execute(
            "SELECT COUNT(*) AS cnt FROM inspections WHERE intake_id=?", (iid,)
        ).fetchone()["cnt"]
        if remaining == 0:
            conn.execute("UPDATE intake_list SET status='대기' WHERE id=?", (iid,))

    conn.commit()
    conn.close()


def import_assembly_from_excel(excel_filepath):
    """MA 자동출력.xlsm의 DATABASE 시트에서 조립 제품 파츠 분해 정보를 읽어 DB에 임포트.

    DATABASE 시트 구조 (2026-08-23 확인):
      - MA명은 A1, D1, G1, J1, ... (3칸 간격) 에 있고
      - 그 MA의 파츠 자재번호는 바로 오른쪽 열(B, E, H, K, ...)의 1~8행에 있다.
      - 파츠 자재번호는 MA_성적서_최종.xlsx의 시트명과 1:1로 같다.
    (예전 구현은 A/D/G 열 자체를 파츠번호로 읽어서 7322·7311 같은 사내 약칭이 들어갔고,
     그 번호는 자재 마스터에 존재하지 않아 검사가 불가능했다 — 그래서 오른쪽 열로 바로잡음)

    반환: (성공 건수, 에러 메시지 또는 None)"""
    try:
        import openpyxl
    except ImportError:
        return 0, "openpyxl 미설치"

    wb = None
    try:
        wb = openpyxl.load_workbook(excel_filepath, read_only=True, data_only=True)
        ws = wb['DATABASE']
        # 1~8행 전체를 읽어둔다 (행1 = MA명 + 첫 파츠, 행2~8 = 나머지 파츠)
        data = [list(r) for r in ws.iter_rows(min_row=1, max_row=8, values_only=True)]
    except Exception as e:
        return 0, f"엑셀 파일 읽기 실패: {e}"
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass

    if not data or not data[0]:
        return 0, "DATABASE 시트가 비어있어"

    def _cell(row_idx, col_idx):
        row = data[row_idx] if row_idx < len(data) else []
        v = row[col_idx] if col_idx < len(row) else None
        if v is None:
            return ""
        return str(v).strip()

    conn = get_conn()
    cur = conn.cursor()
    imported = 0

    try:
        # A(0)=MA명 / B(1)=파츠,  D(3)=MA명 / E(4)=파츠,  G(6)/H(7) ... 3칸 간격
        for ma_col in range(0, len(data[0]), 3):
            part_col = ma_col + 1
            assembly_no = _cell(0, ma_col)
            if not assembly_no:
                continue

            components = []
            for row_idx in range(0, 8):
                component_no = _cell(row_idx, part_col)
                if component_no:
                    components.append(component_no)
            if not components:
                continue

            # 재임포트 안전: 기존 MA는 파츠를 갈아끼운다(파일이 갱신되면 반영돼야 하므로)
            master = cur.execute(
                "SELECT id FROM assembly_masters WHERE assembly_no=?", (assembly_no,)
            ).fetchone()
            if master:
                assembly_id = master["id"]
                cur.execute("DELETE FROM assembly_components WHERE assembly_id=?", (assembly_id,))
                cur.execute(
                    "UPDATE assembly_masters SET assembly_name=?, component_count=? WHERE id=?",
                    (assembly_no, len(components), assembly_id),
                )
            else:
                cur.execute(
                    "INSERT INTO assembly_masters (assembly_no, assembly_name, component_count) VALUES (?, ?, ?)",
                    (assembly_no, assembly_no, len(components)),
                )
                assembly_id = cur.lastrowid

            for order, component_no in enumerate(components, start=1):
                # 파츠명은 자재 마스터(성적서에서 뽑은 품명)를 정본으로 쓴다 — 여기선 번호만 저장
                cur.execute(
                    "INSERT INTO assembly_components (assembly_id, component_order, component_no, component_name) "
                    "VALUES (?, ?, ?, ?)",
                    (assembly_id, order, component_no, ""),
                )

            imported += 1

        conn.commit()
    finally:
        conn.close()

    return imported, None


# ---------- 통합BOM 계층 정보 (assembly_masters와 무관, 조회/필터 전용) ----------

_BOM_ALLOWED_UNITS = {"pc", "set", "ea"}


def import_bom_from_excel(filepath):
    """통합BOM 엑셀("통합BOM" 시트)에서 자재별 계층(모델/Lv/상위품목코드) 정보를 읽어
    material_bom_links에 재삽입한다.

    시트 구조(2026-09-09 확인, 헤더 2행/데이터 3행부터):
      1~5=Lv1~Lv5(그 중 하나만 값 있음)  6=모델명  7=Rev(안씀)  8=품목코드
      9=규격사양(SPEC, bom_name으로 참고 저장 — materials.material_name은 안 건드림)
      10=소요량  11=단위  12=1대당누적  13=상위품목코드  14=구분  15=적용모델수(안씀)  16~=모델별 매트릭스(안씀)

    구분이 '완제품'인 행은 제외. 단위가 Pc/SET/EA(대소문자·공백 무시)가 아니면 제외.

    2026-09-22부터 전량 재삽입이 아니다 — manually_edited=1인 행(사용자가 "자재 찾기"에서
    직접 수정한 행, update_bom_link() 참고)은 보존하고, (material_no, source_row_no)를
    매칭키로 써서 중복 삽입도 방지한다. 원본 엑셀 행순서가 바뀌면 이 매칭이 깨져 중복이
    생길 수 있다는 게 알려진 한계다 — 화면의 '✎ 수정됨' 배지로 사람이 인지할 수 있게만
    해뒀다(자동 재매칭은 하지 않음).
    반환: {"imported", "skipped_unit", "skipped_no_code", "skipped_finished_good",
           "skipped_no_level", "preserved_edited"}"""
    import openpyxl

    wb = openpyxl.load_workbook(filepath, data_only=True)
    try:
        ws = wb["통합BOM"]
        rows_raw = list(ws.iter_rows(min_row=3, values_only=True))
    finally:
        wb.close()

    summary = {"imported": 0, "skipped_unit": 0, "skipped_no_code": 0,
               "skipped_finished_good": 0, "skipped_no_level": 0, "preserved_edited": 0}
    rows_to_insert = []

    for row_no, row in enumerate(rows_raw, start=3):
        if row is None or all(v is None for v in row):
            continue
        row = list(row) + [None] * max(0, 15 - len(row))  # 15열 미만 방어

        kind = str(row[13]).strip() if row[13] is not None else None
        if kind == "완제품":
            summary["skipped_finished_good"] += 1
            continue

        unit_raw = row[10]
        unit_norm = str(unit_raw).strip().lower() if unit_raw is not None else ""
        if unit_norm not in _BOM_ALLOWED_UNITS:
            summary["skipped_unit"] += 1
            continue

        material_no = str(row[7]).strip() if row[7] is not None else ""
        if not material_no:
            summary["skipped_no_code"] += 1
            continue

        level = None
        for lv_idx in range(0, 5):
            v = row[lv_idx]
            if v is not None and str(v).strip() != "":
                level = lv_idx + 1
                break
        if level is None:
            summary["skipped_no_level"] += 1
            continue

        model_name = str(row[5]).strip() if row[5] is not None else ""
        parent_no = str(row[12]).strip() if row[12] is not None else None
        bom_name = str(row[8]).strip() if row[8] is not None else None

        def _num(v):
            if v is None:
                return None
            try:
                return float(v)
            except (TypeError, ValueError):
                return None

        rows_to_insert.append((
            material_no, parent_no or None, model_name, level, kind,
            _num(row[9]), _num(row[11]), str(row[10]).strip() if row[10] is not None else None,
            row_no, bom_name or None,
        ))
        summary["imported"] += 1

    conn = get_conn()
    try:
        edited_keys = {
            (r["material_no"], r["source_row_no"])
            for r in conn.execute(
                "SELECT material_no, source_row_no FROM material_bom_links WHERE manually_edited = 1"
            ).fetchall()
        }

        filtered_rows = []
        for row in rows_to_insert:
            # row 튜플 순서: (material_no, parent_material_no, model_name, level, kind,
            #                 qty_per_parent, qty_per_model, unit, source_row_no, bom_name)
            key = (row[0], row[8])
            if key in edited_keys:
                summary["preserved_edited"] += 1
                summary["imported"] -= 1
                continue
            filtered_rows.append(row)

        conn.execute("DELETE FROM material_bom_links WHERE manually_edited = 0")
        conn.executemany("""
            INSERT INTO material_bom_links
                (material_no, parent_material_no, model_name, level, kind,
                 qty_per_parent, qty_per_model, unit, source_row_no, bom_name)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, filtered_rows)
        conn.commit()
    finally:
        conn.close()

    return summary


def update_bom_link(link_id, model_name, parent_material_no, level):
    """'자재 찾기'에서 BOM 연동 행 하나의 model_name/parent_material_no/level을 사용자가
    직접 수정(2026-09-22). kind는 수정 대상이 아니다(사용자 확정 범위, CLAUDE.md 18절).
    manually_edited=1로 표시돼 다음 import_bom_from_excel() 재임포트 때 이 행은
    덮어쓰기 대상에서 제외된다. 실패 시 ValueError(사람이 읽을 문구)를 낸다."""
    model_name = (model_name or "").strip()
    if not model_name:
        raise ValueError("모델명은 비워둘 수 없어.")
    try:
        level_int = int(level)
    except (TypeError, ValueError):
        raise ValueError("Lv는 1~5 사이 숫자여야 해.")
    if level_int < 1 or level_int > 5:
        raise ValueError("Lv는 1~5 사이여야 해.")
    parent_material_no = (parent_material_no or "").strip() or None

    conn = get_conn()
    try:
        conn.execute("""
            UPDATE material_bom_links
            SET model_name = ?, parent_material_no = ?, level = ?, manually_edited = 1
            WHERE id = ?
        """, (model_name, parent_material_no, level_int, link_id))
        conn.commit()
    finally:
        conn.close()


def get_bom_links_for_material(material_no):
    conn = get_conn()
    rows = conn.execute(
        "SELECT * FROM material_bom_links WHERE material_no = ? ORDER BY model_name, level",
        (material_no,)
    ).fetchall()
    conn.close()
    return rows


def list_bom_model_names():
    conn = get_conn()
    rows = conn.execute(
        "SELECT DISTINCT model_name FROM material_bom_links ORDER BY model_name"
    ).fetchall()
    conn.close()
    return [r["model_name"] for r in rows]


def search_bom_materials(query="", levels=None, models=None, parent_no="", category=None,
                          unregistered_only=False, include=None, exclude=None,
                          major_category=None, middle_category=None):
    """material_bom_links를 materials에 LEFT JOIN해서 조회("자재 찾기" 화면 전용).
    search_materials()와는 완전히 별개 함수(책임 분리).
    include/exclude: 포함/제외 다중 단어(자재번호·자재명만 대상 — 확정된 결정).
    _include_exclude_clause() 참고."""
    conn = get_conn()
    where = ["1=1"]
    params = []

    if unregistered_only:
        where.append("m.material_no IS NULL")

    query = (query or "").strip()
    if query:
        where.append("(b.material_no LIKE ? OR m.material_name LIKE ?)")
        like = f"%{query}%"
        params += [like, like]

    levels = [l for l in (levels or []) if str(l).strip()]
    if levels:
        placeholders = ",".join("?" * len(levels))
        where.append(f"b.level IN ({placeholders})")
        params += [int(l) for l in levels]

    models = [mo for mo in (models or []) if mo]
    if models:
        placeholders = ",".join("?" * len(models))
        where.append(f"b.model_name IN ({placeholders})")
        params += models

    parent_no = (parent_no or "").strip()
    if parent_no:
        where.append("b.parent_material_no LIKE ?")
        params.append(f"%{parent_no}%")

    category = (category or "").strip()
    if category:
        where.append("m.category = ?")
        params.append(category)

    major_category = (major_category or "").strip()
    middle_category = (middle_category or "").strip()
    if major_category == UNASSIGNED_CLASSIFICATION or middle_category == UNASSIGNED_CLASSIFICATION:
        where.append("mc.material_no IS NULL")
    else:
        if major_category:
            where.append("mc.major_category = ?")
            params.append(major_category)
        if middle_category:
            where.append("mc.middle_category = ?")
            params.append(middle_category)

    ie_expr, ie_params = _include_exclude_clause(["b.material_no", "m.material_name"], include, exclude)
    if ie_expr:
        where.append(ie_expr)
        params += ie_params

    sql = f"""
        SELECT b.id, b.material_no, m.material_name, b.bom_name, b.level, b.model_name,
               b.parent_material_no, b.kind, b.qty_per_parent, b.qty_per_model, b.unit, m.category,
               mc.major_category, mc.middle_category,
               b.manually_edited
        FROM material_bom_links b
        LEFT JOIN materials m ON m.material_no = b.material_no
        LEFT JOIN material_classification_map mc ON mc.material_no = b.material_no
        WHERE {" AND ".join(where)}
        ORDER BY b.material_no, b.model_name, b.level
    """
    rows = conn.execute(sql, params).fetchall()
    conn.close()
    return rows


def delete_bom_links_bulk(link_ids):
    """'자재 찾기'에서 선택한 BOM 연동 행만 삭제. 자재(materials)나 규격은 안 건드림 —
    material_bom_links는 조회/필터 전용 참고 데이터라 잘못 들어온 행만 지워도 안전하다."""
    if not link_ids:
        return
    conn = get_conn()
    placeholders = ",".join("?" * len(link_ids))
    conn.execute(f"DELETE FROM material_bom_links WHERE id IN ({placeholders})", link_ids)
    conn.commit()
    conn.close()


def count_unregistered_bom_materials():
    """material_bom_links에는 있는데 materials에는 없는 자재번호 수(자재 찾기 화면 상단 안내용)."""
    conn = get_conn()
    n = conn.execute("""
        SELECT COUNT(DISTINCT b.material_no) FROM material_bom_links b
        LEFT JOIN materials m ON m.material_no = b.material_no
        WHERE m.material_no IS NULL
    """).fetchone()[0]
    conn.close()
    return n


def get_assembly_by_no(assembly_no):
    """MA 조립 제품 정보 조회. 반환: (assembly_id, 파츠 리스트) 또는 (None, [])"""
    conn = get_conn()
    cur = conn.cursor()

    master = cur.execute(
        "SELECT id FROM assembly_masters WHERE assembly_no=?",
        (assembly_no,)
    ).fetchone()

    if not master:
        conn.close()
        return None, []

    assembly_id = master['id']
    components = cur.execute(
        "SELECT component_order, component_no, component_name FROM assembly_components WHERE assembly_id=? ORDER BY component_order",
        (assembly_id,)
    ).fetchall()

    conn.close()
    return assembly_id, [dict(c) for c in components]


def list_all_assemblies(query=None, search_by="all"):
    """조립품 목록 + 각 조립품의 파츠 수. 자재 마스터에 없는 파츠 수(missing)도 같이 센다.

    query: 검색어. search_by: 'assembly_no'(조립품번호) / 'component_no'(파츠번호) / 'all'(둘 다).
    """
    conn = get_conn()
    sql = """
        SELECT m.*,
               (SELECT COUNT(*) FROM assembly_components c WHERE c.assembly_id = m.id) AS part_count,
               (SELECT COUNT(*) FROM assembly_components c
                 WHERE c.assembly_id = m.id
                   AND NOT EXISTS (SELECT 1 FROM materials mt WHERE mt.material_no = c.component_no)
               ) AS missing_count
          FROM assembly_masters m
    """
    params = []
    if query:
        like = f"%{query}%"
        component_match = ("EXISTS (SELECT 1 FROM assembly_components c "
                            "WHERE c.assembly_id = m.id AND c.component_no LIKE ?)")
        if search_by == "component_no":
            sql += f" WHERE {component_match}"
            params.append(like)
        elif search_by == "assembly_no":
            sql += " WHERE m.assembly_no LIKE ?"
            params.append(like)
        else:
            sql += f" WHERE m.assembly_no LIKE ? OR {component_match}"
            params.extend([like, like])
    sql += " ORDER BY m.assembly_no"
    rows = conn.execute(sql, params).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def get_assembly(assembly_id):
    """조립품 1건 + 파츠 목록(자재명·규격수 포함). 반환: (master, parts) 또는 (None, [])"""
    conn = get_conn()
    master = conn.execute("SELECT * FROM assembly_masters WHERE id=?", (assembly_id,)).fetchone()
    if not master:
        conn.close()
        return None, []
    parts = conn.execute("""
        SELECT c.*,
               mt.material_name,
               (SELECT COUNT(*) FROM specs s WHERE s.material_no = c.component_no) AS spec_count
          FROM assembly_components c
          LEFT JOIN materials mt ON mt.material_no = c.component_no
         WHERE c.assembly_id = ?
         ORDER BY c.component_order
    """, (assembly_id,)).fetchall()
    conn.close()
    return dict(master), [dict(p) for p in parts]


def save_assembly(assembly_no, component_nos, assembly_id=None):
    """조립품 하나를 통째로 저장(신규 등록 또는 수정).
    component_nos: 파츠 자재번호 리스트(순서대로). 파츠는 항상 전체 교체된다.
    반환: (assembly_id, 에러메시지 또는 None)"""
    assembly_no = (assembly_no or "").strip()
    if not assembly_no:
        return None, "조립품 이름(또는 대표 자재번호)을 입력해줘."

    seen, parts = set(), []
    for no in component_nos:
        no = (no or "").strip()
        if no and no not in seen:      # 같은 파츠를 두 번 넣는 건 막는다
            seen.add(no)
            parts.append(no)
    if not parts:
        return None, "파츠 자재번호를 하나 이상 입력해줘."

    conn = get_conn()
    try:
        dup = conn.execute(
            "SELECT id FROM assembly_masters WHERE assembly_no=? AND id IS NOT ?",
            (assembly_no, assembly_id),
        ).fetchone()
        if dup:
            return None, f"'{assembly_no}' 이름의 조립품이 이미 있어."

        if assembly_id:
            conn.execute("UPDATE assembly_masters SET assembly_no=?, assembly_name=?, component_count=? WHERE id=?",
                         (assembly_no, assembly_no, len(parts), assembly_id))
            conn.execute("DELETE FROM assembly_components WHERE assembly_id=?", (assembly_id,))
        else:
            cur = conn.execute(
                "INSERT INTO assembly_masters (assembly_no, assembly_name, component_count) VALUES (?, ?, ?)",
                (assembly_no, assembly_no, len(parts)))
            assembly_id = cur.lastrowid

        for order, no in enumerate(parts, start=1):
            conn.execute("INSERT INTO assembly_components (assembly_id, component_order, component_no, component_name) "
                         "VALUES (?, ?, ?, ?)", (assembly_id, order, no, ""))
        conn.commit()
        return assembly_id, None
    finally:
        conn.close()


def delete_assembly(assembly_id):
    conn = get_conn()
    try:
        conn.execute("DELETE FROM assembly_components WHERE assembly_id=?", (assembly_id,))
        conn.execute("DELETE FROM assembly_masters WHERE id=?", (assembly_id,))
        conn.commit()
    finally:
        conn.close()


def delete_assemblies_bulk(assembly_ids):
    """조립품 여러 개를 한번에 삭제 — 조립품 관리 화면 선택삭제용."""
    if not assembly_ids:
        return
    conn = get_conn()
    placeholders = ",".join("?" for _ in assembly_ids)
    conn.execute(f"DELETE FROM assembly_components WHERE assembly_id IN ({placeholders})", assembly_ids)
    conn.execute(f"DELETE FROM assembly_masters WHERE id IN ({placeholders})", assembly_ids)
    conn.commit()
    conn.close()


# ---------- 검사 입력 임시저장 (서버 보관) ----------

def active_inspection_for_intake(intake_id):
    """이 입고 건에 이미 살아있는 성적서가 있으면 돌려준다(재검사로 대체된 건 제외).

    같은 입고 건에 성적서가 여러 개 생기면 대시보드에서 같은 로트 수량이 중복 집계된다.
    (실제로 등록 버튼 연타/새로고침으로 1초 간격 3건이 생긴 사례가 있었다)
    """
    if not intake_id:
        return None
    conn = get_conn()
    row = conn.execute("""SELECT * FROM inspections
                           WHERE intake_id = ? AND status != 'superseded'
                           ORDER BY id DESC LIMIT 1""", (intake_id,)).fetchone()
    conn.close()
    return row


def save_inspection_draft(intake_id, payload_json, user_id=None, username=None):
    """검사 입력 중간값을 서버에 저장. 같은 입고 건은 항상 덮어쓴다."""
    conn = get_conn()
    try:
        conn.execute("""
            INSERT INTO inspection_drafts (intake_id, user_id, username, payload, updated_at)
                 VALUES (?, ?, ?, ?, datetime('now','localtime'))
            ON CONFLICT(intake_id) DO UPDATE SET
                 user_id=excluded.user_id, username=excluded.username,
                 payload=excluded.payload, updated_at=excluded.updated_at
        """, (intake_id, user_id, username, payload_json))
        conn.commit()
    finally:
        conn.close()


def get_inspection_draft(intake_id):
    conn = get_conn()
    row = conn.execute("SELECT * FROM inspection_drafts WHERE intake_id=?", (intake_id,)).fetchone()
    conn.close()
    return row


def delete_inspection_draft(intake_id):
    conn = get_conn()
    conn.execute("DELETE FROM inspection_drafts WHERE intake_id=?", (intake_id,))
    conn.commit()
    conn.close()


# ---------- 4M 변경점 ----------

CHANGE_TYPES = [
    ("Man",      "사람 (작업자·교대 변경)"),
    ("Machine",  "설비 (금형·장비 교체)"),
    ("Material", "자재 (원자재·공급처 변경)"),
    ("Method",   "방법 (공정·조건 변경)"),
]


def add_change_point(supplier, material_no, change_type, change_date, description, reported_by):
    conn = get_conn()
    try:
        cur = conn.execute("""
            INSERT INTO change_points (supplier, material_no, change_type, change_date, description, reported_by)
                 VALUES (?, ?, ?, ?, ?, ?)
        """, (supplier, material_no or None, change_type, change_date, description, reported_by))
        conn.commit()
        return cur.lastrowid
    finally:
        conn.close()


def list_change_points(supplier=None, material_no=None, limit=200):
    conn = get_conn()
    sql = "SELECT * FROM change_points WHERE 1=1"
    params = []
    if supplier:
        sql += " AND supplier = ?"
        params.append(supplier)
    if material_no:
        sql += " AND (material_no = ? OR material_no IS NULL)"
        params.append(material_no)
    sql += " ORDER BY change_date DESC, id DESC LIMIT ?"
    params.append(limit)
    rows = conn.execute(sql, params).fetchall()
    conn.close()
    return rows


def delete_change_point(cp_id):
    conn = get_conn()
    conn.execute("DELETE FROM change_points WHERE id=?", (cp_id,))
    conn.commit()
    conn.close()


def recent_change_points_for(supplier, material_no, within_days=90):
    """이 업체·자재에 최근 변경점이 있었는지 — 검사 화면 경고용."""
    if not supplier:
        return []
    conn = get_conn()
    rows = conn.execute("""
        SELECT * FROM change_points
         WHERE supplier = ?
           AND (material_no IS NULL OR material_no = ?)
           AND date(change_date) >= date('now', 'localtime', ?)
         ORDER BY change_date DESC
    """, (supplier, material_no, f"-{int(within_days)} days")).fetchall()
    conn.close()
    return rows


# ---------- 출고 완제품 S/N ----------

def serial_exists(serial_no):
    conn = get_conn()
    row = conn.execute(
        "SELECT 1 FROM finished_goods_serials WHERE serial_no=?", (serial_no,)).fetchone()
    conn.close()
    return row is not None


def unit_alias_lot_no(serial_no):
    """제품검사 LOT에 완제품으로 등록된 S/N(정규화 비교, 현재·과거 별칭 모두)이면 그 LOT 번호, 아니면 None."""
    key = normalize_serial_key(serial_no)
    if not key:
        return None
    conn = get_conn()
    try:
        row = conn.execute(
            "SELECT l.lot_no FROM unit_serial_aliases a JOIN finished_units u ON u.id=a.unit_id "
            "JOIN production_lots l ON l.id=u.lot_id WHERE a.serial_key=?", (key,)).fetchone()
    finally:
        conn.close()
    return row["lot_no"] if row else None


@_write_retry
def create_serial(serial_no, issued_by):
    """완제품 S/N을 발급 이력에 등록. 이미 발급된 S/N이거나 제품검사 LOT에 등록된 S/N이면 재발급을 막는다(ValueError).
    별칭 확인과 INSERT를 BEGIN IMMEDIATE 한 트랜잭션에 묶어, LOT S/N 추가(add_units)와 동시에 실행돼도 직렬화된다."""
    key = normalize_serial_key(serial_no)
    conn = get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        if conn.execute("SELECT 1 FROM finished_goods_serials WHERE serial_no=?", (serial_no,)).fetchone():
            raise ValueError(f"'{serial_no}'는 이미 발급된 S/N이야.")
        row = conn.execute(
            "SELECT l.lot_no FROM unit_serial_aliases a JOIN finished_units u ON u.id=a.unit_id "
            "JOIN production_lots l ON l.id=u.lot_id WHERE a.serial_key=?", (key,)).fetchone() if key else None
        if row:
            raise ValueError(f"'{serial_no}'는 이미 제품검사 LOT({row['lot_no']})에 등록된 S/N이야.")
        cur = conn.execute(
            "INSERT INTO finished_goods_serials (serial_no, issued_by) VALUES (?, ?)",
            (serial_no, issued_by))
        conn.commit()
        return cur.lastrowid
    except sqlite3.IntegrityError:
        conn.rollback()
        raise ValueError(f"'{serial_no}'는 이미 발급된 S/N이야.")
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@_write_retry
def create_serials_bulk(serial_nos, issued_by):
    """여러 S/N을 한 번에 발급 — 이미 발급된 것·LOT에 등록된 것·제출 목록 내 중복(정규화 비교)은 조용히
    건너뛰고 나머지는 계속 진행한다(붙여넣기 특성상 한 줄 실패로 전체를 막으면 안 되기 때문).
    별칭 조회와 INSERT는 BEGIN IMMEDIATE 한 트랜잭션. 반환: (added: [str], skipped: [str])."""
    conn = get_conn()
    added, skipped, seen = [], [], set()
    try:
        conn.execute("BEGIN IMMEDIATE")
        aliased = set()
        ks = [k for k in (normalize_serial_key(s) for s in serial_nos if (s or "").strip()) if k]
        for i in range(0, len(ks), 500):
            chunk = ks[i:i + 500]
            aliased.update(r[0] for r in conn.execute(
                "SELECT serial_key FROM unit_serial_aliases WHERE serial_key IN (%s)" % ",".join("?" * len(chunk)),
                chunk).fetchall())
        for sn in serial_nos:
            sn = (sn or "").strip()
            if not sn:
                continue
            k = normalize_serial_key(sn)
            if k in seen or k in aliased:
                skipped.append(sn)
                continue
            seen.add(k)
            try:
                conn.execute(
                    "INSERT INTO finished_goods_serials (serial_no, issued_by) VALUES (?, ?)",
                    (sn, issued_by))
                added.append(sn)
            except sqlite3.IntegrityError:
                skipped.append(sn)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return added, skipped


def delete_serial(serial_id):
    conn = get_conn()
    conn.execute("DELETE FROM finished_goods_serials WHERE id=?", (serial_id,))
    conn.commit()
    conn.close()


def delete_serials_bulk(serial_ids):
    """S/N 발급 이력을 여러 건 한 번에 삭제."""
    conn = get_conn()
    conn.executemany("DELETE FROM finished_goods_serials WHERE id=?",
                      [(sid,) for sid in serial_ids])
    conn.commit()
    conn.close()


def list_serials(query=None, limit=50):
    """S/N 발급 이력 최신순. query가 있으면 부분일치로 거른다."""
    conn = get_conn()
    sql = "SELECT * FROM finished_goods_serials WHERE 1=1"
    params = []
    if query:
        sql += " AND serial_no LIKE ?"
        params.append(f"%{query}%")
    sql += " ORDER BY id DESC LIMIT ?"
    params.append(limit)
    rows = conn.execute(sql, params).fetchall()
    conn.close()
    return rows


def find_outbound_item_batch(serial_no, exclude_batch_id=None):
    """이 S/N이 이미 다른 출고 배치에 쓰였으면 그 batch_id, 아니면 None.
    exclude_batch_id(현재 편집 중인 배치)는 '다른 배치'로 치지 않는다."""
    conn = get_conn()
    sql = "SELECT batch_id FROM outbound_items WHERE serial_no=?"
    params = [serial_no]
    if exclude_batch_id is not None:
        sql += " AND batch_id != ?"
        params.append(exclude_batch_id)
    row = conn.execute(sql + " ORDER BY id DESC LIMIT 1", params).fetchone()
    conn.close()
    return row["batch_id"] if row else None


# ---------- 출고 배치/항목/사진 ----------

def create_outbound_batch(customer, ship_date, handler, created_by, round_no=None):
    conn = get_conn()
    cur = conn.execute("""
        INSERT INTO outbound_batches (customer, ship_date, handler, created_by, round_no)
        VALUES (?, ?, ?, ?, ?)
    """, (customer, ship_date, handler, created_by, round_no))
    conn.commit()
    batch_id = cur.lastrowid
    conn.close()
    return batch_id


def update_outbound_batch(batch_id, customer, ship_date, handler, round_no=None):
    conn = get_conn()
    conn.execute("""
        UPDATE outbound_batches SET customer=?, ship_date=?, handler=?, round_no=?,
               updated_at=datetime('now','localtime')
        WHERE id=?
    """, (customer, ship_date, handler, round_no, batch_id))
    conn.commit()
    conn.close()


def confirm_outbound_batch(batch_id, confirmed_by, signature_path=None):
    """확인 시 이 배치의 항목 추가/수정/삭제·배치정보 수정을 잠근다(2026-09-15 변경 —
    예전엔 '확인은 순수 기록용, 잠금 아님'이 설계문서 확정사항이었으나 사용자가 명시적으로
    잠금 방식으로 바꿔달라고 요청해 뒤집었다). 실제 잠금 강제는 app.py의
    _outbound_batch_lock_response()가 각 수정 라우트 앞단에서 한다 — 이 함수 자체는
    여전히 단순 UPDATE다. 2026-09-21: 최종결정권자 서명 경로(confirm_signature)도 같이
    저장한다(사용자 확정 — 출고 확인도 승인/NCR/업체성적표처럼 서명이 필수인 최종 결정으로
    바뀌었다)."""
    conn = get_conn()
    conn.execute("""
        UPDATE outbound_batches SET confirmed_by=?, confirmed_at=datetime('now','localtime'),
               confirm_signature=?
        WHERE id=?
    """, (confirmed_by, signature_path, batch_id))
    conn.commit()
    conn.close()


def revoke_outbound_batch_confirm(batch_id):
    """출고 확인을 취소하고 잠금을 풀어준다. 2026-09-21: 서명이 생기면서 8-2-10절과
    같은 원칙 적용 — 확인이 취소됐는데 확인자 서명이 남아 있으면 안 된다. 성적서
    승인 회수와 같은 관례로 DB 포인터만 NULL로 하고 실제 PNG 파일은 지우지 않는다."""
    conn = get_conn()
    conn.execute("""
        UPDATE outbound_batches SET confirmed_by=NULL, confirmed_at=NULL, confirm_signature=NULL
        WHERE id=?
    """, (batch_id,))
    conn.commit()
    conn.close()


def get_outbound_batch(batch_id):
    conn = get_conn()
    row = conn.execute("SELECT * FROM outbound_batches WHERE id=?", (batch_id,)).fetchone()
    conn.close()
    return row


def count_outbound_pending_batches():
    """출고 승인대기 배치 수 — 미확인 배치 중 '완료된' 것만 센다(2026-09-28 사용자
    확정: 일반 수입검사가 '검사완료 제출' 전엔 승인목록에 안 뜨는 것과 같은 원칙).
    완료 판정은 outbound_batch_confirm()의 실제 게이트 조건과 반드시 같아야 하므로
    새로 계산하지 않고 outbound_batch_has_unplanned_items()/
    outbound_batch_has_unscanned_planned_items()를 그대로 재사용한다(8-1절 —
    같은 판정 기준을 두 곳에서 따로 정의하지 않음). 항목이 0건인 빈 배치는 계획이
    없어도 '완료'로 안 친다."""
    conn = get_conn()
    ids = [r[0] for r in conn.execute("""
        SELECT b.id FROM outbound_batches b
         WHERE b.confirmed_at IS NULL
           AND EXISTS (SELECT 1 FROM outbound_items i WHERE i.batch_id = b.id)
    """).fetchall()]
    conn.close()
    return sum(1 for bid in ids
               if not outbound_batch_has_unplanned_items(bid)
               and not outbound_batch_has_unscanned_planned_items(bid))


def list_outbound_batches(query=None, limit=200):
    conn = get_conn()
    body_required = 1 if outbound_body_photo_enabled() else 0
    sql = """
        SELECT b.*, COUNT(i.id) AS item_count,
               (SELECT COUNT(*) FROM outbound_planned_items p WHERE p.batch_id = b.id) AS planned_count,
               (SELECT COUNT(*) FROM outbound_items i2
                 WHERE i2.batch_id = b.id
                   AND NOT EXISTS (SELECT 1 FROM outbound_planned_items p2
                                    WHERE p2.batch_id = i2.batch_id AND p2.serial_no = i2.serial_no)
               ) AS unplanned_count,
               (SELECT COUNT(*) FROM outbound_planned_items p3
                 WHERE p3.batch_id = b.id
                   AND NOT EXISTS (SELECT 1 FROM outbound_items i3
                                    WHERE i3.batch_id = p3.batch_id AND i3.serial_no = p3.serial_no)
               ) AS missing_count,
               -- 2026-09-23 신규: 계획 대비 "확인됨(confirmed)" 개수. 정의는
               -- outbound_plan_progress()의 confirmed 판정과 반드시 일치시킨다
               -- (8-1절 원칙 — 새 정의 금지).
               (SELECT COUNT(*) FROM outbound_planned_items p4
                 WHERE p4.batch_id = b.id
                   AND EXISTS (
                       SELECT 1 FROM outbound_items i4c
                        WHERE i4c.batch_id = p4.batch_id AND i4c.serial_no = p4.serial_no
                          AND i4c.check_tie IS NOT NULL AND i4c.check_qr IS NOT NULL
                          AND i4c.check_wrap IS NOT NULL AND i4c.check_rst IS NOT NULL
                          AND i4c.check_sticker IS NOT NULL AND i4c.check_paint IS NOT NULL
                          AND i4c.check_access IS NOT NULL AND i4c.check_cable IS NOT NULL
                          AND i4c.check_indicator IS NOT NULL
                   )
                   AND (
                       -- 2026-10-08: 제어함(제품명 기준)은 인디케이터 사진 면제 — outbound_plan_progress와 동일
                       EXISTS (
                           SELECT 1 FROM outbound_items i4n
                            WHERE i4n.batch_id = p4.batch_id AND i4n.serial_no = p4.serial_no
                              AND i4n.product_name LIKE '%제어함%'
                       )
                       OR EXISTS (
                           SELECT 1 FROM outbound_items i4p
                            JOIN outbound_item_photos ph4i ON ph4i.item_id = i4p.id
                            WHERE i4p.batch_id = p4.batch_id AND i4p.serial_no = p4.serial_no
                              AND ph4i.kind = 'indicator'
                       )
                   )
                   AND (
                       ? = 0
                       OR EXISTS (
                           SELECT 1 FROM outbound_items i4b
                            JOIN outbound_item_photos ph4b ON ph4b.item_id = i4b.id
                            WHERE i4b.batch_id = p4.batch_id AND i4b.serial_no = p4.serial_no
                              AND ph4b.kind = 'body'
                       )
                   )
               ) AS confirmed_count,
               -- 2026-09-23 신규: 계획이 없는 자유등록 배치용 — 항목(row) 자신이
               -- confirmed 조건을 만족하는지 카운트.
               (SELECT COUNT(*) FROM outbound_items i6
                 WHERE i6.batch_id = b.id
                   AND i6.check_tie IS NOT NULL AND i6.check_qr IS NOT NULL
                   AND i6.check_wrap IS NOT NULL AND i6.check_rst IS NOT NULL
                   AND i6.check_sticker IS NOT NULL AND i6.check_paint IS NOT NULL
                   AND i6.check_access IS NOT NULL AND i6.check_cable IS NOT NULL
                   AND i6.check_indicator IS NOT NULL
                   AND (i6.product_name LIKE '%제어함%'
                        OR EXISTS (SELECT 1 FROM outbound_item_photos ph6
                                    WHERE ph6.item_id = i6.id AND ph6.kind = 'indicator'))
                   AND (
                       ? = 0
                       OR EXISTS (SELECT 1 FROM outbound_item_photos ph6b
                                   WHERE ph6b.item_id = i6.id AND ph6b.kind = 'body')
                   )
               ) AS item_confirmed_count,
               (SELECT GROUP_CONCAT(DISTINCT i4.inspected_by) FROM outbound_items i4
                 WHERE i4.batch_id = b.id AND i4.inspected_by IS NOT NULL AND i4.inspected_by != ''
               ) AS inspectors
          FROM outbound_batches b
          LEFT JOIN outbound_items i ON i.batch_id = b.id
         WHERE 1=1
    """
    params = [body_required, body_required]
    if query:
        sql += " AND (b.customer LIKE ? OR b.handler LIKE ?)"
        params += [f"%{query}%", f"%{query}%"]
    sql += " GROUP BY b.id ORDER BY b.id DESC LIMIT ?"
    params.append(limit)
    rows = conn.execute(sql, params).fetchall()
    conn.close()
    return rows


def delete_outbound_batches(batch_ids):
    """차수(배치)를 통째로 삭제한다. 배치 안의 스캔 항목·사진(DB row)·계획 S/N·
    QR 라벨 출력 이력까지 전부 같이 지운다 — FK 제약(PRAGMA foreign_keys=ON) 때문에
    자식 테이블부터 지워야 한다. 발급된 S/N 자체(finished_goods_serials)는 배치와
    무관한 별개 이력이라 손대지 않는다. 확인 완료(잠금) 여부는 확인하지 않는다 —
    '삭제'는 '수정'과 다른 개념이라 _outbound_batch_lock_response()의 잠금 대상이 아니다.
    반환: 지워진 사진의 실제 파일명 목록(호출부가 디스크에서도 지울 것 —
    delete_outbound_item()과 같은 계약)."""
    conn = get_conn()
    photo_paths = []
    for batch_id in batch_ids:
        rows = conn.execute("""
            SELECT file_path FROM outbound_item_photos
             WHERE item_id IN (SELECT id FROM outbound_items WHERE batch_id=?)
        """, (batch_id,)).fetchall()
        photo_paths += [r["file_path"] for r in rows]
        conn.execute("""
            DELETE FROM outbound_item_photos
             WHERE item_id IN (SELECT id FROM outbound_items WHERE batch_id=?)
        """, (batch_id,))
        conn.execute("DELETE FROM outbound_items WHERE batch_id=?", (batch_id,))
        conn.execute("DELETE FROM outbound_planned_items WHERE batch_id=?", (batch_id,))
        conn.execute("DELETE FROM outbound_qr_exports WHERE batch_id=?", (batch_id,))
        conn.execute("DELETE FROM outbound_batches WHERE id=?", (batch_id,))
    conn.commit()
    conn.close()
    return photo_paths


def add_outbound_item(batch_id, serial_no, product_name, quantity, inspected_by=None):
    conn = get_conn()
    cur = conn.execute("""
        INSERT INTO outbound_items (batch_id, serial_no, product_name, quantity, inspected_by)
        VALUES (?, ?, ?, ?, ?)
    """, (batch_id, serial_no, product_name, quantity, inspected_by))
    conn.commit()
    item_id = cur.lastrowid
    conn.close()
    return item_id


def get_outbound_item(item_id):
    conn = get_conn()
    row = conn.execute("SELECT * FROM outbound_items WHERE id=?", (item_id,)).fetchone()
    conn.close()
    return row


def list_outbound_items(batch_id):
    """배치의 항목 + 각 항목의 사진 목록을 함께 반환.
    반환: [{...item 컬럼.., "photos": [{"id":.., "file_path":.., "uploaded_at":..}, ...]}, ...]"""
    conn = get_conn()
    items = conn.execute(
        "SELECT * FROM outbound_items WHERE batch_id=? ORDER BY id", (batch_id,)).fetchall()
    result = []
    for it in items:
        photos = conn.execute(
            "SELECT * FROM outbound_item_photos WHERE item_id=? ORDER BY id",
            (it["id"],)).fetchall()
        row = dict(it)
        row["photos"] = [dict(p) for p in photos]
        row["result_auto"] = compute_outbound_item_auto_result(row)
        row["result_effective"] = outbound_item_effective_result(row)
        result.append(row)
    conn.close()
    return result


def update_outbound_item(item_id, product_name, quantity):
    conn = get_conn()
    conn.execute(
        "UPDATE outbound_items SET product_name=?, quantity=? WHERE id=?",
        (product_name, quantity, item_id))
    conn.commit()
    conn.close()


def delete_outbound_item(item_id):
    """항목 삭제. 첨부 사진의 실제 파일명 목록을 반환하니 호출부가 파일도 지울 것
    (FK 제약이 있으므로 사진 행 먼저 지우고 항목을 지운다)."""
    conn = get_conn()
    photo_paths = [r["file_path"] for r in conn.execute(
        "SELECT file_path FROM outbound_item_photos WHERE item_id=?", (item_id,)).fetchall()]
    conn.execute("DELETE FROM outbound_item_photos WHERE item_id=?", (item_id,))
    conn.execute("DELETE FROM outbound_items WHERE id=?", (item_id,))
    conn.commit()
    conn.close()
    return photo_paths


def add_outbound_item_photo(item_id, file_path, kind="indicator"):
    conn = get_conn()
    cur = conn.execute(
        "INSERT INTO outbound_item_photos (item_id, file_path, kind) VALUES (?, ?, ?)",
        (item_id, file_path, kind))
    conn.commit()
    photo_id = cur.lastrowid
    conn.close()
    return photo_id


def update_outbound_item_check(item_id, field, value):
    """field는 OUTBOUND_CHECK_FIELDS 안에 있어야 하고 value는 OUTBOUND_RESULT_VALUES
    안에 있어야 한다 — field가 SQL 컬럼명으로 f-string에 그대로 들어가므로, 호출부
    (app.py 라우트)가 반드시 먼저 화이트리스트 검증을 해야 한다. 여기서도 다시 한 번
    검증한다(방어 2중화, 라우트가 실수로 빠뜨려도 SQL 인젝션 경로가 안 되게)."""
    if field not in OUTBOUND_CHECK_FIELDS:
        raise ValueError(f"허용되지 않은 검사항목 필드: {field}")
    if value not in OUTBOUND_RESULT_VALUES:
        raise ValueError(f"허용되지 않은 판정값: {value}")
    conn = get_conn()
    conn.execute(f"UPDATE outbound_items SET {field}=? WHERE id=?", (value, item_id))
    conn.commit()
    conn.close()


def set_outbound_item_result_override(item_id, value):
    """value: 'PASS'/'FAIL'/'SPECIAL' 중 하나 또는 None(오버라이드 해제 → 자동판정 복귀)."""
    if value is not None and value not in OUTBOUND_RESULT_VALUES:
        raise ValueError(f"허용되지 않은 판정값: {value}")
    conn = get_conn()
    conn.execute("UPDATE outbound_items SET result_override=? WHERE id=?", (value, item_id))
    conn.commit()
    conn.close()


def outbound_body_photo_enabled():
    """관리자가 켠 경우에만 True. 기본값은 꺼짐(새 기능이므로 admin이 명시적으로
    켜기 전까지는 스캔 화면에 본체사진 입력 UI를 노출하지 않는다)."""
    return get_setting("outbound_body_photo_enabled", "0") == "1"


def get_outbound_item_photo(photo_id):
    conn = get_conn()
    row = conn.execute("SELECT * FROM outbound_item_photos WHERE id=?", (photo_id,)).fetchone()
    conn.close()
    return row


def delete_outbound_item_photo(photo_id):
    """사진 1장 삭제. 실제 파일명을 반환(호출부가 파일 삭제)."""
    conn = get_conn()
    row = conn.execute(
        "SELECT file_path FROM outbound_item_photos WHERE id=?", (photo_id,)).fetchone()
    conn.execute("DELETE FROM outbound_item_photos WHERE id=?", (photo_id,))
    conn.commit()
    conn.close()
    return row["file_path"] if row else None


# ---------- 2026-09-15 확장: 모델명 자동분류 규칙 ----------

_OUTBOUND_RULE_TABLES = {
    "voltage": "outbound_rule_voltage",
    "suffix":  "outbound_rule_suffix",
    "pcode":   "outbound_rule_pcode",
}

# 리클로저 본체: CKMR{전압코드}K####USA####{접미사}(({P}P))?
_CKMR_RE = re.compile(r"^CKMR(\d)K\d{4}USA\d{4}([A-Z0-9]*)(?:\((\d+)P\))?$")
# 제어함: CKCB####-{꼬리표}
_CKCB_RE = re.compile(r"^CKCB\d+-(.+)$")


def list_classify_rules(kind):
    """kind: 'voltage'/'suffix'/'pcode'. code 오름차순으로 반환."""
    table = _OUTBOUND_RULE_TABLES[kind]
    conn = get_conn()
    rows = conn.execute(f"SELECT code, label FROM {table} ORDER BY code").fetchall()
    conn.close()
    return rows


def upsert_classify_rule(kind, code, label):
    """code가 이미 있으면 라벨만 갱신, 없으면 새로 추가 — 화면의 추가/수정 폼 하나로 겸용."""
    table = _OUTBOUND_RULE_TABLES[kind]
    conn = get_conn()
    conn.execute(
        f"INSERT INTO {table} (code, label) VALUES (?, ?) "
        f"ON CONFLICT(code) DO UPDATE SET label=excluded.label",
        (code, label))
    conn.commit()
    conn.close()


def delete_classify_rule(kind, code):
    table = _OUTBOUND_RULE_TABLES[kind]
    conn = get_conn()
    conn.execute(f"DELETE FROM {table} WHERE code=?", (code,))
    conn.commit()
    conn.close()


def delete_classify_rules_bulk(kind, codes):
    """kind('voltage'/'suffix'/'pcode') 하나의 매핑표에서 여러 code를 한 번에 삭제."""
    table = _OUTBOUND_RULE_TABLES[kind]
    conn = get_conn()
    conn.executemany(f"DELETE FROM {table} WHERE code=?", [(c,) for c in codes])
    conn.commit()
    conn.close()


def get_classify_rules():
    """세 매핑표를 한 번에 읽어 classify_serial_no()에 넘길 dict로 반환한다.
    여러 S/N을 한 번에 분류할 때(QR 라벨 엑셀 출력 등) 매번 DB를 다시 열지 않도록
    호출부가 이 함수를 한 번만 불러서 재사용할 수 있게 만든 캐시 겸용 함수다."""
    conn = get_conn()
    voltage = {r["code"]: r["label"] for r in
               conn.execute("SELECT code, label FROM outbound_rule_voltage").fetchall()}
    suffix_rows = conn.execute("SELECT code, label FROM outbound_rule_suffix").fetchall()
    pcode = {r["code"]: r["label"] for r in
             conn.execute("SELECT code, label FROM outbound_rule_pcode").fetchall()}
    conn.close()
    # 접미사는 긴 코드부터 검사해야 한다 — "HAT3"가 "HA"/"H"보다 먼저 매칭돼야
    # "CKMR8K0803USA6622HAT3"의 접미사가 "HA"로 잘못 분류되는 걸 막을 수 있다.
    suffix = sorted(((r["code"], r["label"]) for r in suffix_rows), key=lambda kv: -len(kv[0]))
    return {"voltage": voltage, "suffix": suffix, "pcode": pcode}


def _match_ckmr(serial_no, rules):
    """CKMR 리클로저 본체 패턴 매칭 + 접미사 코드 결정. classify_serial_no()와
    resolve_sticker_reference()가 공유한다(8-1절 원칙, 2026-09-21 추출).
    매칭 실패(정규식 자체가 안 맞음) 시 None. 성공 시 dict:
      volt_code: 전압코드 1자리 문자열
      suffix_raw: S/N에 실제로 붙은 접미사 원문(빈 문자열 가능)
      suffix_code: 매핑표에서 매칭된 코드. suffix_raw가 없으면 "", 매핑표에
                   없는 접미사면 None(=인식 실패), 매칭되면 그 코드 문자열
      pcode: 괄호 안 숫자 문자열 또는 None"""
    m = _CKMR_RE.match((serial_no or "").strip())
    if not m:
        return None
    volt_code, suffix_raw, pcode = m.group(1), m.group(2), m.group(3)
    suffix_code = ""
    if suffix_raw:
        suffix_code = None
        for code, label in rules["suffix"]:
            if suffix_raw.startswith(code):
                suffix_code = code
                break
    return {"volt_code": volt_code, "suffix_raw": suffix_raw,
            "suffix_code": suffix_code, "pcode": pcode}


# 2026-09-21 확정값(사용자 확인 완료) — 스티커 부착 기준 참고자료의 21개 CKMR 모델과
# 기존 출고 분류 규칙(전압코드/접미사코드/P코드)의 매칭표. resolve_sticker_reference()가
# S/N에서 뽑아낸 (volt_code, suffix_code, pcode)와 이 표의 값을 비교해서 모델을 찾는다.
# 이 표를 벗어난 model_no는 import_sticker_reference_excel()이 새로 발견했을 때
# 코드 3개를 NULL로 남기고, 관리자가 /outbound/rules 화면에서 수동으로 채운다.
STICKER_MODEL_MATCH_KEYS = {
    "CKMR1070": ("7", "S", None),
    "CKMR1080": ("7", "S", "19"),
    "CKMR1200": ("8", "S", "19"),
    "CKMR1280": ("8", "S", None),
    "CKMR3000": ("7", "HT3", None),
    "CKMR3010": ("7", "HAT3", None),
    "CKMR3050": ("7", "HT1", None),
    "CKMR3060": ("7", "HAT1", None),
    "CKMR3080": ("7", "HT3", "32"),
    "CKMR3500": ("8", "HAT3", None),
    "CKMR3510": ("8", "HT3", None),
    "CKMR3580": ("8", "HT1", None),
    "CKMR3590": ("8", "HT3", "32"),
    "CKMR4000": ("7", "H", None),
    "CKMR4020": ("7", "H", "42"),
    "CKMR4050": ("7", "HA", None),
    "CKMR6050": ("7", "VA", None),
    "CKMR8800": ("8", "H", None),
    "CKMR8830": ("8", "H", "42"),
    "CKMR8880": ("8", "HA", None),
    "CKMR9880": ("9", "H", None),
}

STICKER_REF_NOT_CKMR = "not_ckmr"
STICKER_REF_UNRECOGNIZED_CODE = "unrecognized_code"
STICKER_REF_NO_MODEL = "no_registered_model"


def ensure_sticker_reference_match_seed():
    """STICKER_MODEL_MATCH_KEYS 21건을 sticker_reference_models에 심는다. items/image_path는
    비워두고(관리자 업로드로만 채움, Task 6) 코드 3개만 채워서 resolve_sticker_reference()가
    업로드 이전에도 모델 식별은 가능하게 한다. INSERT OR IGNORE라 이미 있는 행(업로드로
    새로 생겼거나 관리자가 수동으로 채운 행 포함)은 절대 덮어쓰지 않는다. 삭제 기능이
    있으므로(Task 6) 매 재시작마다 다시 실행돼도 무해하게 설계 — 별도 settings 플래그 불필요."""
    conn = get_conn()
    for model_no, (volt, suf, pc) in STICKER_MODEL_MATCH_KEYS.items():
        conn.execute(
            "INSERT OR IGNORE INTO sticker_reference_models "
            "(model_no, voltage_code, suffix_code, pcode) VALUES (?, ?, ?, ?)",
            (model_no, volt, suf, pc))
    conn.commit()
    conn.close()


def bom_model_label(model_no):
    """material_bom_links.model_name(예: "CKMR1070 본체 15kV 단상")에서 이 model_no로
    시작하는 한글 표시명을 찾는다(18절 — assembly_masters와 무관). 못 찾으면 model_no
    자체를 돌려준다(값 유실 방지, 8-3절과 같은 원칙). 표시명은 별도 컬럼에 저장하지
    않고 항상 이 함수로 조회한다 — material_bom_links가 갱신되면 자동으로 최신 반영됨."""
    conn = get_conn()
    row = conn.execute(
        "SELECT model_name FROM material_bom_links "
        "WHERE model_name = ? OR model_name LIKE ? ORDER BY model_name LIMIT 1",
        (model_no, model_no + " %")).fetchone()
    conn.close()
    return row["model_name"] if row else model_no


def _sticker_reference_items(model_no):
    conn = get_conn()
    rows = conn.execute(
        "SELECT no, label, attach_face, qty FROM sticker_reference_items "
        "WHERE model_no=? ORDER BY sort_order", (model_no,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def _sticker_reference_model_dict(row):
    """sticker_reference_models 행(sqlite3.Row) -> API/화면 공용 dict.
    resolve/search/list 세 곳 모두 이 함수 하나로 통일한다(8-1절)."""
    model_no = row["model_no"]
    return {
        "model_no": model_no,
        "display_name": bom_model_label(model_no),
        "voltage_code": row["voltage_code"] or "",
        "suffix_code": row["suffix_code"] or "",
        "pcode": row["pcode"] or "",
        "image_url": f"/static/sticker_reference/{row['image_path']}" if row["image_path"] else None,
        "has_image": bool(row["image_path"]),
        "items": _sticker_reference_items(model_no),
        "item_count": len(_sticker_reference_items(model_no)),
    }


def resolve_sticker_reference(serial_no, rules=None):
    """S/N으로 스티커 부착 기준 모델을 자동 매칭. 반환:
      성공: {"ok": True, "model": {...}}  (_sticker_reference_model_dict 형태)
      실패: {"ok": False, "reason": STICKER_REF_*, "message": "..."}"""
    serial_no = (serial_no or "").strip()
    if rules is None:
        rules = get_classify_rules()

    match = _match_ckmr(serial_no, rules)
    if match is None:
        return {"ok": False, "reason": STICKER_REF_NOT_CKMR,
                "message": "이 항목은 리클로저 본체가 아니라서 자동 매칭 대상이 아니야."}

    volt_label = rules["voltage"].get(match["volt_code"])
    if volt_label is None or (match["suffix_raw"] and match["suffix_code"] is None):
        return {"ok": False, "reason": STICKER_REF_UNRECOGNIZED_CODE,
                "message": "S/N의 전압코드 또는 접미사를 인식하지 못했어."}

    conn = get_conn()
    row = conn.execute(
        "SELECT * FROM sticker_reference_models WHERE voltage_code=? AND suffix_code=? AND pcode IS ?",
        (match["volt_code"], match["suffix_code"] or "", match["pcode"])).fetchone()
    conn.close()
    if row is None:
        label = classify_serial_no(serial_no, rules) or volt_label
        return {"ok": False, "reason": STICKER_REF_NO_MODEL,
                "message": f"'{label}' 조합은 아직 등록된 도면이 없어."}

    return {"ok": True, "model": _sticker_reference_model_dict(row)}


def search_sticker_reference_models(query):
    """model_no 또는 BOM 한글 표시명으로 검색(자동매칭 실패 시 폴백 UI용). 행이 최대
    수십 건뿐이라 전체 스캔 후 파이썬에서 필터링한다(과설계 방지)."""
    conn = get_conn()
    rows = conn.execute("SELECT * FROM sticker_reference_models ORDER BY model_no").fetchall()
    conn.close()
    q = (query or "").strip().lower()
    results = []
    for row in rows:
        d = _sticker_reference_model_dict(row)
        if not q or q in row["model_no"].lower() or q in d["display_name"].lower():
            results.append(d)
    return results


def list_sticker_reference_models():
    """관리 화면(/outbound/rules 카드)용 전체 목록. _sticker_reference_model_dict()를
    그대로 재사용 — search와 다른 shape을 새로 만들지 않는다(8-1절)."""
    conn = get_conn()
    rows = conn.execute("SELECT * FROM sticker_reference_models ORDER BY model_no").fetchall()
    conn.close()
    return [_sticker_reference_model_dict(r) for r in rows]


def set_sticker_reference_codes(model_no, voltage_code, suffix_code, pcode):
    """관리자가 수동으로 매칭 코드 3개를 채우거나 고칠 때(신규 발견 모델 대응).
    빈 문자열은 NULL로 저장(=아직 미분류)."""
    conn = get_conn()
    conn.execute(
        "UPDATE sticker_reference_models SET voltage_code=?, suffix_code=?, pcode=?, "
        "updated_at=datetime('now','localtime') WHERE model_no=?",
        (voltage_code or None, suffix_code or None, pcode or None, model_no))
    conn.commit()
    conn.close()


def create_sticker_reference_model(model_no, voltage_code="", suffix_code="", pcode=""):
    """개별 신규 모델 등록 — model_no만 갖고 빈 모델(항목 없음, 이미지 없음)을 하나 만든다.
    이미 있는 model_no면 아무것도 안 하고 False를 돌려준다(엑셀 재업로드/코드수정 폼으로
    유도 — upsert로 조용히 덮어쓰면 기존 항목표를 실수로 날릴 위험이 있어서 막음).
    표시명은 별도 저장 안 함 — bom_model_label()이 항상 조회해서 계산한다."""
    model_no = (model_no or "").strip()
    if not model_no:
        return False
    conn = get_conn()
    exists = conn.execute(
        "SELECT 1 FROM sticker_reference_models WHERE model_no=?", (model_no,)).fetchone()
    if exists:
        conn.close()
        return False
    conn.execute(
        "INSERT INTO sticker_reference_models (model_no, voltage_code, suffix_code, pcode) "
        "VALUES (?, ?, ?, ?)",
        (model_no, (voltage_code or "").strip() or None,
         (suffix_code or "").strip() or None, (pcode or "").strip() or None))
    conn.commit()
    conn.close()
    return True


def delete_sticker_reference_model(model_no):
    """모델 삭제 — sticker_reference_items(자식)부터 지우고 sticker_reference_models(부모)를
    지운다(SQLite는 기본 FK cascade가 없어 수동으로 순서를 지켜야 함). 이미지 파일은
    여기서 안 지운다 — 파일 경로는 app.py(STICKER_REFERENCE_DIR)만 알고 있으므로 실제
    os.remove는 호출부(app.py) 책임. 반환: 실제로 지웠으면 True, 없던 model_no면 False."""
    conn = get_conn()
    exists = conn.execute(
        "SELECT 1 FROM sticker_reference_models WHERE model_no=?", (model_no,)).fetchone()
    if not exists:
        conn.close()
        return False
    conn.execute("DELETE FROM sticker_reference_items WHERE model_no=?", (model_no,))
    conn.execute("DELETE FROM sticker_reference_models WHERE model_no=?", (model_no,))
    conn.commit()
    conn.close()
    return True


def import_sticker_reference_from_excel(filepath, image_dir):
    """스티커 부착 기준 참고파일("모델별 도면" 시트)을 파싱해서 sticker_reference_models/
    sticker_reference_items를 갱신하고 참고이미지를 image_dir에 저장한다.

    블록 구조(34행 간격, 첫 블록은 엑셀 1행부터):
      오프셋0: A=모델번호(CKMRxxxx)
      오프셋2~5: A=No. B=스티커명 C=부착 면 D=수량 (4행, 라벨 비어있으면 스킵)
      오프셋7: 참고이미지 1장 앵커(있으면 저장, 없으면 기존 이미지 보존)

    voltage_code/suffix_code/pcode는 이 파일에서 읽지 않는다 — 기존 행이 이미 채워져
    있으면 절대 덮어쓰지 않고(관리자 수동편집 보존), 새 model_no면 STICKER_MODEL_MATCH_KEYS에
    있는 값을 쓰거나 없으면 NULL로 남긴다.

    반환: {"imported": 갱신된 모델 수, "no_image": 이미지를 못 찾은 모델 수,
           "new_unmatched": 21개표에 없는 새 모델 수(코드 미배정, 관리자가 채워야 함)}"""
    import openpyxl
    os.makedirs(image_dir, exist_ok=True)

    wb = openpyxl.load_workbook(filepath, data_only=True)
    try:
        ws = wb["모델별 도면"]
        images_by_row = {}
        for img in ws._images:
            images_by_row[img.anchor._from.row] = img

        summary = {"imported": 0, "no_image": 0, "new_unmatched": 0}
        conn = get_conn()
        cur = conn.cursor()

        block_i = 0
        while block_i < 200:
            start0 = block_i * 34
            model_no = ws.cell(row=start0 + 1, column=1).value
            if model_no is None or str(model_no).strip() == "":
                break
            model_no = str(model_no).strip()

            items = []
            for r_off in range(2, 6):
                row1 = start0 + r_off + 1
                no_val = ws.cell(row=row1, column=1).value
                label = ws.cell(row=row1, column=2).value
                face = ws.cell(row=row1, column=3).value
                qty = ws.cell(row=row1, column=4).value
                if label is None or str(label).strip() == "":
                    continue
                items.append((
                    str(no_val).strip() if no_val is not None else "",
                    str(label).strip(),
                    str(face).strip() if face is not None else "",
                    str(qty).strip() if qty is not None else ""))

            existing = cur.execute(
                "SELECT voltage_code, suffix_code, pcode, image_path FROM sticker_reference_models WHERE model_no=?",
                (model_no,)).fetchone()

            if existing and (existing["voltage_code"] or existing["suffix_code"] or existing["pcode"]):
                voltage_code, suffix_code, pcode = existing["voltage_code"], existing["suffix_code"], existing["pcode"]
            else:
                keys = STICKER_MODEL_MATCH_KEYS.get(model_no)
                if keys:
                    voltage_code, suffix_code, pcode = keys
                else:
                    voltage_code, suffix_code, pcode = None, None, None
                    summary["new_unmatched"] += 1

            image_path = existing["image_path"] if existing else None
            img = images_by_row.get(start0 + 7)
            if img is not None:
                try:
                    img_bytes = img._data()
                except Exception:
                    img_bytes = None
                if img_bytes:
                    # 경로 이탈 방지 — model_no는 엑셀 셀 값이라 슬래시/역슬래시/".." 등이
                    # 섞여 들어올 수 있으므로 파일명에 안전한 문자만 남긴다(quality-watcher 지적).
                    safe_model_no = re.sub(r"[^A-Za-z0-9_-]", "", model_no)
                    fname = f"{safe_model_no}.png"
                    with open(os.path.join(image_dir, fname), "wb") as f:
                        f.write(img_bytes)
                    image_path = fname
            if image_path is None:
                summary["no_image"] += 1

            cur.execute("""
                INSERT INTO sticker_reference_models
                    (model_no, voltage_code, suffix_code, pcode, image_path, updated_at)
                VALUES (?, ?, ?, ?, ?, datetime('now','localtime'))
                ON CONFLICT(model_no) DO UPDATE SET
                    voltage_code=excluded.voltage_code, suffix_code=excluded.suffix_code,
                    pcode=excluded.pcode, image_path=excluded.image_path, updated_at=excluded.updated_at
            """, (model_no, voltage_code, suffix_code, pcode, image_path))

            cur.execute("DELETE FROM sticker_reference_items WHERE model_no=?", (model_no,))
            for order, (no_val, label, face, qty) in enumerate(items, start=1):
                cur.execute(
                    "INSERT INTO sticker_reference_items (model_no, no, label, attach_face, qty, sort_order) "
                    "VALUES (?, ?, ?, ?, ?, ?)", (model_no, no_val, label, face, qty, order))

            summary["imported"] += 1
            block_i += 1

        conn.commit()
        conn.close()
    finally:
        wb.close()

    return summary


def classify_serial_no(serial_no, rules=None):
    """S/N 문자열로 제품 분류 라벨을 계산한다.

    - 리클로저 본체: CKMR{전압코드}K####USA####{접미사}(({P}P))?
      예: CKMR7K0902USA5771H(42P) -> "15kV 일반 수평 (42P)"
          CKMR8K0803USA6622HAT3(32P) -> "27kV 트리플 3핸들 앵글 (32P) 155V"
          (P코드가 특수값표에 있으면 "(NNP)" 뒤에 그 특수값도 덧붙인다)
          CKMR9K0912USA8775H -> "38kV 일반 수평"
    - 제어함: CKCB####-{꼬리표}  예: CKCB2758-RA -> "제어함 - RA"
    - 둘 다 매칭 안 되거나, 매칭은 됐는데 전압코드/접미사가 매핑표에 없으면 None(미분류).

    rules를 안 넘기면 이 함수가 직접 get_classify_rules()로 조회한다(단건 호출용).
    여러 건을 한 번에 분류할 때는 호출부가 get_classify_rules()를 한 번만 불러 넘길 것
    (매번 DB를 다시 여는 낭비를 막기 위함)."""
    serial_no = (serial_no or "").strip()
    if not serial_no:
        return None
    if rules is None:
        rules = get_classify_rules()

    match = _match_ckmr(serial_no, rules)
    if match:
        volt_label = rules["voltage"].get(match["volt_code"])
        if volt_label is None:
            return None
        if match["suffix_raw"] and match["suffix_code"] is None:
            return None  # 접미사가 있는데 매핑표에 없으면 분류 불가(미확정 접미사)
        suffix_label = None
        if match["suffix_code"]:
            for code, label in rules["suffix"]:
                if code == match["suffix_code"]:
                    suffix_label = label
                    break
        parts = [volt_label]
        if suffix_label:
            parts.append(suffix_label)
        label = " ".join(parts)
        pcode = match["pcode"]
        if pcode:
            label += f" ({pcode}P)"
            special = rules["pcode"].get(pcode)
            if special:
                label += f" {special}"
        return label

    m2 = _CKCB_RE.match(serial_no)
    if m2:
        return f"제어함 - {m2.group(1)}"

    return None


# ---------- 2026-09-15 확장: 차수(=배치) 계획 항목 ----------

def add_planned_items_bulk(batch_id, serial_nos):
    """차수 계획 S/N을 한 번에 등록. 빈 값과, 이미 이 배치에 등록된 S/N은 건너뛴다.
    반환: 실제로 새로 추가된 건수."""
    conn = get_conn()
    seen = {r["serial_no"] for r in conn.execute(
        "SELECT serial_no FROM outbound_planned_items WHERE batch_id=?", (batch_id,)).fetchall()}
    added = 0
    for sn in serial_nos:
        sn = (sn or "").strip()
        if not sn or sn in seen:
            continue
        seen.add(sn)
        conn.execute(
            "INSERT INTO outbound_planned_items (batch_id, serial_no) VALUES (?, ?)",
            (batch_id, sn))
        added += 1
    conn.commit()
    conn.close()
    return added


def list_planned_items(batch_id):
    conn = get_conn()
    rows = conn.execute(
        "SELECT * FROM outbound_planned_items WHERE batch_id=? ORDER BY id", (batch_id,)).fetchall()
    conn.close()
    return rows


def outbound_plan_progress(batch_id):
    """차수 하나의 "계획 대비 스캔 진행상황"을 계산한다.

    반환: {"rows": [{"serial_no", "model_label", "status"}, ...],
           "summary": {"planned_total", "matched", "extra"}}
    status: 'confirmed'(스캔됨 + 필요 사진 전부 있음 + 9개 품질확인항목 전부 선택됨) /
    'incomplete'(스캔은 됐는데 사진 또는 품질확인이 부족함) / 'pending'(아직 안 스캔됨).

    본체사진이 "필요 사진"에 들어가는지는 이 함수를 호출하는 시점의
    outbound_body_photo_enabled() 값을 따른다(토글이 꺼져 있으면 인디케이터
    사진만 있으면 됨 — 애초에 입력받지 않는 항목을 조건에 넣으면 영원히 확인
    불가능해지기 때문). 품질확인 5종은 2026-09-15 신규 — 5개 항목이 전부 선택된
    항목이 하나라도 있으면 checks_ok로 본다(같은 S/N이 여러 번 스캔된 드문 경우,
    사진 병합과 같은 방식으로 처리)."""
    items = list_outbound_items(batch_id)
    planned = list_planned_items(batch_id)
    body_required = outbound_body_photo_enabled()
    rules = get_classify_rules()

    items_by_serial = {}
    for it in items:
        items_by_serial.setdefault(it["serial_no"], []).append(it)
    planned_serials = {p["serial_no"] for p in planned}

    rows = []
    matched = 0
    for p in planned:
        its = items_by_serial.get(p["serial_no"]) or []
        scanned = bool(its)
        photos_ok = False
        checks_ok = False
        model_label = classify_serial_no(p["serial_no"], rules=rules)
        if scanned:
            kinds = {ph.get("kind") or "indicator" for it in its for ph in (it.get("photos") or [])}
            # 2026-10-07: 제어함은 인디케이터 사진을 찍지 않으므로 필수 조건에서 뺀다
            # 입력된 제품명 기준(화면 JS와 동일) — 자동분류 라벨은 보지 않는다
            control = any("제어함" in (it.get("product_name") or "") for it in its)
            photos_ok = (control or "indicator" in kinds) and (not body_required or "body" in kinds)
            checks_ok = any(compute_outbound_item_auto_result(it) is not None for it in its)
        if scanned and photos_ok and checks_ok:
            status = "confirmed"
            matched += 1
        elif scanned:
            status = "incomplete"
        else:
            status = "pending"
        rows.append({
            "serial_no": p["serial_no"],
            "model_label": model_label,
            "status": status,
            # 미비 사유 — 화면 배지 문구("사진 미비"/"품질확인 미비")용
            "missing_photos": scanned and not photos_ok,
            "missing_checks": scanned and not checks_ok,
        })

    extra = sum(1 for it in items if it["serial_no"] not in planned_serials)
    return {"rows": rows, "summary": {"planned_total": len(planned), "matched": matched, "extra": extra}}


def planned_item_exists(batch_id, serial_no):
    conn = get_conn()
    row = conn.execute(
        "SELECT 1 FROM outbound_planned_items WHERE batch_id=? AND serial_no=?",
        (batch_id, serial_no)).fetchone()
    conn.close()
    return row is not None


def outbound_batch_has_unplanned_items(batch_id):
    """차수 계획이 있는데(planned_items 1건 이상) 그 계획에 없는 항목이 이 배치에
    하나라도 있으면 True. 계획 자체가 없는 배치(자유 등록)는 애초에 "계획에 안
    맞는다"는 개념이 성립하지 않으므로 False — outbound_item_add()의 추가 차단
    로직과 같은 전제를 쓴다(2026-09-16, 사용자 확정: 계획 불일치 항목은 추가 자체를
    막고, 이미 들어와 있는 계획외 항목이 있으면 출고 확인도 막는다). outbound_batch_confirm()
    라우트가 이 함수로 확인 처리를 게이트한다."""
    conn = get_conn()
    planned_count = conn.execute(
        "SELECT COUNT(*) FROM outbound_planned_items WHERE batch_id=?", (batch_id,)).fetchone()[0]
    if planned_count == 0:
        conn.close()
        return False
    unplanned_count = conn.execute("""
        SELECT COUNT(*) FROM outbound_items i
         WHERE i.batch_id=? AND NOT EXISTS (
             SELECT 1 FROM outbound_planned_items p WHERE p.batch_id=i.batch_id AND p.serial_no=i.serial_no)
    """, (batch_id,)).fetchone()[0]
    conn.close()
    return unplanned_count > 0


def outbound_batch_has_unscanned_planned_items(batch_id):
    """차수 계획(outbound_planned_items)에 등록된 S/N 중 아직 스캔(outbound_items)되지
    않은 게 하나라도 있으면 True. 계획 자체가 없으면(자유 등록 배치) False —
    outbound_batch_has_unplanned_items()와 같은 전제(2026-09-21 사용자 확정: '계획에는
    있는데 아직 안 채워진' 반대 방향도 출고 확인을 막는다). outbound_batch_confirm()
    라우트가 이 함수로 확인 처리를 게이트한다."""
    conn = get_conn()
    planned_count = conn.execute(
        "SELECT COUNT(*) FROM outbound_planned_items WHERE batch_id=?", (batch_id,)).fetchone()[0]
    if planned_count == 0:
        conn.close()
        return False
    missing_count = conn.execute("""
        SELECT COUNT(*) FROM outbound_planned_items p
         WHERE p.batch_id=? AND NOT EXISTS (
             SELECT 1 FROM outbound_items i WHERE i.batch_id=p.batch_id AND i.serial_no=p.serial_no)
    """, (batch_id,)).fetchone()[0]
    conn.close()
    return missing_count > 0


def delete_planned_item(planned_id):
    conn = get_conn()
    conn.execute("DELETE FROM outbound_planned_items WHERE id=?", (planned_id,))
    conn.commit()
    conn.close()


# ---------- 2026-09-15 확장: QR 라벨 출력 이력 ----------

def record_qr_export(batch_id, generated_by, item_count):
    conn = get_conn()
    cur = conn.execute("""
        INSERT INTO outbound_qr_exports (batch_id, generated_by, item_count)
        VALUES (?, ?, ?)
    """, (batch_id, generated_by, item_count))
    conn.commit()
    export_id = cur.lastrowid
    conn.close()
    return export_id


def list_qr_exports(query=None, limit=200):
    conn = get_conn()
    sql = """
        SELECT e.*, b.round_no, b.customer, b.ship_date
          FROM outbound_qr_exports e
          JOIN outbound_batches b ON b.id = e.batch_id
         WHERE 1=1
    """
    params = []
    if query:
        sql += " AND (b.customer LIKE ? OR b.round_no LIKE ?)"
        params += [f"%{query}%", f"%{query}%"]
    sql += " ORDER BY e.id DESC LIMIT ?"
    params.append(limit)
    rows = conn.execute(sql, params).fetchall()
    conn.close()
    return rows


def delete_qr_exports(export_ids):
    """QR 라벨 출력 이력을 여러 건 한 번에 삭제. 실제 출력 파일은 디스크에 저장 안 하고
    매번 즉석 생성해서 다운로드만 시키는 방식이라 DB row만 지우면 된다."""
    conn = get_conn()
    conn.executemany("DELETE FROM outbound_qr_exports WHERE id=?",
                      [(eid,) for eid in export_ids])
    conn.commit()
    conn.close()


# ---------- 성적서 위변조 검증 ----------

def set_inspection_hashes(inspection_id, content_hash=None, pdf_hash=None):
    """None으로 준 값은 기존 값을 그대로 둔다(승인 시 content, 출력 시 pdf 를 따로 채우기 위함)."""
    conn = get_conn()
    if content_hash is not None:
        conn.execute("UPDATE inspections SET content_hash=? WHERE id=?", (content_hash, inspection_id))
    if pdf_hash is not None:
        conn.execute("UPDATE inspections SET pdf_hash=? WHERE id=?", (pdf_hash, inspection_id))
    conn.commit()
    conn.close()


# ---------- 품질 현황 집계 (대시보드 / 보고서) ----------

PERIOD_TYPES = [
    ("daily",     "일간"),
    ("weekly",    "주간"),
    ("monthly",   "월간"),
    ("quarterly", "분기"),
    ("half",      "반기"),
    ("yearly",    "연간"),
]


def _period_key(date_str, period_type):
    """검사일을 기간 유형에 맞는 묶음 키로 바꾼다. 못 읽는 날짜는 None."""
    from datetime import date as _date
    text = str(date_str or "")[:10]
    try:
        y, m, d = int(text[0:4]), int(text[5:7]), int(text[8:10])
        dt = _date(y, m, d)
    except (ValueError, IndexError):
        return None
    if period_type == "daily":
        return f"{dt:%Y-%m-%d}"
    if period_type == "weekly":
        iso_y, iso_w, _ = dt.isocalendar()
        return f"{iso_y}-W{iso_w:02d}"
    if period_type == "monthly":
        return f"{dt:%Y-%m}"
    if period_type == "quarterly":
        return f"{dt.year}-Q{(dt.month - 1) // 3 + 1}"
    if period_type == "half":
        return f"{dt.year}-H{1 if dt.month <= 6 else 2}"
    return str(dt.year)


def _lot_state(status, approval_type, overall_result=None):
    """로트 하나의 최종 상태. 불량률은 '판정이 확정된 것'만으로 계산한다.

    superseded(재검사로 대체된 옛 성적서)는 집계에서 아예 빼야 한다.
    같은 입고 건을 재검사하면 성적서가 하나 더 생기는데, 옛 건까지 세면
    같은 로트의 수량이 두 번 잡혀서 검사 수량이 부풀려진다.

    overall_result 인자(2026-09-14, 항목 판정 수동 오버라이드 기능과 함께 도입):
    - 넘기면(집계·통계용 호출) 항목 오버라이드가 반영된 로트 상태를 돌려준다 —
      '합격승인(normal)'인데 오버라이드로 overall_result가 '합격'이 아니게 됐으면 '불합격'으로 재분류.
    - 안 넘기면(기본값 None, 공식 출력물용 호출) 예전 그대로 status+approval_type만 본다.
      공식 출력물(커스텀 자유양식 성적서 등)을 만드는 호출부는 절대 이 인자를 넘기면 안 된다.
    - 특채/불합격확정 로트는 오버라이드가 나중에 뭐가 되든 재분류하지 않는다
      (이미 서명받은 최종 결정을 오버라이드로 되돌리지 않기 위함).
    """
    if status == "superseded":
        return "대체됨"
    if status == "approved":
        if approval_type == "failed":
            return "불합격"
        if approval_type == "special":
            return "특채"
        if overall_result is not None and overall_result != "합격":
            return "불합격"
        return "합격"
    return "미결"      # pending / rejected


def _ncr_extra_stats(rows, manual_ncr_lot_qtys):
    """②③(NCR포함 불량률/건수율, 2026-09-15 사용자 확정)에 얹을 'NCR만 있는 로트'의
    수량 합·건수를 계산한다. quality_report()에서만 쓰는 헬퍼.

    이미 baseline(①, 불합격 확정)에서 잡힌 로트는 여기서 또 세면 이중집계이므로
    제외한다(8-2-9절 '재검사한 옛 성적서 집계 제외'와 같은 정신). 합격/특채/미결
    로트인데 발송·확인(status가 'sent' 또는 'confirmed') NCR이 딸려있으면 '실질
    불량이 이미 보고된 로트'로 보고 분자·분모 양쪽에 더한다.

    rows: quality_report()가 이미 superseded 제외 필터를 거친 inspections dict 리스트.
          각 dict에 'ncr_sent_count'(이 로트에 연결된 sent/confirmed NCR 개수,
          0 이상 정수), 'status', 'approval_type', 'overall_result', 'quantity',
          'id' 키가 있어야 한다.
    manual_ncr_lot_qtys: 성적서 미연결(inspection_id IS NULL) sent/confirmed NCR의
          lot_qty 문자열 리스트(예: "12대", "200"). 숫자 추출 실패한 건은 이 함수
          안에서 스킵한다.

    반환: (추가수량, 추가건수)
    """
    qty = 0
    cnt = 0
    for r in rows:
        if not r["ncr_sent_count"]:
            continue
        state = _lot_state(r["status"], r["approval_type"], r["overall_result"])
        if state == "불합격":
            continue  # 이미 ①(baseline)에서 분자·분모 둘 다 잡혀있음 — 이중집계 방지
        qty += int(r["quantity"] or 0)
        cnt += 1
    for lot_qty in manual_ncr_lot_qtys:
        m = re.search(r"\d+", str(lot_qty or ""))
        if not m:
            continue  # 수량 파싱 실패 — 근거 없는 숫자를 만들지 않기 위해 건수도 같이 스킵
        qty += int(m.group())
        cnt += 1
    return qty, cnt


def _recorded_sample_count(measured_value):
    """저장된 측정값 문자열에서 실제 기록 개수만 센다.

    과거 검사에 AQL 스냅샷이 없을 때 현재 specs로 표본수를 다시 계산하면 당시 이력이
    바뀌므로, 외관 항목에 실제로 남아 있는 값만 근거로 사용한다.
    """
    return len([v for v in str(measured_value or "").split(",") if v.strip()])


def _preferred_quality_document(rows):
    """확정/발송 문서를 먼저, 같은 우선순위에서는 최신 id를 고른다."""
    if not rows:
        return None
    return max(rows, key=lambda r: (
        1 if (r.get("status") or "") in ("confirmed", "sent") else 0,
        int(r.get("id") or 0),
    ))


def quality_status_export_rows(start_date, end_date, supplier=None, po_number=None,
                               material=None, states=None, major_category=None,
                               middle_category=None):
    """품질현황 엑셀 전용 원본 행을 DB에서 직접 조회한다.

    화면용 ``quality_report()``와 분리해 과거 검사 표본수를 현재 규격으로 재계산하지 않는다.
    일반 표본수는 외관 항목의 저장 스냅샷, 구형 검사는 외관 항목의 실제 기록 개수만 쓴다.
    불량 문서 수량은 NCR → 개선요청서 → 일반 외관 근거 순서로 적용한다.
    """
    def as_list(value):
        if value is None:
            return []
        if isinstance(value, (list, tuple, set)):
            return [str(v).strip() for v in value if str(v).strip()]
        value = str(value).strip()
        return [value] if value else []

    suppliers = as_list(supplier)
    po_numbers = as_list(po_number)
    states = as_list(states)
    conn = get_conn()
    sql = """
        SELECT i.*, mc.major_category, mc.middle_category
          FROM inspections i
          LEFT JOIN material_classification_map mc ON mc.material_no = i.material_no
         WHERE i.status != 'superseded'
    """
    params = []
    if start_date:
        sql += " AND i.inspect_date >= ?"; params.append(start_date)
    if end_date:
        sql += " AND i.inspect_date <= ?"; params.append(end_date)
    if suppliers:
        sql += f" AND i.supplier IN ({','.join('?' * len(suppliers))})"; params += suppliers
    if po_numbers:
        sql += f" AND i.po_number IN ({','.join('?' * len(po_numbers))})"; params += po_numbers
    if material:
        sql += " AND (i.material_no LIKE ? OR i.material_name LIKE ?)"
        params += [f"%{material}%", f"%{material}%"]
    major_category = (major_category or "").strip()
    middle_category = (middle_category or "").strip()
    if major_category == UNASSIGNED_CLASSIFICATION or middle_category == UNASSIGNED_CLASSIFICATION:
        sql += " AND mc.material_no IS NULL"
    else:
        if major_category:
            sql += " AND mc.major_category = ?"; params.append(major_category)
        if middle_category:
            sql += " AND mc.middle_category = ?"; params.append(middle_category)
    sql += " ORDER BY i.inspect_date, i.id"
    lots = [dict(r) for r in conn.execute(sql, params).fetchall()]
    for row in lots:
        row["판정"] = _lot_state(row["status"], row["approval_type"], row["overall_result"])
    if states:
        lots = [row for row in lots if row["판정"] in states]

    ids = [row["id"] for row in lots]
    exterior_by_inspection = {}
    ncr_by_inspection = {}
    improvement_by_inspection = {}
    if ids:
        placeholders = ",".join("?" * len(ids))
        exterior_rows = conn.execute(f"""
            SELECT ii.id, ii.inspection_id, ii.measured_value, ii.aql_snapshot,
                   ii.sample_qty_snapshot, ii.aql_snapshot_set, ii.item_name,
                   COALESCE(s.spec_display, '') AS spec_display,
                   COALESCE(s.item_order, 999999) AS spec_item_order
              FROM inspection_items ii
              JOIN inspections i ON i.id = ii.inspection_id
              LEFT JOIN specs s
                ON s.material_no = COALESCE(ii.part_material_no, i.material_no)
               AND s.item_name = ii.item_name
             WHERE ii.inspection_id IN ({placeholders})
               AND (TRIM(ii.item_name) LIKE '%외관%'
                    OR TRIM(COALESCE(s.spec_display, '')) LIKE '%외관%')
             ORDER BY ii.inspection_id,
                      CASE WHEN TRIM(ii.item_name) = '외관' THEN 0
                           WHEN TRIM(COALESCE(s.spec_display, '')) = '외관' THEN 1
                           ELSE 2 END,
                      COALESCE(s.item_order, 999999),
                      ii.id
        """, ids).fetchall()
        for item in exterior_rows:
            exterior_by_inspection.setdefault(item["inspection_id"], dict(item))

        for item in conn.execute(f"""
            SELECT id, inspection_id, ncr_no, status, sample_qty, defect_qty,
                   defect_description, special_note
              FROM ncr
             WHERE inspection_id IN ({placeholders})
             ORDER BY id
        """, ids).fetchall():
            ncr_by_inspection.setdefault(item["inspection_id"], []).append(dict(item))
        for item in conn.execute(f"""
            SELECT id, inspection_id, request_no, status, sample_qty, defect_qty,
                   request_detail
              FROM improvement_requests
             WHERE inspection_id IN ({placeholders})
             ORDER BY id
        """, ids).fetchall():
            improvement_by_inspection.setdefault(item["inspection_id"], []).append(dict(item))
    conn.close()

    for row in lots:
        exterior = exterior_by_inspection.get(row["id"])
        if exterior and exterior.get("aql_snapshot_set") == 1:
            general_sample_qty = exterior.get("sample_qty_snapshot")
            general_source = "외관 AQL 저장 스냅샷"
        elif exterior:
            recorded = _recorded_sample_count(exterior.get("measured_value"))
            general_sample_qty = recorded or None
            general_source = "외관 실제 기록" if recorded else ""
        else:
            general_sample_qty = None
            general_source = ""

        ncr = _preferred_quality_document(ncr_by_inspection.get(row["id"], []))
        improvement = _preferred_quality_document(improvement_by_inspection.get(row["id"], []))
        sample_qty = (ncr.get("sample_qty") if ncr and ncr.get("sample_qty") is not None else
                      improvement.get("sample_qty") if improvement and improvement.get("sample_qty") is not None else
                      general_sample_qty)
        if ncr and ncr.get("defect_qty") is not None:
            defect_qty = ncr["defect_qty"]
        elif improvement and improvement.get("defect_qty") is not None:
            defect_qty = improvement["defect_qty"]
        elif row["판정"] in ("합격", "미결"):
            defect_qty = 0
        else:
            defect_qty = None

        if ncr and ncr.get("sample_qty") is not None:
            sample_source = f"NCR {ncr.get('ncr_no') or ''}".strip()
        elif improvement and improvement.get("sample_qty") is not None:
            sample_source = f"개선요청서 {improvement.get('request_no') or ''}".strip()
        else:
            sample_source = general_source

        # ``UNASSIGNED_CLASSIFICATION`` is a filter sentinel and must never leak
        # into user-facing exports.
        major = row.get("major_category") or "미지정"
        middle = row.get("middle_category") or "미지정"
        aql_snapshot = exterior.get("aql_snapshot") if exterior else None
        if major in ("하우징", "기구물") or middle in ("고저항", "VI"):
            inspection_method = "전수 검사"
        elif str(aql_snapshot or "").strip() == "체크검사":
            inspection_method = "체크 검사"
        elif "전수" in str(aql_snapshot or ""):
            inspection_method = "전수 검사"
        elif exterior and (aql_snapshot is not None or general_sample_qty is not None):
            inspection_method = "샘플링 검사"
        else:
            inspection_method = ""

        special_parts = []
        if row["판정"] == "특채":
            if ncr:
                for value in (ncr.get("special_note"), ncr.get("defect_description")):
                    if value and value not in special_parts:
                        special_parts.append(value)
            if improvement and improvement.get("request_detail"):
                special_parts.append(improvement["request_detail"])
        note_parts = []
        for label, value in (("검사자", row.get("remark_inspector")),
                             ("검토자", row.get("remark_manager")),
                             ("최종승인자", row.get("remark_approver"))):
            if value:
                note_parts.append(f"{label}: {value}")

        row.update({
            "major_category": major,
            "middle_category": middle,
            "exterior_aql_snapshot": aql_snapshot,
            "sample_qty": sample_qty,
            "defect_qty": defect_qty,
            "sample_source": sample_source,
            "inspection_method": inspection_method,
            "selected_ncr": ncr,
            "selected_improvement": improvement,
            "special_content": "\n".join(special_parts),
            "notes": "\n".join(note_parts),
        })
    return lots


LOT_STATES = ["합격", "특채", "불합격", "미결"]   # 화면 상태 필터에 쓰는 값


def daily_status(day=None):
    """금일(또는 지정일) 현황 — 홈 요약보다 자세한 하루치 상황판.

    "아침에 출근해서 무엇부터 해야 하는가"에 답하는 화면용 데이터.
    입고→검사→승인→출력 흐름이 각 단계에서 얼마나 밀려 있는지, 오늘 누가 얼마나 했는지,
    그리고 지금 손대야 할 일(미조치)이 뭔지를 한 번에 모아준다.
    """
    from datetime import date as _date
    day = day or _date.today().isoformat()
    conn = get_conn()

    def one(sql, params=()):
        return conn.execute(sql, params).fetchone()[0] or 0

    def rows(sql, params=()):
        return [dict(r) for r in conn.execute(sql, params).fetchall()]

    # ── 오늘 들어온 것 ──
    intake_today = rows("""SELECT * FROM intake_list WHERE date(created_at)=?
                            ORDER BY id DESC""", (day,))
    intake_qty = sum(int(r["quantity"] or 0) for r in intake_today)

    # 업체별로 몇 건·몇 개 들어왔는지 + 그 업체의 실제 입고 행 목록도 같이 담아둔다
    # (2026-09-02: 예전엔 요약표와 전체 목록 두 섹션이 따로였는데, 업체별 접기/펼치기
    # 하나로 합치는 게 스크롤도 덜하고 "어느 업체에서 뭐가 왔는지"가 한눈에 보인다.)
    by_supplier = {}
    for r in intake_today:
        sup = r["supplier"] or "(미입력)"
        b = by_supplier.setdefault(sup, {"업체": sup, "건수": 0, "수량": 0, "목록": []})
        b["건수"] += 1
        b["수량"] += int(r["quantity"] or 0)
        b["목록"].append(r)
    by_supplier_list = sorted(by_supplier.values(), key=lambda x: -x["건수"])

    # ── 오늘 검사한 것 ── 제출(created_at)이 아니라 실제 검사한 날짜(inspect_date) 기준으로 집계한다.
    # 검사를 그날 했어도 제출은 다른 날 할 수 있어서, created_at으로 걸러면 검사자별 실적이 어긋난다.
    inspected = rows("""SELECT * FROM inspections WHERE date(inspect_date)=? AND status!='superseded'
                         ORDER BY id DESC""", (day,))
    insp_qty = sum(int(r["quantity"] or 0) for r in inspected)
    defect_today = [r for r in inspected if (r["overall_result"] or "") not in ("합격", "")]

    # ── 단계별 밀린 일 (오늘 것만이 아니라 '지금 쌓여 있는 것') ──
    backlog = {
        "검사대기": one("SELECT COUNT(*) FROM intake_list WHERE status='대기'"),
        "승인대기": one("SELECT COUNT(*) FROM inspections WHERE status='pending'"),
        "출력대기": one("""SELECT COUNT(*) FROM inspections
                            WHERE status='approved' AND (pdf_path IS NULL OR pdf_path='')"""),
        "반려": one("SELECT COUNT(*) FROM inspections WHERE status='rejected'"),
    }

    # ── 조치가 남아 있는 것 (이게 진짜 '해야 할 일') ──
    todo = {
        # 반려됐으니 다시 검사해야 하는 건
        "재검사필요": rows("""SELECT id, material_no, material_name, supplier, reject_reason, inspect_date
                                FROM inspections WHERE status='rejected' ORDER BY id DESC"""),
        # 불합격 확정인데 아직 부적합 통보서를 안 쓴 건
        "통보서작성필요": rows("""SELECT i.id, i.material_no, i.material_name, i.supplier, i.inspect_date
                                    FROM inspections i
                                   WHERE i.status='approved' AND i.approval_type='failed'
                                     AND NOT EXISTS (SELECT 1 FROM ncr n WHERE n.inspection_id=i.id)
                                   ORDER BY i.id DESC"""),
        # 통보서는 썼는데 아직 확인(승인) 안 된 건
        "통보서확인필요": rows("""SELECT id, ncr_no, material_no, supplier, issued_by, created_at
                                    FROM ncr WHERE status='draft' ORDER BY id DESC"""),
        # 확인은 됐는데 아직 협력사로 발송 안 한 건
        "통보서발송필요": rows("""SELECT id, ncr_no, material_no, supplier, confirmed_by, confirmed_at
                                    FROM ncr WHERE status='confirmed' ORDER BY id DESC"""),
    }

    # ── 오늘 누가 얼마나 했나 ──
    by_person = {}
    for r in inspected:
        p = by_person.setdefault(r["inspector"] or "(미입력)",
                                 {"검사자": r["inspector"] or "(미입력)", "건수": 0,
                                  "수량": 0, "총초": 0, "불량": 0})
        p["건수"] += 1
        p["수량"] += int(r["quantity"] or 0)
        p["총초"] += int(r["total_time_sec"] or 0) or int(r["actual_time_sec"] or 0)
        if (r["overall_result"] or "") not in ("합격", ""):
            p["불량"] += 1
    for p in by_person.values():
        p["평균초"] = int(round(p["총초"] / p["건수"])) if p["건수"] else 0

    # ── 오늘 승인/불합격 결정 ──
    decided = rows("""SELECT * FROM inspections
                       WHERE date(approved_at)=? ORDER BY approved_at DESC""", (day,))

    conn.close()

    # 검사 진행률 = 오늘 입고분 중 몇 건이나 검사가 끝났나
    done_pos = len([r for r in intake_today if r["status"] == "검사완료"])
    progress = round(done_pos / len(intake_today) * 100) if intake_today else 0

    return {
        "날짜": day,
        "입고": {"건수": len(intake_today), "수량": intake_qty,
                 "검사완료": done_pos, "진행률": progress, "목록": intake_today,
                 "업체별": by_supplier_list},
        "검사": {"건수": len(inspected), "수량": insp_qty,
                 "불량건수": len(defect_today), "목록": inspected, "불량목록": defect_today},
        "결정": {"건수": len(decided),
                 "합격": len([r for r in decided if r["approval_type"] == "normal"]),
                 "특채": len([r for r in decided if r["approval_type"] == "special"]),
                 "불합격": len([r for r in decided if r["approval_type"] == "failed"])},
        "밀린일": backlog,
        "해야할일": todo,
        "인원별": sorted(by_person.values(), key=lambda x: -x["건수"]),
    }


def quality_report(start_date, end_date, period_type="monthly",
                   supplier=None, po_number=None, material=None, states=None, category=None,
                   major_category=None, middle_category=None):
    """품질 현황 집계 — 대시보드·보고서·내보내기가 전부 이 함수 하나를 쓴다.

    불량률은 **수량 기준**이다(사용자 확정):
        불량률 = 불합격 확정 수량 / 판정 확정 수량 × 100
        PPM    = 불량률 × 10,000
    '판정 확정'은 합격 + 특채 + 불합격이고, 승인 대기·반려 건은 분모에서 뺀다
    (아직 결과가 아니므로). 미결 수량은 따로 보여준다.

    material 인자는 자재번호와 제품명 양쪽에서 부분일치로 찾는다(화면 입력 하나로 둘 다 커버).

    supplier / po_number 는 문자열 하나 또는 **리스트**를 받는다(화면에서 여러 개 선택 가능).
    states 는 ['합격','불합격', ...] 처럼 로트 상태로 거를 때 쓴다(비우면 전체).
    """
    def _as_list(v):
        if v is None:
            return []
        if isinstance(v, (list, tuple, set)):
            return [str(x).strip() for x in v if str(x).strip()]
        return [v.strip()] if str(v).strip() else []

    suppliers = _as_list(supplier)
    po_numbers = _as_list(po_number)
    states = _as_list(states)

    conn = get_conn()

    sql = """SELECT i.*, mc.major_category, mc.middle_category,
                    (SELECT COUNT(*) FROM ncr n WHERE n.inspection_id = i.id) AS ncr_count,
                    (SELECT n.id FROM ncr n WHERE n.inspection_id = i.id
                      ORDER BY n.id DESC LIMIT 1) AS ncr_id,
                    (SELECT COUNT(*) FROM ncr n WHERE n.inspection_id = i.id
                      AND n.status IN ('sent', 'confirmed')) AS ncr_sent_count
               FROM inspections i
               LEFT JOIN materials mt ON mt.material_no = i.material_no
               LEFT JOIN material_classification_map mc ON mc.material_no = i.material_no
              WHERE 1=1"""
    params = []
    if start_date:
        sql += " AND i.inspect_date >= ?"; params.append(start_date)
    if end_date:
        sql += " AND i.inspect_date <= ?"; params.append(end_date)
    if suppliers:
        sql += f" AND i.supplier IN ({','.join('?' * len(suppliers))})"; params += suppliers
    if po_numbers:
        sql += f" AND i.po_number IN ({','.join('?' * len(po_numbers))})"; params += po_numbers
    if material:
        sql += " AND (i.material_no LIKE ? OR i.material_name LIKE ?)"
        params += [f"%{material}%", f"%{material}%"]
    if category:
        sql += " AND mt.category = ?"; params.append(category)
    major_category = (major_category or "").strip()
    middle_category = (middle_category or "").strip()
    if major_category == UNASSIGNED_CLASSIFICATION or middle_category == UNASSIGNED_CLASSIFICATION:
        sql += " AND mc.material_no IS NULL"
    else:
        if major_category:
            sql += " AND mc.major_category = ?"; params.append(major_category)
        if middle_category:
            sql += " AND mc.middle_category = ?"; params.append(middle_category)
    # 재검사로 대체된 옛 성적서는 집계에서 제외 (같은 로트가 두 번 잡히는 것 방지)
    sql += " AND i.status != 'superseded'"
    sql += " ORDER BY i.inspect_date DESC, i.id DESC"

    rows = [dict(r) for r in conn.execute(sql, params).fetchall()]

    # 로트 상태(합격/특채/불합격/미결) 필터는 status+approval_type 조합이라 SQL이 지저분해진다.
    # 판정 규칙이 _lot_state() 한 곳에만 있도록 여기서 걸러낸다.
    if states:
        rows = [r for r in rows
                if _lot_state(r["status"], r["approval_type"], r["overall_result"]) in states]

    ids = [r["id"] for r in rows]

    # 불량 항목 — 어떤 검사항목에서 많이 터지는지
    defect_items = []
    if ids:
        ph = ",".join("?" * len(ids))
        defect_items = [dict(r) for r in conn.execute(f"""
            SELECT ii.item_name, ii.result, ii.measured_value,
                   i.material_no, i.material_name, i.supplier, i.id AS inspection_id,
                   COALESCE(s.spec_display, ii.item_name) AS spec_display
              FROM inspection_items ii
              JOIN inspections i ON i.id = ii.inspection_id
              LEFT JOIN specs s ON s.material_no = COALESCE(ii.part_material_no, i.material_no)
                               AND s.item_name = ii.item_name
             WHERE ii.inspection_id IN ({ph})
               AND ii.result NOT IN ('합격', '미측정', '')
               AND ii.result IS NOT NULL
        """, ids).fetchall()]

    # 기간 내 4M 변경점 — 불량률 변화의 원인 후보
    cp_sql = "SELECT * FROM change_points WHERE 1=1"
    cp_params = []
    if start_date:
        cp_sql += " AND change_date >= ?"; cp_params.append(start_date)
    if end_date:
        cp_sql += " AND change_date <= ?"; cp_params.append(end_date)
    if suppliers:
        cp_sql += f" AND supplier IN ({','.join('?' * len(suppliers))})"; cp_params += suppliers
    cp_sql += " ORDER BY change_date DESC"
    change_points = [dict(r) for r in conn.execute(cp_sql, cp_params).fetchall()]

    # 실제 측정 표본수 — inspection_id별 measured_value 콤마수 최댓값
    # (값이 3개면 콤마 2개 = 길이차 + 1 = 3)
    sample_count_map = {}
    if ids:
        ph2 = ",".join("?" * len(ids))
        for sc_row in conn.execute(f"""
            SELECT inspection_id,
                   MAX(LENGTH(measured_value) - LENGTH(REPLACE(measured_value, ',', '')) + 1) AS cnt
              FROM inspection_items
             WHERE inspection_id IN ({ph2})
               AND measured_value IS NOT NULL AND TRIM(measured_value) != ''
             GROUP BY inspection_id
        """, ids).fetchall():
            sample_count_map[sc_row[0]] = int(sc_row[1])

    # NCR 집계용 데이터 — issued_date 기준, supplier 필터 적용
    ncr_sql = """
        SELECT supplier, issued_date,
               COALESCE(occurrence_type, '입고검사') AS occurrence_type,
               defect_type
          FROM ncr
         WHERE issued_date IS NOT NULL
    """
    ncr_params = []
    if start_date:
        ncr_sql += " AND issued_date >= ?"; ncr_params.append(start_date)
    if end_date:
        ncr_sql += " AND issued_date <= ?"; ncr_params.append(end_date)
    if suppliers:
        ncr_sql += f" AND supplier IN ({','.join('?' * len(suppliers))})"; ncr_params += suppliers
    ncr_rows = [dict(r) for r in conn.execute(ncr_sql, ncr_params).fetchall()]

    # 개선요청서 집계용 데이터 — NCR과 병렬(issued_date 기준, supplier 필터 적용)
    imp_sql = "SELECT supplier, issued_date FROM improvement_requests WHERE issued_date IS NOT NULL"
    imp_params = []
    if start_date:
        imp_sql += " AND issued_date >= ?"; imp_params.append(start_date)
    if end_date:
        imp_sql += " AND issued_date <= ?"; imp_params.append(end_date)
    if suppliers:
        imp_sql += f" AND supplier IN ({','.join('?' * len(suppliers))})"; imp_params += suppliers
    imp_rows = [dict(r) for r in conn.execute(imp_sql, imp_params).fetchall()]

    # 수기입력(성적서 미연결) sent NCR — ②③(NCR포함 불량률/건수율)용 lot_qty만 뽑는다.
    manual_ncr_sql = "SELECT lot_qty FROM ncr WHERE inspection_id IS NULL AND status IN ('sent', 'confirmed')"
    manual_ncr_params = []
    if start_date:
        manual_ncr_sql += " AND issued_date >= ?"; manual_ncr_params.append(start_date)
    if end_date:
        manual_ncr_sql += " AND issued_date <= ?"; manual_ncr_params.append(end_date)
    if suppliers:
        manual_ncr_sql += f" AND supplier IN ({','.join('?' * len(suppliers))})"; manual_ncr_params += suppliers
    manual_ncr_lot_qtys = [r[0] for r in conn.execute(manual_ncr_sql, manual_ncr_params).fetchall()]

    conn.close()

    # ---- 집계 ----
    def blank():
        return {"로트": 0, "수량": 0, "합격수량": 0, "특채수량": 0,
                "불합격수량": 0, "미결수량": 0, "불합격로트": 0, "합격로트": 0,
                "검사표본수": 0, "표본불량수": 0, "특채건수": 0}

    # 로트(=성적서 1건) 단위 판정/표본통계 — 세부 내보내기(자재별 상세시트)에서도 그대로 쓴다.
    # 비고의 "검사수량 N개 중 M개 불량"에서 M(표본불량수)만 뽑고 N은 버린 채, 측정값
    # 콤마개수로 따로 계산한 검사표본수로 나누면 표본불량률이 100%를 넘는 사고가 났다
    # (2026-10-03 실사용자 확인) — 비고의 N은 측정값 표본수와 무관하게 "그 항목의 QMS AQL
    # 샘플수량"을 가리킬 수 있는 별개 개념이라, M은 반드시 같은 문구에 적힌 N과 짝지어
    # 계산해야 한다. 비고에 이 문구가 있으면 그 N을 검사표본수로 쓰고, 없을 때만 기존처럼
    # 측정값 기반 sample_count_map으로 대체한다.
    for r in rows:
        r["판정"] = _lot_state(r["status"], r["approval_type"], r["overall_result"])
        remark_n = remark_m = 0
        for m in _DEFECT_RE.finditer(r.get("remark_inspector") or ""):
            remark_n += int(m.group(1))
            remark_m += int(m.group(2))
        r["표본불량수"] = remark_m
        r["검사표본수"] = remark_n if remark_n else sample_count_map.get(r["id"], 0)

    def add(acc, r):
        qty = int(r["quantity"] or 0)
        state = r["판정"]
        acc["로트"] += 1
        acc["수량"] += qty
        acc["검사표본수"] += r["검사표본수"]
        acc["표본불량수"] += r["표본불량수"]
        if state == "합격":
            acc["합격수량"] += qty
            acc["합격로트"] += 1
        elif state == "특채":
            acc["특채수량"] += qty
            acc["특채건수"] += 1
        elif state == "불합격":
            acc["불합격수량"] += qty
            acc["불합격로트"] += 1
        else:
            acc["미결수량"] += qty

    def finish(acc):
        confirmed = acc["합격수량"] + acc["특채수량"] + acc["불합격수량"]
        acc["확정수량"] = confirmed
        acc["판정확정로트"] = acc["합격로트"] + acc["특채건수"] + acc["불합격로트"]
        acc["불량률"] = round(acc["불합격수량"] / confirmed * 100, 3) if confirmed else 0.0
        acc["PPM"] = int(round(acc["불합격수량"] / confirmed * 1_000_000)) if confirmed else 0
        # 특채까지 포함한 '규격 이탈률' — 참고용
        out = acc["불합격수량"] + acc["특채수량"]
        acc["규격이탈률"] = round(out / confirmed * 100, 3) if confirmed else 0.0
        # 표본 기준 불량률 (None = 표본데이터 없음, 0%와 구분)
        samp = acc["검사표본수"]
        acc["표본불량률"] = round(acc["표본불량수"] / samp * 100, 3) if samp else None
        acc["표본PPM"] = int(round(acc["표본불량수"] / samp * 1_000_000)) if samp else None
        return acc

    summary = blank()
    by_period, by_supplier, by_material = {}, {}, {}
    for r in rows:
        add(summary, r)
        pk = _period_key(r["inspect_date"], period_type) or "(날짜없음)"
        add(by_period.setdefault(pk, blank()), r)
        add(by_supplier.setdefault(r["supplier"] or "(미입력)", blank()), r)
        mk = r["material_no"] or "(미입력)"
        acc_m = by_material.setdefault(mk, blank())
        acc_m["자재명"] = r["material_name"]
        acc_m["대분류"] = r.get("major_category") or "미지정"
        acc_m["중분류"] = r.get("middle_category") or "미지정"
        add(acc_m, r)

    finish(summary)

    # ---- ②③ NCR포함 불량률/건수율 (2026-09-15 사용자 확정) ----
    ncr_extra_qty, ncr_extra_cnt = _ncr_extra_stats(rows, manual_ncr_lot_qtys)
    summary["NCR추가수량"] = ncr_extra_qty
    summary["NCR추가로트"] = ncr_extra_cnt

    confirmed_ncr = summary["확정수량"] + ncr_extra_qty
    defect_ncr = summary["불합격수량"] + ncr_extra_qty
    summary["불량률_NCR포함"] = round(defect_ncr / confirmed_ncr * 100, 3) if confirmed_ncr else 0.0
    summary["PPM_NCR포함"] = int(round(defect_ncr / confirmed_ncr * 1_000_000)) if confirmed_ncr else 0

    confirmed_cnt_ncr = summary["판정확정로트"] + ncr_extra_cnt
    defect_cnt_ncr = summary["불합격로트"] + ncr_extra_cnt
    summary["불량건수율_NCR포함"] = round(defect_cnt_ncr / confirmed_cnt_ncr * 100, 3) if confirmed_cnt_ncr else 0.0

    period_list = []
    for k in sorted(by_period.keys()):
        period_list.append({"구간": k, **finish(by_period[k])})

    supplier_list = []
    for k, v in by_supplier.items():
        supplier_list.append({"업체": k, **finish(v)})
    supplier_list.sort(key=lambda x: (-x["불합격수량"], -x["수량"]))

    material_list = []
    for k, v in by_material.items():
        name = v.pop("자재명", None)
        major = v.pop("대분류", "미지정")
        middle = v.pop("중분류", "미지정")
        material_list.append({"자재번호": k, "자재명": name,
                              "대분류": major, "중분류": middle, **finish(v)})
    material_list.sort(key=lambda x: (-x["불합격수량"], -x["수량"]))

    # 불량 항목 순위
    item_rank = {}
    for d in defect_items:
        key = (d["material_no"], d["item_name"])
        e = item_rank.setdefault(key, {
            "자재번호": d["material_no"], "자재명": d["material_name"],
            "항목": d["item_name"], "규격": d["spec_display"], "발생건수": 0, "업체": set(),
        })
        e["발생건수"] += 1
        if d["supplier"]:
            e["업체"].add(d["supplier"])
    top_items = sorted(item_rank.values(), key=lambda x: -x["발생건수"])
    for e in top_items:
        e["업체"] = ", ".join(sorted(e["업체"]))

    # 검사 소요시간 — 인력 산정 근거
    time_rows = [r for r in rows if (r["total_time_sec"] or r["actual_time_sec"])]
    per_person = {}
    total_sec = 0
    for r in time_rows:
        sec = int(r["total_time_sec"] or 0) or int(r["actual_time_sec"] or 0)
        total_sec += sec
        p = per_person.setdefault(r["inspector"] or "(미입력)", {"검사자": r["inspector"] or "(미입력)",
                                                              "건수": 0, "총초": 0})
        p["건수"] += 1
        p["총초"] += sec
    for p in per_person.values():
        p["평균초"] = int(round(p["총초"] / p["건수"])) if p["건수"] else 0
    time_stats = {
        "측정건수": len(time_rows),
        "총초": total_sec,
        "평균초": int(round(total_sec / len(time_rows))) if time_rows else 0,
        "인원별": sorted(per_person.values(), key=lambda x: -x["건수"]),
    }

    # ---- NCR 집계 ----
    # 최상위 요약
    ncr_total = len(ncr_rows)
    ncr_후발 = sum(1 for n in ncr_rows if n["occurrence_type"] == "사후")
    ncr_불량유형별: dict = {}
    for n in ncr_rows:
        dt = n["defect_type"] or "(미분류)"
        ncr_불량유형별[dt] = ncr_불량유형별.get(dt, 0) + 1

    # 기간별 NCR 건수 — issued_date 기준으로 period_key 계산
    ncr_by_period: dict = {}
    for n in ncr_rows:
        pk = _period_key(n["issued_date"], period_type) or "(날짜없음)"
        entry = ncr_by_period.setdefault(pk, {"ncr_건수": 0, "사후불량_건수": 0})
        entry["ncr_건수"] += 1
        if n["occurrence_type"] == "사후":
            entry["사후불량_건수"] += 1

    # period_list에 NCR 건수 병합
    for item in period_list:
        pk = item["구간"]
        ncr_entry = ncr_by_period.get(pk, {"ncr_건수": 0, "사후불량_건수": 0})
        item["ncr_건수"] = ncr_entry["ncr_건수"]
        item["사후불량_건수"] = ncr_entry["사후불량_건수"]

    # 업체별 NCR 건수
    ncr_by_supplier: dict = {}
    for n in ncr_rows:
        sup = n["supplier"] or "(미입력)"
        entry = ncr_by_supplier.setdefault(sup, {"ncr_건수": 0, "사후불량_건수": 0})
        entry["ncr_건수"] += 1
        if n["occurrence_type"] == "사후":
            entry["사후불량_건수"] += 1

    # supplier_list에 NCR 건수 병합
    for item in supplier_list:
        sup = item["업체"]
        ncr_entry = ncr_by_supplier.get(sup, {"ncr_건수": 0, "사후불량_건수": 0})
        item["ncr_건수"] = ncr_entry["ncr_건수"]
        item["사후불량_건수"] = ncr_entry["사후불량_건수"]

    # 요약에 NCR 건수 추가
    summary["ncr_건수"] = ncr_total
    summary["사후불량_건수"] = ncr_후발
    summary["불량유형별"] = ncr_불량유형별

    # supplier_ncr_rank — NCR 건수 내림차순
    sup_ncr_map: dict = {}
    for n in ncr_rows:
        sup = n["supplier"] or "(미입력)"
        sup_ncr_map[sup] = sup_ncr_map.get(sup, 0) + 1
    # 불합격수량은 supplier_list에서 가져온다
    sup_불합격_map = {item["업체"]: item["불합격수량"] for item in supplier_list}
    supplier_ncr_rank = sorted(
        [{"name": s, "ncr_건수": cnt, "불합격수량": sup_불합격_map.get(s, 0)}
         for s, cnt in sup_ncr_map.items()],
        key=lambda x: -x["ncr_건수"],
    )

    # ---- 개선요청서 집계 (NCR과 병렬) ----
    summary["개선요청_건수"] = len(imp_rows)

    imp_by_period: dict = {}
    for n in imp_rows:
        pk = _period_key(n["issued_date"], period_type) or "(날짜없음)"
        imp_by_period[pk] = imp_by_period.get(pk, 0) + 1
    for item in period_list:
        item["개선요청_건수"] = imp_by_period.get(item["구간"], 0)

    imp_by_supplier: dict = {}
    for n in imp_rows:
        sup = n["supplier"] or "(미입력)"
        imp_by_supplier[sup] = imp_by_supplier.get(sup, 0) + 1
    for item in supplier_list:
        item["개선요청_건수"] = imp_by_supplier.get(item["업체"], 0)

    supplier_improvement_rank = sorted(
        [{"name": s, "개선요청_건수": cnt} for s, cnt in imp_by_supplier.items()],
        key=lambda x: -x["개선요청_건수"],
    )

    # material_deviation_rank — 특채건수 내림차순 (특채건수 > 0인 자재만)
    material_deviation_rank = sorted(
        [{"material_no": item["자재번호"], "material_name": item["자재명"],
          "특채건수": item["특채건수"]}
         for item in material_list if item.get("특채건수", 0) > 0],
        key=lambda x: -x["특채건수"],
    )

    # supplier_deviation_rank — 특채건수 내림차순 (특채건수 > 0인 업체만)
    supplier_deviation_rank = sorted(
        [{"name": item["업체"], "특채건수": item["특채건수"]}
         for item in supplier_list if item.get("특채건수", 0) > 0],
        key=lambda x: -x["특채건수"],
    )

    return {
        "기간": {"시작": start_date, "종료": end_date, "유형": period_type},
        "필터": {"업체": ", ".join(suppliers) if suppliers else "전체",
                 "로트번호": ", ".join(po_numbers) if po_numbers else "전체",
                 "자재/제품명": material or "전체",
                 "대분류": ("미지정" if major_category == UNASSIGNED_CLASSIFICATION else major_category) or "전체",
                 "중분류": ("미지정" if middle_category == UNASSIGNED_CLASSIFICATION else middle_category) or "전체",
                 "판정상태": ", ".join(states) if states else "전체"},
        "불량률기준": "수량기준 (불합격 확정수량 ÷ 판정 확정수량)",
        "불량률기준_NCR포함": ("수량기준 ((불합격확정수량 + NCR만있는로트 수량) ÷ "
                          "(판정확정수량 + NCR만있는로트 수량)). NCR만있는로트 = "
                          "불합격확정이 아닌 로트(합격/특채/미결) 중 발송·확인(status가 "
                          "sent 또는 confirmed)된 NCR이 연결된 것 + 수기입력 발송·확인 "
                          "NCR(lot_qty 숫자 인식 성공분)."),
        "불량건수율기준_NCR포함": "건수기준 ((불합격로트 + NCR만있는로트) ÷ (판정확정로트 + NCR만있는로트))",
        "요약": summary,
        "기간별": period_list,
        "업체별": supplier_list,
        "자재별": material_list,
        "불량항목순위": top_items,
        "소요시간": time_stats,
        "변경점": change_points,
        "성적서목록": rows,
        "supplier_ncr_rank": supplier_ncr_rank,
        "supplier_improvement_rank": supplier_improvement_rank,
        "material_deviation_rank": material_deviation_rank,
        "supplier_deviation_rank": supplier_deviation_rank,
    }


_HL_DEVIATION_MIN = 5.0     # 규격이탈률(%) 이 이상이어야 특이사항 후보로 본다
_HL_MIN_CONFIRMED = 10      # 확정수량이 이 미만이면 표본이 너무 작아 노이즈로 보고 제외한다
_HL_SUPPLIER_SLOTS = 3      # 특이사항 중 업체 몫(수량 영향이 큰 순)
_HL_MATERIAL_SLOTS = 2      # 특이사항 중 자재 몫


def quality_highlights(report):
    """quality_report() 결과를 임원도 바로 읽을 수 있는 문장으로 요약한다(규칙 기반 — AI 아님).

    대시보드 화면과 JSON/엑셀 내보내기가 이 함수 하나만 쓴다(집계는 한 곳에서만 하는
    quality_report()의 원칙을 요약 문장에도 그대로 적용) — 화면과 발표자료가 서로
    다른 이야기를 하지 않게 하기 위함이다.

    특이사항은 '규격이탈률이 높은 것'이 아니라 '실제 영향 수량(특채+불합격)이 큰 것'을
    우선 노출한다 — 작은 로트 하나가 100% 이탈인 것보다, 큰 로트에서 이탈이 반복되는
    업체가 경영진에게 더 중요한 신호이기 때문. 업체 몫과 자재 몫을 미리 나눠서(3+2),
    업체 단위 이슈가 개별 부품 이슈에 전부 밀려나지 않게 한다.
    """
    s = report["요약"]
    confirmed = s["확정수량"]
    if confirmed == 0:
        if s["로트"] == 0:
            headline = "이 조건에 해당하는 검사 기록이 없습니다."
        else:
            headline = f"이 기간 {s['로트']}개 로트 중 판정이 확정된 건이 아직 없습니다(전부 승인 대기·반려)."
        return {"headline": headline, "톤": "정보없음", "특이사항": [], "표본참고": None}

    def _candidates(rows, kind, name_key):
        out = []
        for r in rows:
            if r["확정수량"] < _HL_MIN_CONFIRMED or r["규격이탈률"] < _HL_DEVIATION_MIN:
                continue
            name = r[name_key]
            if kind == "자재" and r.get("자재명"):
                name = f"{name} ({r['자재명']})"
            impact = r["특채수량"] + r["불합격수량"]
            out.append({
                "구분": kind, "이름": name, "규격이탈률": r["규격이탈률"], "불량률": r["불량률"],
                "특채수량": r["특채수량"], "불합격수량": r["불합격수량"], "확정수량": r["확정수량"],
                "_impact": impact,
            })
        out.sort(key=lambda x: (-x["_impact"], -x["규격이탈률"]))
        return out

    sup_cands = _candidates(report["업체별"], "업체", "업체")
    mat_cands = _candidates(report["자재별"], "자재", "자재번호")

    top = sup_cands[:_HL_SUPPLIER_SLOTS] + mat_cands[:_HL_MATERIAL_SLOTS]
    total_slots = _HL_SUPPLIER_SLOTS + _HL_MATERIAL_SLOTS
    remaining = total_slots - len(top)
    if remaining > 0:
        used_sup = min(len(sup_cands), _HL_SUPPLIER_SLOTS)
        used_mat = min(len(mat_cands), _HL_MATERIAL_SLOTS)
        top += sup_cands[used_sup:used_sup + remaining]
        remaining = total_slots - len(top)
        if remaining > 0:
            top += mat_cands[used_mat:used_mat + remaining]
    top.sort(key=lambda x: (-x["_impact"], -x["규격이탈률"]))

    for c in top:
        parts = []
        if c["불합격수량"]:
            parts.append(f"불합격 {c['불합격수량']:,}개")
        if c["특채수량"]:
            parts.append(f"특채 {c['특채수량']:,}개")
        detail = ", ".join(parts) if parts else "규격 이탈"
        c["메시지"] = f"{c['이름']} — 규격이탈률 {c['규격이탈률']}%(확정 {c['확정수량']:,}개 중 {detail})"
        c.pop("_impact", None)

    pass_rate = round(s["합격수량"] / confirmed * 100, 1)
    headline = (f"이번 기간 판정 확정 {confirmed:,}개 중 합격 {s['합격수량']:,}개({pass_rate}%), "
                f"불합격 {s['불합격수량']:,}개, 특채 {s['특채수량']:,}개입니다.")
    if top:
        names = ", ".join(c["이름"] for c in top[:3])
        suffix = " 등" if len(top) > 3 else ""
        headline += f" 다만 {names}{suffix}에서 규격을 반복적으로 벗어나 특채·불합격 처리된 사례가 있어 확인이 필요합니다."
    else:
        headline += " 규격 이탈이 두드러지는 업체·자재는 없습니다."

    sample_note = None
    sample_defect_rate = s.get("표본불량률")  # 이 필드가 없던 시절(구버전) 성적표 스냅샷 대응
    if sample_defect_rate is not None and (sample_defect_rate - s["불량률"]) >= 2:
        sample_note = (f"참고: 개별 표본 기준으로는 불량 표본 비율이 {sample_defect_rate}%로 나타납니다. "
                        "이 수치는 특채로 넘어간 물량의 표본 불량까지 포함하며, "
                        "위 불량률·PPM에는 특채가 들어가지 않습니다.")

    if s["불량률"] >= 3 or any(c["규격이탈률"] >= 20 for c in top):
        tone = "경고"
    elif top or s["불량률"] >= 1:
        tone = "주의"
    else:
        tone = "양호"

    return {"headline": headline, "톤": tone, "특이사항": top, "표본참고": sample_note}


def upsert_supplier_report(supplier, period, start_date, end_date, payload_json, created_by):
    """업체 월간 성적표 생성/갱신. 이미 승인·발송된 건은 덮어쓰지 않는다."""
    conn = get_conn()
    try:
        existing = conn.execute(
            "SELECT id, status FROM supplier_reports WHERE supplier=? AND period=?",
            (supplier, period)).fetchone()
        if existing and existing["status"] != "draft":
            return existing["id"], f"이미 {existing['status']} 상태라 다시 만들 수 없어."
        if existing:
            conn.execute("""UPDATE supplier_reports
                               SET start_date=?, end_date=?, payload=?, created_by=?,
                                   created_at=datetime('now','localtime')
                             WHERE id=?""",
                         (start_date, end_date, payload_json, created_by, existing["id"]))
            rid = existing["id"]
        else:
            cur = conn.execute("""INSERT INTO supplier_reports
                                   (supplier, period, start_date, end_date, payload, created_by)
                                   VALUES (?, ?, ?, ?, ?, ?)""",
                               (supplier, period, start_date, end_date, payload_json, created_by))
            rid = cur.lastrowid
        conn.commit()
        return rid, None
    finally:
        conn.close()


def get_supplier_report_trend(supplier, upto_period_exclusive, limit=5):
    """최근 성적표 목록 — 월별 추이 차트용. upto_period_exclusive 는 포함하지 않는다
    (같은 달 draft가 자기 자신을 중복으로 끌고 오는 걸 막기 위해 반드시 '<')."""
    import json
    conn = get_conn()
    rows = conn.execute(
        """SELECT period, payload FROM supplier_reports
            WHERE supplier = ? AND period < ?
            ORDER BY period DESC LIMIT ?""",
        (supplier, upto_period_exclusive, limit)).fetchall()
    conn.close()
    out = []
    for row in reversed(rows):
        try:
            data = json.loads(row["payload"])
            s = data["요약"]
            out.append({"기간": row["period"], "수량": s["수량"], "불합격수량": s["불합격수량"],
                        "특채수량": s["특채수량"], "합격수량": s["합격수량"],
                        "불량률": s["불량률"], "규격이탈률": s["규격이탈률"], "PPM": s["PPM"]})
        except (ValueError, KeyError, TypeError):
            continue
    return out


def list_supplier_reports(status=None):
    conn = get_conn()
    sql = "SELECT * FROM supplier_reports"
    params = []
    if status:
        sql += " WHERE status = ?"; params.append(status)
    sql += " ORDER BY period DESC, supplier"
    rows = conn.execute(sql, params).fetchall()
    conn.close()
    return rows


def get_supplier_report(report_id):
    conn = get_conn()
    row = conn.execute("SELECT * FROM supplier_reports WHERE id=?", (report_id,)).fetchone()
    conn.close()
    return row


def approve_supplier_report(report_id, approved_by, signature_path):
    conn = get_conn()
    conn.execute("""UPDATE supplier_reports
                       SET status='approved', approved_by=?, approve_signature=?,
                           approved_at=datetime('now','localtime')
                     WHERE id=? AND status='draft'""",
                 (approved_by, signature_path, report_id))
    conn.commit()
    conn.close()


def mark_supplier_report_sent(report_id, sent_to):
    conn = get_conn()
    conn.execute("""UPDATE supplier_reports
                       SET status='sent', sent_to=?, sent_at=datetime('now','localtime')
                     WHERE id=? AND status='approved'""",
                 (sent_to, report_id))
    conn.commit()
    conn.close()


def delete_supplier_report(report_id):
    conn = get_conn()
    conn.execute("DELETE FROM supplier_reports WHERE id=? AND status='draft'", (report_id,))
    conn.commit()
    conn.close()


def delete_supplier_reports_admin(report_ids):
    """admin 전용: 상태와 무관하게 성적표 삭제."""
    if not report_ids:
        return
    conn = get_conn()
    placeholders = ",".join("?" * len(report_ids))
    conn.execute(f"DELETE FROM supplier_reports WHERE id IN ({placeholders})", report_ids)
    conn.commit()
    conn.close()


def delete_ncrs(ncr_ids):
    """부적합 통보서 일괄 삭제 (photos는 JSON으로 인라인 저장돼 있어서 부수 테이블 없음)."""
    if not ncr_ids:
        return
    conn = get_conn()
    placeholders = ",".join("?" * len(ncr_ids))
    conn.execute(f"DELETE FROM ncr WHERE id IN ({placeholders})", ncr_ids)
    conn.commit()
    conn.close()


def process_capability(material_no, item_name=None, min_samples=5):
    """공정능력(Cp/Cpk) 계산 — 규격 하한·상한과 실제 측정값으로.

    Cpk = min( (USL-μ)/3σ , (μ-LSL)/3σ )  — 단측 규격이면 있는 쪽만 본다.
    구간 표시(우수/양호/주의/부족)는 1.33/1.67 기준이고, 실제 요구 기준은 거래처·품목마다
    다르니 확인해야 한다(CLAUDE.md 8-2-12절 원칙과 동일 — 업종 단정 금지).
    표본이 적으면 숫자가 튀므로 min_samples 미만은 계산하지 않고 표시만 한다.
    """
    import statistics
    conn = get_conn()
    specs = conn.execute("""
        SELECT item_name, spec_display, lower_limit, upper_limit
          FROM specs
         WHERE material_no = ? AND judge_type = 'numeric'
           AND (lower_limit IS NOT NULL OR upper_limit IS NOT NULL)
           AND (? IS NULL OR item_name = ?)
         ORDER BY item_order
    """, (material_no, item_name, item_name)).fetchall()

    out = []
    for sp in specs:
        raw = conn.execute("""
            SELECT ii.measured_value
              FROM inspection_items ii
              JOIN inspections i ON i.id = ii.inspection_id
             WHERE ii.item_name = ?
               AND COALESCE(ii.part_material_no, i.material_no) = ?
               AND ii.measured_value IS NOT NULL AND ii.measured_value != ''
        """, (sp["item_name"], material_no)).fetchall()

        values = []
        for r in raw:
            for tok in str(r["measured_value"]).split(","):
                tok = tok.strip()
                if not tok:
                    continue
                try:
                    values.append(float(tok))
                except ValueError:
                    pass

        entry = {
            "항목": sp["item_name"], "규격": sp["spec_display"],
            "하한": sp["lower_limit"], "상한": sp["upper_limit"],
            "표본수": len(values), "평균": None, "표준편차": None,
            "Cp": None, "Cpk": None, "판정": None,
        }
        if len(values) >= max(2, min_samples):
            mu = statistics.fmean(values)
            sigma = statistics.stdev(values)
            entry["평균"] = round(mu, 4)
            entry["표준편차"] = round(sigma, 4)
            if sigma > 0:
                lsl, usl = sp["lower_limit"], sp["upper_limit"]
                cpu = (usl - mu) / (3 * sigma) if usl is not None else None
                cpl = (mu - lsl) / (3 * sigma) if lsl is not None else None
                cands = [c for c in (cpu, cpl) if c is not None]
                if cands:
                    entry["Cpk"] = round(min(cands), 3)
                if lsl is not None and usl is not None:
                    entry["Cp"] = round((usl - lsl) / (6 * sigma), 3)
                if entry["Cpk"] is not None:
                    entry["판정"] = ("우수" if entry["Cpk"] >= 1.67 else
                                    "양호" if entry["Cpk"] >= 1.33 else
                                    "주의" if entry["Cpk"] >= 1.0 else "부족")
            else:
                entry["판정"] = "산포없음"     # 측정값이 전부 같음 (분해능 부족 의심)
        else:
            entry["판정"] = "표본부족"
        out.append(entry)

    conn.close()
    return out


# ---------- 마이그레이션 ----------

def ensure_ncr_columns_migration():
    """ncr 테이블에 occurrence_type / defect_type 컬럼 추가 (멱등).

    - occurrence_type: '입고검사' | '사후'. NULL은 '입고검사'로 간주.
    - defect_type: '치수불량'|'외관불량'|'기능불량'|'재질불량'|'수량불량'|'기타'. NULL 허용.
    """
    if get_setting("ncr_columns_migrated_20260908") == "1":
        return
    conn = get_conn()
    existing = [r[1] for r in conn.execute("PRAGMA table_info('ncr')").fetchall()]
    if "occurrence_type" not in existing:
        conn.execute("ALTER TABLE ncr ADD COLUMN occurrence_type TEXT DEFAULT '입고검사'")
    if "defect_type" not in existing:
        conn.execute("ALTER TABLE ncr ADD COLUMN defect_type TEXT")
    conn.commit()
    conn.close()
    set_setting("ncr_columns_migrated_20260908", "1")


def ensure_defect_types_seed_20260908():
    """불량 유형 마스터 최초 시드 — 폼에 하드코딩돼 있던 6종을 마스터 테이블로 옮긴다(멱등)."""
    if get_setting("defect_types_seeded_20260908") == "1":
        return
    for name in ["치수불량", "외관불량", "기능불량", "재질불량", "수량불량", "기타"]:
        add_defect_type(name)
    set_setting("defect_types_seeded_20260908", "1")


# ---------- 불량 유형 마스터 (NCR 폼 드롭다운, 생성/수정/삭제 팝업에서 씀) ----------

def list_defect_types():
    """등록된 불량 유형명 전체 (가나다순)."""
    conn = get_conn()
    rows = conn.execute("SELECT name FROM defect_types ORDER BY name").fetchall()
    conn.close()
    return [r["name"] for r in rows]


def list_improvement_defect_categories():
    """개선요청서 불량유형 체크박스 목록 — list_defect_types()(NCR과 공용 마스터)를
    그대로 쓰되 '기타'만 맨 끝으로 옮긴다(2026-09-16 사용자 확정: "개선요청서에
    있는 불량 유형들을 기존에 등록되어있는 불량유형의 내용으로 통일" — 예전엔
    이 앱 두 곳(app.py/report_builder.py)에 NCR 마스터와 이름도 항목 구성도 다른
    5종이 따로 하드코딩돼 있었다). '기타'는 옆에 항상 자유입력칸이 붙는 항목이라
    마스터의 가나다순 그대로 두면 중간에 섞여 나와 어색해서 여기서만 보정 —
    app.py(폼 렌더/검증)와 report_builder.py(엑셀 생성) 둘 다 이 함수 하나를 쓴다
    (CLAUDE.md 8-1절)."""
    cats = list_defect_types()
    if "기타" not in cats:
        return cats
    return [c for c in cats if c != "기타"] + ["기타"]


def ensure_improvement_defect_categories_migration_20260916():
    """개선요청서 불량유형을 NCR 마스터로 통일하기 전(2026-09-16 이전)에 만들어진
    기존 레코드는 `defect_categories`에 옛 5종("치수 불량"/"표면 결함"/
    "기능 부적합"/"포장 손상") 문자열이 그대로 저장돼 있다 — 새 마스터
    (치수불량/외관불량/기능불량/재질불량/수량불량/기타, 공백 없음)와 문자열이
    달라서 엑셀을 다시 발행해도 체크박스가 하나도 안 켜져 보이는 문제가 실제로
    있었다(사용자가 실제 발행 PDF로 발견). 저장된 문자열을 새 이름으로 1회
    치환한다(멱등). '포장 손상'은 새 마스터에 대응하는 항목이 없어서 그냥
    제거한다(사용자가 원하면 나중에 새 항목으로 등록하면 됨, 임의로 지어내지 않음)."""
    if get_setting("improvement_defect_categories_migrated_20260916") == "1":
        return
    mapping = {
        "치수 불량": "치수불량",
        "표면 결함": "외관불량",
        "기능 부적합": "기능불량",
        "포장 손상": None,  # 새 마스터에 대응 항목 없음 — 제거
    }
    conn = get_conn()
    rows = conn.execute(
        "SELECT id, defect_categories FROM improvement_requests WHERE defect_categories IS NOT NULL"
    ).fetchall()
    for row in rows:
        old_cats = [c.strip() for c in (row["defect_categories"] or "").split(",") if c.strip()]
        new_cats = []
        changed = False
        for c in old_cats:
            if c in mapping:
                changed = True
                mapped = mapping[c]
                if mapped:
                    new_cats.append(mapped)
            else:
                new_cats.append(c)
        if changed:
            conn.execute("UPDATE improvement_requests SET defect_categories=? WHERE id=?",
                        (",".join(new_cats), row["id"]))
    conn.commit()
    conn.close()
    set_setting("improvement_defect_categories_migrated_20260916", "1")


def ensure_material_category_merge_20260916():
    """자재 분류 마스터에 '★완전포장제품'과 '★완포장 제품'이 표기만 다른 같은
    분류로 두 개 나뉘어 있던 걸(2026-09-10 BOM 자동분류 때 키워드 매칭이 두 형태를
    각각 새 분류로 등록해버린 것으로 추정) 사용자 요청으로 '★완포장 제품' 하나로
    합친다. materials.category가 옛 이름인 행을 새 이름으로 바꾸고, 마스터에서
    옛 이름 항목을 제거(새 이름 항목이 마스터에 없으면 새로 등록). 1회성, 멱등."""
    if get_setting("material_category_merge_20260916") == "1":
        return
    OLD_NAME = "★완전포장제품"
    NEW_NAME = "★완포장 제품"
    add_material_category(NEW_NAME)
    conn = get_conn()
    conn.execute("UPDATE materials SET category = ? WHERE category = ?", (NEW_NAME, OLD_NAME))
    conn.execute("DELETE FROM material_categories WHERE name = ?", (OLD_NAME,))
    conn.commit()
    conn.close()
    set_setting("material_category_merge_20260916", "1")


def add_defect_type(name):
    """불량 유형명을 마스터에 등록. 이미 있으면(공백/대소문자 무시) 기존 정본 표기를
    반환. name이 빈 값이면 None. (material_categories.add_material_category와 동일 관례)"""
    name = (name or "").strip()[:30]  # 프론트 maxlength=30과 맞춰 서버에서도 강제
    if not name:
        return None
    conn = get_conn()
    for row in conn.execute("SELECT name FROM defect_types").fetchall():
        if row["name"].strip().casefold() == name.casefold():
            conn.close()
            return row["name"]
    conn.execute("INSERT INTO defect_types (name) VALUES (?)", (name,))
    conn.commit()
    conn.close()
    return name


def rename_defect_type(old_name, new_name):
    """불량 유형명 변경 — 이미 발행된 NCR에 쓰인 값도 새 이름으로 같이 갱신한다
    (과거 통계·목록에서 라벨이 끊기지 않게, rename_material()과 동일한 사고방식).
    반환: (성공여부, 에러메시지)."""
    old_name = (old_name or "").strip()
    new_name = (new_name or "").strip()[:30]  # 프론트 maxlength=30과 맞춰 서버에서도 강제
    if not new_name:
        return False, "이름을 입력해줘."
    conn = get_conn()
    if new_name.casefold() != old_name.casefold():
        dup = conn.execute("SELECT 1 FROM defect_types WHERE name = ?", (new_name,)).fetchone()
        if dup:
            conn.close()
            return False, f"'{new_name}'은 이미 있는 이름이야."
    existing = conn.execute("SELECT 1 FROM defect_types WHERE name = ?", (old_name,)).fetchone()
    if not existing:
        conn.close()
        return False, "존재하지 않는 항목이야."
    conn.execute("UPDATE defect_types SET name = ? WHERE name = ?", (new_name, old_name))
    conn.execute("UPDATE ncr SET defect_type = ? WHERE defect_type = ?", (new_name, old_name))
    conn.commit()
    conn.close()
    return True, None


def delete_defect_type(name):
    """불량 유형을 마스터에서만 제거한다 — FK가 없어서 이미 발행된 NCR의 defect_type
    값은 그대로 남는다(자재 분류 삭제와 동일한 방식). 목록/드롭다운에서만 안 보이게 된다."""
    conn = get_conn()
    conn.execute("DELETE FROM defect_types WHERE name = ?", ((name or "").strip(),))
    conn.commit()
    conn.close()


# ---------- 데이터 점검 (관리자 전용) ----------

def data_health_report():
    """자재·규격 데이터에서 검사·성적서를 망가뜨릴 수 있는 것들을 찾아서 모아준다.
    각 항목: {key, title, desc, severity('high'/'mid'/'low'), rows:[...], columns:[...]}"""
    conn = get_conn()

    def q(sql, params=()):
        return [dict(r) for r in conn.execute(sql, params).fetchall()]

    checks = []

    checks.append({
        "key": "no_limit",
        "severity": "high",
        "title": "규격(하한·상한)이 비어 있는 숫자 항목",
        "desc": "판정 기준이 없어서 측정값을 넣어도 합격/불합격을 낼 수 없어. 검사는 되지만 승인이 막혀.",
        "columns": ["자재번호", "항목", "규격 표기", "검사방법"],
        "rows": q("""SELECT material_no, item_name, spec_display, inspect_method
                       FROM specs
                      WHERE judge_type='numeric' AND lower_limit IS NULL AND upper_limit IS NULL
                      ORDER BY material_no, item_order"""),
    })

    checks.append({
        "key": "no_spec",
        "severity": "high",
        "title": "규격 항목이 하나도 없는 자재",
        "desc": "입고돼도 검사를 시작할 수 없어. 자재 관리에서 항목을 등록해줘.",
        "columns": ["자재번호", "자재명", "도면"],
        "rows": q("""SELECT m.material_no, m.material_name,
                            CASE WHEN m.drawing_file IS NOT NULL AND m.drawing_file!='' THEN '있음' ELSE '없음' END AS drawing
                       FROM materials m
                      WHERE NOT EXISTS (SELECT 1 FROM specs s WHERE s.material_no = m.material_no)
                      ORDER BY m.material_no"""),
    })

    checks.append({
        "key": "bad_material_no",
        "severity": "mid",
        "title": "자재번호가 이상한 자재",
        "desc": "자재번호 자리에 한글 품명이나 '-' 같은 값이 들어가 있어. "
                "도면번호 자동계산(A+번호)과 도면 파일 매칭이 안 돼.",
        "columns": ["자재번호", "자재명", "규격 수", "검사 이력"],
        "rows": q("""SELECT m.material_no, m.material_name,
                            (SELECT COUNT(*) FROM specs s WHERE s.material_no=m.material_no) AS spec_count,
                            (SELECT COUNT(*) FROM inspections i WHERE i.material_no=m.material_no) AS insp_count
                       FROM materials m
                      WHERE LENGTH(TRIM(m.material_no)) <= 2
                         OR m.material_no GLOB '*[가-힣]*'
                      ORDER BY m.material_no"""),
    })

    checks.append({
        "key": "assembly_missing_part",
        "severity": "high",
        "title": "조립품 파츠인데 자재로 등록 안 된 번호",
        "desc": "입고 때 조립품이 펼쳐지면 이 번호로 입고 줄이 생기는데, 자재가 없어서 검사를 못 해.",
        "columns": ["조립품", "순서", "파츠 자재번호"],
        "rows": q("""SELECT am.assembly_no, ac.component_order, ac.component_no
                       FROM assembly_components ac
                       JOIN assembly_masters am ON am.id = ac.assembly_id
                      WHERE NOT EXISTS (SELECT 1 FROM materials m WHERE m.material_no = ac.component_no)
                      ORDER BY am.assembly_no, ac.component_order"""),
    })

    gauge_count = conn.execute("SELECT COUNT(*) FROM gauge_master").fetchone()[0]
    checks.append({
        "key": "gauge_master_empty",
        "severity": "high" if gauge_count == 0 else "low",
        "title": "계측기 마스터 비어 있음",
        "desc": "등록된 계측기가 없으면 교정 유효기간 임박(D-15) 경고가 절대 뜨지 않는다. "
                "'경고 없음 = 문제 없음'으로 오해하기 쉬워서 위험하다.",
        "columns": ["상태"],
        "rows": ([{"상태": "계측기가 하나도 등록돼 있지 않아 — 설정 → 계측기 관리에서 등록해줘"}]
                 if gauge_count == 0 else []),
    })

    checks.append({
        "key": "no_integrity_hash",
        "severity": "low",
        "title": "무결성 검증 기준값이 없는 승인 성적서",
        "desc": "위변조 검증 기능이 생기기 전에 승인된 건이라 '승인 당시 그대로인지' 확인할 수 없다. "
                "새로 승인되는 건부터는 자동으로 기준값이 저장된다.",
        "columns": ["성적서 번호", "자재번호", "업체", "검사일"],
        "rows": q("""SELECT id, material_no, supplier, inspect_date
                       FROM inspections
                      WHERE status='approved' AND (content_hash IS NULL OR content_hash='')
                      ORDER BY id DESC"""),
    })

    checks.append({
        "key": "no_drawing",
        "severity": "low",
        "title": "도면이 연결 안 된 자재",
        "desc": "검사 화면에서 도면을 못 봐. 도면 파일명이 자재번호와 같으면 자동으로 연결돼.",
        "columns": ["자재번호", "자재명"],
        "rows": q("""SELECT material_no, material_name FROM materials
                      WHERE drawing_file IS NULL OR drawing_file=''
                      ORDER BY material_no"""),
    })

    checks.append({
        "key": "dup_spec",
        "severity": "low",
        "title": "같은 자재 안에 규격 표기가 똑같은 항목",
        "desc": "성적서에 같은 내용이 두 줄로 나와. 일부러 나눈 거면 그대로 둬도 돼.",
        "columns": ["자재번호", "항목들", "규격 표기"],
        "rows": q("""SELECT material_no,
                            GROUP_CONCAT(item_name, ', ') AS item_names,
                            spec_display
                       FROM specs
                      WHERE spec_display IS NOT NULL AND TRIM(spec_display) != ''
                      GROUP BY material_no, TRIM(spec_display)
                     HAVING COUNT(*) > 1
                      ORDER BY material_no"""),
    })

    checks.append({
        "key": "star_mismatch",
        "severity": "low",
        "title": "중요항목(*) 표시와 AQL이 안 맞는 항목",
        "desc": "원칙은 'AQL 0.65 = 중요항목(*)'인데, *는 붙어 있고 AQL은 0.65가 아니야. "
                "성적서에는 *가 그대로 나가니 AQL 쪽을 확인해줘.",
        "columns": ["자재번호", "항목", "AQL", "규격 표기"],
        "rows": q("""SELECT material_no, item_name, aql, spec_display
                       FROM specs
                      WHERE TRIM(item_name) LIKE '*%' AND (aql IS NULL OR CAST(aql AS TEXT) != '0.65')
                      ORDER BY material_no"""),
    })

    conn.close()
    return checks
